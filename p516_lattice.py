"""P5.16 convention change C5 -- the general (node, year) investment lattice (pure; no Pyomo import).

Authority: P5_16_CONVENTION_CHANGE_SPEC_v4.md section 1 C5; PLANNER_BRIEF_2026-09-13.md Addendum 88 R3 (the
neighbourhood, authoritative wording) and the author's decisions 4-5 of 2026-10-30; STEP4_DFO_METHOD.md 1.1-1.2, 5.4.
Gate: p516_lattice_gate.py (G-CC-4, zero solves).

DECISION VECTOR. z in Z^{2|N||Y|}: per (node n, representative year y) a pair (zP, zE) of granules
GRANULE_P_MVA = 0.25 MVA and GRANULE_E_MWH = 0.5 MWh (the values of p515_s47_phase_b_record.GRANULE_P_MVA /
GRANULE_E_MWH, l.186-187; the gate asserts equality). Order: node-major (instance node order), then year (instance
order), then (zP, zE) -- the ordering of p56a_oracle.vector_to_candidate. Nodes and years are never literals here: they
come from the constructor (`LatticeInstance`) or from a case file (`load_instance` / `load_instance_from_case_dict`).

ADMISSIBILITY (STEP4 1.2 + spec v4 C5; mirrors the rows of shared_energy_storage_data._build_master_problem):
  - sign: zP >= 0, zE >= 0 per cohort (the master's NonNegativeReals domain);
  - zP = 0 <=> zE = 0 per cohort;
  - duration per cohort: min_ratio * P <= E <= max_ratio * P (master `energy_storage_power_to_energy_factor`; 2-4 h
    with the committed params file, i.e. zP <= zE <= 2 zP on the lattice);
  - cumulative rated energy per node and block: sum over cohorts y_inv alive in block y of E_{n,y_inv} <= max_capacity
    (master `energy_storage_maximum_capacity` on `es_e_rated`, with the cohort window of `rated_e_capacity`:
    tcal_norm = round(t_cal / Years[y_inv]), alive in blocks y_inv .. min(y_inv + tcal_norm, |Y|) - 1);
  - budget I(x) <= B, I(x) = sum_n sum_y sum_m (1+d)^-(y-y1) * omega_m * (c^S_{m,y} P_{n,y} + c^E_{m,y} E_{n,y})
    (master `energy_storage_investment` / `investment_cost`, evaluated in the master's own term order).
  `reasons(z)` keeps the short-circuit semantics of p515_s47_phase_b_record.Lattice.reasons: a cohort with a sign or
  P = 0 <=> E = 0 failure is not checked further and suppresses its node's cumulative check; the budget is evaluated
  only when no structural reason exists. `family_violations(z)` evaluates every family independently (oracle use).
  Reason classes: 'bound' (sign, cumulative max_capacity), 'duration' (P = 0 <=> E = 0, E/P window), 'budget'.

CANONICAL FORM / KEY SCHEMA v2: {'schema': 'v2', 'nodes': {'5': {'2025': [s, e], ...}, ...}} -- floats, string keys,
ascending node and year order, zero entries included. `candidate_key` = p515_s44_campaign_harness.candidate_key
(sha256 of the compact key-sorted JSON). x = 0 ALIAS: in the old single-cohort schema x = 0 had one representative per
investment year and was canonicalised to investment_year = y1 (STEP4 5.4); in schema v2 there is no year field, so x = 0
has exactly one canonical form (all entries [0.0, 0.0]) and one key, labelled 'x0'; every old x = 0 representation
maps to it. Old keys are not translated (spec v1 C5).

NEIGHBOURHOOD N(z) (Addendum 88 R3.1-R3.2), listed in this deterministic order:
  (a) coordinate steps +zP, -zP, +zE, -zE on every ACTIVE cohort (node order, then year order);
  (b) activation of every inactive cohort to (0.25, 0.5) and (0.25, 1.0); deactivation of every active smallest cohort
      (0.25 MVA with E in {0.5, 1.0}) to (0, 0);
  (c) one-block shifts of every active cohort to y - 1 then y + 1 at the same (P, E), added to the target cohort.
  Raw size <= 2|N||Y| + 5A (= 30 + 5A on the five-block instance: 4A + 2(15 - A) + A + 2A), asserted before return.
  Inadmissible points are dropped and counted by reason class; coordinate steps on inactive cohorts are never listed.
RESOLUTION BY BOUND (R3.3): `resolution_by_bound(points, f_lb, margin, ub)` -- pure; x' resolved iff
  F_LB(x') - margin >= UB(x_hat).
RUN-RULE CONSTANTS (R3.4): COMPLETION_CAP = 60 (stop for review above it), N_MAX = 120 new evaluations per run.
"""

import hashlib
import json
import os

from p515_s44_campaign_harness import candidate_key as _harness_candidate_key  # stdlib only at import

GRANULE_P_MVA = 0.25  # = p515_s47_phase_b_record.GRANULE_P_MVA (l.186)
GRANULE_E_MWH = 0.5   # = p515_s47_phase_b_record.GRANULE_E_MWH (l.187)
KEY_SCHEMA = 'v2'
X0_LABEL = 'x0'
COMPLETION_CAP = 60   # Addendum 88 R3.4
N_MAX = 120           # Addendum 88 R3.4
ACTIVATION_PAIRS = ((1, 1), (1, 2))      # (0.25 MVA, 0.5 MWh), (0.25 MVA, 1.0 MWh) -- R3.1(b)
SMALLEST_COHORTS = frozenset(ACTIVATION_PAIRS)  # deactivated to (0, 0) -- R3.1(b)
REASON_CLASSES = ('bound', 'duration', 'budget')
COST_SHEETS = {'scenarios': 'Scenarios', 'power': 'Investment Cost, Power', 'energy': 'Investment Cost, Energy'}


def candidate_key(canonical):
    """sha256 of the canonical form's compact key-sorted JSON -- p515_s44_campaign_harness.candidate_key, reused."""
    return _harness_candidate_key(canonical)


def sha256_file(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b''):
            digest.update(chunk)
    return digest.hexdigest()


# ======================================================================================================================
#  the instance (data only)
# ======================================================================================================================
class LatticeInstance:
    """Everything the lattice needs, passed in. `years` ascending; `block_widths[y]` = the case file's Years[y];
    `prob_market` = omega_m; `cost_power[m][y]`, `cost_energy[m][y]` = undiscounted EUR/MVA, EUR/MWh."""

    def __init__(self, nodes, years, block_widths, t_cal_years, max_capacity_mwh, min_ratio, max_ratio, budget_eur,
                 discount_factor, prob_market, cost_power, cost_energy, sources=None):
        self.nodes = tuple(int(n) for n in nodes)
        self.years = tuple(int(y) for y in years)
        if len(set(self.nodes)) != len(self.nodes) or not self.nodes:
            raise ValueError(f'nodes must be distinct and non-empty: {self.nodes}')
        if list(self.years) != sorted(set(self.years)) or not self.years:
            raise ValueError(f'years must be strictly ascending and non-empty: {self.years}')
        self.block_widths = {int(y): int(block_widths[y] if y in block_widths else block_widths[str(y)])
                             for y in self.years}
        self.t_cal_years = float(t_cal_years)
        self.max_capacity_mwh = float(max_capacity_mwh)
        self.min_ratio = float(min_ratio)
        self.max_ratio = float(max_ratio)
        self.budget_eur = float(budget_eur)
        self.discount_factor = float(discount_factor)
        self.prob_market = tuple(float(p) for p in prob_market)
        self.cost_power = {int(m): {int(y): float(v) for y, v in table.items()} for m, table in cost_power.items()}
        self.cost_energy = {int(m): {int(y): float(v) for y, v in table.items()} for m, table in cost_energy.items()}
        for m in range(len(self.prob_market)):
            for y in self.years:
                if y not in self.cost_power.get(m, {}) or y not in self.cost_energy.get(m, {}):
                    raise ValueError(f'cost table lacks scenario {m} year {y}')
        self.sources = dict(sources or {})

    def describe(self):
        return {'nodes': list(self.nodes), 'years': list(self.years),
                'block_widths': {str(y): w for y, w in self.block_widths.items()},
                't_cal_years': self.t_cal_years, 'max_capacity_mwh': self.max_capacity_mwh,
                'min_energy_to_power_ratio': self.min_ratio, 'max_energy_to_power_ratio': self.max_ratio,
                'budget_eur': self.budget_eur, 'discount_factor': self.discount_factor,
                'prob_market': list(self.prob_market),
                'cost_power': {str(m): {str(y): v for y, v in t.items() if y in self.years}
                               for m, t in self.cost_power.items()},
                'cost_energy': {str(m): {str(y): v for y, v in t.items() if y in self.years}
                                for m, t in self.cost_energy.items()},
                'sources': self.sources}


def _read_cost_tables(cost_path, years):
    """The ESS cost workbook read as production reads it (shared_energy_storage_data.py
    `_get_operational_scenarios_info_from_excel_file` / `_get_investment_costs_from_excel_file`, l.1753-1800), with
    openpyxl (cached values): sheet 'Scenarios' row 1: B1 = number of market scenarios, C1.. = probabilities;
    sheets 'Investment Cost, Power' / 'Investment Cost, Energy': row 1 from column B = years, row 1 + i = scenario i."""
    import openpyxl  # local: the module stays import-light
    workbook = openpyxl.load_workbook(cost_path, data_only=True, read_only=True)
    try:
        rows = list(workbook[COST_SHEETS['scenarios']].iter_rows(values_only=True))
        num = int(rows[0][1])
        prob = [float(rows[0][i + 2]) for i in range(num)]
        if round(sum(prob), 2) != 1.00:
            raise ValueError(f'scenario probabilities sum to {sum(prob)}')
        tables = {}
        for kind in ('power', 'energy'):
            rows = list(workbook[COST_SHEETS[kind]].iter_rows(values_only=True))
            header = rows[0]
            tables[kind] = {}
            for i in range(num):
                tables[kind][i] = {}
                for year in years:
                    cols = [j for j in range(1, len(header)) if header[j] is not None and int(header[j]) == year]
                    if not cols:
                        raise ValueError(f'{COST_SHEETS[kind]}: year {year} not found')
                    for j in cols:  # production keeps the last matching column
                        tables[kind][i][year] = float(rows[i + 1][j])
    finally:
        workbook.close()
    return prob, tables['power'], tables['energy']


def load_instance_from_case_dict(case, data_dir, case_source=None):
    """Build the instance from a parsed case dict (e.g. p515_s44_scale_measurement.derive_case output): Years map,
    DistributionNetworks[].connection_node_id (the active DN nodes, case order -- shared_resources_planning.py l.8975,
    9049), DiscountFactor, SharedEnergyStorage.{params_file, data_file} under <data_dir>/SharedESS. Imports no model
    code. Every file read is sha256-recorded in `sources`."""
    years_map = {int(y): int(w) for y, w in case['Years'].items()}
    years = sorted(years_map)
    if list(years_map) != years:
        raise ValueError(f'case Years map is not in ascending order: {list(years_map)}')
    nodes = [int(dn['connection_node_id']) for dn in case['DistributionNetworks']]
    ess_dir = os.path.join(data_dir, 'SharedESS')
    params_path = os.path.join(ess_dir, case['SharedEnergyStorage']['params_file'])
    cost_path = os.path.join(ess_dir, case['SharedEnergyStorage']['data_file'])
    with open(params_path) as handle:
        params = json.load(handle)
    ageing = params.get('ageing') or {}
    t_cal = float(ageing.get('calendar_life_years', 15))  # production default (shared_energy_storage_parameters l.133)
    prob, c_p, c_e = _read_cost_tables(cost_path, years)
    sources = {'case': case_source,
               'params_file': {'path': params_path, 'sha256': sha256_file(params_path),
                               'fields': 'budget, max_capacity, min_energy_to_power_factor, max_energy_to_power_factor, '
                                         'ageing.calendar_life_years'},
               'cost_file': {'path': cost_path, 'sha256': sha256_file(cost_path),
                             'sheets': dict(COST_SHEETS),
                             'cells': ("Scenarios!B1 = number of market scenarios, Scenarios!C1.. = omega_m; "
                                       "'Investment Cost, Power'/'Investment Cost, Energy' row 1 (from column B) = "
                                       "years, row 1 + m = scenario m (EUR/MVA, EUR/MWh, undiscounted)")},
               'num_market_scenarios_case_file': case.get('NumMarketScenarios'),
               'note': ('omega_m is read from the workbook Scenarios sheet, as production does; it is independent '
                        'of the case file NumMarketScenarios')}
    return LatticeInstance(nodes, years, years_map, t_cal, params['max_capacity'], params['min_energy_to_power_factor'],
                           params['max_energy_to_power_factor'], params['budget'], case['DiscountFactor'], prob, c_p, c_e,
                           sources)


def load_instance(case_path):
    """Read a case file (e.g. data/SRP1/SRP1.json) and build the instance; see `load_instance_from_case_dict`."""
    with open(case_path) as handle:
        case = json.load(handle)
    return load_instance_from_case_dict(case, os.path.dirname(os.path.abspath(case_path)),
                                        case_source={'path': case_path, 'sha256': sha256_file(case_path)})


# ======================================================================================================================
#  the lattice (pure)
# ======================================================================================================================
class GeneralLattice:
    """z in Z^{2|N||Y|}; see the module docstring."""

    def __init__(self, instance):
        self.inst = instance
        self.nodes = instance.nodes
        self.years = instance.years
        self.dim = 2 * len(self.nodes) * len(self.years)
        self.cohorts = tuple((n, y) for n in self.nodes for y in self.years)
        # production cohort window (shared_energy_storage_data._build_master_problem l.333-340)
        self.alive_blocks = {}
        for yi, year in enumerate(self.years):
            tcal_norm = round(instance.t_cal_years / instance.block_widths[year])
            self.alive_blocks[year] = tuple(self.years[x] for x in range(yi, min(yi + tcal_norm, len(self.years))))
        # discounted expected unit costs, as the master's coefficients (annualization * omega_m) * c
        self._coef = {}
        for year in self.years:
            annualization = 1 / ((1 + instance.discount_factor) ** (int(year) - int(self.years[0])))
            self._coef[year] = [(annualization * instance.prob_market[m] * instance.cost_power[m][year],
                                 annualization * instance.prob_market[m] * instance.cost_energy[m][year])
                                for m in range(len(instance.prob_market))]

    # ---- coordinates -------------------------------------------------------------------------------------------------
    def x0(self):
        return (0,) * self.dim

    def index(self, node, year):
        return 2 * (self.nodes.index(int(node)) * len(self.years) + self.years.index(int(year)))

    def pair(self, z, node, year):
        i = self.index(node, year)
        return z[i], z[i + 1]

    def with_pair(self, z, node, year, zp, ze):
        z = list(z)
        i = self.index(node, year)
        z[i], z[i + 1] = int(zp), int(ze)
        return tuple(z)

    def pairs(self, z):
        return {(n, y): self.pair(z, n, y) for n, y in self.cohorts}

    def active_cohorts(self, z):
        return [(n, y) for n, y in self.cohorts if self.pair(z, n, y) != (0, 0)]

    def check_z(self, z):
        z = tuple(z)
        if len(z) != self.dim or any(not isinstance(v, int) or isinstance(v, bool) for v in z):
            raise ValueError(f'z must be {self.dim} ints, got {z!r}')
        return z

    # ---- closed forms ------------------------------------------------------------------------------------------------
    def investment_cost(self, z):
        """I(x), the master's `investment_cost` in its own term order (node, year, scenario; power then energy)."""
        total = 0.0
        for n in self.nodes:
            for y in self.years:
                zp, ze = self.pair(z, n, y)
                p, e = GRANULE_P_MVA * zp, GRANULE_E_MWH * ze
                for c_s, c_e in self._coef[y]:
                    total += c_s * p
                    total += c_e * e
        return total

    def cumulative_energy(self, z):
        """{(node, block): sum of E over the cohorts alive in that block} (the master's es_e_rated)."""
        out = {(n, y): 0.0 for n in self.nodes for y in self.years}
        for n in self.nodes:
            for y_inv in self.years:
                e = GRANULE_E_MWH * self.pair(z, n, y_inv)[1]
                for y in self.alive_blocks[y_inv]:
                    out[(n, y)] += e
        return out

    # ---- admissibility -----------------------------------------------------------------------------------------------
    def _cohort_reasons(self, n, y, zp, ze):
        """(reasons, well_formed): sign and P = 0 <=> E = 0 failures stop the cohort's checks (old-Lattice semantics)."""
        p, e = GRANULE_P_MVA * zp, GRANULE_E_MWH * ze
        if zp < 0 or ze < 0:
            return [f'bound: node {n} year {y} negative capacity (P={p:g}, E={e:g})'], False
        if (zp == 0) != (ze == 0):
            return [f'duration: node {n} year {y} P = 0 <=> E = 0 violated (P={p:g}, E={e:g})'], False
        if zp > 0 and not (self.inst.min_ratio * p <= e <= self.inst.max_ratio * p):
            return [f'duration: node {n} year {y} E/P = {e / p:g} h outside '
                    f'[{self.inst.min_ratio:g}, {self.inst.max_ratio:g}] h'], True
        return [], True

    def _cumulative_reason(self, n, cum):
        over = [(y, cum[(n, y)]) for y in self.years if cum[(n, y)] > self.inst.max_capacity_mwh]
        if not over:
            return []
        blocks = ', '.join(f'{y}: {v:g}' for y, v in over)
        return [f'bound: node {n} cumulative rated energy > {self.inst.max_capacity_mwh:g} MWh in blocks [{blocks}]']

    def reasons(self, z):
        """Every violated constraint (empty list = admissible); semantics in the module docstring."""
        z = self.check_z(z)
        out = []
        cum = self.cumulative_energy(z)
        for n in self.nodes:
            node_well_formed = True
            for y in self.years:
                why, well_formed = self._cohort_reasons(n, y, *self.pair(z, n, y))
                out += why
                node_well_formed &= well_formed
            if node_well_formed:
                out += self._cumulative_reason(n, cum)
        if not out:
            i_x = self.investment_cost(z)
            if i_x > self.inst.budget_eur:
                out.append(f'budget: I(x) = {i_x:.2f} > B = {self.inst.budget_eur:g}')
        return out

    def admissible(self, z):
        return not self.reasons(z)

    def family_violations(self, z):
        """Every family evaluated independently (no short-circuit), named after the master rows they mirror."""
        z = self.check_z(z)
        out = {'sign': [], 'p0_iff_e0': [], 'energy_storage_power_to_energy_factor': [],
               'energy_storage_maximum_capacity': [], 'energy_storage_investment': []}
        for n, y in self.cohorts:
            zp, ze = self.pair(z, n, y)
            p, e = GRANULE_P_MVA * zp, GRANULE_E_MWH * ze
            if zp < 0 or ze < 0:
                out['sign'].append((n, y))
            if (zp == 0) != (ze == 0):
                out['p0_iff_e0'].append((n, y))
            if not (e >= self.inst.min_ratio * p) or not (e <= self.inst.max_ratio * p):
                out['energy_storage_power_to_energy_factor'].append((n, y))
        for (n, y), v in self.cumulative_energy(z).items():
            if v > self.inst.max_capacity_mwh:
                out['energy_storage_maximum_capacity'].append((n, y))
        if self.investment_cost(z) > self.inst.budget_eur:
            out['energy_storage_investment'].append('budget')
        return out

    @staticmethod
    def reason_classes(reasons):
        return sorted(r.split(':', 1)[0] for r in reasons)

    # ---- canonical form and keys -------------------------------------------------------------------------------------
    def canonical_of_z(self, z):
        z = self.check_z(z)
        return {'schema': KEY_SCHEMA,
                'nodes': {str(n): {str(y): [GRANULE_P_MVA * self.pair(z, n, y)[0] + 0.0,
                                            GRANULE_E_MWH * self.pair(z, n, y)[1] + 0.0]
                                   for y in sorted(self.years)} for n in sorted(self.nodes)}}

    def z_of_canonical(self, canonical):
        """Inverse of `canonical_of_z`; None when off-lattice, wrong schema, or nodes/years not exactly the instance's."""
        if canonical.get('schema') != KEY_SCHEMA or set(canonical) != {'schema', 'nodes'}:
            return None
        nodes = canonical['nodes']
        if set(nodes) != {str(n) for n in self.nodes}:
            return None
        z = list(self.x0())
        for n in self.nodes:
            per_year = nodes[str(n)]
            if set(per_year) != {str(y) for y in self.years}:
                return None
            for y in self.years:
                s_val, e_val = per_year[str(y)]
                zp, ze = s_val / GRANULE_P_MVA, e_val / GRANULE_E_MWH
                if zp != int(zp) or ze != int(ze):
                    return None
                i = self.index(n, y)
                z[i], z[i + 1] = int(zp), int(ze)
        return tuple(z)

    def canonicalize(self, candidate):
        """Canonical v2 form of {node: {year: (s, e)}} (missing entries are zero) or of a v2 canonical (idempotent).
        Refuses off-lattice values, unknown nodes/years, negative or non-finite values."""
        if isinstance(candidate, dict) and candidate.get('schema') == KEY_SCHEMA:
            candidate = candidate['nodes']
        z = list(self.x0())
        for n_key, per_year in candidate.items():
            n = int(n_key)
            if n not in self.nodes:
                raise ValueError(f'node {n} not in the instance {self.nodes}')
            for y_key, value in per_year.items():
                y = int(y_key)
                if y not in self.years:
                    raise ValueError(f'year {y} not in the instance {self.years}')
                s_val, e_val = float(value[0]), float(value[1])
                for name, v in (('s', s_val), ('e', e_val)):
                    if not (v == v) or v in (float('inf'), float('-inf')) or v < 0.0:
                        raise ValueError(f'node {n} year {y}: invalid {name}={v}')
                zp, ze = s_val / GRANULE_P_MVA, e_val / GRANULE_E_MWH
                if zp != int(zp) or ze != int(ze):
                    raise ValueError(f'node {n} year {y}: ({s_val}, {e_val}) is off the lattice')
                i = self.index(n, y)
                z[i], z[i + 1] = int(zp), int(ze)
        return self.canonical_of_z(tuple(z))

    def key(self, z):
        return candidate_key(self.canonical_of_z(z))

    def label(self, z):
        z = self.check_z(z)
        active = self.active_cohorts(z)
        if not active:
            return X0_LABEL
        return '__'.join(f'n{n}_y{y}_p{GRANULE_P_MVA * self.pair(z, n, y)[0]:g}_e{GRANULE_E_MWH * self.pair(z, n, y)[1]:g}'
                         for n, y in active)

    def production_investment(self, z):
        """{node: {year: {'s': MVA, 'e': MWh}}} -- the shape of candidate_solution['investment']."""
        return {n: {y: {'s': GRANULE_P_MVA * self.pair(z, n, y)[0], 'e': GRANULE_E_MWH * self.pair(z, n, y)[1]}
                    for y in self.years} for n in self.nodes}

    # ---- the neighbourhood (Addendum 88 R3) --------------------------------------------------------------------------
    def size_bound(self, n_active):
        """|N| <= 2|N||Y| + 5A (R3.2: 4A coordinate steps + 2(|N||Y| - A) activations + <= A deactivations + <= 2A
        shifts); 30 + 5A on a 3-node five-block instance."""
        return 2 * len(self.cohorts) + 5 * n_active

    def neighbourhood(self, z_inc):
        z_inc = self.check_z(z_inc)
        if self.reasons(z_inc):
            raise ValueError(f'incumbent is inadmissible: {self.reasons(z_inc)}')
        active = self.active_cohorts(z_inc)
        raw = []
        # (a) coordinate steps on active cohorts
        for n, y in active:
            zp, ze = self.pair(z_inc, n, y)
            for coord, delta in (('zP', +1), ('zP', -1), ('zE', +1), ('zE', -1)):
                new = (zp + delta, ze) if coord == 'zP' else (zp, ze + delta)
                raw.append(('a', f'coord:{coord}{delta:+d}:n{n}:y{y}', self.with_pair(z_inc, n, y, *new)))
        # (b) activations of inactive cohorts, then deactivations of active smallest cohorts
        for n, y in self.cohorts:
            if (n, y) in active:
                continue
            for zp, ze in ACTIVATION_PAIRS:
                raw.append(('b', f'activate:n{n}:y{y}:p{GRANULE_P_MVA * zp:g}_e{GRANULE_E_MWH * ze:g}',
                            self.with_pair(z_inc, n, y, zp, ze)))
        for n, y in active:
            if self.pair(z_inc, n, y) in SMALLEST_COHORTS:
                raw.append(('b', f'deactivate:n{n}:y{y}', self.with_pair(z_inc, n, y, 0, 0)))
        # (c) one-block shifts, merged additively into the target cohort
        for n, y in active:
            yi = self.years.index(y)
            zp, ze = self.pair(z_inc, n, y)
            for t in (yi - 1, yi + 1):
                if not 0 <= t < len(self.years):
                    continue
                target = self.years[t]
                tp, te = self.pair(z_inc, n, target)
                new = self.with_pair(self.with_pair(z_inc, n, y, 0, 0), n, target, tp + zp, te + ze)
                raw.append(('c', f'shift:n{n}:y{y}->y{target}' + (':merged' if (tp, te) != (0, 0) else ''), new))
        bound = self.size_bound(len(active))
        assert len(raw) <= bound, f'|N| raw {len(raw)} > closed-form bound {bound}'
        points, dropped, by_class, seen, duplicates = [], [], {}, {}, []
        for part, move, z in raw:
            why = self.reasons(z)
            entry = {'part': part, 'move': move, 'z': list(z), 'label': self.label(z), 'key': self.key(z),
                     'admissible': not why, 'reasons': why}
            if why:
                dropped.append(entry)
                for cls in sorted(set(self.reason_classes(why))):
                    by_class[cls] = by_class.get(cls, 0) + 1
                continue
            if entry['key'] in seen or z == z_inc:
                duplicates.append({'move': move, 'same_as': seen.get(entry['key'], 'incumbent')})
                continue
            seen[entry['key']] = move
            entry['I_x_eur'] = self.investment_cost(z)
            points.append(entry)
        return {'incumbent': list(z_inc), 'incumbent_label': self.label(z_inc), 'incumbent_key': self.key(z_inc),
                'A': len(active), 'size_bound': bound, 'n_raw': len(raw),
                'n_raw_by_part': {p: sum(1 for r in raw if r[0] == p) for p in ('a', 'b', 'c')},
                'n_admissible': len(points), 'n_dropped': len(dropped), 'dropped_by_reason_class': by_class,
                'duplicates': duplicates, 'over_completion_cap': len(points) > COMPLETION_CAP,
                'completion_cap': COMPLETION_CAP, 'n_max': N_MAX, 'points': points, 'dropped': dropped}


def resolution_by_bound(points, f_lb, margin, ub):
    """Addendum 88 R3.3 (pure). `points`: neighbourhood entries (each with 'key'); `f_lb`: {key: F_LB(x')} or a
    callable key -> F_LB (None = no bound); `margin`: a number, or {key: margin} / callable; `ub`: UB(x_hat).
    x' is resolved without evaluation iff F_LB(x') - margin(x') >= ub. Returns the split and the report sentence."""
    lb_of = f_lb if callable(f_lb) else (lambda k: f_lb.get(k))
    if callable(margin):
        margin_of = margin
    elif isinstance(margin, dict):
        margin_of = margin.get
    else:
        margin_of = (lambda k: margin)
    resolved, unresolved = [], []
    for point in points:
        k = point['key']
        lb, m = lb_of(k), margin_of(k)
        ok = lb is not None and m is not None and (lb - m) >= ub
        (resolved if ok else unresolved).append({'key': k, 'move': point.get('move'), 'F_LB': lb, 'margin': m,
                                                 'F_LB_minus_margin': None if lb is None or m is None else lb - m,
                                                 'UB': ub})
    return {'k': len(resolved), 'n': len(points), 'resolved': resolved, 'unresolved': unresolved,
            'statement': f'resolved by the lower bound: {len(resolved)} of {len(points)}'}
