"""P5.15 Addendum 68 order, Planner task W169 -- CLOSURE SLACK AND ESSO P-NET SLACK AT THE CERTIFIED POINT, EVERY
EVALUATION IN T1-T11, FROM THE RECORDS. ZERO SOLVES, NO MODEL LOADS. A NEW FILE; no production module, record, frozen
JSON, manuscript file or committed artifact is edited.

AUTHORITY. PLANNER_BRIEF_2026-09-13.md Addendum 68; STEP6_ROUND1_CORRECTIONS.md section A.3 / A.4 and section C (W169);
TASKS.md (Addendum 68 order). The two manuscript comments this answers are read (only) from the Overleaf clone
manuscript/6a67305f25e8348fb71380c3/main.tex at Overleaf 42794d4 (sha256 pinned below), l. 747 and l. 857.

THE TWO QUANTITIES (Pyomo names located in the code at run time, file:line written to the output):
  * CLOSURE SLACK s+ / s- = `slack_shared_es_soc_final_up` / `slack_shared_es_soc_final_down` (network.py, Var over
    (e, s_m, s_o); only the first scenario pair is referenced -- `sess_na_scenario` / `sess_row_is_duplicate`), in the
    TSO and every DSO block model, in the closure row of `sess_soc_final_rule` (model_construction_helpers.py),
    bounded [0, 0.05 E^Av + 1e-5] p.u. by `configure_shared_ess_operational_state` ([0, 0] for a zero-capacity unit),
    penalised baseMVA * PENALTY_SHARED_ESS_BALANCE * (s+ + s-) in the block objective
    (`shared_ess_day_balance_slack_penalty`). Units: p.u. energy (x baseMVA = MWh).
  * ESSO P-NET SLACK sigma+ / sigma- = `slack_es_pnet_up` / `slack_es_pnet_down` (shared_energy_storage_data.py, Var
    over (y, d, t) of each node's ESSO model), in `energy_storage_operation_agg`
    (es_pnet = sum_cohorts(pch - pdch) + sigma+ - sigma-), penalised PENALTY_ESSO_SLACK * (sigma+ + sigma-) in the ESSO
    objective. Units: MW (the ESSO model is in MW / MVA / MWh, not p.u.).

THE CELLS. Every evaluation the frozen Step 6 tables use (frozen_step6_tables_v1_590088fe.json): the 49 T2 rows
(tables.cells) and the 12 certificates appended to T2 for T4 / T5 / T10 (tables.cells_appended_w160) -- these are
also every cell T1, T3, T7, T8, T9 and T10 read; the six T6 benchmark arm starts (w116 NRF r2 arms; the T6
"coordinated" row is the T2 reference ref:7aa017f0); the T11 3 x 3 pair (x = 0 and the unit, w91 pair) and the 3 x 3
x = 0 continuation (w98) whose record the T11 post-hoc descent D was fitted on.

THE POINT. Certified cell: k* (frozen `k_star`). Uncertified cell: the terminal / cap cycle (frozen `end_cycle`),
labelled. T11: the cell's own certification cycle (production exit). w98 continuation: its terminal cycle, labelled
(not a certified point of any tabulated value). T6 arms: the arm's evaluation (no ADMM cycles).

ORDER OF WORK (CLAUDE.md rule eleven -- capture path asserted before extracting).
  1. INVENTORY, per cell, of every place the two quantities could be captured: every file of the eval directory
     (committed or not; JSON parsed and key-scanned when <= 40 MB, JSONL first records key-scanned, logs <= 50 MB
     pattern-counted, workbooks' sheet names read with openpyxl, pickles LISTED, never opened), the campaign
     directory's top-level files, and the esso_capture directory (the file of cycle P per node).
  2. CAPTURE PLAN: per cell and quantity, the single source used and why, written to capture_plan.json BEFORE any
     value is extracted; every planned file must exist (precondition fault otherwise) and its sha256 is recorded in
     manifest_inputs_sha256.json then (re-hashed after extraction, must be unchanged).
  3. EXTRACTION from the planned sources only. Committed records are primary; an uncommitted file is read only where it
     carries what no committed record does, and is labelled "uncommitted". A zero is never inferred from absence: a
     quantity no record carries at P is reported "not carried at P", with the pickles that may hold it listed
     ("only in pickle -- not read") or, where no pickle holds cycle P either, "no record or pickle holds cycle P".

DEFINITIONS (stated in every output).
  * closure, per block (TSO|y|d or DSO|n|y|d): component_levels_terminal.json blocks[block].unweighted.
    shared_ess_day_balance_slack is sum_s omega_s * baseMVA * PENALTY_SHARED_ESS_BALANCE * sum_e (s+ + s-) (the
    scenario-free copy; sum_s omega_s = 1), so  sum_e (s+ + s-) [MWh] = value [EUR] / PENALTY_SHARED_ESS_BALANCE
    [EUR/MWh]. A DSO block holds one unit (its node's); a TSO block holds every unit, so its value is the SUM over the
    units; a per-unit upper bound there is value + (n_active - 1) * 2 * FLOOR, FLOOR = 1e-8 p.u. * baseMVA MWh (IPOPT
    bound_relax_factor 1e-8, honor_original_bounds no: data/SRP1/Results/P515S48/negative_penalty_check/) -- a
    zero-capacity unit's pair has bounds [0, 0].
  * closure RELATIVE = (s+ + s-) / E^Av(unit, block year), E^Av = evaluation_record.json storage_per_node[n].
    published_available_capacity_terminal[year].e_available (MWh; the available energy the agent published, which
    sets the slack bound); a TSO block uses the per-unit bound and the smallest E^Av of its active units.
  * sigma, per (e, y, d, t): esso_capture/<label>/node<e>_cycle<P>.jsonl slack_pnet_up + slack_pnet_down (MW; written
    for active cohort-periods only); sigma RELATIVE = sigma / S(e, y), S = the cohort s_max of that row
    (es_s_rated_per_unit), cross-checked against published s_available. The committed record carries only the SUM
    over every node and (y, d, t): component_levels_terminal.json esso_feasibility_violation_D3
    (= SharedEnergyStorageData.get_feasibility_violation); from it a per-element upper bound D3 + (N - 1) * 2e-8 MW
    (same relaxed-bound floor, N = 3 nodes x |Y| x |D| x 24) is reported beside, labelled.
  * PREDICTION (expert, Addendum 68): both slacks are zero to solver tolerance (<= 1e-6 relative) at every certified
    point. Scored per quantity over the certified points whose records carry it, in two readings: SIGNED (the slack
    sum <= 1e-6 x E^Av, resp. S -- the inequality as written) and ABSOLUTE (|slack sum| <= 1e-6 x E^Av, resp. S;
    the solver's bound relaxation leaves the sums slightly NEGATIVE, so the absolute reading measures that residue).
  * the min(pch, pdch) detector (Addendum 3 item 1, remedy (h)): `_get_esso_complementarity_diagnostics` /
    `_complementarity_ratio_for_model` (shared_energy_storage_data.py), called after every successful ESSO solve in
    the production ESSO solve path; its per-solve outcome is recorded by the campaign harness in
    g_<label>.json esso_complementarity_diagnostics_by_round and leak_classification_<label>.jsonl. Reported at P per
    node: complementarity_ratio_max = max over active cohort-periods of min(pch, pdch) / s_max, and its argmax.

GUARDS. SolveProfileGuard(permitted=()) installed BEFORE any other project import and verified at exactly 0;
pickle.load / pickle.loads blocked for the whole run, counters verified at 0. git reads are `git rev-parse` /
`git ls-files` / `git status --porcelain` / `git grep` only.

MODE (repo root, canonical interpreter; attached, both streams captured; mkdir refuses an existing directory):
    mkdir data/SRP1/Results/P515S53/w169_slack_inventory && \\
    /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s53_w169_slack_inventory.py \\
        > data/SRP1/Results/P515S53/w169_slack_inventory/launch.log 2>&1
    /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_gate_result_bool_typing_test.py \\
        --out data/SRP1/Results/P515S53/w169_slack_inventory/w169_bool_typing_test.json \\
        > data/SRP1/Results/P515S53/w169_slack_inventory/w169_bool_typing_test.log 2>&1
    /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s53_w169_slack_inventory.py --post-run
  --out-dir DIR   trial runs only (absolute or repo-relative); the directory must hold nothing but launch.log.
Exit: 0 = written, integrity checks hold, guards at 0 (a failed prediction or an uncovered cell is a FINDING and never
changes the exit code); 3 = written, an integrity check failed (listed in the JSON and the log); 1 = precondition or
guard fault.
"""
import argparse
import glob
import hashlib
import json
import os
import pickle
import subprocess
import sys
import time
from datetime import datetime, timezone

REPO = os.path.dirname(os.path.abspath(__file__))
if REPO not in sys.path:
    sys.path.insert(0, REPO)
os.chdir(REPO)

from p513_solve_profile_guard import SolveProfileGuard  # noqa: E402

GUARD = SolveProfileGuard(permitted=(), label='P5.15 W169 slack inventory (never solves)').install()

PICKLE_COUNTS = {'load': 0, 'loads': 0}


def _blocked_load(*_a, **_k):
    PICKLE_COUNTS['load'] += 1
    raise RuntimeError('W169: pickle.load called -- no model loads are permitted')


def _blocked_loads(*_a, **_k):
    PICKLE_COUNTS['loads'] += 1
    raise RuntimeError('W169: pickle.loads called -- no model loads are permitted')


pickle.load, pickle.loads = _blocked_load, _blocked_loads

import gate_result_io as GRIO  # noqa: E402
import definitions as DEF  # noqa: E402
from model_construction_helpers import shared_ess_capacity_is_inactive  # noqa: E402

# ======================================================================================================================
#  constants
# ======================================================================================================================
SCRIPT_REL = os.path.basename(__file__)
RES = os.path.join('data', 'SRP1', 'Results')
S53 = os.path.join(RES, 'P515S53')
FZ_REL = os.path.join(S53, 'w160_step6_frozen', 'frozen_step6_tables_v1_590088fe.json')
FZ_SHA = '590088fe6b364c265c97998c5491edea7ca70baad0dbe2ba5150273656d9b6f4'
CLONE_REL = os.path.join('manuscript', '6a67305f25e8348fb71380c3')
MAIN_REL = os.path.join(CLONE_REL, 'main.tex')
MAIN_SHA = '7effd8983df50bdce94e1bc2de07d95d79098d82ce6e0fa201ce711b515230b7'
CLONE_HEAD_EXPECTED = '42794d4'
CONFIRM_LINES = {747: 'closure', 857: 'sigma'}
CASE_SRP1_REL = os.path.join('data', 'SRP1', 'SRP1.json')
CASE_3X3_REL = os.path.join(S53, 'w89_3x3', 'instance', 'SRP1__s53_3x3.json')
INSTANCE_3X3_RECORD_REL = os.path.join(S53, 'w89_3x3', 'instance', 'instance_record.json')
ESS_PARAMS_REL = os.path.join('data', 'SRP1', 'SharedESS', 'SRP1_ESS_Params.json')
NET_PARAMS_RELS = [os.path.join('data', 'SRP1', n, f'{n}_params.json') for n in ('case9', 'case33_1', 'case33_2', 'case33_3')]
NET_2025_RELS = [os.path.join('data', 'SRP1', n, f'{n}_2025.json') for n in ('case9', 'case33_1', 'case33_2', 'case33_3')]
S48_REL = os.path.join(RES, 'P515S48', 'negative_penalty_check', 'negative_penalty_check.json')

T11_EVALS = [
    ('t11_3x3_x0', os.path.join(S53, 'w90_3x3', 'campaign_s53_w91_3x3_pair', 'evals', 'f6e9cd53fdbb8ee8_x0'),
     'certified_point', 'T11: 3 x 3 x = 0 cell of the value V (w91 pair)'),
    ('t11_3x3_unit', os.path.join(S53, 'w90_3x3', 'campaign_s53_w91_3x3_pair', 'evals', 'c82522f470b35b58_n7_4h_e1'),
     'certified_point', 'T11: 3 x 3 unit cell (n7_4h_e1) of the value V (w91 pair)'),
    ('t11_3x3_x0_continuation', os.path.join(S53, 'w98_continuation', 'campaign_s53_w98_x0_continuation_r2', 'evals',
                                             '25b92ae0f1f2c02e_x0_cont'),
     'continuation_terminal', 'T11: 3 x 3 x = 0 continuation (w98), the record the post-hoc descent D was fitted on'),
]
T6_ARMS = [(f't6_{arm}_{start}', os.path.join(S53, 'w116_benchmark_nrf', f'nrf_arm_{arm}_{start}_r2'), arm, start)
           for arm in ('passive', 'price_taker') for start in ('cold', 'perturbed', 'warm_from_certified')]

THRESH = 1e-6                       # the prediction's relative tolerance
BOUND_RELAX = 1e-8                  # IPOPT bound_relax_factor in force (S48: not overridden; default)
BASE_MVA = 100.0                    # asserted against the case network files below
FLOOR_MWH_PER_VAR = BOUND_RELAX * BASE_MVA   # closure slack floor per variable, MWh
FLOOR_MW_PER_VAR = BOUND_RELAX                # ESSO slack floor per variable, MW (ESSO model unit)

SCAN_KEYS = ('shared_ess_day_balance_slack', 'esso_feasibility_violation', 'slack_pnet_up', 'slack_pnet_down',
             'slack_es_pnet', 'soc_final', 'slack_shared_es_soc', 'complementarity_ratio_max', 'esso_violation',
             'relaxation_slacks', 'relaxation_variables', 'slack_penalties')
LOG_PATTERNS = ('Shared ESS complementarity diagnostics', 'ESSO violation', 'soc_final', '[SLACK COMPONENTS]')
XLSX_SHEETS = ('Relaxation Slacks TSO, DSOs', 'Slacks operation, aggregated')
JSON_SCAN_MAX = 40 * 1024 * 1024
LOG_SCAN_MAX = 50 * 1024 * 1024

OUT_DIR = os.path.join(S53, 'w169_slack_inventory')
OUT_NAMES = {'plan': 'capture_plan.json', 'inventory': 'inventory.json', 'summary': 'w169_slack_inventory.json',
             'md': 'w169_slack_inventory.md', 'inputs_man': 'manifest_inputs_sha256.json', 'run': 'run_record.json',
             'man': 'manifest_sha256.json', 'post': 'manifest_post_run_sha256.json',
             'typing_json': 'w169_bool_typing_test.json', 'typing_log': 'w169_bool_typing_test.log',
             'log': 'launch.log'}
CELLS_SUB = 'cells'

# code locations: (file, exact needle, what). Located at run time; a needle not found is an integrity failure.
CODE_NEEDLES = [
    ('network.py', 'model.slack_shared_es_soc_final_up = pe.Var(', 'closure slack s+ declared (e, s_m, s_o)'),
    ('network.py', 'model.slack_shared_es_soc_final_down = pe.Var(', 'closure slack s- declared'),
    ('model_construction_helpers.py', 'def sess_soc_final_rule(', 'closure row'),
    ('model_construction_helpers.py',
     'return m.shared_es_soc[e, s_m0, s_o0, final_p] == final_soc + m.slack_shared_es_soc_final_up[e, s_m0, s_o0]',
     'closure row with the slack pair: SoC_T = 0.5 E + s+ - s- (scenario-free copy)'),
    ('model_construction_helpers.py', 'final_soc = m.shared_es_e_rated_fixed[e] * ENERGY_STORAGE_RELATIVE_INIT_SOC',
     'closure target 0.5 E (E = the published capacity Param)'),
    ('model_construction_helpers.py', 'def sess_row_is_duplicate(', 'non-first scenario pairs skipped'),
    ('model_construction_helpers.py', 'ESS_DAY_BALANCE_SLACK_FRACTION = 0.05', 'closure slack bound fraction'),
    ('model_construction_helpers.py',
     'slack_ub = 0.0 if inactive else e_capacity * ESS_DAY_BALANCE_SLACK_FRACTION + EQUALITY_TOLERANCE',
     'closure slack upper bound, re-applied per candidate ([0, 0] when inactive)'),
    ('model_construction_helpers.py', 'def shared_ess_day_balance_slack_penalty(', 'closure penalty'),
    ('model_construction_helpers.py',
     'total += base * PENALTY_SHARED_ESS_BALANCE * (model.slack_shared_es_soc_final_up[e, s_m0, s_o0]',
     'closure penalty baseMVA * 1e3 * (s+ + s-)'),
    ('definitions.py', 'PENALTY_SHARED_ESS_BALANCE = 1e3', 'closure penalty weight (EUR/MWh)'),
    ('definitions.py', 'PENALTY_ESSO_SLACK = 1e3', 'ESSO P-net slack penalty weight'),
    ('definitions.py', 'EQUALITY_TOLERANCE = 1e-5', 'closure slack bound offset (p.u.)'),
    ('shared_resources_planning.py',
     "components['shared_ess_day_balance_slack'] += probability * pe.value(shared_ess_day_balance_slack_penalty(",
     'per-block closure level (the field component_levels_terminal.json records)'),
    ('shared_energy_storage_data.py',
     'model.slack_es_pnet_up = pe.Var(model.years, model.days, model.periods, domain=pe.NonNegativeReals',
     'ESSO P-net slack sigma+ declared (y, d, t)'),
    ('shared_energy_storage_data.py',
     'model.slack_es_pnet_down = pe.Var(model.years, model.days, model.periods, domain=pe.NonNegativeReals',
     'ESSO P-net slack sigma- declared'),
    ('shared_energy_storage_data.py',
     'model.energy_storage_operation_agg.add(model.es_pnet[y, d, p] == agg_pnet + model.slack_es_pnet_up[y, d, p]',
     'P-net row: es_pnet = sum_cohorts(pch - pdch) + sigma+ - sigma-'),
    ('shared_energy_storage_data.py',
     'slack_penalty += PENALTY_ESSO_SLACK * (model.slack_es_pnet_up[y_inv, d, p] + model.slack_es_pnet_down[y_inv, d, p])',
     'sigma penalty in the ESSO objective'),
    ('shared_energy_storage_data.py', 'def get_feasibility_violation(self, models):',
     'sum of sigma over nodes and (y, d, t) (component_levels_terminal esso_feasibility_violation_D3)'),
    ('shared_energy_storage_data.py', "ESSO_TOL_OVERRIDES = {'tol': 1e-10, 'acceptable_tol': 1e-9}",
     'Addendum 3 remedy (h) tolerance override (tightened by Addendum 8)'),
    ('shared_energy_storage_data.py', 'def _complementarity_ratio_for_model(model):',
     'min(pch, pdch) / s_max detector (ratio form)'),
    ('shared_energy_storage_data.py', 'def _get_esso_complementarity_diagnostics(model, node_id, log_path):',
     'per-solve detector + barrier-identity estimate'),
    ('shared_energy_storage_data.py', 'complementarity_diagnostics = _get_esso_complementarity_diagnostics(',
     'detector CALLED after every successful ESSO solve (production ESSO solve path)'),
    ('shared_energy_storage_data.py', 'complementarity_diagnostics_sink.append(complementarity_diagnostics)',
     'detector outcome appended to the sink the campaign harness records'),
    ('shared_energy_storage_data.py', 'complementarity_diagnostics_sink=self.esso_complementarity_diagnostics,',
     'sink wired by the production ESSO optimize entry point'),
    ('p515_g_g1_g4_admm_gates.py', 'def write_component_levels_terminal(planning, sed, models, rows, report, out_dir, label):',
     'writer of component_levels_terminal.json (from the final models)'),
    ('p515_g_g1_g4_admm_gates.py', "'esso_feasibility_violation_D3': esso_feasibility_violation,",
     'esso_feasibility_violation_D3 written'),
    ('p515_g_g1_g4_admm_gates.py', "'slack_pnet_up': slack_pnet_up, 'slack_pnet_down': slack_pnet_down,",
     'esso_capture per-element sigma written'),
    ('p515_g_g1_g4_admm_gates.py', 'def _capture_esso_solve(sed, models, node_diag, esso_capture_dir, cycle_label, hook_state):',
     'esso_capture writer (active cohort-periods only)'),
    ('model_construction_helpers.py', 'def shared_ess_capacity_is_inactive(s_capacity, e_capacity):',
     'zero-capacity test (used here, in p.u., to decide which units are active)'),
    ('definitions.py', 'SHARED_ESS_ZERO_CAPACITY_TOLERANCE = 1e-10', 'its tolerance (p.u.)'),
]

FAILED = []


def _log(msg):
    print(f'{datetime.now(timezone.utc).strftime("%H:%M:%S")} {msg}', flush=True)


def _fail(msg):
    FAILED.append(msg)
    _log(f'[W169 INTEGRITY FAIL] {msg}')


def _sha_file(path):
    h = hashlib.sha256()
    with open(os.path.join(REPO, path), 'rb') as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def _git(*args, cwd=REPO):
    r = subprocess.run(['git', *args], cwd=cwd, capture_output=True, text=True)
    return r.stdout.strip()


def _write_x(path, text):
    with open(path, 'x', encoding='utf-8') as handle:
        handle.write(text)


def _write_json(path, obj):
    _write_x(path, GRIO.dumps(obj, indent=1, sort_keys=False, ensure_ascii=False) + '\n')


def _load_json(rel):
    with open(os.path.join(REPO, rel), encoding='utf-8') as handle:
        return json.load(handle)


def _rel(path):
    return os.path.relpath(os.path.join(REPO, path), REPO)


def _cyc(x):
    """Cycle labels are written as ints, '150' or zero-padded '069' depending on the harness; compare as ints."""
    sx = str(x)
    return int(sx) if sx.isdigit() else sx


def _locate(rel, needle):
    with open(os.path.join(REPO, rel), encoding='utf-8') as handle:
        for i, line in enumerate(handle, 1):
            if needle in line:
                return f'{rel}:{i}'
    return None


# ----------------------------------------------------------------------------------------------------------------------
#  git status of the results tree (one call each, then lookups)
# ----------------------------------------------------------------------------------------------------------------------
class GitIndex:
    def __init__(self):
        out = _git('ls-files', '-s', '--', RES)
        self.blob = {}
        for line in out.split('\n'):
            if not line:
                continue
            meta, path = line.split('\t', 1)
            self.blob[path] = meta.split()[1]
        dirty = _git('status', '--porcelain', '--untracked-files=no', '--', RES)
        self.dirty = set()
        for line in dirty.split('\n'):
            if line.strip():
                self.dirty.add(line[3:].strip())

    def status(self, rel):
        rel = _rel(rel)
        tracked = rel in self.blob
        return {'tracked': tracked, 'clean': (tracked and rel not in self.dirty) if tracked else None,
                'blob': self.blob.get(rel), 'label': ('committed' if tracked and rel not in self.dirty else
                                                      'committed-but-modified' if tracked else 'uncommitted')}


# ======================================================================================================================
#  cells
# ======================================================================================================================
def build_cells(fz):
    t = fz['tables']
    cells = []
    coordinated_src = t['benchmark'].get('coordinated_source', '')
    for group, rows in (('T2', t['cells']), ('T2 appended', t['cells_appended_w160'])):
        for name, c in rows.items():
            certified = c.get('status') == 'certified'
            point = c.get('k_star') if certified else c.get('end_cycle')
            tables = ['T2'] if group == 'T2' else [c.get('row_origin'), 'T2 (appended by W160)']
            if c.get('eval_key', '')[:16] and c['eval_key'][:16] in coordinated_src:
                tables.append('T6 (coordinated row)')
            cells.append({
                'cell': name, 'group': group, 'tables': tables, 'item': c.get('item'),
                'status_frozen': c.get('status'), 'superseded': bool(c.get('superseded')),
                'certified_point': certified,
                'point': point, 'point_kind': 'k*' if certified else 'terminal (uncertified; frozen end_cycle)',
                'eval_key': c.get('eval_key'), 'candidate_key': c.get('candidate_key'),
                'candidate_canonical': c.get('candidate_canonical'),
                'certifying_spec': c.get('certifying_spec') or c.get('source'), 'case': 'SRP1', 'kind': 'admm',
            })
    for name, rel, role, note in T11_EVALS:
        er = _load_json(os.path.join(rel, 'evaluation_record.json'))
        cells.append({
            'cell': name, 'group': 'T11', 'tables': ['T11'], 'item': note, 'status_frozen': None,
            'status_record': er.get('status'), 'superseded': False,
            'certified_point': role == 'certified_point',
            'point': er.get('certification_cycle'),
            'point_kind': ('certification cycle (3 x 3 production exit; the T11 value is read here)'
                           if role == 'certified_point' else
                           'continuation terminal (post-hoc D source; not a certified point of a tabulated value)'),
            'eval_key': er.get('eval_key'), 'candidate_key': er.get('candidate_key'),
            'candidate_canonical': er.get('candidate_canonical'), 'certifying_spec': er.get('campaign_spec_path'),
            'case': '3x3', 'kind': 'admm', 'eval_dir': rel,
        })
    for name, rel, arm, start in T6_ARMS:
        cells.append({
            'cell': name, 'group': 'T6', 'tables': ['T6'], 'item': f'uncoordinated NRF arm {arm}, start {start}',
            'status_frozen': None, 'superseded': False, 'certified_point': False, 'point': None,
            'point_kind': 'uncoordinated arm evaluation (no ADMM cycles; phase A, phase C beside)',
            'eval_key': None, 'candidate_key': None, 'candidate_canonical': None,
            'certifying_spec': None, 'case': 'SRP1', 'kind': 'arm', 'arm_dir': rel, 'arm': arm, 'start': start,
        })
    return cells


def resolve_eval_dirs(cells):
    all_dirs = sorted(os.path.dirname(p) for p in glob.glob(os.path.join(S53, '**', 'evals', '*', 'evaluation_record.json'),
                                                           recursive=True))
    for c in cells:
        if c['kind'] != 'admm' or c.get('eval_dir'):
            continue
        ek = c['eval_key'] or ''
        hits = [d for d in all_dirs if os.path.basename(d).startswith(ek[:16] + '_')]
        if len(hits) != 1:
            _fail(f"{c['cell']}: {len(hits)} eval directories for eval key {ek[:16]} ({hits})")
            c['eval_dir'] = hits[0] if hits else None
            continue
        c['eval_dir'] = hits[0]
    for c in cells:
        if c['kind'] != 'admm' or not c.get('eval_dir'):
            continue
        er = _load_json(os.path.join(c['eval_dir'], 'evaluation_record.json'))
        c['eval_record_status'] = er.get('status')
        c['eval_record_cycles_run'] = er.get('cycles_run')
        if er.get('eval_key') != c['eval_key']:
            _fail(f"{c['cell']}: evaluation_record eval_key {er.get('eval_key')} != {c['eval_key']}")
        if c['candidate_key'] and er.get('candidate_key') != c['candidate_key']:
            _fail(f"{c['cell']}: evaluation_record candidate_key {er.get('candidate_key')} != {c['candidate_key']}")
        if c['candidate_key'] is None:
            c['candidate_key'] = er.get('candidate_key')
        if c['candidate_canonical'] is None:
            c['candidate_canonical'] = er.get('candidate_canonical')
        if c['point'] is None:
            _fail(f"{c['cell']}: no point (k* / end) in the frozen or evaluation record")


# ======================================================================================================================
#  inventory
# ======================================================================================================================
def _key_scan(obj, found, path='', depth=0):
    if depth > 12:
        return
    if isinstance(obj, dict):
        for k, v in obj.items():
            ks = str(k)
            for key in SCAN_KEYS:
                if key in ks.lower():
                    rec = found.setdefault(key, {'count': 0, 'example_path': f'{path}/{ks}'})
                    rec['count'] += 1
            _key_scan(v, found, f'{path}/{"*" if ("|" in ks or ks.lstrip("-").isdigit() or ks.startswith("(")) else ks}',
                      depth + 1)
    elif isinstance(obj, list):
        for v in obj[:50]:
            _key_scan(v, found, f'{path}[*]', depth + 1)


def _scan_file(rel, gi):
    size = os.path.getsize(os.path.join(REPO, rel))
    st = gi.status(rel)
    rec = {'path': rel, 'size_bytes': size, **st, 'scan': None, 'keys_found': {}}
    ext = os.path.splitext(rel)[1].lower()
    try:
        if ext == '.json':
            if size <= JSON_SCAN_MAX:
                obj = _load_json(rel)
                found = {}
                _key_scan(obj, found)
                rec['keys_found'] = found
                rec['scan'] = 'json parsed in full; keys scanned'
                if os.path.basename(rel) == 'component_levels_terminal.json':
                    rec['cycles_run'] = obj.get('cycles_run')
            else:
                rec['scan'] = f'not scanned (json larger than {JSON_SCAN_MAX} bytes)'
        elif ext == '.jsonl':
            found = {}
            n = 0
            with open(os.path.join(REPO, rel), encoding='utf-8') as handle:
                for line in handle:
                    if not line.strip():
                        continue
                    _key_scan(json.loads(line), found)
                    n += 1
                    if n >= 3:
                        break
            rec['keys_found'] = found
            rec['scan'] = f'first {n} records parsed; keys scanned (records are homogeneous)'
        elif ext in ('.log', '.txt'):
            if size <= LOG_SCAN_MAX:
                counts = {p: 0 for p in LOG_PATTERNS}
                with open(os.path.join(REPO, rel), encoding='utf-8', errors='replace') as handle:
                    for line in handle:
                        for p in LOG_PATTERNS:
                            if p in line:
                                counts[p] += 1
                rec['log_pattern_counts'] = counts
                rec['scan'] = 'log scanned in full for the patterns'
            else:
                rec['scan'] = f'not scanned (log larger than {LOG_SCAN_MAX} bytes)'
        elif ext == '.xlsx':
            import openpyxl
            wb = openpyxl.load_workbook(os.path.join(REPO, rel), read_only=True)
            rec['sheets'] = list(wb.sheetnames)
            rec['carries_sheets'] = [s for s in XLSX_SHEETS if s in wb.sheetnames]
            wb.close()
            rec['scan'] = 'workbook sheet names read (openpyxl, read-only)'
        elif ext == '.pkl':
            rec['scan'] = 'pickle: LISTED, NOT OPENED (pickle blocked)'
        else:
            rec['scan'] = f'not scanned (type {ext or "none"})'
    except Exception as exc:  # a scan error is recorded, never hidden
        rec['scan'] = f'scan error: {type(exc).__name__}: {exc}'
    return rec


def inventory_cell(c, gi):
    if c['kind'] == 'arm':
        files = sorted(os.path.join(c['arm_dir'], f) for f in os.listdir(c['arm_dir'])
                       if os.path.isfile(os.path.join(c['arm_dir'], f)))
        return {'eval_dir': c['arm_dir'], 'files': [_scan_file(f, gi) for f in files], 'campaign_files': [],
                'esso_capture': None, 'pickles': []}
    ed = c['eval_dir']
    files = []
    for dp, dns, fns in os.walk(ed):
        dns[:] = sorted(d for d in dns if d != 'esso_capture')
        for f in sorted(fns):
            files.append(os.path.relpath(os.path.join(dp, f), REPO))
    camp = os.path.dirname(os.path.dirname(ed))
    camp_files = sorted(os.path.join(camp, f) for f in os.listdir(camp) if os.path.isfile(os.path.join(camp, f)))
    inv = {'eval_dir': ed, 'campaign_dir': camp, 'files': [_scan_file(f, gi) for f in files],
           'campaign_files': [_scan_file(f, gi) for f in camp_files]}
    caps = sorted(glob.glob(os.path.join(ed, 'esso_capture', '*')))
    cap = {'dirs': [os.path.relpath(x, REPO) for x in caps]}
    if len(caps) == 1:
        names = sorted(os.listdir(caps[0]))
        cap['n_files'] = len(names)
        cap['label'] = os.path.basename(caps[0])
        p = c['point']
        cap['files_at_point'] = {}
        for n in (5, 7, 9):
            rel = os.path.join(caps[0], f'node{n}_cycle{p:03d}.jsonl')
            if os.path.exists(os.path.join(REPO, rel)):
                st = gi.status(rel)
                with open(os.path.join(REPO, rel), encoding='utf-8') as handle:
                    first = handle.readline()
                cap['files_at_point'][str(n)] = {'path': rel, **st,
                                                 'first_row_keys': sorted(json.loads(first).keys()) if first.strip() else []}
            else:
                cap['files_at_point'][str(n)] = None
        cap['status_sample'] = gi.status(os.path.join(caps[0], names[0])) if names else None
    inv['esso_capture'] = cap
    inv['pickles'] = [f['path'] for f in inv['files'] if f['path'].endswith('.pkl')]
    return inv


def _find(inv, basename_prefix, ext=None):
    out = []
    for f in inv['files']:
        b = os.path.basename(f['path'])
        if b.startswith(basename_prefix) and (ext is None or b.endswith(ext)):
            out.append(f)
    return out


def plan_cell(c, inv):
    """The capture plan: one source per quantity, decided BEFORE extraction from the inventory alone."""
    plan = {'cell': c['cell'], 'point': c['point'], 'point_kind': c['point_kind'], 'sources': {}, 'not_carried': {},
            'records_not_carrying': [], 'pickles': []}
    if c['kind'] == 'arm':
        arm_json = [f for f in inv['files'] if f['path'].endswith(f"nrf_arm_{c['arm']}_{c['start']}_r2.json")]
        if len(arm_json) != 1:
            _fail(f"{c['cell']}: arm record not found uniquely")
            return plan
        a = arm_json[0]
        plan['sources']['closure_total_weighted'] = {
            'path': a['path'], 'label': a['label'],
            'fields': ['phase_A.evaluation.recourse_components.detector_components.shared_ess_day_balance_slack',
                       'phase_C_sequential_pass.evaluation.recourse_components.detector_components.shared_ess_day_balance_slack'],
            'why': 'the arm record carries the closure level only as the block-weighted TOTAL over every block'}
        plan['sources']['solve_kinds'] = {'path': a['path'], 'label': a['label'], 'fields': ['solve_accounting.by_kind'],
                                          'why': 'establishes whether an ESSO model is solved in the arm at all'}
        plan['not_carried']['closure_per_block'] = 'arm record: block_components carry one scalar per block, not the closure level'
        for f in inv['files']:
            if f['path'] != a['path']:
                plan['records_not_carrying'].append({'path': f['path'], 'label': f['label'],
                                                     'keys_found': f['keys_found'], 'scan': f['scan']})
        return plan
    p = c['point']
    cl = _find(inv, 'component_levels_terminal.json')
    er = _find(inv, 'evaluation_record.json')
    g = [f for f in inv['files'] if os.path.basename(f['path']).startswith('g_') and f['path'].endswith('.json')]
    lk = _find(inv, 'leak_classification_', '.jsonl')
    xl = [f for f in inv['files'] if f['path'].endswith('.xlsx')]
    pc = _find(inv, 'post_certification.json')
    if len(er) != 1:
        _fail(f"{c['cell']}: evaluation_record.json not unique")
        return plan
    plan['sources']['capacities'] = {
        'path': er[0]['path'], 'label': er[0]['label'],
        'fields': ['storage_per_node.<n>.published_available_capacity_terminal.<year>.{e_available,s_available}',
                   'storage_per_node.<n>.has_storage', 'ageing_trajectory_terminal (y -> block_year cross-check)'],
        'at': 'terminal cycle of the run', 'why': 'E^Av and S for the relative values'}
    cl_at_p = len(cl) == 1 and cl[0].get('cycles_run') == p
    if cl_at_p:
        plan['sources']['closure'] = {
            'path': cl[0]['path'], 'label': cl[0]['label'],
            'fields': ['blocks.<block>.unweighted.shared_ess_day_balance_slack', 'cycles_run'],
            'at': f'cycle {p} (cycles_run == P)', 'why': 'per-block closure level at P'}
        plan['sources']['sigma_sum'] = {
            'path': cl[0]['path'], 'label': cl[0]['label'], 'fields': ['esso_feasibility_violation_D3'],
            'at': f'cycle {p}', 'why': 'sum of sigma over every node and (y, d, t) at P (the only committed sigma field)'}
    else:
        plan['not_carried']['closure'] = (
            f'no record carries the closure level at cycle {p}: component_levels_terminal.json is at cycle '
            f"{cl[0].get('cycles_run') if cl else None}; per-cycle records carry only the aggregate slack_penalties "
            '(recourse_blocks_all.jsonl / creep_diagnostic_per_cycle.jsonl: voltage + flexibility + closure summed); '
            'ess_schedule_per_cycle.jsonl (uncommitted, where present) carries the network copies\' charge / discharge '
            'per cycle, from which only the SIGNED s+ - s- could be reconstructed through the SoC recursion -- not done '
            '(a derived value, not the variable)')
        plan['not_carried']['sigma_sum'] = f'component_levels_terminal.json is not at cycle {p}'
        if cl:
            plan['sources']['closure_terminal_beside'] = {
                'path': cl[0]['path'], 'label': cl[0]['label'],
                'fields': ['blocks.<block>.unweighted.shared_ess_day_balance_slack', 'esso_feasibility_violation_D3'],
                'at': f"cycle {cl[0].get('cycles_run')} (terminal; NOT the point)", 'why': 'reported beside, labelled'}
    cap = inv['esso_capture'] or {}
    fap = cap.get('files_at_point') or {}
    present = {n: v for n, v in fap.items() if v}
    if present:
        plan['sources']['sigma_elementwise'] = {
            'paths': {n: v['path'] for n, v in present.items()},
            'labels': {n: v['label'] for n, v in present.items()},
            'fields': ['slack_pnet_up', 'slack_pnet_down', 's_max', 'y', 'd', 'p', 'y_inv'],
            'at': f'cycle {p}', 'why': 'per-element sigma at P (rows exist for active cohort-periods only)'}
    else:
        plan['not_carried']['sigma_elementwise'] = f'no esso_capture file for cycle {p}'
    if len(g) == 1:
        plan['sources']['detector'] = {'path': g[0]['path'], 'label': g[0]['label'],
                                       'fields': [f'esso_complementarity_diagnostics_by_round.per_round[round == {p}]',
                                                  'esso_complementarity_diagnostics_by_round.per_node_summary',
                                                  'rule_eleven_checklist.s44_campaign_configuration_checks.persistent_workers_off'],
                                       'why': 'the min(pch, pdch) detector outcome at P per node'}
    else:
        plan['not_carried']['detector'] = f'{len(g)} g_*.json files'
    nf = _find(inv, 'network_failures_', '.jsonl')
    if len(nf) == 1:
        plan['sources']['network_failures'] = {'path': nf[0]['path'], 'label': nf[0]['label'],
                                               'fields': [f'records with cycle == {p}: block, primary_termination, class'],
                                               'why': 'context: which block solves at P were recoveries (not a slack value)'}
    if len(lk) == 1:
        plan['sources']['detector_crosscheck'] = {'path': lk[0]['path'], 'label': lk[0]['label'],
                                                  'fields': [f'records with cycle == "{p}": complementarity_ratio_max'],
                                                  'why': 'cross-check of the g record'}
    for x in xl:
        if cl and cl[0].get('cycles_run') == p and set(XLSX_SHEETS) <= set(x.get('carries_sheets') or []):
            plan['sources'].setdefault('xlsx_crosscheck', {
                'path': x['path'], 'label': x['label'],
                'fields': ['Relaxation Slacks TSO, DSOs: "Shared Energy Storage, soc_final" (signed s+ - s-, MWh)',
                           'Slacks operation, aggregated: "Pnet, up" / "Pnet, down" (every node, MW)'],
                'at': 'terminal (== P)', 'why': 'per-unit signed closure and every-node sigma (cross-check)'})
    # pickles: what they hold and at which cycle (post_certification.json names the persisted model's cycle)
    pk_cycle = None
    if pc:
        pcd = _load_json(pc[0]['path'])
        pk_cycle = (pcd.get('certification') or {}).get('certification_cycle')
    for pth in inv['pickles']:
        b = os.path.basename(pth)
        holds = ('the persisted certified models (TSO / DSO / ESSO) at cycle '
                 f'{pk_cycle} (post_certification.json)' if b == 'certified_models.pkl' else
                 'ESSO models at the run terminal' if b.startswith('esso_models') else
                 'FrozenSMOPF fixture (a frozen single-block solve, not this cell\'s point)' if 'FrozenSMOPF' in pth else
                 'unknown content')
        plan['pickles'].append({'path': pth, 'holds': holds, 'read': False})
    plan['pickle_cycle'] = pk_cycle
    if 'closure' in plan['not_carried']:
        plan['not_carried']['closure'] += ('; pickles: ' + ('certified_models.pkl holds cycle '
                                                            f'{pk_cycle}, not {p} -- no record or pickle holds cycle {p}'
                                                            if pk_cycle is not None and pk_cycle != p else
                                                            'only in pickle -- not read'))
    used = set()
    for s in plan['sources'].values():
        if 'path' in s:
            used.add(s['path'])
        for v in (s.get('paths') or {}).values():
            used.add(v)
    for f in inv['files'] + inv['campaign_files']:
        if f['path'] in used or f['path'].endswith('.pkl'):
            continue
        plan['records_not_carrying'].append({'path': f['path'], 'label': f['label'], 'scan': f['scan'],
                                             'keys_found': f['keys_found'],
                                             'log_pattern_counts': f.get('log_pattern_counts'),
                                             'carries_sheets': f.get('carries_sheets')})
    return plan


def plan_paths(plan):
    out = []
    for s in plan['sources'].values():
        if 'path' in s:
            out.append(s['path'])
        out.extend((s.get('paths') or {}).values())
    return sorted(set(out))


# ======================================================================================================================
#  extraction
# ======================================================================================================================
def _case_orders(case):
    rel = CASE_SRP1_REL if case == 'SRP1' else CASE_3X3_REL
    d = _load_json(rel)
    return [str(y) for y in d['Years']], list(d['Days'])


def _capacities(er):
    spn = er['storage_per_node']
    caps = {}
    for n, v in spn.items():
        caps[str(n)] = {'has_storage': bool(v.get('has_storage')),
                        'e': {str(y): float(x['e_available']) for y, x in v['published_available_capacity_terminal'].items()},
                        's': {str(y): float(x['s_available']) for y, x in v['published_available_capacity_terminal'].items()}}
    return caps


def _active(caps, n, year):
    """Production's own zero-capacity test (model_construction_helpers.shared_ess_capacity_is_inactive), applied in
    p.u. as production applies it (network_data.py divides the published MVA / MWh by baseMVA)."""
    s_ = caps[n]['s'].get(year, 0.0)
    e_ = caps[n]['e'].get(year, 0.0)
    return not shared_ess_capacity_is_inactive(s_ / BASE_MVA, e_ / BASE_MVA)


def extract_closure(c, plan, caps, years):
    src = plan['sources']['closure']
    cl = _load_json(src['path'])
    if cl.get('cycles_run') != c['point']:
        _fail(f"{c['cell']}: component_levels cycles_run {cl.get('cycles_run')} != P {c['point']}")
    blocks = []
    w_total = 0.0
    for key, b in cl['blocks'].items():
        parts = key.split('|')
        if parts[0] == 'TSO':
            units, year, day = sorted(caps.keys(), key=int), parts[1], parts[2]
        else:
            units, year, day = [parts[1]], parts[2], parts[3]
        v_eur = float(b['unweighted']['shared_ess_day_balance_slack'])
        w_total += float(b['weighted']['shared_ess_day_balance_slack'])
        s_mwh = v_eur / DEF.PENALTY_SHARED_ESS_BALANCE
        active = [u for u in units if _active(caps, u, year)]
        e_min = min(caps[u]['e'][year] for u in active) if active else None
        per_unit_bound = s_mwh + max(len(active) - 1, 0) * 2 * FLOOR_MWH_PER_VAR
        blocks.append({'block': key, 'year': year, 'day': day, 'units': units, 'active_units': active,
                       'level_eur_unweighted': v_eur, 'sum_s_mwh': s_mwh, 'per_unit_upper_bound_mwh': per_unit_bound,
                       'e_av_min_active_mwh': e_min,
                       'per_unit_exact': len(active) == 1,
                       'frac_signed': (per_unit_bound / e_min) if e_min else None,
                       'frac_abs': (abs(s_mwh) / e_min) if (e_min and len(active) == 1) else None,
                       'frac_abs_of_multi_unit_sum': (abs(s_mwh) / e_min) if (e_min and len(active) > 1) else None})
    er = _load_json(plan['sources']['capacities']['path'])
    rec_total = (er.get('component_decomposition_totals_weighted') or {}).get('shared_ess_day_balance_slack')
    if rec_total is not None and abs(rec_total - w_total) > 1e-9 * max(1.0, abs(rec_total)):
        _fail(f"{c['cell']}: weighted closure total {w_total} != evaluation_record {rec_total}")
    act = [b for b in blocks if b['active_units']]
    ina = [b for b in blocks if not b['active_units']]
    mx = max(blocks, key=lambda b: b['sum_s_mwh'])
    out = {
        'source': {'path': src['path'], 'label': src['label'], 'sha256': None,
                   'field': 'blocks.<block>.unweighted.shared_ess_day_balance_slack / PENALTY_SHARED_ESS_BALANCE'},
        'point_of_record': cl.get('cycles_run'),
        'n_blocks': len(blocks), 'n_blocks_with_active_unit': len(act), 'n_blocks_without_active_unit': len(ina),
        'max_sum_s_mwh': mx['sum_s_mwh'], 'argmax_block': mx['block'], 'argmax_active_units': mx['active_units'],
        'min_sum_s_mwh': min(b['sum_s_mwh'] for b in blocks),
        'n_blocks_positive': sum(1 for b in blocks if b['sum_s_mwh'] > 0.0),
        'max_positive_s_mwh': max([b['sum_s_mwh'] for b in blocks if b['sum_s_mwh'] > 0.0], default=None),
        'inactive_only_blocks_all_exactly_zero': all(b['sum_s_mwh'] == 0.0 for b in ina) if ina else None,
        'inactive_only_blocks_max_abs_mwh': max((abs(b['sum_s_mwh']) for b in ina), default=None),
        'weighted_total_eur_record': rec_total, 'weighted_total_eur_from_blocks': w_total,
        'blocks': blocks,
    }
    if act:
        ms = max(act, key=lambda b: b['frac_signed'])
        mact = max(act, key=lambda b: b['sum_s_mwh'])
        exact = [b for b in act if b['per_unit_exact']]
        multi = [b for b in act if not b['per_unit_exact']]
        ma = max(exact, key=lambda b: b['frac_abs']) if exact else None
        out.update({'max_sum_s_mwh_active_blocks': mact['sum_s_mwh'], 'argmax_active_block': mact['block'],
                    'min_sum_s_mwh_active_blocks': min(b['sum_s_mwh'] for b in act),
                    'max_frac_signed': ms['frac_signed'], 'argmax_frac_signed_block': ms['block'],
                    'argmax_frac_signed_e_av_mwh': ms['e_av_min_active_mwh'],
                    'max_frac_abs': ma['frac_abs'] if ma else None,
                    'argmax_frac_abs_block': ma['block'] if ma else None,
                    'argmax_frac_abs_e_av_mwh': ma['e_av_min_active_mwh'] if ma else None,
                    'abs_reading_scope': 'blocks with exactly one active unit (DSO blocks; TSO blocks of one active unit, '
                                         'the zero-capacity pairs being bounded [0, 0])',
                    'n_multi_unit_tso_blocks': len(multi),
                    'max_frac_abs_of_multi_unit_sum': max((b['frac_abs_of_multi_unit_sum'] for b in multi), default=None),
                    'max_per_unit_upper_bound_mwh': max(b['per_unit_upper_bound_mwh'] for b in act),
                    'every_active_block_at_or_above_relaxed_floor': all(
                        b['sum_s_mwh'] >= -2 * len(b['active_units']) * FLOOR_MWH_PER_VAR * (1 + 1e-9) for b in act),
                    'min_ratio_to_relaxed_floor': min(b['sum_s_mwh'] / (-2 * len(b['active_units']) * FLOOR_MWH_PER_VAR)
                                                      for b in act)})
    else:
        out.update({'max_frac_signed': None, 'max_frac_abs': None,
                    'note_no_active_unit': 'every unit has zero capacity (x = 0): every closure pair has bounds [0, 0]; '
                                           'the relative value is undefined; the absolute level is reported'})
    return out


def extract_closure_terminal_beside(c, plan, caps):
    src = plan['sources']['closure_terminal_beside']
    cl = _load_json(src['path'])
    vals = []
    for key, b in cl['blocks'].items():
        parts = key.split('|')
        units, year = (sorted(caps.keys(), key=int), parts[1]) if parts[0] == 'TSO' else ([parts[1]], parts[2])
        act = [u for u in units if _active(caps, u, year)]
        e_min = min(caps[u]['e'][year] for u in act) if act else None
        v = float(b['unweighted']['shared_ess_day_balance_slack']) / DEF.PENALTY_SHARED_ESS_BALANCE
        vals.append((v, key, act, e_min))
    av = [x for x in vals if x[2]]
    mx = max(av) if av else max(vals)
    fr = [(x[0] + max(len(x[2]) - 1, 0) * 2 * FLOOR_MWH_PER_VAR) / x[3] for x in av]
    return {'path': src['path'], 'label': src['label'], 'at_cycle': cl.get('cycles_run'),
            'NOT_THE_POINT': True, 'max_sum_s_mwh_active_blocks': mx[0], 'argmax_block': mx[1],
            'min_sum_s_mwh': min(vals)[0], 'n_blocks_positive': sum(1 for x in vals if x[0] > 0.0),
            'max_frac_signed': max(fr) if fr else None,
            'esso_feasibility_violation_D3': cl.get('esso_feasibility_violation_D3')}


def extract_sigma(c, plan, caps, years, days):
    src = plan['sources']['sigma_elementwise']
    per_node = {}
    elements = []
    p = c['point']
    for n, rel in sorted(src['paths'].items()):
        rows = []
        with open(os.path.join(REPO, rel), encoding='utf-8') as handle:
            for line in handle:
                if line.strip():
                    rows.append(json.loads(line))
        if caps[n]['has_storage'] and not rows:
            _fail(f"{c['cell']}: active node {n} has an empty esso_capture file at cycle {p}")
        if not caps[n]['has_storage'] and rows:
            _fail(f"{c['cell']}: inactive node {n} has esso_capture rows at cycle {p}")
        by_ydp = {}
        for r in rows:
            if _cyc(r.get('cycle')) != _cyc(p):
                _fail(f"{c['cell']}: capture row cycle {r.get('cycle')} != {p} in {rel}")
            if r.get('slack_pnet_up') is None or r.get('slack_pnet_down') is None:
                _fail(f"{c['cell']}: sigma None in {rel} (slacks disabled?)")
                continue
            k = (r['y'], r['d'], r['p'])
            sig = (float(r['slack_pnet_up']), float(r['slack_pnet_down']))
            if k in by_ydp:
                if by_ydp[k]['sig'] != sig:
                    _fail(f"{c['cell']}: sigma differs across cohorts at {k} in {rel}")
                by_ydp[k]['s'] += float(r['s_max'])
                by_ydp[k]['n_cohorts'] += 1
            else:
                by_ydp[k] = {'sig': sig, 's': float(r['s_max']), 'n_cohorts': 1}
        node_elems = []
        for (y, d, t), v in by_ydp.items():
            yl, dl = years[y], days[d]
            s_pub = caps[n]['s'].get(yl)
            if s_pub is not None and abs(s_pub - v['s']) > 1e-9 * max(1.0, s_pub):
                _fail(f"{c['cell']}: node {n} {yl}: capture S {v['s']} != published s_available {s_pub}")
            sig = v['sig'][0] + v['sig'][1]
            node_elems.append({'e': int(n), 'y': y, 'year': yl, 'd': d, 'day': dl, 't': t, 'sigma_up': v['sig'][0],
                               'sigma_down': v['sig'][1], 'sigma_sum_mw': sig, 'S_mva': v['s'],
                               'frac_signed': sig / v['s'] if v['s'] else None,
                               'frac_abs': abs(sig) / v['s'] if v['s'] else None})
        elements.extend(node_elems)
        n_years_with_s = sum(1 for yl in years if _active(caps, n, yl))
        per_node[n] = {'path': rel, 'n_rows': len(rows), 'n_elements': len(node_elems),
                       'expected_elements': n_years_with_s * len(days) * 24,
                       'years_with_capacity': [yl for yl in years if _active(caps, n, yl)],
                       'max_sigma_mw': max((e['sigma_sum_mw'] for e in node_elems), default=None),
                       'min_sigma_mw': min((e['sigma_sum_mw'] for e in node_elems), default=None)}
        if len(node_elems) != n_years_with_s * len(days) * 24:
            _fail(f"{c['cell']}: node {n}: {len(node_elems)} sigma elements != {n_years_with_s * len(days) * 24} "
                  '(years with capacity x days x 24)')
    out = {'source_paths': src['paths'], 'labels': src['labels'], 'per_node': per_node, 'n_elements': len(elements)}
    if elements:
        ms = max(elements, key=lambda e: e['sigma_sum_mw'])
        fs = [e for e in elements if e['frac_signed'] is not None]
        mfs = max(fs, key=lambda e: e['frac_signed'])
        mfa = max(fs, key=lambda e: e['frac_abs'])
        out.update({'max_sigma_mw': ms['sigma_sum_mw'], 'argmax': {k: ms[k] for k in ('e', 'year', 'day', 't', 'y', 'd')},
                    'min_sigma_mw': min(e['sigma_sum_mw'] for e in elements),
                    'n_positive': sum(1 for e in elements if e['sigma_sum_mw'] > 0.0),
                    'max_positive_sigma_mw': max([e['sigma_sum_mw'] for e in elements if e['sigma_sum_mw'] > 0.0],
                                                 default=None),
                    'max_frac_signed': mfs['frac_signed'],
                    'argmax_frac_signed': {k: mfs[k] for k in ('e', 'year', 'day', 't', 'S_mva', 'sigma_sum_mw')},
                    'max_frac_abs': mfa['frac_abs'],
                    'argmax_frac_abs': {k: mfa[k] for k in ('e', 'year', 'day', 't', 'S_mva', 'sigma_sum_mw')},
                    'sum_captured_mw': sum(e['sigma_sum_mw'] for e in elements),
                    'every_element_at_or_above_relaxed_floor': all(
                        e['sigma_sum_mw'] >= -2 * FLOOR_MW_PER_VAR * (1 + 1e-9) for e in elements),
                    'max_minus_min_mw': max(e['sigma_sum_mw'] for e in elements) - min(e['sigma_sum_mw'] for e in elements)})
    else:
        out.update({'max_sigma_mw': None, 'note_no_active_unit': 'no active cohort-period: no per-element rows exist'})
    return out


def extract_sigma_sum(c, plan, caps, years, days, sig_el):
    src = plan['sources']['sigma_sum']
    cl = _load_json(src['path'])
    d3 = cl.get('esso_feasibility_violation_D3')
    n_nodes = len(caps)
    n_total = n_nodes * len(years) * len(days) * 24
    out = {'path': src['path'], 'label': src['label'], 'field': 'esso_feasibility_violation_D3', 'D3_mw': d3,
           'N_elements_all_nodes': n_total, 'floor_per_element_mw': -2 * FLOOR_MW_PER_VAR,
           'D3_over_N_mean_mw': d3 / n_total if d3 is not None else None,
           'per_element_upper_bound_mw': (d3 + (n_total - 1) * 2 * FLOOR_MW_PER_VAR) if d3 is not None else None,
           'bound_assumption': 'every sigma+ and sigma- >= -1e-8 MW (IPOPT bound_relax_factor 1e-8; S48)'}
    active_s = [caps[n]['s'][y] for n in caps for y in years if _active(caps, n, y)]
    out['per_element_upper_bound_frac_of_min_S'] = (out['per_element_upper_bound_mw'] / min(active_s)
                                                     if active_s and d3 is not None else None)
    if sig_el and sig_el.get('n_elements') is not None and d3 is not None:
        captured = sig_el.get('sum_captured_mw') or 0.0
        n_cap = sig_el['n_elements']
        n_inact = n_total - n_cap
        out['inactive_nodes_total_mw'] = d3 - captured
        out['inactive_nodes_N'] = n_inact
        out['inactive_nodes_mean_mw'] = (d3 - captured) / n_inact if n_inact else None
        out['inactive_nodes_per_element_upper_bound_mw'] = ((d3 - captured) + (n_inact - 1) * 2 * FLOOR_MW_PER_VAR
                                                            if n_inact else None)
    return out


def extract_xlsx(c, plan):
    import openpyxl
    src = plan['sources']['xlsx_crosscheck']
    wb = openpyxl.load_workbook(os.path.join(REPO, src['path']), read_only=True)
    ws = wb['Relaxation Slacks TSO, DSOs']
    rows = list(ws.iter_rows(values_only=True))
    head = rows[0]
    qi = head.index('Quantity')
    soc = []
    for r in rows[1:]:
        if r[qi] == 'Shared Energy Storage, soc_final':
            vals = [v for v in r[qi + 3:] if isinstance(v, (int, float))]
            soc.append({'operator': r[0], 'adn_node': r[1], 'resource': r[2], 'year': r[3], 'day': r[4],
                        'net_s_mwh': vals[0] if vals else None})
    ws2 = wb['Slacks operation, aggregated']
    rows2 = list(ws2.iter_rows(values_only=True))
    sig = {}
    for r in rows2[1:]:
        key = (r[0], r[1], r[2])
        for t, v in enumerate(r[4:]):
            if isinstance(v, (int, float)):
                sig.setdefault(key + (t,), 0.0)
                sig[key + (t,)] += float(v)
    wb.close()
    nets = [s['net_s_mwh'] for s in soc if s['net_s_mwh'] is not None]
    return {'path': src['path'], 'label': src['label'], 'at': 'terminal (== P)',
            'closure_signed_net_rows': len(soc),
            'closure_signed_net_max_abs_mwh': max((abs(x) for x in nets), default=None),
            'closure_signed_net_values_distinct': sorted(set(nets))[:10],
            'sigma_elements_all_nodes': len(sig),
            'sigma_max_mw': max(sig.values()) if sig else None, 'sigma_min_mw': min(sig.values()) if sig else None,
            'note': 'the workbook carries s+ - s- (network.py writes up - down), not s+ + s-; sigma for EVERY node'}


def extract_detector(c, plan):
    p = c['point']
    g = _load_json(plan['sources']['detector']['path'])
    br = g.get('esso_complementarity_diagnostics_by_round') or {}
    pr = [r for r in br.get('per_round', []) if r.get('round') == p]
    out = {'path': plan['sources']['detector']['path'], 'label': plan['sources']['detector']['label'],
           'round_found': len(pr) == 1, 'per_node_at_point': {}, 'per_node_summary_all_rounds': br.get('per_node_summary'),
           'grouping_clean': br.get('grouping_clean'), 'expected_total_solves': br.get('expected_total_solves'),
           'observed_total_solves': br.get('observed_total_solves'),
           'persistent_workers_off': (((g.get('rule_eleven_checklist') or {}).get('s44_campaign_configuration_checks')
                                       or {}).get('persistent_workers_off'))}
    if len(pr) != 1:
        _fail(f"{c['cell']}: detector round {p} found {len(pr)} times in {out['path']}")
        return out
    for e in pr[0]['entries']:
        out['per_node_at_point'][str(e['node_id'])] = {
            k: e.get(k) for k in ('complementarity_ratio_max', 'complementarity_ratio_argmax', 'n_active_cohort_periods',
                                  'spurious_throughput_measured', 'spurious_throughput_bound', 'mu_final', 's_obj',
                                  'parse_reason', 'log_path')}
    if 'detector_crosscheck' in plan['sources']:
        rel = plan['sources']['detector_crosscheck']['path']
        lk = {}
        with open(os.path.join(REPO, rel), encoding='utf-8') as handle:
            for line in handle:
                if line.strip():
                    r = json.loads(line)
                    if _cyc(r.get('cycle')) == _cyc(p):
                        lk[str(r['node_id'])] = r.get('complementarity_ratio_max')
        agree = all(lk.get(n) == v['complementarity_ratio_max'] for n, v in out['per_node_at_point'].items())
        out['leak_classification_crosscheck'] = {'path': rel, 'values': lk, 'agree_exactly': agree}
        if not agree:
            _fail(f"{c['cell']}: detector g vs leak_classification disagree at {p}: {lk}")
    vals = [v['complementarity_ratio_max'] for v in out['per_node_at_point'].values()
            if v['complementarity_ratio_max'] is not None]
    out['max_ratio_at_point'] = max(vals) if vals else None
    return out


def extract_network_failures(c, plan):
    rel = plan['sources']['network_failures']['path']
    out = []
    with open(os.path.join(REPO, rel), encoding='utf-8') as handle:
        for line in handle:
            if line.strip():
                r = json.loads(line)
                if _cyc(r.get('cycle')) == _cyc(c['point']):
                    blk = (f"TSO|{r.get('year')}|{r.get('day')}" if r.get('agent') == 'TSO' else
                           f"DSO|{r.get('node_id')}|{r.get('year')}|{r.get('day')}")
                    out.append({'block': blk, 'primary_termination': r.get('primary_termination'),
                                'termination': r.get('termination'), 'class': r.get('class')})
    return {'path': rel, 'label': plan['sources']['network_failures']['label'], 'at_point': out}


def extract_arm(c, plan):
    a = _load_json(plan['sources']['closure_total_weighted']['path'])
    pa = a['phase_A']['evaluation']['recourse_components']['detector_components']['shared_ess_day_balance_slack']
    pcx = (a.get('phase_C_sequential_pass') or {}).get('evaluation') or {}
    pcv = ((pcx.get('recourse_components') or {}).get('detector_components') or {}).get('shared_ess_day_balance_slack')
    kinds = sorted((a.get('solve_accounting') or {}).get('by_kind', {}).keys())
    inst = a.get('instance') or {}
    return {'closure_total_weighted_eur_phase_A': pa, 'closure_total_weighted_eur_phase_C': pcv,
            'closure_total_weighted_mwh_equiv_phase_A': pa / DEF.PENALTY_SHARED_ESS_BALANCE if pa is not None else None,
            'solve_kinds': kinds, 'esso_solved': any('ESS' in k.upper() for k in kinds),
            'instance': {'candidate_key': inst.get('candidate_key'), 'canonical_candidate': inst.get('canonical_candidate'),
                         'label': inst.get('label')},
            'arm_record': plan['sources']['closure_total_weighted']['path'],
            'arm_label': a.get('arm'), 'start_label': a.get('start')}


# ======================================================================================================================
#  checks on the configuration the values rest on
# ======================================================================================================================
def configuration_evidence():
    ev = {'ess_params': {}, 'network_params': {}, 'baseMVA': {}}
    ess = _load_json(ESS_PARAMS_REL)
    ev['ess_params'] = {'path': ESS_PARAMS_REL, 'slacks': ess.get('slacks'), 'sha256': _sha_file(ESS_PARAMS_REL),
                        'clean_vs_HEAD': _git('status', '--porcelain', '--', ESS_PARAMS_REL) == ''}
    for rel in NET_PARAMS_RELS:
        d = _load_json(rel)
        ev['network_params'][rel] = {'slacks.shared_ess': (d.get('slacks') or {}).get('shared_ess'),
                                     'shared_ess_model': d.get('shared_ess_model'),
                                     'clean_vs_HEAD': _git('status', '--porcelain', '--', rel) == ''}
        if not ((d.get('slacks') or {}).get('shared_ess') or {}).get('day_balance'):
            _fail(f'{rel}: slacks.shared_ess.day_balance is not true')
    for rel in NET_2025_RELS:
        b = _load_json(rel).get('baseMVA')
        ev['baseMVA'][rel] = b
        if b != BASE_MVA:
            _fail(f'{rel}: baseMVA {b} != {BASE_MVA}')
    s48 = _load_json(S48_REL)
    ev['bound_relaxation_prior_art'] = {
        'path': S48_REL, 'sha256': _sha_file(S48_REL),
        'IPOPT_bound_relax_factor_default': s48.get('IPOPT_bound_relax_factor_default'),
        'verdict': s48.get('verdict')}
    rec = _load_json(INSTANCE_3X3_RECORD_REL)
    txt = json.dumps(rec)
    ev['probabilities_3x3_all_one_third'] = ('"probabilities_3x3_all_one_third": true' in txt)
    ev['bound_relax_search'] = {
        'scope': 'git grep -l over tracked *.py, data/SRP1/*.json and data/SRP1/*/*.json (case and parameter files) '
                 'for "bound_relax_factor" / "honor_original_bounds"',
        'hits': [h for h in _git('grep', '-l', '-e', 'bound_relax_factor', '-e', 'honor_original_bounds', '--',
                                 '*.py', ':(glob)data/SRP1/*.json', ':(glob)data/SRP1/*/*.json').split('\n') if h]}
    det = _git('grep', '-n', '-e', 'esso_complementarity_diagnostics', '-e', 'complementarity_ratio_max', '-e',
               'get_complementarity_violation_ratio', '--', '*.py', ':!p[0-9]*', ':!data/*')
    ev['detector_consumers_search'] = {
        'scope': 'git grep -n over tracked *.py outside the p<digit>* harnesses and data/ for '
                 '"esso_complementarity_diagnostics" / "complementarity_ratio_max" / "get_complementarity_violation_ratio"',
        'hits': [h for h in det.split('\n') if h],
        'reading': 'definition, call after each ESSO solve, log line, sink append, sink wiring, the persistent-worker '
                   'sink extend; no hit compares the value with a threshold or branches on it'}
    return ev


# ======================================================================================================================
#  markdown
# ======================================================================================================================
def _f(x, fmt='{:.3e}'):
    return '—' if x is None else (fmt.format(x) if isinstance(x, (int, float)) else str(x))


def md_summary(summary, cells_out):
    L = []
    L.append('# W169 — closure slack and ESSO P-net slack at the certified point, every T1–T11 evaluation')
    L.append('')
    L.append(f"Script `{SCRIPT_REL}`; zero solves (guard verified 0), pickle blocked (0 calls). Frozen tables "
             f"`{FZ_REL}` (sha256 {FZ_SHA[:8]}). Manuscript comments read at Overleaf {CLONE_HEAD_EXPECTED} "
             f"(main.tex sha256 {MAIN_SHA[:8]}) l. 747 and l. 857.")
    L.append('')
    L.append('## Definitions')
    L.append('')
    for k, v in summary['definitions'].items():
        L.append(f'- **{k}**: {v}')
    L.append('')
    L.append('## Prediction (expert, Addendum 68) against the outcome')
    L.append('')
    pr = summary['prediction']
    L.append(f"Prediction: {pr['statement']}")
    L.append('')
    L.append('| quantity | reading | certified points | scored | source of the score | pass | fail | max relative (argmax cell) |')
    L.append('|---|---|---:|---:|---|---:|---:|---|')
    for row in pr['rows']:
        L.append(f"| {row['quantity']} | {row['reading']} | {row['n_certified_points']} | {row['n_scored']} | "
                 f"{row['source']} | {row['n_pass']} | {row['n_fail']} | {_f(row['max_relative'])} ({row['argmax_cell']}) |")
    L.append('')
    L.append(f"Outcome: {pr['outcome_text']}")
    L.append('')
    L.append('## Coverage — every evaluation')
    L.append('')
    L.append('Closure: max over the blocks holding an active unit of s⁺+s⁻ (MWh; a TSO block = the sum over its '
             'units); for x = 0 cells (no active unit) the max over all blocks. Relative signed = per-unit bound / E^Av '
             '(max over blocks); relative abs = |s⁺+s⁻| / E^Av over the blocks with exactly one active unit. '
             'σ: max over (e, y, d, t) of σ⁺+σ⁻ (MW) and σ / S. C = committed record, U = uncommitted file '
             '(read-only), — = not carried at P. Every value: the cell\'s point (k* when certified).')
    L.append('')
    L.append('| cell | table | point | eval key / candidate | closure src | max s⁺+s⁻ (MWh) | argmax block | rel. signed | rel. abs | '
             'σ src | max σ (MW) | argmax (e, y, d, t) | σ/S signed | σ/S abs | D3 sum (C) | detector max ratio at P |')
    L.append('|---|---|---|---|---|---:|---|---:|---:|---|---:|---|---:|---:|---:|---:|')
    for co in cells_out:
        c = co['cell_meta']
        cl = co.get('closure') or {}
        sg = co.get('sigma_elementwise') or {}
        ss = co.get('sigma_sum') or {}
        det = co.get('detector') or {}
        inst = (f"{(c.get('eval_key') or '')[:16]} / {(c.get('candidate_key') or '')[:8]}" if c['kind'] == 'admm'
                else f"arm / {(((co.get('arm') or {}).get('instance') or {}).get('candidate_key') or '')[:8]}")
        if c['kind'] == 'arm':
            arm = co.get('arm') or {}
            L.append(f"| {c['cell']} | T6 | arm evaluation | {inst} | C (weighted total) | "
                     f"total {_f(arm.get('closure_total_weighted_mwh_equiv_phase_A'))} | — | — | — | "
                     f"n/a (no ESSO solved: {arm.get('solve_kinds')}) | — | — | — | — | — | — |")
            continue
        argb = cl.get('argmax_active_block', cl.get('argmax_block', '—'))
        clmax = cl.get('max_sum_s_mwh_active_blocks', cl.get('max_sum_s_mwh'))
        ag = sg.get('argmax') or {}
        agt = f"({ag.get('e')}, {ag.get('year')}, {ag.get('day')}, {ag.get('t')})" if ag else '—'
        cls = co['closure_src_flag']
        sgs = co['sigma_src_flag']
        L.append(f"| {c['cell']} | {', '.join(t for t in c['tables'] if t)} | {c['point_kind'].split(' ')[0]} {c['point']} | "
                 f"{inst} | {cls} | {_f(clmax)} | {argb} | {_f(cl.get('max_frac_signed'))} | "
                 f"{_f(cl.get('max_frac_abs'))} | {sgs} | {_f(sg.get('max_sigma_mw'))} | {agt} | "
                 f"{_f(sg.get('max_frac_signed'))} | {_f(sg.get('max_frac_abs'))} | {_f(ss.get('D3_mw'))} | "
                 f"{_f(det.get('max_ratio_at_point'))} |")
    L.append('')
    L.append('## Cells not covered at P, and why')
    L.append('')
    any_nc = False
    for co in cells_out:
        if co.get('not_carried'):
            any_nc = True
            for q, why in co['not_carried'].items():
                L.append(f"- **{co['cell_meta']['cell']}** — {q}: {why}")
            if co.get('closure_terminal_beside'):
                b = co['closure_terminal_beside']
                L.append(f"  - beside (terminal cycle {b['at_cycle']}, NOT the point, {b['label']}): max s⁺+s⁻ over "
                         f"active blocks {_f(b['max_sum_s_mwh_active_blocks'])} MWh at {b['argmax_block']} (relative "
                         f"signed {_f(b['max_frac_signed'])}); D3 {_f(b['esso_feasibility_violation_D3'])} MW")
    if not any_nc:
        L.append('- none')
    L.append('')
    L.append('## Which records carry the quantities (every ADMM cell)')
    L.append('')
    for k, v in summary['records_carrying'].items():
        L.append(f'- **{k}**: {v}')
    L.append('')
    L.append('## The min(pch, pdch) detector (l. 857)')
    L.append('')
    for k, v in summary['detector'].items():
        L.append(f'- **{k}**: {v}')
    L.append('')
    L.append('## Code locations (found by exact text at run time)')
    L.append('')
    for cl in summary['code_locations']:
        L.append(f"- `{cl['location']}` — {cl['what']}")
    L.append('')
    L.append('## Sentences for the two CONFIRM comments (scoped to this evidence)')
    L.append('')
    for k, v in summary['manuscript_sentences'].items():
        L.append(f'- **{k}**: {v}')
    L.append('')
    L.append('## Integrity')
    L.append('')
    L.append(f"Integrity failures: {len(summary['integrity_failures'])}" +
             (''.join(f'\n- {x}' for x in summary['integrity_failures']) if summary['integrity_failures'] else ''))
    L.append('')
    return '\n'.join(L) + '\n'


# ======================================================================================================================
#  post-run
# ======================================================================================================================
def manifest(paths):
    return {p: _sha_file(p) for p in sorted(paths)}


def post_run(out_dir):
    names = [os.path.join(out_dir, n) for n in (OUT_NAMES['log'], OUT_NAMES['typing_json'], OUT_NAMES['typing_log'],
                                                OUT_NAMES['man'])]
    for rel in names:
        if not os.path.exists(os.path.join(REPO, rel)):
            _log(f'[W169 post-run PRECONDITION FAILED] {rel} missing')
            sys.exit(1)
    man = manifest(names)
    _write_json(os.path.join(REPO, out_dir, OUT_NAMES['post']), man)
    _log(f'[W169 post-run] wrote {OUT_NAMES["post"]}: ' + ', '.join(f'{os.path.basename(k)} {v[:8]}' for k, v in man.items()))
    sys.exit(0)


# ======================================================================================================================
#  main
# ======================================================================================================================
def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--post-run', action='store_true')
    ap.add_argument('--out-dir', default=None)
    args = ap.parse_args()
    out_rel = args.out_dir or OUT_DIR
    out_abs = out_rel if os.path.isabs(out_rel) else os.path.join(REPO, out_rel)
    if args.post_run:
        post_run(out_rel)
    t0 = time.time()
    _log(f'[W169] start; HEAD {_git("rev-parse", "HEAD")}; out {out_rel}')
    # ---- preconditions --------------------------------------------------------------------------------------------
    pre = []
    if not os.path.isdir(out_abs):
        pre.append(f'{out_rel} missing (mkdir it before the launch; the launch log goes there)')
    else:
        extra = [f for f in os.listdir(out_abs) if f != OUT_NAMES['log']]
        if extra:
            pre.append(f'{out_rel} holds {extra} (write-once: refuse)')
    if _sha_file(FZ_REL) != FZ_SHA:
        pre.append(f'frozen JSON sha256 {_sha_file(FZ_REL)} != {FZ_SHA}')
    if _git('status', '--porcelain', '--', FZ_REL) != '':
        pre.append('frozen JSON not clean vs HEAD')
    if _sha_file(MAIN_REL) != MAIN_SHA:
        pre.append(f'main.tex sha256 {_sha_file(MAIN_REL)} != {MAIN_SHA}')
    clone_head = _git('rev-parse', '--short=7', 'HEAD', cwd=os.path.join(REPO, CLONE_REL))
    if clone_head != CLONE_HEAD_EXPECTED:
        pre.append(f'clone HEAD {clone_head} != {CLONE_HEAD_EXPECTED}')
    if _git('status', '--porcelain', '--', 'main.tex', cwd=os.path.join(REPO, CLONE_REL)) != '':
        pre.append('main.tex modified in the clone')
    with open(os.path.join(REPO, MAIN_REL), encoding='utf-8') as handle:
        main_lines = handle.read().split('\n')
    confirm = {}
    for ln, what in CONFIRM_LINES.items():
        txt = main_lines[ln - 1] + ' ' + main_lines[ln]
        confirm[str(ln)] = {'what': what, 'text': (main_lines[ln - 1] + '\n' + main_lines[ln]).strip()}
        if '[CONFIRM — W169]' not in txt:
            pre.append(f'main.tex l. {ln}: no "[CONFIRM — W169]" comment')
    if pre:
        for p in pre:
            _log(f'[W169 PRECONDITION FAILED] {p}')
        sys.exit(1)
    # ---- cells ----------------------------------------------------------------------------------------------------
    fz = _load_json(FZ_REL)
    cells = build_cells(fz)
    resolve_eval_dirs(cells)
    _log(f"[W169] cells: {len(cells)} (T2 {sum(1 for c in cells if c['group'] == 'T2')}, appended "
         f"{sum(1 for c in cells if c['group'] == 'T2 appended')}, T11 {sum(1 for c in cells if c['group'] == 'T11')}, "
         f"T6 arms {sum(1 for c in cells if c['group'] == 'T6')})")
    # ---- 1. inventory -----------------------------------------------------------------------------------------------
    gi = GitIndex()
    inventories = {}
    for c in cells:
        inventories[c['cell']] = inventory_cell(c, gi)
    _log('[W169] inventory built')
    # ---- 2. capture plan, asserted BEFORE extraction ---------------------------------------------------------------
    plans = {c['cell']: plan_cell(c, inventories[c['cell']]) for c in cells}
    missing = []
    planned_files = set()
    for c in cells:
        for pth in plan_paths(plans[c['cell']]):
            planned_files.add(pth)
            if not os.path.exists(os.path.join(REPO, pth)):
                missing.append(f"{c['cell']}: planned source missing: {pth}")
    for c in cells:
        pl = plans[c['cell']]
        for q in (['closure_total_weighted', 'solve_kinds'] if c['kind'] == 'arm' else ['capacities', 'detector']):
            if q not in pl['sources']:
                missing.append(f"{c['cell']}: no source for the required {q}")
        if c['kind'] == 'admm' and 'closure' not in pl['sources'] and 'closure' not in pl['not_carried']:
            missing.append(f"{c['cell']}: closure neither planned nor declared not carried")
        if c['kind'] == 'admm' and 'sigma_elementwise' not in pl['sources'] and 'sigma_elementwise' not in pl['not_carried']:
            missing.append(f"{c['cell']}: sigma neither planned nor declared not carried")
    if missing:
        for m in missing:
            _log(f'[W169 CAPTURE-PATH ASSERTION FAILED] {m}')
        sys.exit(1)
    inputs_fixed = [FZ_REL, MAIN_REL, CASE_SRP1_REL, CASE_3X3_REL, INSTANCE_3X3_RECORD_REL, ESS_PARAMS_REL, S48_REL,
                    SCRIPT_REL, 'network.py', 'model_construction_helpers.py', 'shared_energy_storage_data.py',
                    'shared_resources_planning.py', 'definitions.py', 'p515_g_g1_g4_admm_gates.py',
                    'p513_solve_profile_guard.py', 'gate_result_io.py'] + NET_PARAMS_RELS + NET_2025_RELS
    inputs_all = sorted(set(inputs_fixed) | planned_files)
    inputs_man_before = manifest(inputs_all)
    status_of = {p: gi.status(p) for p in sorted(planned_files)}
    _write_json(os.path.join(out_abs, OUT_NAMES['inputs_man']),
                {'written': 'BEFORE extraction (capture plan asserted)', 'files': inputs_man_before,
                 'planned_source_status': status_of})
    _write_json(os.path.join(out_abs, OUT_NAMES['plan']),
                {'written': 'BEFORE extraction; every planned file exists; nothing below is read except these sources',
                 'plans': plans})
    _log(f'[W169] capture plan asserted and written ({len(planned_files)} planned source files); extracting')
    # ---- 3. extraction ---------------------------------------------------------------------------------------------
    cells_out = []
    for c in cells:
        pl = plans[c['cell']]
        co = {'cell_meta': c, 'plan_ref': f"{OUT_NAMES['plan']} plans.{c['cell']}", 'not_carried': pl['not_carried'],
              'pickles': pl['pickles'], 'records_not_carrying': pl['records_not_carrying']}
        if c['kind'] == 'arm':
            co['arm'] = extract_arm(c, pl)
            co['closure_src_flag'] = 'C (total)'
            co['sigma_src_flag'] = 'n/a'
            cells_out.append(co)
            continue
        years, days = _case_orders(c['case'])
        er = _load_json(pl['sources']['capacities']['path'])
        caps = _capacities(er)
        # year-index cross-check against the ageing trajectory the record carries
        for n, nd in ((er.get('ageing_trajectory_terminal') or {}).get('nodes') or {}).items():
            for cell_ in nd.get('cells', []):
                if years[cell_['y']] != str(cell_['block_year']):
                    _fail(f"{c['cell']}: year index {cell_['y']} -> {years[cell_['y']]} != record {cell_['block_year']}")
        co['capacities'] = {'path': pl['sources']['capacities']['path'], 'label': pl['sources']['capacities']['label'],
                            'at': ('terminal == P' if c.get('eval_record_cycles_run') == c['point'] else
                                   f"terminal {c.get('eval_record_cycles_run')} != P {c['point']} (E^Av and S at the "
                                   'terminal used for the relative value; labelled)'),
                            'nodes': caps}
        co['index_orders'] = {'years': years, 'days': days,
                              'source': CASE_SRP1_REL if c['case'] == 'SRP1' else CASE_3X3_REL}
        if 'closure' in pl['sources']:
            co['closure'] = extract_closure(c, pl, caps, years)
        if 'closure_terminal_beside' in pl['sources']:
            co['closure_terminal_beside'] = extract_closure_terminal_beside(c, pl, caps)
        if 'sigma_elementwise' in pl['sources']:
            co['sigma_elementwise'] = extract_sigma(c, pl, caps, years, days)
        if 'sigma_sum' in pl['sources']:
            co['sigma_sum'] = extract_sigma_sum(c, pl, caps, years, days, co.get('sigma_elementwise'))
        if 'xlsx_crosscheck' in pl['sources']:
            co['xlsx_crosscheck'] = extract_xlsx(c, pl)
        co['detector'] = extract_detector(c, pl)
        if 'network_failures' in pl['sources']:
            co['network_failures'] = extract_network_failures(c, pl)
            if co.get('closure'):
                fb = {x['block']: x for x in co['network_failures']['at_point']}
                for key in ('argmax_frac_signed_block', 'argmax_active_block'):
                    b = co['closure'].get(key)
                    co['closure'][key + '_solve_at_point'] = (fb[b] if b in fb else 'no failure record at P (primary '
                                                              'solve accepted)') if b else None
        clab = (pl['sources'].get('closure') or {}).get('label')
        co['closure_src_flag'] = {'committed': 'C', 'uncommitted': 'U'}.get(clab, '—' if clab is None else clab)
        slabs = set(((pl['sources'].get('sigma_elementwise') or {}).get('labels') or {}).values())
        co['sigma_src_flag'] = ('U' if slabs == {'uncommitted'} else 'C' if slabs == {'committed'} else
                                '—' if not slabs else '/'.join(sorted(slabs)))
        cells_out.append(co)
    # sha256 of every source a value was read from (from the before-manifest; re-verified below)
    for co in cells_out:
        if co.get('closure'):
            co['closure']['source']['sha256'] = inputs_man_before[co['closure']['source']['path']]
        if co.get('sigma_elementwise'):
            co['sigma_elementwise']['source_sha256'] = {n: inputs_man_before[p_]
                                                        for n, p_ in co['sigma_elementwise']['source_paths'].items()}
        if co.get('sigma_sum'):
            co['sigma_sum']['sha256'] = inputs_man_before[co['sigma_sum']['path']]
        if co.get('detector'):
            co['detector']['sha256'] = inputs_man_before[co['detector']['path']]
        if co.get('capacities'):
            co['capacities']['sha256'] = inputs_man_before[co['capacities']['path']]
        if co.get('arm'):
            co['arm']['arm_record_sha256'] = inputs_man_before[co['arm']['arm_record']]
    inputs_man_after = manifest(inputs_all)
    changed = [p for p in inputs_all if inputs_man_after[p] != inputs_man_before[p]]
    if changed:
        _fail(f'inputs changed during the run: {changed}')
    # ---- code locations --------------------------------------------------------------------------------------------
    code_locations = []
    for rel, needle, what in CODE_NEEDLES:
        loc = _locate(rel, needle)
        if loc is None:
            _fail(f'code needle not found in {rel}: {needle[:80]}')
        code_locations.append({'location': loc, 'needle': needle, 'what': what})
    conf = configuration_evidence()
    # ---- scoring ----------------------------------------------------------------------------------------------------
    cert = [co for co in cells_out if co['cell_meta']['certified_point']]
    rows = []

    def _score(quantity, reading, getter, source_desc, pool):
        scored = [(co, getter(co)) for co in pool if getter(co) is not None]
        n_pass = sum(1 for _, v in scored if v <= THRESH)
        mx = max(scored, key=lambda t: t[1]) if scored else (None, None)
        rows.append({'quantity': quantity, 'reading': reading, 'n_certified_points': len(pool), 'n_scored': len(scored),
                     'source': source_desc, 'n_pass': n_pass, 'n_fail': len(scored) - n_pass,
                     'max_relative': mx[1], 'argmax_cell': mx[0]['cell_meta']['cell'] if mx[0] else None,
                     'failing_cells': [co['cell_meta']['cell'] + (' (superseded certificate)'
                                                                  if co['cell_meta'].get('superseded') else '')
                                       for co, v in scored if v > THRESH],
                     'unscored_cells': [co['cell_meta']['cell'] for co in pool if getter(co) is None]})

    _score('closure (s+ + s-) / E^Av', 'signed', lambda co: (co.get('closure') or {}).get('max_frac_signed'),
           'committed component_levels_terminal.json at P', cert)
    _score('closure (s+ + s-) / E^Av', 'absolute', lambda co: (co.get('closure') or {}).get('max_frac_abs'),
           'committed component_levels_terminal.json at P', cert)
    _score('sigma (sigma+ + sigma-) / S', 'signed', lambda co: (co.get('sigma_elementwise') or {}).get('max_frac_signed'),
           'UNCOMMITTED esso_capture file at P (per element)', cert)
    _score('sigma (sigma+ + sigma-) / S', 'absolute', lambda co: (co.get('sigma_elementwise') or {}).get('max_frac_abs'),
           'UNCOMMITTED esso_capture file at P (per element)', cert)
    _score('sigma (sigma+ + sigma-) / S', 'signed, committed bound',
           lambda co: (co.get('sigma_sum') or {}).get('per_element_upper_bound_frac_of_min_S'),
           'committed D3 sum at P -> per-element upper bound (relaxed-bound floor) / min active S', cert)
    no_active_cert = [co['cell_meta']['cell'] for co in cert
                      if co.get('closure') and co['closure'].get('max_frac_signed') is None]
    zero_ok = all(co['closure'].get('inactive_only_blocks_all_exactly_zero') in (True, None)
                  for co in cells_out if co.get('closure'))
    signed_fail = any(r['n_fail'] for r in rows if r['reading'] == 'signed')
    abs_rows = [r for r in rows if r['reading'] == 'absolute']
    outcome_text = (
        f"closure scored at {rows[0]['n_scored']} / {len(cert)} certified points (committed), sigma per element at "
        f"{rows[2]['n_scored']} / {len(cert)} (uncommitted esso_capture), sigma by the committed sum bound at "
        f"{rows[4]['n_scored']} / {len(cert)}. SIGNED reading: "
        + ('HOLDS at every scored point' if not signed_fail else 'FAILS at ' + '; '.join(
            f"{r['quantity']}: {r['failing_cells']}" for r in rows if r['reading'] == 'signed' and r['n_fail']))
        + '. ABSOLUTE reading: ' + '; '.join(f"{r['quantity']}: {r['n_pass']} pass / {r['n_fail']} fail (max "
                                              f"{_f(r['max_relative'])} at {r['argmax_cell']})" for r in abs_rows)
        + f". Certified points with no active unit (x = 0; relative closure undefined, absolute level reported): "
          f"{no_active_cert}; every block without an active unit is exactly 0.0 in every scored cell: {zero_ok}.")
    # all-evaluation activity for the l. 747 sentence
    all_admm = [co for co in cells_out if co['cell_meta']['kind'] == 'admm']
    cl_cov = [co for co in all_admm if co.get('closure')]
    cl_pos_rel = [co['cell_meta']['cell'] for co in cl_cov
                  if co['closure'].get('max_frac_signed') is not None and co['closure']['max_frac_signed'] > THRESH]
    cl_rel_vals = [co['closure']['max_frac_signed'] for co in cl_cov if co['closure'].get('max_frac_signed') is not None]
    cl_pos_any = [co['cell_meta']['cell'] for co in cl_cov if co['closure']['n_blocks_positive']]
    max_pos = max((co['closure']['max_positive_s_mwh'] for co in cl_cov if co['closure']['max_positive_s_mwh']),
                  default=None)
    arms_zero = all((co['arm'] or {}).get('closure_total_weighted_eur_phase_A') == 0.0 for co in cells_out if co.get('arm'))
    sg_cov = [co for co in cert if co.get('sigma_elementwise') and co['sigma_elementwise'].get('max_frac_signed') is not None]
    sg_pos = [co['cell_meta']['cell'] for co in sg_cov if co['sigma_elementwise']['max_frac_signed'] > THRESH]
    n_closure_all = len(all_admm)
    cl_max_pos_rel = max(((co['closure']['max_frac_signed'], co['cell_meta']['cell']) for co in cl_cov
                          if co['closure'].get('max_frac_signed') is not None), default=(None, None))
    not_cov = [co['cell_meta']['cell'] for co in all_admm if not co.get('closure')]
    scope_747 = (f"[Scope: {len(cl_cov)} of {n_closure_all} ADMM evaluations carried at their point by a committed "
                 f"record (component_levels_terminal.json); not carried at the point: {not_cov}; the six T6 arms are at "
                 f"x = 0, every closure pair bounded [0, 0], recorded weighted total exactly 0.0: {arms_zero}.]")
    sup = {co['cell_meta']['cell'] for co in cl_cov if co['cell_meta'].get('superseded')}
    cl_pos_rel_ns = [x for x in cl_pos_rel if x not in sup]
    cl_max_ns = max(((co['closure']['max_frac_signed'], co['cell_meta']['cell']) for co in cl_cov
                     if co['closure'].get('max_frac_signed') is not None and co['cell_meta']['cell'] not in sup),
                    default=(None, None))
    pos_detail = {co['cell_meta']['cell']: {'superseded': bool(co['cell_meta'].get('superseded')),
                                            'block': co['closure'].get('argmax_frac_signed_block'),
                                            'relative': co['closure'].get('max_frac_signed'),
                                            's_mwh': co['closure'].get('max_sum_s_mwh_active_blocks'),
                                            'solve_at_point': co['closure'].get('argmax_frac_signed_block_solve_at_point')}
                  for co in cl_cov if co['cell_meta']['cell'] in cl_pos_rel}
    floor_ok = all(co['closure'].get('every_active_block_at_or_above_relaxed_floor', True) for co in cl_cov)
    if not cl_pos_rel:
        s747 = ('The closure slack was inactive in every evaluation reported in this paper. '
                f'(Evidence: s+ + s- <= {_f(cl_max_pos_rel[0])} x E^Av at every block, signed; the pairs sit at their '
                'lower bound within the solver\'s bound relaxation.) ' + scope_747)
    else:
        s747 = (f"The strict sentence is NOT supported at 1e-6 of E^Av: s+ + s- exceeds 1e-6 x E^Av at {cl_pos_rel} "
                f"(largest {_f(cl_max_pos_rel[0])} x E^Av = {_f(max_pos)} MWh, at {cl_max_pos_rel[1]}). A sentence the "
                f"evidence supports: 'The closure slack was at most {_f(max_pos, '{:.1e}')} MWh "
                f"({_f(cl_max_pos_rel[0], '{:.0e}')} of the available energy) in every evaluation reported in this "
                "paper.' Detail: " + json.dumps(pos_detail) + '. Excluding superseded certificates: '
                + (f"no block exceeds 1e-6 x E^Av (signed max {_f(cl_max_ns[0])} at {cl_max_ns[1]}), and the strict "
                   "sentence holds for every evaluation the tables use as a current certificate. "
                   if not cl_pos_rel_ns else f"exceeds at {cl_pos_rel_ns}. ")
                + f"Every negative block level is at or above the solver's relaxed-bound floor: {floor_ok}. " + scope_747)
    sg_max = max((co['sigma_elementwise']['max_frac_signed'] for co in sg_cov), default=None)
    sg_abs = max((co['sigma_elementwise']['max_frac_abs'] for co in sg_cov), default=None)
    if sg_cov and not sg_pos:
        s857 = ('The slack pair sigma was inactive at every certified point: sigma+ + sigma- stood at its lower bound to '
                f'solver tolerance (signed maximum {_f(sg_max)}, absolute maximum {_f(sg_abs)} of the converter '
                'rating). The post-solve check of min(P^Ch, P^Dch) / S runs after every ESSO solve in the production '
                'path and its value is recorded for every solve.')
    elif sg_pos:
        s857 = f'NOT SUPPORTED as worded: sigma exceeds 1e-6 x S at {sg_pos}.'
    else:
        s857 = 'no per-element evidence'
    s857 += (f" [Scope: per-element sigma from UNCOMMITTED esso_capture files at {len(sg_cov)} of {len(cert)} certified "
             f"points (the rest have no active unit: x = 0); the committed record carries only the sum over all "
             f"elements (esso_feasibility_violation_D3), which bounds every element at the same level under the "
             f"relaxed-bound floor.]")
    sentences = {'l. 747 (closure)': s747, 'l. 857 (sigma)': s857}
    summary = {
        'stage': 'P5.15 Addendum 68 order, W169 (zero solves, no model loads)',
        'script': SCRIPT_REL, 'git_HEAD': _git('rev-parse', 'HEAD'),
        'frozen_tables': {'path': FZ_REL, 'sha256': FZ_SHA},
        'manuscript_comments': {'clone': CLONE_REL, 'clone_head': CLONE_HEAD_EXPECTED, 'main_tex_sha256': MAIN_SHA,
                                'lines': confirm},
        'definitions': {
            'closure slack': 's+ / s- = slack_shared_es_soc_final_up / _down (network block models; scenario-free copy)',
            'closure level per block': 'component_levels_terminal.json blocks[b].unweighted.shared_ess_day_balance_slack '
                                       '/ PENALTY_SHARED_ESS_BALANCE = sum over the block\'s units of (s+ + s-) in MWh',
            'closure relative': '(s+ + s-) / E^Av(unit, year); TSO blocks: per-unit bound (sum + (n_active - 1) x 2 x '
                                f'{FLOOR_MWH_PER_VAR:g} MWh) / smallest active E^Av; E^Av = published e_available',
            'sigma': 'sigma+ / sigma- = slack_es_pnet_up / _down (ESSO, per node and (y, d, t), MW)',
            'sigma relative': '(sigma+ + sigma-) / S(e, y), S = the cohort converter rating s_max (MVA)',
            'threshold': f'{THRESH:g} (the prediction\'s "1e-6 relative")',
            'signed reading': 'max relative <= 1e-6 (the slack sum may be negative: IPOPT bound relaxation)',
            'absolute reading': 'max |relative| <= 1e-6',
            'point': 'certified: k*; uncertified: terminal (frozen end_cycle); T11: certification cycle; '
                     'w98: continuation terminal; arms: arm evaluation',
        },
        'n_cells': len(cells_out), 'n_certified_points': len(cert),
        'prediction': {'statement': 'both slacks are zero to solver tolerance (<= 1e-6 relative) at every certified '
                                    'point (expert, Addendum 68)', 'rows': rows, 'outcome_text': outcome_text},
        'closure_activity_all_evaluations': {
            'n_admm_evaluations': n_closure_all, 'n_carried_at_point': len(cl_cov),
            'cells_above_threshold_signed': cl_pos_rel, 'cells_with_any_positive_block': cl_pos_any,
            'max_positive_block_level_mwh': max_pos, 't6_arms_total_exactly_zero': arms_zero,
            'cells_above_threshold_signed_excluding_superseded': cl_pos_rel_ns,
            'max_relative_signed_excluding_superseded': cl_max_ns, 'above_threshold_detail': pos_detail,
            'every_negative_level_within_relaxed_bound_floor': floor_ok,
            'sigma_every_element_within_relaxed_bound_floor': all(
                co['sigma_elementwise'].get('every_element_at_or_above_relaxed_floor', True)
                for co in cells_out if co.get('sigma_elementwise'))},
        'records_carrying': {
            'closure at P (exact, s+ + s- per block)': 'component_levels_terminal.json (committed) blocks.*.unweighted.'
                                                       'shared_ess_day_balance_slack -- only when its cycles_run == P',
            'closure (signed s+ - s- per unit, terminal)': 'results/*.xlsx "Relaxation Slacks TSO, DSOs" -- only where a '
                                                           'workbook exists (uncommitted)',
            'closure, weighted total': 'evaluation_record.json component_decomposition_totals_weighted.'
                                       'shared_ess_day_balance_slack (committed; total over blocks)',
            'closure NOT carried': 'per-cycle records (recourse_blocks_all.jsonl, creep_diagnostic_per_cycle.jsonl) '
                                   'carry only slack_penalties = voltage + flexibility + closure summed; '
                                   'ess_schedule_per_cycle.jsonl (uncommitted) carries pch / pdch, not s+-',
            'sigma per element at P': 'esso_capture/<label>/node<n>_cycle<P>.jsonl (UNCOMMITTED; active cohort-periods only)',
            'sigma sum at P': 'component_levels_terminal.json esso_feasibility_violation_D3 (committed; sum over every '
                              'node and (y, d, t)) -- only when cycles_run == P',
            'sigma, every node, terminal': 'results/*.xlsx "Slacks operation, aggregated" (uncommitted; where present)',
            'pickles (never read)': 'certified_models.pkl (persisted models at post_certification.json certification_cycle), '
                                    'esso_models_<label>.pkl (terminal ESSO models), results/FrozenSMOPF/*.pkl (fixtures)',
        },
        'detector': {
            'exists in the production path': 'yes -- see code_locations: `_get_esso_complementarity_diagnostics` is called '
                                             'in the production ESSO solve routine after every successful ESSO solve '
                                             '(node_id set, solver ipopt) and its result appended to the sink wired by '
                                             'SharedEnergyStorageData.optimize; persistent workers were off in every '
                                             'scored run (g record), so this sequential path is the one that ran',
            'what it measures': 'complementarity_ratio_max = max over active cohort-periods of min(pch, pdch) / s_max '
                                '(plus the barrier-identity estimate of spurious throughput); it is LOGGED and RECORDED; '
                                'no production code compares it with a threshold (scope: configuration_evidence.'
                                'detector_consumers_search)',
            'recorded per evaluation': 'g_<label>.json esso_complementarity_diagnostics_by_round (committed) and '
                                       'leak_classification_<label>.jsonl (committed), per ESSO solve and cycle; '
                                       'child_stdout.log lines "[INFO] Shared ESS complementarity diagnostics"',
            'outcomes at P': 'per cell in cells/<cell>.json detector.per_node_at_point; max over cells in the table above',
            'max ratio at P over certified points': max((co['detector']['max_ratio_at_point'] for co in cert
                                                         if co.get('detector') and co['detector'].get('max_ratio_at_point')
                                                         is not None), default=None),
            'argmax cell': max(((co['detector']['max_ratio_at_point'], co['cell_meta']['cell']) for co in cert
                                if co.get('detector') and co['detector'].get('max_ratio_at_point') is not None),
                               default=(None, None))[1],
        },
        'code_locations': code_locations,
        'configuration_evidence': conf,
        'manuscript_sentences': sentences,
        'integrity_failures': FAILED,
    }
    # ---- guards ----------------------------------------------------------------------------------------------------
    gfail = GUARD.verify(0)
    if gfail or PICKLE_COUNTS['load'] or PICKLE_COUNTS['loads']:
        _log(f'[W169 GUARD FAULT] {gfail} pickle {PICKLE_COUNTS}')
        sys.exit(1)
    run = {'stage': summary['stage'], 'script': SCRIPT_REL, 'script_sha256': _sha_file(SCRIPT_REL),
           'git_HEAD': _git('rev-parse', 'HEAD'), 'wall_s': round(time.time() - t0, 2),
           'interpreter': sys.executable, 'python': sys.version.split()[0],
           'guard': {'permitted': [], 'counts': GUARD.counts, 'verify_0_failures': gfail}, 'pickle_counts': PICKLE_COUNTS,
           'inputs_unchanged': not changed, 'integrity_failures': FAILED,
           'finished_utc': datetime.now(timezone.utc).isoformat()}
    # ---- write -----------------------------------------------------------------------------------------------------
    os.makedirs(os.path.join(out_abs, CELLS_SUB), exist_ok=False)
    for co in cells_out:
        safe = co['cell_meta']['cell'].replace(':', '_')
        _write_json(os.path.join(out_abs, CELLS_SUB, f'{safe}.json'), co)
    _write_json(os.path.join(out_abs, OUT_NAMES['inventory']), inventories)
    _write_json(os.path.join(out_abs, OUT_NAMES['summary']), summary)
    _write_x(os.path.join(out_abs, OUT_NAMES['md']), md_summary(summary, cells_out))
    _write_json(os.path.join(out_abs, OUT_NAMES['run']), run)
    written = [os.path.relpath(os.path.join(dp, f), REPO) for dp, _d, fs in os.walk(out_abs) for f in fs
               if f not in (OUT_NAMES['log'], OUT_NAMES['man'])]
    _write_json(os.path.join(out_abs, OUT_NAMES['man']), manifest(written))
    _log(f'[W169] {outcome_text}')
    _log(f'[W169] wrote {len(written)} files + manifest; integrity failures {len(FAILED)}; wall {time.time() - t0:.1f}s')
    sys.exit(3 if FAILED else 0)


if __name__ == '__main__':
    main()
