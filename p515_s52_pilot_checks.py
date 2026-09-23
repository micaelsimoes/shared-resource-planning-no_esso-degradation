"""
P5.15 Addendum 39 (task W47) -- ZERO-SOLVE checks for the multi-scenario pilot's harness extension, the
scenario-indexing audit fixes and the manuscript-output capture, BEFORE the pilot runs.

Nothing here solves. `SolveProfileGuard(permitted=())` is installed before any production import and
`verify(0)`-ed exactly at the end (0 solves, 0 process launches, 0 blocked). Production's model builders run with
`.optimize` intercepted (`p515_s44_scale_measurement.Interceptor`, the scale harness's own stub: records the call
and returns "no result"), so production's construction code runs and nothing reaches a solver.

CHECKS
 A  HARNESS BACK-COMPATIBILITY OF IDENTITIES. Every entry of every TRACKED committed campaign spec
    (data/SRP1/Results/**/campaign_spec_*.json carrying `eval_key`) recomputes its eval key under the modified
    `evaluation_key` with its own declarations -- the two new options absent -- to the committed value.
 B  SPEC FORMAT. A synthetic freeze WITHOUT the new options carries neither `configuration.derived_instance` nor
    any entry `interface_deviation_premium`; WITH them, both are recorded and the eval key differs from the
    undeclared key and from the premium-only / instance-only keys. The validators refuse malformed declarations.
 C  THE PILOT CHILD'S PRE-RUN PATH ON THE PILOT INSTANCE, in the order `_child_real` runs it:
    `install_derived_instance` (hash, production reader, checksum, `install_baseline`), `instance_investment_years`,
    the floor-row precheck (`p515_s40_polish_gap._build_floor_rows`), `_construct_arm_planning` with the unit
    candidate, and the configuration hook (`_config_hook_factory`: D checks, AA declaration, derived-instance
    checks, the premium at alpha = 0.50, the ESS ageing verification and read-back on probe ESSO models of the
    5-year ESSO).
 D  THE MULTI-SCENARIO CAPTURE AND THE WORKBOOK ON THE PILOT'S OWN BUILT MODELS: production's ADMM model builders
    (`create_admm_variables`, `create_distribution_networks_models(premium_alpha = 0.50)`,
    `create_transmission_network_model`, `create_shared_energy_storage_model`, `_prepare_*_objectives_for_admm`)
    with `.optimize` intercepted; every free variable put on a deterministic, box-feasible, NON-TRIVIAL point
    (W39's `_randomize_point` rule: a value from sha256(seed | name) inside the variable's bounds -- nothing is
    solved or feasible; the point exists so the identities have something to disagree about); then
    `multiscenario_terminal_capture` (every identity within the declared tolerance, the row 18 read-back on every
    DSO block), `write_multiscenario_terminal` and `write_operational_workbook` (production's writer) to scratch.
 E  THE TWO SCENARIO-INDEXING FIXES, demonstrated on those models:
    E1 `_process_scenario_dispersion_results` (the workbook's 'Scenario Dispersion' sheet) reads the shared-ESS
       schedule at the non-anticipative copy: with the expectation Var equal to that copy and the unwired copies at
       their post-solve value (their initial 0 -- they are in no row), the fixed reading reports ZERO dispersion,
       while the pre-fix reading (production's own `_weighted_scenario_dispersion` on the per-scenario copies)
       reports the spurious one.
    E2 `_s31c_interface_detail` flexibility volumes: with the scenario-free TSO interface delta set, the fixed sum
       is sum_t |delta| (probability-weighted over the scenario keys), where the pre-fix sum was n_scenarios x it.
 F  SRP1 INVARIANCE OF E1 / E2, structurally: at SRP1 every network's market and operation probability vectors
    are exactly [1.0] (so omega = 1.0 * 1.0 = 1.0 and 1.0 * x == x bit for bit), and a built SRP1 block has the
    single scenario pair (0, 0), which IS `sess_na_scenario` -- the fixed readings read the same VarData the
    pre-fix readings read.
 D2 THE CHILD POST-RUN HOOK'S OTHER WRITERS on those models (zero solves): the S31C settlement detail and the
    component levels (the penalty table), the interface-voltage writer, the ESSO capture of `run_admm_arm`, the
    harness's ageing trajectory, the ESS ageing read-back on CLONES, `get_updated_capacities`.
 D3 THE POST-CERTIFICATION PATH UP TO THE FIRST SOLVE: `_persist_certified_models` (the ~1 GB pickle: size, time,
    memory), `per_block_base_objectives`, `hull_entries_with_esso` and `apply_hull_bounds` on the 5-year pilot
    (W37 exercised them at 2x2 on ONE year). The polish solves themselves are not run.
 G  The harness's own rule-eleven checklists pass (`assert_record_capture_paths`,
    `assert_post_certification_capture_paths`).
 H  The terminal-phase lock (`acquire_terminal_phase_lock` / `release_terminal_phase_lock`): held; a second
    acquisition while a LIVE pid holds it waits and then runs without the lock (recorded); release removes only the
    holder's own lock; a lock whose pid is dead is taken over (recorded).

EXACT COMMAND (repo root, canonical interpreter, attached, both streams captured):
  /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s52_pilot_checks.py --label r1 \\
      --scratch <dir outside the repo> > data/SRP1/Results/P515S52/pilot_checks_r1_launch.log 2>&1
OUTPUT (write-once): data/SRP1/Results/P515S52/pilot_checks/<label>/{pilot_checks.json, manifest_sha256.json}; the
dry-test capture JSON and workbook go to --scratch and are hash-recorded only (placeholder values, not results).
Exit 0 when every check passes, 1 otherwise.
"""

import argparse
import hashlib
import json
import os
import subprocess
import sys
import time
import traceback
from datetime import datetime, timezone

REPO = os.path.dirname(os.path.abspath(__file__))
if REPO not in sys.path:
    sys.path.insert(0, REPO)
os.chdir(REPO)

from p513_solve_profile_guard import SolveProfileGuard  # noqa: E402

GUARD = SolveProfileGuard(permitted=(), label='P5.15 W47 pilot checks (zero solves)').install()

import pyomo.environ as pe  # noqa: E402
import p515_s44_campaign_harness as H  # noqa: E402
import p515_s44_scale_measurement as S  # noqa: E402
import p515_g_g1_g4_admm_gates as G  # noqa: E402
import p56a_oracle as O  # noqa: E402
import shared_resources_planning as srp  # noqa: E402
import model_construction_helpers as MCH  # noqa: E402

SCHEMA = 'p515_s52_pilot_checks_v1'
STAGE = 'P5.15 Addendum 39 W47 -- zero-solve checks for the multi-scenario pilot'
OUT_ROOT_REL = os.path.join('data', 'SRP1', 'Results', 'P515S52', 'pilot_checks')
ALPHA = 0.50
PREMIUM = {'alpha': ALPHA, 'floor': None}
CASE_FILE_AA = {'enabled': True, 'memory': 5, 'regularization': 1e-10, 'reject_policy': 'keep_memory'}
ESS_AGEING_BASELINE = {'calendar_life_years': 15, 'cycle_life_nominal': 10000, 'depth_of_discharge_nominal': 0.8,
                       'minimum_soh': 0.7, 'calendar_retention_per_year': 0.985,
                       'calibration': {'status': 'ACTIVE', 'cycles_n': 10000, 'reference_dod_d': 0.8,
                                       'eol_retention_r': 0.8}}
ESS_AGEING_LABEL = 'BASELINE (C2 + phi_cal 0.985 + soh_min 0.70)'
UNIT_MAP = {5: (0.0, 0.0), 7: (0.25, 1.0), 9: (0.0, 0.0)}
EXPECTED_CASE_SHA256 = '7ecff44a874892187d1dd2e3d5ed0664a2f90a4bfa4cdb820a4abcbbd828f949'
EXPECTED_SCENARIO_CHECKSUM = '53b4bea4142001617282a08541079952564903b6d84db066f90d71a650b7d563'
SEED = 'W47'
SIGMA_STATE_PLACEHOLDER = {'sigma_fixed': 93635360.0, 'sigma_computed': 33764631.272205874,
                           'admm_diagnostics': [], 'solver_recovery_diagnostics': []}


def _utc():
    return datetime.now(timezone.utc).isoformat()


class _MemoryMarks:
    """Current RSS (psutil) and peak RSS (ru_maxrss, bytes on macOS) at each stage boundary -- a zero-solve data
    point for the pilot child's memory (the same models, the capture and the workbook)."""

    def __init__(self):
        self.t0 = time.time()
        self.marks = []

    def mark(self, stage):
        import psutil
        import resource
        self.marks.append({'stage': stage, 't_s': round(time.time() - self.t0, 3),
                           'rss_bytes': psutil.Process().memory_info().rss,
                           'ru_maxrss_bytes': resource.getrusage(resource.RUSAGE_SELF).ru_maxrss})
        _log(f"memory at '{stage}': rss {self.marks[-1]['rss_bytes'] / (1 << 30):.2f} GiB, peak "
             f"{self.marks[-1]['ru_maxrss_bytes'] / (1 << 30):.2f} GiB")


MEM = _MemoryMarks()


def _log(msg):
    print(f'{datetime.now(timezone.utc).strftime("%H:%M:%S")} [W47-checks] {msg}', flush=True)


def _git(args):
    return subprocess.run(['git', *args], cwd=REPO, capture_output=True, text=True).stdout.strip()


def _randomize_point(model, seed):
    """W39's rule (p515_s51_row18_zero_solve_checks._randomize_point, replicated because that module installs its
    own guard at import): every unfixed variable gets a deterministic value from sha256(seed|name), inside its
    bounds. Nothing is solved or feasible."""
    n_set = 0
    for component in model.component_objects(pe.Var, active=None):
        for data in component.values():
            if data.fixed:
                continue
            digest = hashlib.sha256(f'{seed}|{data.name}'.encode()).digest()
            unit = int.from_bytes(digest[:8], 'big') / float(1 << 64)
            lower, upper = data.lb, data.ub
            if lower is not None and upper is not None:
                value = lower + unit * (upper - lower)
            elif lower is not None:
                value = lower + 0.1 + 0.9 * unit
            elif upper is not None:
                value = upper - (0.1 + 0.9 * unit)
            else:
                value = 0.1 + 0.9 * unit
            data.set_value(value)
            n_set += 1
    return n_set


# ======================================================================================================================
#  A / B -- harness identities and format
# ======================================================================================================================
def check_a_keys():
    files = [f for f in _git(['ls-files', '--', 'data/SRP1/Results']).splitlines()
             if os.path.basename(f).startswith('campaign_spec_') and f.endswith('.json')]
    per_file, n_entries, mismatches = {}, 0, []
    for rel in sorted(files):
        with open(os.path.join(REPO, rel)) as handle:
            spec = json.load(handle)
        cfg = spec.get('configuration') or {}
        n_file = 0
        for entry in spec.get('candidates') or []:
            if 'eval_key' not in entry:
                continue
            recomputed = H.evaluation_key(entry['key'], entry.get('overrides') or {},
                                          case_file_aa=cfg.get('case_file_anderson_acceleration'),
                                          model_variant=entry.get('model_variant'),
                                          ess_ageing_baseline=cfg.get('ess_ageing_baseline'),
                                          flex_price_multiplier=entry.get('flex_price_multiplier'))
            n_file += 1
            if recomputed != entry['eval_key']:
                mismatches.append({'file': rel, 'label': entry.get('label'), 'committed': entry['eval_key'],
                                   'recomputed': recomputed})
        per_file[rel] = n_file
        n_entries += n_file
    return {'n_spec_files': len(files), 'n_entries_with_eval_key': n_entries, 'per_file': per_file,
            'mismatches': mismatches, 'pass': n_entries > 0 and not mismatches}


def check_b_format(scratch, derived):
    out = {}
    nodes = {5: (0.0, 0.0), 7: (0.25, 1.0), 9: (0.0, 0.0)}
    base_cfg = {'name': 'W47 synthetic', 'arm_label': 's39_D', 'overrides': {},
                'case_file_anderson_acceleration': dict(CASE_FILE_AA),
                'ess_ageing_baseline': dict(ESS_AGEING_BASELINE), 'ess_ageing_baseline_label': ESS_AGEING_LABEL}
    root_plain = os.path.join(scratch, 'synthetic_freeze_plain')
    _p, _s, plain = H.freeze_campaign_spec(root_plain, 'w47_synthetic_plain', [('u', nodes, {'investment_year': 2025})],
                                           configuration=dict(base_cfg), cap=2, concurrency=1, authority=['W47 checks'])
    root_decl = os.path.join(scratch, 'synthetic_freeze_declared')
    _p2, _s2, decl = H.freeze_campaign_spec(
        root_decl, 'w47_synthetic_declared',
        [('u', nodes, {'investment_year': 2025, 'interface_deviation_premium': dict(PREMIUM)})],
        configuration=dict(base_cfg, derived_instance=derived), cap=2, concurrency=1, authority=['W47 checks'])
    key = H.candidate_key(H.canonical_candidate(nodes, investment_year=2025))
    k_plain = H.evaluation_key(key, {}, case_file_aa=CASE_FILE_AA, ess_ageing_baseline=ESS_AGEING_BASELINE)
    k_both = H.evaluation_key(key, {}, case_file_aa=CASE_FILE_AA, ess_ageing_baseline=ESS_AGEING_BASELINE,
                              derived_instance=derived, interface_deviation_premium=PREMIUM)
    k_inst = H.evaluation_key(key, {}, case_file_aa=CASE_FILE_AA, ess_ageing_baseline=ESS_AGEING_BASELINE,
                              derived_instance=derived)
    k_prem = H.evaluation_key(key, {}, case_file_aa=CASE_FILE_AA, ess_ageing_baseline=ESS_AGEING_BASELINE,
                              interface_deviation_premium=PREMIUM)
    k_alpha0 = H.evaluation_key(key, {}, case_file_aa=CASE_FILE_AA, ess_ageing_baseline=ESS_AGEING_BASELINE,
                                derived_instance=derived, interface_deviation_premium={'alpha': 0.0, 'floor': None})
    out['plain_spec_has_no_new_keys'] = ('derived_instance' not in plain['configuration']
                                         and all('interface_deviation_premium' not in e for e in plain['candidates']))
    out['plain_eval_key_is_the_undeclared_formula'] = plain['candidates'][0]['eval_key'] == k_plain
    out['declared_spec_records_both'] = (decl['configuration'].get('derived_instance') == derived
                                         and decl['candidates'][0].get('interface_deviation_premium') == PREMIUM)
    out['declared_eval_key_recomputes'] = decl['candidates'][0]['eval_key'] == k_both
    out['keys_all_distinct'] = len({k_plain, k_both, k_inst, k_prem, k_alpha0}) == 5
    refusals = {}
    bad_derived = [dict(derived, instance_label='srp1'), dict(derived, case_sha256='xyz'),
                   {k: v for k, v in derived.items() if k != 'scenario_checksum'}, dict(derived, extra=1),
                   dict(derived, case_path='/abs/path.json')]
    for i, bad in enumerate(bad_derived):
        try:
            H.validate_derived_instance(bad)
            refusals[f'derived_{i}'] = False
        except ValueError:
            refusals[f'derived_{i}'] = True
    for i, bad in enumerate([{'alpha': -0.1, 'floor': None}, {'alpha': True, 'floor': None}, {'alpha': 0.5},
                             {'alpha': 0.5, 'floor': None, 'x': 1}, {'alpha': float('nan'), 'floor': None}]):
        try:
            H.validate_interface_deviation_premium(bad)
            refusals[f'premium_{i}'] = False
        except ValueError:
            refusals[f'premium_{i}'] = True
    try:
        H.freeze_campaign_spec(os.path.join(scratch, 'synthetic_freeze_badsha'), 'w47_badsha',
                               [('u', nodes, {'investment_year': 2025})],
                               configuration=dict(base_cfg, derived_instance=dict(derived, case_sha256='0' * 64)),
                               cap=2, concurrency=1, authority=['W47 checks'])
        refusals['freeze_refuses_wrong_case_sha256'] = False
    except ValueError:
        refusals['freeze_refuses_wrong_case_sha256'] = True
    out['refusals'] = refusals
    out['all_refused'] = all(refusals.values())
    out['keys'] = {'plain': k_plain, 'instance_and_premium': k_both, 'instance_only': k_inst,
                   'premium_only': k_prem, 'instance_alpha0': k_alpha0}
    out['pass'] = all(v for k, v in out.items() if isinstance(v, bool))
    return out


# ======================================================================================================================
#  C -- the child's pre-run path on the pilot instance
# ======================================================================================================================
def derive_pilot_case(out_dir):
    case, spec, changes = S.derive_case('paper', {'num_market_scenarios': 2, 'num_operation_scenarios': 2})
    case_dir = os.path.join(out_dir, 'case')
    os.makedirs(case_dir)
    path = os.path.join(case_dir, 'SRP1__s52_pilot_2x2.json')
    with open(path, 'x') as handle:
        handle.write(json.dumps(case, indent='\t'))
    source = os.path.join(REPO, 'data', 'SRP1', 'SRP1.json')
    return {'instance_label': 's52_pilot_2x2', 'case_path': os.path.relpath(path, REPO),
            'case_sha256': H.sha256_file(path), 'scenario_checksum': EXPECTED_SCENARIO_CHECKSUM,
            'source_case_path': os.path.relpath(source, REPO), 'source_case_sha256': H.sha256_file(source),
            'changes_vs_source': changes}


def check_c_child_prerun(derived, scratch, label):
    from p515_s40_polish_gap import _build_floor_rows
    out = {}
    probe_dir = os.path.join(scratch, 'child_probe')
    os.makedirs(probe_dir)
    t0 = time.time()
    installed = H.install_derived_instance(derived, probe_dir)
    MEM.mark('C: derived instance installed (the baseline planning in memory)')
    out['install'] = {k: v for k, v in installed.items() if k != 'planning_dimensions'}
    out['install']['wall_s'] = time.time() - t0
    out['installed_label_and_checksum'] = (installed['instance_label'] == derived['instance_label']
                                           and installed['scenario_checksum_in_child'] == derived['scenario_checksum'])
    out['case_sha256_equals_w45'] = installed['case_sha256_in_child'] == EXPECTED_CASE_SHA256
    years = H.instance_investment_years()
    out['instance_investment_years'] = years
    out['investment_year_2025_in_instance'] = 2025 in years and len(years) == 5
    ids = {'run': f'p515s52_checks_{label}_run', 'precheck': f'p515s52_checks_{label}_precheck'}
    for eid in ids.values():
        if os.path.exists(os.path.join(G.O.WORK_DIR, eid)):
            raise RuntimeError(f'working dir id already used: {eid}')
    _cc, floor_rows, floor_counts = _build_floor_rows(ids['precheck'])
    out['floor_rows_counts_by_node'] = {str(k): v for k, v in floor_counts.items()}
    report = {}
    planning, sed, candidate = G._construct_arm_planning(
        's52checks', os.path.join(scratch, 'arm'), report, investment_map=UNIT_MAP, eval_id=ids['run'],
        num_max_iters_override=500, apply_rho=False, investment_year=2025)
    spec_like = {'configuration': {'overrides': {}, 'case_file_anderson_acceleration': dict(CASE_FILE_AA),
                                   'ess_ageing_baseline': dict(ESS_AGEING_BASELINE),
                                   'ess_ageing_baseline_label': ESS_AGEING_LABEL,
                                   'ess_params_file': {'path': H.ESS_PARAMS_FILE_REL,
                                                       'sha256': H.sha256_file(os.path.join(REPO, H.ESS_PARAMS_FILE_REL))},
                                   'derived_instance': derived},
                 'cap': 500, 'required_consecutive_cycles': 10}
    holder = {}
    hook = H._config_hook_factory(spec_like, holder, overrides={}, investment_year=2025,
                                  expected_floor_rows=floor_rows, interface_deviation_premium=PREMIUM)
    MEM.mark('C: floor precheck + arm planning constructed (fresh_planning deepcopy)')
    report.setdefault('rule_eleven_checklist', {})
    hook(planning=planning, sed=sed, candidate=candidate, report=report)
    MEM.mark('C: configuration hook (probe ESSO read-back) done')
    out['configuration_checks'] = holder.get('configuration_checks')
    out['derived_instance_checks'] = holder.get('derived_instance_checks')
    out['premium_applied'] = holder.get('interface_deviation_premium_applied')
    verified = holder.get('ess_ageing_verified_pre_run') or {}
    out['ess_ageing_checks'] = verified.get('checks')
    out['ess_ageing_readback_all_match'] = (verified.get('readback_pre_run') or {}).get('all_match')
    out['ess_ageing_floor_rows_identical'] = verified.get('floor_rows_identical_to_precheck')
    out['premium_in_force'] = dict(planning.params.admm.interface_deviation_premium)
    out['pass'] = bool(out['installed_label_and_checksum'] and out['case_sha256_equals_w45']
                       and out['investment_year_2025_in_instance'] and all((out['configuration_checks'] or {}).values())
                       and all((out['derived_instance_checks'] or {}).values())
                       and (out['premium_applied'] or {}).get('took_effect')
                       and out['premium_in_force'].get('alpha') == ALPHA
                       and all((out['ess_ageing_checks'] or {}).values()) and out['ess_ageing_readback_all_match']
                       and out['ess_ageing_floor_rows_identical'])
    return out, planning, sed, candidate


# ======================================================================================================================
#  D / E -- the capture and the workbook on the pilot's own built models; the two fixes
# ======================================================================================================================
def build_models(planning, sed, candidate):
    interceptor = S.Interceptor()
    tn = planning.transmission_network
    tn.optimize = interceptor.network(tn, 'tso')
    for dn in planning.distribution_networks.values():
        dn.optimize = interceptor.network(dn, 'dso')
    sed.optimize = interceptor.esso(sed)
    premium = planning.params.admm.interface_deviation_premium
    t0 = time.time()
    try:
        cv, _dv = srp.create_admm_variables(planning)
        dso_models, res_dso = srp.create_distribution_networks_models(
            planning.distribution_networks, cv, candidate['total_capacity'], parallel_execution=False,
            premium_alpha=premium['alpha'], premium_floor=premium['floor'])
        tso_model, res_tso = srp.create_transmission_network_model(planning, cv, candidate['total_capacity'])
        esso_model, res_esso = srp.create_shared_energy_storage_model(sed, cv, candidate['investment'])
        srp._prepare_distribution_objectives_for_admm(planning.distribution_networks, dso_models)
        srp._prepare_transmission_objectives_for_admm(tn, tso_model)
    finally:
        del tn.optimize
        for dn in planning.distribution_networks.values():
            del dn.optimize
        del sed.optimize
    models = {'tso': tso_model, 'dso': dso_models, 'esso': esso_model}
    MEM.mark('D: ADMM models built (80 network blocks + 3 ESSO, no pristine clones)')
    calls = interceptor.counts()
    return models, cv, {'tso': res_tso, 'dso': res_dso, 'esso': res_esso}, {
        'interceptor_calls': calls, 'build_wall_s': time.time() - t0,
        'pass': calls == {'dso': len(planning.distribution_networks), 'tso': 1, 'esso_init': 1}}


def _all_blocks(planning, models):
    tn = planning.transmission_network
    for y in tn.years:
        for d in tn.days:
            yield 'TSO', None, tn, y, d, models['tso'][y][d]
    for n, dn in planning.distribution_networks.items():
        for y in dn.years:
            for d in dn.days:
                yield 'DSO', n, dn, y, d, models['dso'][n][y][d]


def check_d_capture(planning, models, results, scratch):
    out = {}
    n_set = 0
    for _k, _n, _h, _y, _d, model in _all_blocks(planning, models):
        n_set += _randomize_point(model, SEED)
    for node_id, model in models['esso'].items():
        n_set += _randomize_point(model, f'{SEED}|esso|{node_id}')
    out['n_variables_set'] = n_set
    MEM.mark('D: point set on every block')
    capture_dir = os.path.join(scratch, 'dry_capture')
    os.makedirs(capture_dir)
    written = H.write_multiscenario_terminal(planning, models, SIGMA_STATE_PLACEHOLDER, capture_dir, premium=PREMIUM)
    summary = written
    MEM.mark('D: multiscenario_terminal written')
    out['capture_wall_s'] = written['runtime_s']
    out['summary_checks'] = summary['checks']
    out['summary_all_checks_pass_is_python_bool'] = summary['all_checks_pass'] is True
    out['identity_worst_rel_diffs'] = summary['identity_worst_rel_diffs']
    out['n_blocks'] = summary['n_blocks']
    out['n_dso_blocks_row18_wired'] = summary['n_dso_blocks_row18_wired']
    out['all_dso_keys'] = sorted(summary['all_dso'])
    out['per_dso_nodes'] = sorted(summary['per_dso'])
    out['settlement_identity_keys'] = sorted(summary['settlement_identity'])
    out['placeholder_values_note'] = ('values are a deterministic non-solved point; the numbers are not results, only '
                                      'the identities and the code paths are checked')
    # the W46 formulas (p515_s51_coordinated_decomposition, pure functions, nothing installed at import) applied
    # to the WRITTEN capture must reproduce its aggregates -- one definition of E|d| and sum omega d^2
    import p515_s51_coordinated_decomposition as DEC
    with open(os.path.join(REPO, written['path'])) as handle:
        captured = json.load(handle)
    dec_blocks, cov_rel, weights_seen = {}, 0.0, set()
    for key, blk in captured['blocks'].items():
        disp = blk.get('dispersion')
        if not disp:
            continue
        dec = DEC.block_decomposition({k: v['d_p_mw'] for k, v in disp['per_scenario_d'].items()},
                                      disp['probabilities'], disp['pi_by_market_by_hour'], disp['pibar_by_hour'])
        dec_blocks['DSO:' + ':'.join(key.split('|')[1:])] = dec
        settle = blk['settlement']
        # DEC's covariance is the market part on d; the settlement deviation is on p_int = pbar + d, the pbar
        # terms cancel because sum_m omega_m (pi_m - pibar) = 0; units: DEC MW x price, the capture x baseMVA
        # is already in MW -> identical up to rounding at weight 1
        cov_rel = max(cov_rel, abs(dec['totals']['covariance'] * settle['settlement_weight'] - settle['deviation'])
                      / max(abs(settle['total']), abs(settle['contracted']), 1.0))
        weights_seen.add(settle['settlement_weight'])
    agg = DEC.aggregate(dec_blocks)
    mine = captured['summary']['all_dso']
    out['w46_formula_cross_check'] = {
        'n_blocks': agg['n_blocks'],
        'sum_omega_d2_dec': agg['all_dso']['sum_omega_d2'], 'sum_omega_d2_capture': mine['sum_omega_d2_p_mw2h_sum_over_blocks'],
        'E_abs_d_dec': agg['all_dso']['E_abs_d'], 'E_abs_d_capture': mine['E_abs_d_p_mwh_sum_over_blocks'],
        'max_covariance_rel_diff_vs_settlement_scale': cov_rel,
        'market_share_of_sum_omega_d2': agg['all_dso']['market_share_of_sum_omega_d2'],
        'settlement_weights_seen_on_dso_blocks': sorted(weights_seen)}
    w46 = out['w46_formula_cross_check']
    w46['pass'] = bool(agg['n_blocks'] == 60
                       and abs(w46['sum_omega_d2_dec'] - w46['sum_omega_d2_capture']) <= 1e-9 * abs(w46['sum_omega_d2_dec'])
                       and abs(w46['E_abs_d_dec'] - w46['E_abs_d_capture']) <= 1e-9 * abs(w46['E_abs_d_dec'])
                       and cov_rel <= 1e-9 and weights_seen == {1.0})
    out['write_multiscenario_terminal'] = {k: written.get(k) for k in ('status', 'path', 'sha256', 'runtime_s')}
    out['write_multiscenario_terminal']['size_bytes'] = os.path.getsize(os.path.join(REPO, written['path']))
    out['sigma_record_placeholder'] = written['sigma_calibration']
    planning.results_dir = os.path.join(scratch, 'dry_workbook')
    os.makedirs(planning.results_dir)
    wb = H.write_operational_workbook(planning, models, results, [], SIGMA_STATE_PLACEHOLDER, 1.0)
    MEM.mark('D: operational workbook written')
    out['write_operational_workbook'] = wb
    try:
        from openpyxl import load_workbook
        sheets = load_workbook(os.path.join(REPO, wb['path']), read_only=True).sheetnames
    except Exception as error:  # noqa: BLE001
        sheets = f'{type(error).__name__}: {error}'
    out['workbook_sheets'] = sheets
    out['pass'] = bool(summary['all_checks_pass'] is True and w46['pass'] and summary['n_dso_blocks_row18_wired'] == 60
                       and summary['n_blocks'] == 80 and written['status'] == 'written'
                       and wb['status'] == 'written' and isinstance(sheets, list) and 'Scenario Dispersion' in sheets)
    return out


def check_d2_writers(planning, sed, models, scratch):
    import p514_n_instrumented_cstar as N
    out_dir = os.path.join(scratch, 'dry_writers')
    os.makedirs(out_dir)
    t0 = time.time()
    s31c_path = G.write_interface_settlement_detail_s31c(planning, sed, models, [], {}, out_dir, 'dry')
    volt_path = G.write_interface_voltage_terminal(planning, models, out_dir, 'dry', cycle=None)
    with open(os.path.join(out_dir, 'component_levels_terminal.json')) as handle:
        cl = json.load(handle)
    esso_capture = N.capture_esso(models['esso'], sed)
    trajectory = H.ageing_trajectory_terminal(models['esso'], sed)
    readback = H.ess_ageing_readback_models(models['esso'], sed, ESS_AGEING_BASELINE, 2025, clone=True)
    caps = sed.get_updated_capacities(models['esso'])
    MEM.mark('D2: post-run writers done')
    out = {'wall_s': time.time() - t0,
           'component_levels_n_blocks': len(cl.get('blocks') or {}),
           'component_levels_recourse_keys_present': all(k in (cl.get('recourse_components') or {}) for k in (
               'gross_operational_cost', 'interface_settlement_deviation_total', 'interface_settlement_identity_residual',
               'voltage_pin_total')),
           's31c_written': os.path.isfile(os.path.join(out_dir, 'interface_settlement_detail_s31c.json')),
           'voltage_written': os.path.isfile(volt_path),
           'esso_capture_nodes': sorted(str(k) for k in esso_capture),
           'ageing_trajectory_nodes': sorted((trajectory.get('nodes') or {}).keys()),
           'ageing_trajectory_n_cells_node7': len(((trajectory.get('nodes') or {}).get('7') or {}).get('cells') or []),
           'ess_ageing_readback_on_clones_all_match': readback.get('all_match'),
           'updated_capacity_nodes': sorted(str(k) for k in caps),
           'paths': {'s31c': os.path.abspath(os.path.join(REPO, s31c_path)) if not os.path.isabs(s31c_path) else s31c_path,
                     'voltage': volt_path}}
    out['pass'] = bool(out['component_levels_n_blocks'] == 80 and out['component_levels_recourse_keys_present']
                       and out['s31c_written'] and out['voltage_written'] and len(out['esso_capture_nodes']) == 3
                       and out['ageing_trajectory_nodes'] == ['5', '7', '9'] and out['ess_ageing_readback_on_clones_all_match']
                       and len(out['updated_capacity_nodes']) == 3)
    return out


def check_d3_post_certification(planning, models, cv, scratch):
    import p515_s41_hull_polish as HP
    import p515_s42_exact_fix_rerun as EF
    out = {}
    persist_dir = os.path.join(scratch, 'dry_persist')
    os.makedirs(persist_dir)
    t0 = time.time()
    persisted = EF._persist_certified_models(models, persist_dir)
    out['persist'] = dict(persisted, wall_s=time.time() - t0)
    MEM.mark('D3: certified models pickled')
    t0 = time.time()
    before = O.per_block_base_objectives(planning, models)
    hull = HP.hull_entries_with_esso(planning, models, cv)
    descriptors = HP.apply_hull_bounds(planning, models, hull)
    out['hull'] = {'wall_s': time.time() - t0, 'n_per_block_base_objectives': len(before), 'n_hull_entries': len(hull),
                   'n_descriptors': len(descriptors),
                   'channels': sorted({d.get('channel') for d in descriptors if isinstance(d, dict)})}
    MEM.mark('D3: hull bounds applied (no solve)')
    out['pass'] = bool(persisted.get('size_bytes') and len(before) == 80 and len(hull) > 0 and len(descriptors) > 0)
    return out


def check_h_lock(scratch):
    lock = os.path.join(scratch, 'terminal_lock_test', H.TERMINAL_PHASE_LOCK_NAME)
    os.makedirs(os.path.dirname(lock))
    first = H.acquire_terminal_phase_lock(lock)
    second = H.acquire_terminal_phase_lock(lock, timeout_s=0.3, poll_s=0.1)
    released = H.release_terminal_phase_lock(lock, first)
    gone = not os.path.exists(lock)
    dead = subprocess.Popen(['true'])
    dead.wait()
    with open(lock, 'w') as handle:
        json.dump({'pid': dead.pid, 'utc': _utc()}, handle)
    third = H.acquire_terminal_phase_lock(lock, timeout_s=5.0, poll_s=0.1)
    released_third = H.release_terminal_phase_lock(lock, third)
    out = {'first': first, 'second': second, 'released_first': released, 'lock_removed': gone, 'third': third,
           'released_third': released_third}
    out['pass'] = bool(first['status'] == 'held' and second['status'] == 'timeout_ran_without_lock'
                       and second['held_by'].get('pid') == os.getpid() and released and gone
                       and third['status'] == 'held' and len(third['stale_locks_taken_over']) == 1
                       and released_third and not os.path.exists(lock))
    return out


def check_e_fixes(planning, models):
    out = {}
    # E1 -- the workbook's shared-ESS scenario dispersion
    for kind, node_id, holder, y, d, model in _all_blocks(planning, models):
        s_m0, s_o0 = MCH.sess_na_scenario(model)
        for e in model.shared_energy_storages:
            for p in model.periods:
                for s_m in model.scenarios_market:
                    for s_o in model.scenarios_operation:
                        if (s_m, s_o) != (s_m0, s_o0):
                            model.shared_es_pnet[e, s_m, s_o, p].set_value(0.0)   # unwired: their post-solve value
                            model.shared_es_qnet[e, s_m, s_o, p].set_value(0.0)
                na_p = pe.value(model.shared_es_pnet[e, s_m0, s_o0, p])
                na_q = pe.value(model.shared_es_qnet[e, s_m0, s_o0, p])
                if kind == 'TSO':
                    model.expected_shared_ess_p[e, p].set_value(na_p)
                    model.expected_shared_ess_q[e, p].set_value(na_q)
        if kind == 'DSO':
            ref = holder.network[y][d].get_reference_node_id()
            idx = holder.network[y][d].get_shared_energy_storage_idx(ref)
            for p in model.periods:
                model.expected_shared_ess_p[p].set_value(pe.value(model.shared_es_pnet[idx, s_m0, s_o0, p]))
                model.expected_shared_ess_q[p].set_value(pe.value(model.shared_es_qnet[idx, s_m0, s_o0, p]))
    records = srp._process_scenario_dispersion_results(planning, models['tso'], models['dso'])
    fixed_max = max(r['maximum'] for r in records if r['quantity'].startswith('Shared ESS'))
    legacy_max, n_legacy = 0.0, 0
    for kind, node_id, holder, y, d, model in _all_blocks(planning, models):
        network = holder.network[y][d]
        if kind == 'TSO':
            for dn in model.active_distribution_networks:
                e = network.get_shared_energy_storage_idx(holder.active_distribution_network_nodes[dn])
                disp = srp._weighted_scenario_dispersion(
                    model, network, lambda p, e=e: model.expected_shared_ess_p[e, p],
                    lambda s_m, s_o, p, e=e: model.shared_es_pnet[e, s_m, s_o, p], scale=network.baseMVA)
                legacy_max = max(legacy_max, max(disp['maximum_absolute_deviation']))
                n_legacy += 1
        else:
            e = network.get_shared_energy_storage_idx(network.get_reference_node_id())
            disp = srp._weighted_scenario_dispersion(
                model, network, lambda p: model.expected_shared_ess_p[p],
                lambda s_m, s_o, p, e=e: model.shared_es_pnet[e, s_m, s_o, p], scale=network.baseMVA)
            legacy_max = max(legacy_max, max(disp['maximum_absolute_deviation']))
            n_legacy += 1
    out['E1'] = {'fixed_reading_max_shared_ess_dispersion': fixed_max,
                 'pre_fix_reading_max_shared_ess_dispersion': legacy_max, 'n_blocks_pre_fix_reading': n_legacy,
                 'n_records': len(records),
                 'pass': fixed_max == 0.0 and legacy_max > 0.0}
    # E2 -- s31c flexibility volumes (TSO interface delta, scenario-free)
    tn = planning.transmission_network
    expected_sum = {}
    for y in tn.years:
        for d in tn.days:
            model = models['tso'][y][d]
            network = tn.network[y][d]
            s_m0, s_o0 = MCH.sess_na_scenario(model)
            for node_id in planning.distribution_networks:
                dn = tn.active_distribution_network_nodes.index(node_id)
                for p in model.periods:
                    value = 0.001 * (1 + p) * (1 if p % 2 else -1)
                    model.interface_delta_p[dn, s_m0, s_o0, p].set_value(value)
                    model.interface_delta_q[dn, s_m0, s_o0, p].set_value(0.5 * value)
                    expected_sum[node_id] = expected_sum.get(node_id, 0.0) + abs(value) * network.baseMVA
    detail = G._s31c_interface_detail(planning, models)
    n_scen = planning.num_market_scenarios * tn.num_oper_scenarios
    e2 = {}
    ok = True
    for node_id, vols in detail['flexibility_volumes_per_dso'].items():
        got = vols['sum_abs_delta_p_mw']
        rel = abs(got - expected_sum[node_id]) / max(expected_sum[node_id], 1e-12)
        e2[str(node_id)] = {'fixed_sum_abs_delta_p_mw': got, 'sum_t_abs_delta_mw': expected_sum[node_id],
                            'pre_fix_would_be_n_scenarios_times': n_scen * expected_sum[node_id], 'rel_diff': rel}
        ok = ok and rel <= 1e-12
    out['E2'] = {'per_dso': e2, 'n_scenarios': n_scen, 'pass': ok and n_scen == 4}
    out['pass'] = out['E1']['pass'] and out['E2']['pass']
    return out


# ======================================================================================================================
#  F -- SRP1 invariance, structural
# ======================================================================================================================
def check_f_srp1(scratch):
    case, _spec, _changes = S.derive_case('srp1', {})
    case_dir = os.path.join(scratch, 'srp1_case')
    os.makedirs(case_dir)
    path = os.path.join(case_dir, 'SRP1__srp1.json')
    with open(path, 'w') as handle:
        json.dump(case, handle, indent='\t')
    read_dir = os.path.join(scratch, 'srp1_read')
    planning = S.read_planning_from_derived_case({'derived_case': {'path': os.path.relpath(path, REPO)}}, read_dir,
                                                 H._NoStageLog())
    probs_exact = True
    holders = [planning.transmission_network] + list(planning.distribution_networks.values())
    for holder in holders:
        for y in holder.years:
            for d in holder.days:
                net = holder.network[y][d]
                probs_exact = probs_exact and list(net.prob_market_scenarios) == [1.0] \
                    and list(net.prob_operation_scenarios) == [1.0] \
                    and all(type(v) is float for v in list(net.prob_market_scenarios) + list(net.prob_operation_scenarios))
    tn = planning.transmission_network
    y0, d0 = next(iter(tn.years)), next(iter(tn.days))
    t_model = tn.network[y0][d0].build_model(tn.params)
    dn = planning.distribution_networks[7]
    d_model = dn.network[y0][d0].build_model(dn.params)
    pairs = {kind: (list(m.scenarios_market), list(m.scenarios_operation), MCH.sess_na_scenario(m))
             for kind, m in (('tso', t_model), ('dso7', d_model))}
    na_is_only_pair = all(sm == [0] and so == [0] and na == (0, 0) for sm, so, na in pairs.values())
    return {'scenario_checksum': planning.scenario_metadata['combined_scenario_checksum'],
            'checksum_is_canonical': planning.scenario_metadata['combined_scenario_checksum'] == O.CANONICAL_CHECKSUM,
            'all_probability_vectors_exactly_one_float': probs_exact,
            'block_scenario_pairs_and_na': {k: [v[0], v[1], list(v[2])] for k, v in pairs.items()},
            'na_is_the_only_pair': na_is_only_pair,
            'omega_is_exactly_one': (1.0 * 1.0) == 1.0,
            'pass': probs_exact and na_is_only_pair}


# ======================================================================================================================
#  main
# ======================================================================================================================
def main():
    parser = argparse.ArgumentParser(description=STAGE)
    parser.add_argument('--label', required=True)
    parser.add_argument('--scratch', required=True)
    args = parser.parse_args()
    out_dir = os.path.join(REPO, OUT_ROOT_REL, args.label)
    scratch = os.path.abspath(os.path.join(args.scratch, f'w47_checks_{args.label}'))
    if os.path.exists(out_dir) or os.path.exists(scratch):
        print(f'REFUSED: output exists (write-once): {out_dir} / {scratch}', file=sys.stderr)
        return 2
    if scratch.startswith(REPO + os.sep):
        print('REFUSED: --scratch must be outside the repository', file=sys.stderr)
        return 2
    os.makedirs(out_dir)
    os.makedirs(scratch)
    started = time.time()
    MEM.mark('start (production modules imported)')
    results = {'schema': SCHEMA, 'stage': STAGE, 'label': args.label, 'argv': sys.argv, 'interpreter': sys.executable,
               'script_sha256': H.sha256_file(os.path.abspath(__file__)),
               'harness_sha256': H.sha256_file(H.HARNESS_PATH),
               'production_sha256': {f: H.sha256_file(os.path.join(REPO, f)) for f in (
                   'shared_resources_planning.py', 'p515_g_g1_g4_admm_gates.py', 'model_construction_helpers.py',
                   'network.py', 'p56a_oracle.py', 'p515_s44_scale_measurement.py')},
               'git_head': _git(['rev-parse', 'HEAD']),
               'git_tracked_changes': _git(['status', '--porcelain', '--untracked-files=no']).splitlines(),
               'started_utc': _utc(), 'scratch': scratch, 'checks': {}}
    checks = results['checks']
    try:
        _log('A: eval keys of every committed campaign spec')
        checks['A_keys'] = check_a_keys()
        _log(f"A: {checks['A_keys']['n_entries_with_eval_key']} entries in {checks['A_keys']['n_spec_files']} specs, "
             f"mismatches {len(checks['A_keys']['mismatches'])}")
        derived = derive_pilot_case(scratch)
        results['derived_instance_declaration_used'] = derived
        _log(f"pilot case derived: sha256 {derived['case_sha256']} (W45 {derived['case_sha256'] == EXPECTED_CASE_SHA256})")
        checks['B_format'] = check_b_format(scratch, derived)
        _log(f"B: pass={checks['B_format']['pass']}")
        checks['H_terminal_lock'] = check_h_lock(scratch)
        _log(f"H: pass={checks['H_terminal_lock']['pass']}")
        checks['G_rule_eleven'] = {'record': H.assert_record_capture_paths(),
                                   'post_certification': H.assert_post_certification_capture_paths(), 'pass': True}
        _log('C: the pilot child pre-run path (install, precheck, construct, configuration hook)')
        checks['C_child_prerun'], planning, sed, candidate = check_c_child_prerun(derived, scratch, args.label)
        _log(f"C: pass={checks['C_child_prerun']['pass']}")
        _log('D: intercepted build of the pilot ADMM models (unit candidate, alpha 0.50)')
        models, cv, opt_results, build = build_models(planning, sed, candidate)
        checks['D_build'] = build
        _log(f"D: built in {build['build_wall_s']:.1f}s, interceptor {build['interceptor_calls']}")
        checks['D_capture'] = check_d_capture(planning, models, opt_results, scratch)
        _log(f"D: pass={checks['D_capture']['pass']} worst={checks['D_capture']['identity_worst_rel_diffs']}")
        checks['D2_writers'] = check_d2_writers(planning, sed, models, scratch)
        _log(f"D2: pass={checks['D2_writers']['pass']}")
        checks['E_fixes'] = check_e_fixes(planning, models)
        _log(f"E: pass={checks['E_fixes']['pass']} E1={checks['E_fixes']['E1']}")
        checks['D3_post_certification'] = check_d3_post_certification(planning, models, cv, scratch)
        _log(f"D3: pass={checks['D3_post_certification']['pass']} {checks['D3_post_certification']}")
        del models, planning
        checks['F_srp1_invariance'] = check_f_srp1(scratch)
        _log(f"F: pass={checks['F_srp1_invariance']['pass']}")
    except Exception as error:  # noqa: BLE001 -- recorded, the run fails
        results['error'] = f'{type(error).__name__}: {error}'
        results['traceback'] = traceback.format_exc()
        print(results['traceback'], file=sys.stderr, flush=True)
    GUARD.uninstall()
    guard_failures = GUARD.verify(0)
    results['solve_profile'] = {'declared': 'zero solves', 'counts': dict(GUARD.counts), 'verify_failures': guard_failures}
    results['memory_by_stage'] = MEM.marks
    failing = sorted(k for k, v in checks.items() if not v.get('pass'))
    results['failing_checks'] = failing
    results['all_ok'] = bool(not results.get('error') and not failing and not guard_failures
                             and set(checks) == {'A_keys', 'B_format', 'G_rule_eleven', 'C_child_prerun', 'D_build',
                                                 'D_capture', 'D2_writers', 'E_fixes', 'D3_post_certification',
                                                 'F_srp1_invariance', 'H_terminal_lock'})
    results['finished_utc'] = _utc()
    results['wall_s'] = time.time() - started
    results['scratch_artifacts_sha256'] = {}
    for root, _dirs, files in os.walk(scratch):
        for fname in sorted(files):
            if fname.endswith(('.json', '.xlsx', '.pkl')):
                fpath = os.path.join(root, fname)
                results['scratch_artifacts_sha256'][fpath] = H.sha256_file(fpath)
    path = os.path.join(out_dir, 'pilot_checks.json')
    H._write_once_json(path, results)
    H._write_once_json(os.path.join(out_dir, 'manifest_sha256.json'),
                       {os.path.relpath(path, REPO): H.sha256_file(path)})
    _log(f"all_ok={results['all_ok']} failing={failing} guard={dict(GUARD.counts)} wall={results['wall_s']:.1f}s")
    return 0 if results['all_ok'] else 1


if __name__ == '__main__':
    sys.exit(main())
