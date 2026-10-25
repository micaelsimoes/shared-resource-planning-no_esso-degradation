"""P5.15 Addendum 82 v2, W182 order item A2, Planner task W182b -- NUMBER CHECK OF EVERY NUMBER CHANGED OR MOVED IN THE
49 HUNKS OF `git diff 8f051a2 edd7a63 -- main.tex` (Overleaf clone, read only), AGAINST THE FROZEN TABLES
(frozen_step6_tables_v1_590088fe.json) AND NAMED COMMITTED RECORDS. ZERO SOLVES, NO MODEL LOADS, NO PICKLE.

SCOPE. The lines of main.tex at edd7a63 that the diff adds ('+' lines of `git diff -U0 8f051a2 edd7a63 -- main.tex`:
the 49 default-context hunks and the 176 zero-context hunks have the same '+' lines). Every numeric token on those
lines (the checker's tokenizer: W164 tokenize + W174a number words, `W174.tokenize_v3`) is assigned exactly one status.
A moved number is a token on a '+' line; it is re-rated at its new location against its source (no carry of an earlier
rating), and the line is classified "moved (identical text at 8f051a2 l. N)" or "new or changed"; every moved line
carries a context note saying whether the move changed what the number means.

METHOD (the checker's, W164 -> W171 -> W174 -> W176 -> W177b-1' -> W178b; DECLARATIONS VERSION W182b, new, for edd7a63).
p515_s53_w178b_number_check.py is IMPORTED and NOT edited (through it W177b-1', W177b-1, W177a and the chain to W160 are
imported, each arming its guards); its spec constructors, token evaluator and helpers are used; no earlier checker is
re-run and no earlier output directory is written. Declarations are keyed by (main.tex line at edd7a63, token order);
a declaration whose written text or token count differs from the line is stale (exit 3); a token claimed twice, a line
declared twice or an unassigned token fails the run (exit 3). One tokenizer adjustment, declared: a math-mode
alphanumeric token whose leading digits are a coefficient ("3b" in max{3b, 2 tau}, "2P_" in 2P_max) is split and its
digits rated (the W164 tokenizer would exclude it as an identifier; at 8f051a2 the same coefficients were written
"3\\,b" and were tokenized).

STATUSES (the checker's) and the A2 RATING they map to:
  match -> MATCH; approximate -> MATCH (written at a coarser precision than the record; listed);
  derived -> DERIVED (the formula is in this file, function named in the token's counterpart);
  MISMATCH -> MISMATCH; no counterpart -> UNSOURCED (where it was looked for is stated);
  unchecked (category) / excluded (LaTeX structure, identifiers) -> NOT RATED (not a quantity: notation, pointer,
  language, comment, structure); every one is listed with its category.
New spec kinds: CODE (a constant read from a committed production / harness source file by a pinned regex, with the
lines that use it), BOUND (a written upper bound against a recorded maximum), and derived functions (DER) for 12.2 %,
61.4 h, 230.5 kEUR, the RES capacity sums, the lattice count, the E.21 row lists and the 60-cycle window.

OBJECTIVE CONVENTION. Q = gross_operational_cost, settlement excluded (the frozen primary); net = Q - terminal salvage
credit; value = Q(0) - Q(x); F = Q + I. Every rated monetary token carries its convention (gross / net / investment
cost / unit cost / salvage credit) in the output.

GUARDS. SolveProfileGuard(permitted=()) installed BEFORE any other project import, verified at exactly 0 at the end with
every guard the imported chain arms; pickle.load / pickle.loads blocked for the whole run (counters verified 0). git
reads only, also in the clone (rev-parse / show / diff / status / hash-object without -w). Refuses to run beside another
copy of itself, and unless the output directory exists and holds nothing but launch logs. JSON written only through
gate_result_io. NEGATIVE CONTROL: every rated token (match / approximate / derived) is perturbed in memory (last digit,
exponent or number word; a BOUND token downward) and re-evaluated against the same declaration; every perturbation must
be refused (exit 3 otherwise).

MODE (repo root, canonical interpreter; attached, both streams captured):
    mkdir data/SRP1/Results/P515S53/w182b_number_check && \\
    /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s53_w182b_number_check.py \\
        --overleaf-commit edd7a63 --expect main.tex=b08a3365... (full sha256 of every *.tex of the clone) \\
        > data/SRP1/Results/P515S53/w182b_number_check/launch.log 2>&1
    /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_gate_result_bool_typing_test.py \\
        --out data/SRP1/Results/P515S53/w182b_number_check/w182b_bool_typing_test.json \\
        > data/SRP1/Results/P515S53/w182b_number_check/w182b_bool_typing_test.log 2>&1
    /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s53_w182b_number_check.py \\
        --overleaf-commit edd7a63 --post-run
Exit: 0 = written, every integrity check holds and the guards are at 0 (MISMATCH / UNSOURCED are FINDINGS and never change
the exit code); 3 = written, an integrity check failed (listed); 1 = precondition or guard fault (nothing written).
"""
import argparse
import collections
import copy
import json
import math
import os
import pickle
import re
import subprocess
import sys
import time
from datetime import datetime, timezone

REPO = os.path.dirname(os.path.abspath(__file__))
if REPO not in sys.path:
    sys.path.insert(0, REPO)
os.chdir(REPO)

from p513_solve_profile_guard import SolveProfileGuard  # noqa: E402

GUARD = SolveProfileGuard(permitted=(), label='P5.15 W182b number check at Overleaf edd7a63 (never solves)').install()

PICKLE_COUNTS = {'load': 0, 'loads': 0}
_PICKLE_ORIG = (pickle.load, pickle.loads)


def _blocked_load(*_a, **_k):
    PICKLE_COUNTS['load'] += 1
    raise RuntimeError('W182b: pickle.load called -- no model loads are permitted')


def _blocked_loads(*_a, **_k):
    PICKLE_COUNTS['loads'] += 1
    raise RuntimeError('W182b: pickle.loads called -- no model loads are permitted')


pickle.load, pickle.loads = _blocked_load, _blocked_loads

import openpyxl  # noqa: E402 -- reads the committed operational-data workbooks (growth factors), read only
import gate_result_io as GRIO  # noqa: E402
import p515_s53_w178b_number_check as W178B  # noqa: E402 -- arms its guard; imports W177b-1' -> ... -> W160 (none edited)

pickle.load, pickle.loads = _blocked_load, _blocked_loads  # this script's block, re-installed after the imports

B1R, B1, W177A, W174, W164, W160 = W178B.B1R, W178B.B1, W178B.W177A, W178B.W174, W178B.W164, W178B.W160
GUARDS = W160.W157._dedupe((('w182b_number_check', GUARD),) + tuple(W178B.GUARDS))
_log, _sha, _sha_bytes, _git = W160._log, W160._sha, W160._sha_bytes, W160._git
S, U, P, W3, D = W177A.S, W177A.U, W177A.P, W177A.W3, W177A.D
NOT, SINGLE, LANG, XREF = W177A.NOT, W177A.SINGLE, W177A.LANG, W177A.XREF
DER, COMP, NOC = B1.DER, B1.COMP, B1R.NOC
_cl, _t10, _text, _json = W177A._cl, W177A._t10, W177A._text, W177A._json
_dp, _fmt = W177A._dp, W177A._fmt
UNIT, REF4, C5, C59, C57, BM, AF, YL, WC = B1.UNIT, B1.REF4, B1.C5, B1.C59, B1.C57, B1.BM, B1.AF, B1.YL, B1.WC
M15, M175, M2, G = B1.M15, B1.M175, B1.M2, B1.G
NAN = float('nan')

SCRIPT_REL = os.path.basename(__file__)
S53 = W160.S53
OUT_DIR = os.path.join(S53, 'w182b_number_check')
OUT_NAMES = {'json': 'w182b_number_check.json', 'md': 'w182b_number_check.md', 'man': 'manifest_sha256.json',
             'log': 'launch.log', 'typing_json': 'w182b_bool_typing_test.json', 'typing_log': 'w182b_bool_typing_test.log',
             'post': 'manifest_post_run_sha256.json'}
LAUNCH_LOG_RX = re.compile(r'launch(_attempt\d+)?\.log')

CLONE = W164.DEFAULT_MANUSCRIPT_DIR                       # manuscript/6a67305f25e8348fb71380c3 (read only)
DECLARED_COMMIT = 'edd7a63'
CLONE_HEAD = 'edd7a63a6c2ea06c1db17e3b6941bde8aa6d6838'
BASE_COMMIT = '8f051a2a0cd8b37378d7407699e38c3bf20a71e7'  # the A2 base: the 49 hunks are 8f051a2 -> edd7a63
W181A_COMMIT = 'cd57efcb54399cb652a0ac2dcd64e5d3a78d89fc'  # the HEAD of W181a's committed E.21 row map
EXPECTED_HUNKS = {'default_context': 49, 'zero_context': 176}
TEX_SHA = {
    'main.tex': 'b08a3365607e56f2ff8c56572a9a74bc9ea97d5fb461faa252fab7955ffa8f5b',
    'highlights.tex': '001352eebb3ac647f8fe1f5c2211025816ca5143c74c02a01af731630a8fea08',
    'response_to_reviewers_draft.tex': 'dba1ed69974b6585eac05d07d620553a8c9c6b6fdc0dbf5a3f9c136f8ab31692',
    'cover_letter.tex': '363ad82139c8bbf1873e64c8984934d4e049a67af593fa90936a54f3601d7508',
    'section2_expert_draft.tex': '7e347dc830d1d9a8f0c6892a963a534d63d90846c6508108821e9cfac9bcabd1',
    'main_latexdiff.tex': '18c37907f2e8cbee837f4ae50491ca86688581aa55cc1f23c5e8ac67b7d6f43a',
}
MAIN_LINES_N = 2956
BIB = 'bibliography.bib'
FZ_REL, FZ_SHA = B1R.FZ_REL, B1R.FZ_SHA

CASE_NETS = ('case9', 'case33_1', 'case33_2', 'case33_3')
CASE_YEARS = ('2025', '2028', '2030', '2031', '2034', '2035', '2037')
EXTRA_INPUTS = {
    'W178B_SCRIPT': os.path.basename(W178B.__file__),
    'W181A_JSON': os.path.join(S53, 'w181a_identity_appendixE', 'w181a_identity_appendixE.json'),
    'W180D_JSON': os.path.join(S53, 'w180d_closure_sentence', 'w180d_closure_sentence.json'),
    'W169_JSON': os.path.join(S53, 'w169_slack_inventory', 'w169_slack_inventory.json'),
    'BENCH_SPEC_V5': os.path.join(S53, 'w116_benchmark_nrf', 'frozen_s53_benchmark_spec_v5_bca69f97.json'),
    'CODE_MCH': 'model_construction_helpers.py',
    'CODE_NET': 'network.py',
    'CODE_SC1': 'settling_criterion.py',
    'CODE_SC6': 'settling_criterion_v6.py',
    'CODE_S47PB': 'p515_s47_phase_b_record.py',
    'FIG_LATTICE_PY': os.path.join('manuscript_review', 'fig_lattice_case.py'),
    'W165_SCRIPT': 'p515_s53_w165_step6_figures.py',
    'FIG4_PDF_CLONE': os.path.join(W164.DEFAULT_MANUSCRIPT_DIR, 'fig4_q_versus_cycle.pdf'),
}
for _n in CASE_NETS:
    for _y in CASE_YEARS:
        EXTRA_INPUTS[f'CASE_{_n}_{_y}'] = os.path.join('data', 'SRP1', _n, f'{_n}_{_y}.json')
    if _n != 'case9':
        EXTRA_INPUTS[f'OPDATA_{_n}'] = os.path.join('data', 'SRP1', _n, f'{_n}_operational_data.xlsx')
EXTRA_INPUTS['OPDATA_case9'] = os.path.join('data', 'SRP1', 'case9', 'case9_operational_data.xlsx')
EXTRA_JSON = ('W181A_JSON', 'W180D_JSON', 'W169_JSON', 'BENCH_SPEC_V5') + tuple(
    f'CASE_{_n}_{_y}' for _n in CASE_NETS for _y in CASE_YEARS)
EXTRA_TEXT = ('CODE_MCH', 'CODE_NET', 'CODE_SC1', 'CODE_SC6', 'CODE_S47PB', 'FIG_LATTICE_PY', 'W165_SCRIPT')
STATUSES = ('match', 'approximate', 'derived', 'MISMATCH', 'no counterpart', 'unchecked', 'excluded')
RATING = {'match': 'MATCH', 'approximate': 'MATCH', 'derived': 'DERIVED', 'MISMATCH': 'MISMATCH',
          'no counterpart': 'UNSOURCED', 'unchecked': 'NOT RATED', 'excluded': 'NOT RATED'}


# ======================================================================================================================
#  1. new spec kinds, safe accessors
# ======================================================================================================================
def CODE(w, key, rx, scale=1.0, uses=(), desc=None, conv='n/a'):
    """a constant in a committed source file X.texts[key]: every line matching `rx` (group 1 = the value) must give the same
    value; each regex of `uses` = (key, rx) must match at least one line (where the constant is used)."""
    return dict(w=w, kind='code', key=key, rx=rx, scale=scale, uses=tuple(uses), desc=desc, conv=conv)


def BOUND(w, fn, other, desc):
    """a written upper bound (two tokens, e.g. 4 x 10^-5): the bound is the token's value times `other`; fn(X) -> (recorded
    maximum, its source)."""
    return dict(w=w, kind='bound', fn=fn, other=other, desc=desc)


def CONV(spec, conv):
    spec = dict(spec)
    spec['conv'] = conv
    return spec


def _safe(fn):
    def f(X):
        try:
            v = fn(X)
        except Exception:  # noqa: BLE001 -- an unreproducible source value is a MISMATCH, not a crash
            return NAN
        return NAN if v is None else v
    return f


def SF(w, kind, fn, desc, **o):
    """S with a callable source made safe (None / exception -> NaN -> MISMATCH, never an evaluation error)."""
    return S(w, kind, _safe(fn), desc, **o)


def _code_matches(X, key, rx):
    text = X.texts[key]
    return [(i + 1, m.group(1) if m.groups() else m.group(0)) for i, ln in enumerate(text.split('\n'))
            for m in [re.search(rx, ln)] if m]


def _pnum(s):
    return float(str(s).replace('_', ''))


def evaluate(X, spec, t, line):
    k = spec['kind']
    if k == 'code':
        hits = _code_matches(X, spec['key'], spec['rx'])
        vals = sorted({_pnum(v) for _l, v in hits})
        use_hits = [(u[0], u[1], [ln for ln, _v in _code_matches(X, u[0], u[1])]) for u in spec['uses']]
        uses_ok = all(h[2] for h in use_hits)
        path = X.inputs[spec['key']]['path']
        cp = (f"{path} (sha256 {X.inputs[spec['key']]['sha256'][:8]}) lines {[ln for ln, _v in hits]} /{spec['rx']}/ -> "
              f"{vals}" + ('' if not use_hits else '; used at ' + '; '.join(
                  f"{X.inputs[a]['path']}:{b}" for a, _r, b in use_hits)))
        if len(vals) != 1 or not uses_ok:
            return {'status': 'MISMATCH', 'counterpart': cp, 'counterpart_value': vals, 'rule': 'code constant (pinned regex)',
                    'note': f'{len(vals)} distinct values; uses found {uses_ok}', 'code_lines': [ln for ln, _v in hits]}
        val = vals[0] * spec['scale']
        wv = W177A._written_value(t, line)
        if t['type'] == 'sci':
            ok = abs(val - wv) <= 1e-12 * max(1.0, abs(wv))
            return {'status': 'match' if ok else 'MISMATCH', 'counterpart': cp, 'counterpart_value': val,
                    'value_at_written_precision': f'{val:g}', 'rule': 'code constant (exponent, exact)', 'note': spec.get('desc'),
                    'code_lines': [ln for ln, _v in hits]}
        dp = _dp(t['written']) if t['type'] != 'word' else 0
        shown = _fmt(val, dp)
        ok = float(shown) == abs(wv)
        exact = abs(val - float(shown)) <= 1e-12 * max(1.0, abs(val))
        return {'status': ('match' if exact else 'approximate') if ok else 'MISMATCH', 'counterpart': cp,
                'counterpart_value': val, 'value_at_written_precision': shown,
                'rule': f'code constant x {spec["scale"]:g}, at the written precision ({dp} dp)', 'note': spec.get('desc'),
                'code_lines': [ln for ln, _v in hits]}
    if k == 'bound':
        try:
            rec, src = spec['fn'](X)
        except Exception as e:  # noqa: BLE001
            rec, src = NAN, f'error {type(e).__name__}: {e}'
        wv = W177A._written_value(t, line)
        bound = wv * spec['other']
        ok = rec == rec and rec < bound
        return {'status': 'match' if ok else 'MISMATCH', 'counterpart': src, 'counterpart_value': rec,
                'value_at_written_precision': f'recorded max {rec:.3g} < written bound {bound:.3g}: {ok}',
                'rule': 'written upper bound against the recorded maximum', 'note': spec.get('desc')}
    return B1R.evaluate(X, spec, t, line)


# ======================================================================================================================
#  2. derived numbers and source helpers (every formula of a DERIVED rating is here)
# ======================================================================================================================
def _dval(v, arith, consistent=True):
    return {'value': NAN if v is None else v, 'arithmetic': arith, 'consistent': consistent}


def der_pct_price_taker(X):
    """12.2 %: coordination saving relative to the price-taker benchmark = benefit / Q_price_taker (gross Q, x = 0)."""
    b, q = X.fzp(f'{BM}.benefit'), X.fzp(f'{BM}.arms.price_taker.Q_best')
    d = X.fzp(f'{BM}.decomposition.price_taker_NRF_minus_coordinated_eur')
    return _dval(100.0 * b / q, f"{FZ_REL} tables.benchmark.benefit {b:,.2f} / benchmark.arms.price_taker.Q_best {q:,.2f} "
                               f"x 100 = {100 * b / q:.4f} % (gross Q, settlement excluded; denominator the price-taker "
                               f"cost; benefit = decomposition.price_taker_NRF_minus_coordinated_eur {d:,.2f})",
                 abs(b - d) < 1e-6)


def der_wall_h(X):
    w = X.rec('W153_CERT', 'result.totals.wall_s_total')
    return _dval(w / 3600.0, f"{X.inputs['W153_CERT']['path']} result.totals.wall_s_total {w:,.3f} s / 3600 = {w / 3600:.4f} h")


def der_v_adjusted(X):
    v, dd = X.fzp('three_by_three.rows[0].value'), X.fzp('three_by_three.rows[5].value')
    v6 = X.fzp('three_by_three.rows[6].value')
    r8 = X.fzp('three_by_three.rows[8].value')
    return _dval((v - dd) / 1e3, f"T11 rows[0] V {v:,.2f} - rows[5] D {dd:,.2f} = {v - dd:,.2f} EUR = {(v - dd) / 1e3:.3f} kEUR "
                                 f"(printed 236.7 - 6.2 = 230.5); consistent with rows[8] R = (V - D)/V_single "
                                 f"{(v - dd) / v6:.6f} vs {r8:.6f}", abs((v - dd) / v6 - r8) < 1e-9)


def der_retention80(X):
    phi, T = X.rec('ESS_PARAMS', 'ageing.calendar_retention_per_year'), X.rec('ESS_PARAMS', 'ageing.calendar_life_years')
    return _dval(100.0 * phi ** T, f"SRP1_ESS_Params.json calendar_retention_per_year {phi} ^ calendar_life_years {T} = "
                                   f"{phi ** T:.6f} = {100 * phi ** T:.2f} % ('approximately 80')")


def _c4_entry(X):
    """the C4 arm's calibration as entered (e_c4 cell of the frozen re-settling ext spec v3): cycles_n, reference_dod_d,
    and the arm's eol_retention_r override; with T8's k for C4."""
    cell = X.j['RESETTLE_EXT_V3']['cells']
    cell = cell['e_c4'] if isinstance(cell, dict) else [c for c in cell if c.get('cell') == 'e_c4'][0]
    found = collections.defaultdict(set)

    def walk(o):
        if isinstance(o, dict):
            for kk, vv in o.items():
                if kk in ('cycles_n', 'reference_dod_d', 'eol_retention_r') and isinstance(vv, (int, float)):
                    found[kk].add(float(vv))
                walk(vv)
        elif isinstance(o, list):
            for vv in o:
                walk(vv)
    walk(cell)
    k = W177A._arm(X, 'C4')['k']
    return found, k


def der_c4(which):
    def f(X):
        found, k = _c4_entry(X)
        n = found['cycles_n'].pop() if len(found['cycles_n']) == 1 else None
        dod = found['reference_dod_d'].pop() if len(found['reference_dod_d']) == 1 else None
        ok = n is not None and dod is not None and 0.7 in found['eol_retention_r'] and \
            abs(n * dod / -math.log(0.70) - k) < 1e-6 * k
        v = {'N': n, 'dod_pct': None if dod is None else 100 * dod, 'Nd': None if n is None or dod is None else n * dod,
             'k_ln': k * -math.log(0.70)}[which]
        return _dval(v, f"e_c4 cell of {X.inputs['RESETTLE_EXT_V3']['path']}: cycles_n {n}, reference_dod_d {dod}, "
                        f"eol_retention_r values {sorted(found['eol_retention_r'])} (override 0.7); N x d = "
                        f"{None if n is None or dod is None else n * dod}; T8 C4 k {k:,.4f} x (-ln 0.70) = "
                        f"{k * -math.log(0.70):,.4f}", ok)
    return f


def _bench_q(X):
    return X.fzp(f'{BM}.arms.price_taker.Q_best')


def pos_m(which):
    """the multiplier m of a positive, determinate investment margin of the reference unit (fig2 data multiplier; claim)."""
    def f(X):
        if which == '1.75':
            c, m = X.A64['H:m1.75:value_minus_I'], X.rec(M175[1], M175[2])
        else:
            c, m = X.CL['H:m2:value_minus_I'], X.rec(M2[1], M2[2])
        return m if c['d_gross'] > 0 and c['gross_verdict'] == 'determinate' else None
    return f


def _dn_by_node(X):
    return {d['connection_node_id']: d['name'] for d in X.j['SRP1_JSON']['DistributionNetworks']}


def _gens(X, net, year):
    return {g['gen_id']: g for g in X.j[f'CASE_{net}_{year}']['generators'] if g['type'] in ('WIND', 'PV')}


def res_sum(net_or_node, tech, year):
    """aggregate installed capacity (MW) of `tech` in a network: sum of Pmax over its WIND / PV generators."""
    def f(X):
        net = net_or_node if isinstance(net_or_node, str) else _dn_by_node(X)[net_or_node]
        return sum(float(g['Pmax']) for g in _gens(X, net, year).values() if g['type'] == tech)
    return f


def res_cell(net, gen_id, field, year=None):
    def f(X):
        g = _gens(X, net, year or '2025')[gen_id]
        return {'Pmax': float(g['Pmax']) if year else None, 'bus': g['bus'], 'id': g['gen_id']}[field]
    return f


def _years_equal(X, ya, yb):
    """every WIND / PV generator of the four networks has the same Pmax (and bus, type) in years ya and yb."""
    for net in CASE_NETS:
        a, b = _gens(X, net, ya), _gens(X, net, yb)
        if set(a) != set(b) or any((a[i]['Pmax'], a[i]['bus'], a[i]['type']) != (b[i]['Pmax'], b[i]['bus'], b[i]['type'])
                                   for i in a):
            return False
    return True


YEAR_MAP = {'2025': ('2025', '2028'), '2028': ('2025', '2028'), '2030': ('2030', '2031'), '2031': ('2030', '2031'),
            '2034': ('2030', '2034'), '2035': ('2035', '2037'), '2037': ('2035', '2037')}


def year_map(y):
    """l. 1255: 2025 -> 2028, 2030 -> 2031 and 2034, 2035 -> 2037 (both years of each pair in their instance)."""
    def f(X):
        a, b = YEAR_MAP[y]
        single, multi = X.j['SRP1_JSON']['Years'], X.j['CASE_3X3']['Years']
        ok = _years_equal(X, a, b) and a in single and b in multi and \
            "os.path.join(self.data_dir, self.name, f'{self.name}_{self.year}.json')" in X.texts['CODE_NET']
        if y == '2030':
            ok = ok and _years_equal(X, '2030', '2034') and '2034' in multi
        return int(y) if ok else None
    return f


def growth(which):
    """the ADN demand growth factors (static load / flexible load), per operational-data workbook: the same in the three
    ADNs (network.py reads sheet 'Growth Factors', rows 'Load' / 'Flexible Load')."""
    def f(X):
        vals = set()
        for net in ('case33_1', 'case33_2', 'case33_3'):
            wb = openpyxl.load_workbook(os.path.join(REPO, X.inputs[f'OPDATA_{net}']['path']), read_only=True, data_only=True)
            rows = [list(r) for r in wb['Growth Factors'].iter_rows(values_only=True)]
            wb.close()
            vals.add([r[1] for r in rows if r and r[0] == which][0])
        return 100.0 * vals.pop() if len(vals) == 1 else None
    return f


def tn_growth(X):
    wb = openpyxl.load_workbook(os.path.join(REPO, X.inputs['OPDATA_case9']['path']), read_only=True, data_only=True)
    rows = [list(r) for r in wb['Growth Factors'].iter_rows(values_only=True)]
    wb.close()
    return {r[0]: r[1] for r in rows if r and r[0] in ('Load', 'Flexible Load')}


def n_flex_loads(X):
    vals = {(len(X.j[f'CASE_{n}_2025']['loads']), sum(int(ld.get('fl_reg', 0)) for ld in X.j[f'CASE_{n}_2025']['loads']))
            for n in ('case33_1', 'case33_2', 'case33_3')}
    if len(vals) != 1:
        return None
    n, nf = vals.pop()
    return n if n == nf else None


def _lattice_plans(X):
    ds, de = B1R._lattice(X, 'P'), B1R._lattice(X, 'E')
    p = X.j['ESS_PARAMS']
    lo, hi, emax, budget = p['min_energy_to_power_factor'], p['max_energy_to_power_factor'], p['max_capacity'], p['budget']
    q = X.j['W167_Q4']['as_read_by_production']
    cs, ce = q['expected_power_eur_per_MVA']['2025'], q['expected_energy_eur_per_MWh_production_function']['2025']
    out = []
    for i in range(1, int(round(emax / de)) + 1):
        for j in range(1, int(round(emax / de)) + 1):
            s_, e_ = i * ds, j * de
            if lo * s_ - 1e-9 <= e_ <= hi * s_ + 1e-9 and e_ <= emax + 1e-9 and cs * s_ + ce * e_ <= budget + 1e-6:
                out.append((s_, e_))
    return out, {'dS': ds, 'dE': de, 'E/S': [lo, hi], 'Emax': emax, 'budget': budget, 'cs_2025': cs, 'ce_2025': ce}


def der_lattice(which):
    def f(X):
        plans, prm = _lattice_plans(X)
        smallest = min(plans, key=lambda x: (x[0], x[1]))
        unit_step = [x for x in plans if x[0] <= prm['dS'] + 1e-12 and x[1] <= prm['dE'] + 1e-12]
        v = {'n': len(plans), 'S': smallest[0], 'E': smallest[1], 'unit': len(unit_step)}[which]
        return _dval(v, f"single-node lattice in 2025 from {prm} (A1 campaign LATTICE steps; SRP1_ESS_Params.json; W167 "
                        f"2025 expected unit costs; installation in the base year, discount factor 1): admissible plans "
                        f"{len(plans)} {plans}; smallest {smallest}; plans one unit step from x = 0: {unit_step}")
    return f


def lattice_fig_year(X):
    m = re.search(r'^CS, CE, BUDGET = ([0-9.]+), ([0-9.]+), ([0-9.]+)', X.texts['FIG_LATTICE_PY'], re.M)
    q = X.j['W167_Q4']['as_read_by_production']
    ok = m and abs(float(m.group(1)) - q['expected_power_eur_per_MVA']['2025'] / 1e3) < 1e-5 and \
        abs(float(m.group(2)) - q['expected_energy_eur_per_MWh_production_function']['2025'] / 1e3) < 1e-5 and \
        float(m.group(3)) * 1e3 == X.j['ESS_PARAMS']['budget']
    return 2025 if ok else None


def x0_neighbour_years(X):
    ys = sorted({int(re.match(r'y(\d{4})__', r['label']).group(1)) for r in X.w173['item2']['x0']['rows']})
    return ys


def x0_year(y):
    def f(X):
        ys = x0_neighbour_years(X)
        return y if ys == [2025, 2030] and y in ys else None
    return f


def m2_positive_count(X):
    cl = ('H:m2:value_minus_I', 'I:m2:e2_value_minus_I', 'J:e3_value_minus_I', 'J:e4_value_minus_I', 'J:e5_value_minus_I')
    ok = all(X.cand(X.other(c))['nodes']['7'][1] / X.cand(X.other(c))['nodes']['7'][0] == 4.0 for c in cl)
    return sum(1 for c in cl if X.CL[c]['d_gross'] > 0) if ok else None


def be_quarter(X):
    conv = X.t['break_even']['objective_convention']
    return 4 if 'e* = b + c/4 - p_cost/4' in conv else None


def der_two_tau(X):
    iv = X.t['break_even']['conservative_A64']['intervals']
    tau = X.fz['constants']['TAU']
    ks = sorted({round((v['half_width'] - v['bar']) / tau, 9) for v in iv.values()})
    return _dval(ks[0] if len(ks) == 1 else None, f"T3 conservative_A64.intervals: (half_width - bar) / tau = {ks} for "
                                                  f"{sorted(iv)} (the fitting intervals add 2 tau to the bar)")


def der_window60(X):
    a, b = X.fig4['b.band_window']
    shaded = [i + 1 for i, ln in enumerate(X.texts['W165_SCRIPT'].split('\n'))
              if 'axb.axvspan(blo - 0.5, bhi + 0.5' in ln]
    same_pdf = X.inputs['FIG4_PDF_CLONE']['sha256'] == X.inputs['PDF_FIG4']['sha256']
    return _dval(b - a + 1, f"W165 fig4 data b.band_window [{a}, {b}]: {b} - {a} + 1 = {b - a + 1} cycles (cap {X.fig4['b.cap']}); "
                            f"shaded in panel (b) by {X.inputs['W165_SCRIPT']['path']} lines {shaded} (axvspan of band_window); "
                            f"the clone's fig4_q_versus_cycle.pdf equals the W165 export: {same_pdf}",
                 bool(shaded) and same_pdf)


def der_three_continued(X):
    lc = X.W163['N25']['value']['last_cycles']
    return _dval(len(lc), f"W163 N25 value.last_cycles {lc}: {len(lc)} evaluations continued past a settling-rule certificate")


def der_864(X):
    s = X.j['SRP1_JSON']
    n = len(s['DistributionNetworks']) * len(s['Years']) * len(s['Days']) * s['NumInstants']
    return _dval(n, f"SRP1.json DistributionNetworks {len(s['DistributionNetworks'])} x Years {len(s['Years'])} x Days "
                    f"{len(s['Days'])} x NumInstants {s['NumInstants']} = {n} network-hours")


def _all_same(vals):
    vals = list(vals)
    return vals[0] if vals and all(v == vals[0] for v in vals) else None


def _alpha(X):
    return _all_same([(c.get('interface_deviation_premium') or {}).get('alpha') for c in X.j['SPEC_3X3']['candidates']])


def _w169_max(X):
    return (X.j['W169_JSON']['detector']['max ratio at P over certified points'],
            f"{X.inputs['W169_JSON']['path']} (sha256 {X.inputs['W169_JSON']['sha256'][:8]}) detector.\"max ratio at P over "
            f"certified points\" ({X.j['W169_JSON']['detector']['argmax cell']}): max min(P^Ch, P^Dch) / rating")


def appE_claims_rows_17_20(X):
    """tab:appE_claims rows 17-20 as printed at edd7a63 are the four G claims (7: 0.5/1 and 0.5/2, 2030 and 2035) with the
    gross Delta of the frozen claims."""
    lines = X.main_lines
    a = next(i for i, ln in enumerate(lines) if '\\label{tab:appE_claims}' in ln)
    b = next(i for i in range(a, len(lines)) if lines[i].startswith('\\end{longtable}'))
    rows = {int(m.group(1)): ln for ln in lines[a:b] for m in [re.match(r'^(\d+) & ', ln)] if m}
    want = {17: ('7: 0.5/1, 2030', G[0]), 18: ('7: 0.5/1, 2035', G[1]), 19: ('7: 0.5/2, 2030', G[2]), 20: ('7: 0.5/2, 2035', G[3])}
    ok = all(r in rows and rows[r].split(' & ')[1] == lbl and rows[r].split(' & ')[3] == f"{X.CL[c]['d_gross'] / 1e3:.1f}"
             for r, (lbl, c) in want.items())
    return ok


def e21(X):
    """W181a's committed E.21 row -> cell map, valid only while the E.21 data rows are unchanged since cd57efc."""
    if not X.e21_rows_unchanged:
        return None
    return {int(r): c for r, c in X.j['W181A_JSON']['e21_map'].items()}


def _compress(rows):
    rows = sorted(rows)
    out, i = [], 0
    while i < len(rows):
        j = i
        while j + 1 < len(rows) and rows[j + 1] == rows[j] + 1:
            j += 1
        out += [rows[i], rows[j]] if j - i >= 2 else rows[i:j + 1]
        i = j + 1
    return out


def campaign42_seq(X):
    m = e21(X)
    cells = {c['cell'] for c in X.rec('W153_CERT', 'result.cells')}
    rows = [r for r, c in m.items() if c in cells]
    return _compress(rows) if len(rows) == len(cells) == 42 else None


def seq_tok(fn, i):
    return lambda X: fn(X)[i]


def added_later_rows(X):
    m = e21(X)
    cells = {k for k, c in X.t['cells_appended_w160'].items() if c.get('row_origin') == 'T10' and c['status'] == 'certified'}
    return sorted(r for r, c in m.items() if c in cells)


def counted_rows(X):
    m = e21(X)
    inv = {c: r for r, c in m.items()}
    reg = X.t['at_or_above_0_95_tau']['registry']
    return sorted(inv[r['cell']] for r in reg if r.get('counted'))


def beside_row(X):
    m = e21(X)
    inv = {c: r for r, c in m.items()}
    b = list(X.t['at_or_above_0_95_tau']['beside'])
    c = X.cells[b[0]] if len(b) == 1 else None
    return inv[b[0]] if c and c['branch'] == 'monotone' else None


def beside_range(X):
    b = list(X.t['at_or_above_0_95_tau']['beside'])
    return X.cells[b[0]]['range_over_tau']


N25_ORDER = ('b_2a0ba8b2', 'b_4649234b', 'd_c52e1670')


def continued_row(i):
    def f(X):
        m = e21(X)
        inv = {c: r for r, c in m.items()}
        cells = {k.split(' ')[0] for k in X.W163['N25']['value']['over_tau']}
        return inv[N25_ORDER[i]] if cells == set(N25_ORDER) else None
    return f


def continued_over_tau(cell, sel):
    def f(X):
        o = X.W163['N25']['value']['over_tau']
        ks = [k for k in o if k.startswith(cell) and sel in k]
        return abs(o[ks[0]]) if len(ks) == 1 else None
    return f


def w180d_row16(field):
    def f(X):
        r = X.j['W180D_JSON']['row16']
        ok = r['e21_row']['no'] == 16 and r['cell'] == e21(X)[16] and r['S_rounds_to_printed_2sf'] is True
        if not ok:
            return None
        if field == 'no':
            return 16
        if field == 'S_mant':
            return r['S_mwh'] / 10 ** math.floor(math.log10(r['S_mwh']))
        if field == 'S_exp':
            return 10.0 ** math.floor(math.log10(r['S_mwh']))
        if field == 'F_mant':
            return r['S_over_e_av'] / 10 ** math.floor(math.log10(r['S_over_e_av']))
        if field == 'F_exp':
            return 10.0 ** math.floor(math.log10(r['S_over_e_av']))
    return f


def w180d_relax(X):
    c = X.j['W180D_JSON']['conversion']
    return c['relaxation_mwh_per_slack'] if c['conversion_correct'] is True else None


def ex_lad(rows):
    return lambda X: _all_same([X.cand(X.other(c))['investment_year'] for c in rows])


# ======================================================================================================================
#  3. declarations, version W182b (main.tex at edd7a63; one entry per '+' line that carries non-excluded tokens)
# ======================================================================================================================
SR, C3X = 'SRP1_JSON', 'CASE_3X3'
TWO_PCT = ('rec', 'SRP1_JSON', 'DiscountFactor')
CTX = {}  # line -> context note for moved lines (set below)


def DL(line, *specs, note=None, ctx=None):
    d = D(line, None, *specs, note=note)
    d['file'] = 'main.tex'
    if ctx:
        CTX[line] = ctx
    return d


def ladder6(line, claim, nodes, year, w):
    """tab:res_value_ladder at edd7a63: node(s) [, year], S, E, duration, cost increase F(x) - F(0), threshold."""
    sp = [P(str(n), 'node', claim, [n]) for n in nodes]
    if year != 2025:
        sp.append(P(str(year), 'year', claim))
    sp += [P(w[0], 'S', claim, nodes), P(w[1], 'E', claim, nodes), P(w[2], 'dur', claim, nodes),
           S(w[3], 'k', _cl(claim, 'd_gross'), signed=True), S(w[4], 'k', _cl(claim, 'gross_threshold_or_bar'))]
    return DL(line, *sp)


def gen_row(line, net, gid, w, years=('2025', '2028', '2031', '2034', '2037'), ctx=None):
    sp = [SF(w[0], 'int', res_cell(net, gid, 'id'), f'{net}_2025.json generators gen_id {gid}'),
          SF(w[1], 'int', res_cell(net, gid, 'bus'), f'{net}_2025.json generators[gen_id {gid}].bus')]
    sp += [SF(v, 'raw', res_cell(net, gid, 'Pmax', y), f'{net}_{y}.json generators[gen_id {gid}].Pmax (MW)')
           for v, y in zip(w[2:], years)]
    return DL(line, *sp, ctx=ctx)


def declarations():
    L = []
    add = L.append
    yrs3 = lambda y: SF(y, 'int', B1R._year_in(C3X, y), 'SRP1__s53_3x3.json Years')  # noqa: E731
    yrs1 = lambda y: SF(y, 'int', B1R._year_in(SR, y), 'SRP1.json Years')  # noqa: E731
    n_dn = SF('three', 'int', lambda X: len(X.j[SR]['DistributionNetworks']), 'SRP1.json DistributionNetworks (count)')
    m3 = [SF('3', 'int', lambda X: X.j[C3X]['NumMarketScenarios'], 'SRP1__s53_3x3.json NumMarketScenarios'),
          SF('3', 'int', lambda X: B1R._op_scen(X.j[C3X]), 'SRP1__s53_3x3.json num_operation_scenarios (TN and DNs)')]
    pct122 = DER('12.2', der_pct_price_taker, 'coordination saving relative to the price-taker benchmark (gross Q)')
    meur = DER('90.9', B1.der_meur(f'{BM}.benefit', 'benefit'), 'coordination saving vs the price-taker arrangement')
    two_stage = U('two', 'compound adjective', '"two-stage" (the model class, not a count)')
    two_pct = S('2', 'pct0', TWO_PCT, 'SRP1.json DiscountFactor (x 100)')
    one_inst = SF('one', 'int', B1R._one_install_per_node, 'every frozen cell candidate: one investment_year and one [S, E] '
                                                           'pair per node (no staged additions)')
    year_dims = SF('one', 'int', B1._w173_year_dims, 'W173 box_definition: one year coordinate zY common to all nodes')
    soh070 = S('0.70', 'raw', ('rec', 'ESS_PARAMS', 'ageing.minimum_soh'), 'SRP1_ESS_Params.json minimum_soh')
    soh050 = S('0.50', 'raw', ('rec', 'W155_SOH050', 'reports.e_soh050.minimum_soh'), 'A64 e_soh050 minimum_soh')
    m15 = S('1.5', 'raw', M15, 'fig2 data: the multiplier of the m = 1.5 cells')
    tau2 = W3('2', 'N12', 2.0, 'the gap band tau/2 (settling_criterion_v2.GAP_BOUND = TAU / 2)')

    # ---- abstract, highlights ---------------------------------------------------------------------------------------
    add(DL(145, two_stage))
    add(DL(146, DER('108', B1.der_108, 'buses of the integrated system'),
           SF('15', 'int', lambda X: sum(X.j[SR]['Years'].values()), 'SRP1.json Years: sum of the block lengths'),
           pct122, P('0.25', 'S', UNIT, [7]), P('1', 'E', UNIT, [7]),
           S('64.4', 'k', _cl(REF4, 'd_gross'), note='"baseline ageing yields an investment shortfall of 64.4" = '
                                                     'F(x) - F(0) = -(value - I) of the reference unit, gross'),
           SF('1.75', 'raw', pos_m('1.75'), 'multiplier with a positive, determinate margin (A64 H:m1.75; fig2 data)'),
           SF('2', 'raw', pos_m('2'), 'multiplier with a positive, determinate margin (T1 H:m2; fig2 data)'), *m3))
    add(DL(166, SF('1.75', 'raw', pos_m('1.75'), 'multiplier with a positive, determinate margin'),
           SF('2', 'raw', pos_m('2'), 'multiplier with a positive, determinate margin')))
    add(DL(167, pct122))
    # ---- section 2 --------------------------------------------------------------------------------------------------
    add(DL(462, LANG('two', '"the two scenario sets" (market, operation): the sets enumerated in the formula')))
    add(DL(463, U('0', 'notation', 'u^{0} (a symbol superscript)')))
    add(DL(486, U('two', 'comment (not typeset)', 'commented-out caption of fig:search_generic (figure removed)')))
    add(DL(492, year_dims))
    add(DL(525, CODE('ten', 'CODE_S47PB', r'step over the last (\d+) cycles', 1.0,
                     desc='the settling bar of a search evaluation: RESOLUTION_RULE "max gross step over the last 10 cycles"')))
    add(DL(737, SINGLE('single')))
    add(DL(743, W3('two', 'N73', 2, 'a difference of two certified evaluations'),
           W3('3', 'N28', 3.0, 'DETERMINACY_BAR_FACTOR (coefficient of b; split from the token "3b")'),
           W3('2', 'N29', 2.0, 'DETERMINACY_TAU_MULTIPLE')))
    add(DL(748, SINGLE('single')))
    add(DL(755, U('one', 'model structure', '"one shared-ESS agent per interface node" (the decomposition, not a figure)')))
    add(DL(978, U('0', 'notation', 'sigma >= 0 (non-negativity)')))
    add(DL(996, one_inst))
    add(DL(1061, U('1', 'notation', 'eq:soh_initial SoH = 1 (an equation constant; the equation is rated against the code by '
                                   'A1)')))
    add(DL(1132, U('0', 'notation', 'd >= 0 (non-negativity)')))
    add(DL(1158, LANG('single', '"With a single scenario" (the one-scenario case)')))
    # ---- section 3 opening, 3.1 -------------------------------------------------------------------------------------
    add(DL(1169, SF('9', 'int', lambda X: B1R._n_bus(X, 9), 'case9_2025.json nodes'), n_dn,
           SF('33', 'int', lambda X: B1R._n_bus(X, 33), 'case33_1/2/3_2025.json nodes (each)')))
    add(DL(1170, n_dn))
    add(DL(1196, SF('Two', 'int', B1R._two_instances, 'the instances with results: SRP1.json and its derived 3 x 3 instance '
                                                     '(3 x 3 campaign spec derived_instance, sha256 checked)'),
           SF('15', 'int', lambda X: _all_same([sum(X.j[SR]['Years'].values()), sum(X.j[C3X]['Years'].values())]),
              'sum of the block lengths of both instances')))
    add(DL(1197, SINGLE('single'), S('one', 'int', ('rec', SR, 'NumMarketScenarios'), 'SRP1.json NumMarketScenarios'),
           SF('one', 'int', lambda X: B1R._op_scen(X.j[SR]), 'SRP1.json num_operation_scenarios (TN and DNs)'),
           yrs1('2025'), yrs1('2030'), yrs1('2035'),
           SF('five', 'int', lambda X: _all_same(X.j[SR]['Years'].values()), 'SRP1.json Years: block length (each)')))
    add(DL(1198, SF('three', 'int', lambda X: X.j[C3X]['NumMarketScenarios'], '3 x 3 NumMarketScenarios'),
           SF('three', 'int', lambda X: B1R._op_scen(X.j[C3X]), '3 x 3 num_operation_scenarios'),
           *[yrs3(y) for y in ('2025', '2028', '2031', '2034', '2037')],
           SF('three', 'int', lambda X: _all_same(X.j[C3X]['Years'].values()), '3 x 3 Years: block length (each)')))
    add(DL(1202, SF('four', 'int', lambda X: _all_same([len(X.j[SR]['Days']), len(X.j[C3X]['Days'])]), 'Days (count), both'),
           *[SF(w, 'int', (lambda d: (lambda X: _all_same([X.j[SR]['Days'][d], X.j[C3X]['Days'][d]])))(d),
                f'Days["{d}"] weight, both instances') for w, d in (('92', 'Spring'), ('91', 'Summer'), ('91', 'Autumn'),
                                                                    ('91', 'Winter'))]))
    add(DL(1203, SINGLE('single'), one_inst))
    add(DL(1207, SF('1', 'raw', lambda X: X.j['ESS_PARAMS']['budget'] / 1e6, 'SRP1_ESS_Params.json budget (EUR) / 1e6',
                    conv='investment budget (expected discounted investment cost, EUR)'),
           S('5', 'raw', ('rec', 'ESS_PARAMS', 'max_capacity'), 'SRP1_ESS_Params.json max_capacity (MWh per node)'),
           SF('2', 'pct0', B1R._disc_both, 'DiscountFactor of SRP1.json and SRP1__s53_3x3.json (equal)')))
    add(DL(1211, NOT('0')))
    add(DL(1219, SF('Three', 'int', lambda X: len(X.expected_cost('energy', 2025)[1]),
                    'SRP1_ESS.xlsx sheet Scenarios: investment-cost trajectories (probabilities)')))
    add(DL(1227, yrs1('2025'), yrs1('2030'), yrs1('2035')))
    q4 = 'as_read_by_production'
    for ln, kind, sc, w in ((1229, 'power', 1, ('35.00', '214.17', '168.72', '153.93')),
                            (1230, 'power', 2, ('55.00', '267.53', '224.25', '206.96')),
                            (1231, 'power', 3, ('10.00', '342.17', '278.10', '268.58')),
                            (1233, 'energy', 1, ('35.00', '212.13', '167.11', '152.46')),
                            (1234, 'energy', 2, ('55.00', '264.98', '222.12', '204.99')),
                            (1235, 'energy', 3, ('10.00', '338.91', '275.45', '266.02'))):
        f = 'cost_power_eur_per_MVA' if kind == 'power' else 'cost_energy_eur_per_MWh'
        cv = 'unit investment cost (undiscounted, k EUR per MVA)' if kind == 'power' else \
            'unit investment cost (undiscounted, k EUR per MWh)'
        add(DL(ln, S(w[0], 'pct2', ('rec', 'W167_Q4', f'{q4}.scenario_probabilities[{sc - 1}]'), 'trajectory probability'),
               *[CONV(S(v, 'kmwh', ('rec', 'W167_Q4', f'{q4}.{f}.{sc}.{y}'), f'W167 cost file as read by production, '
                                                                            f'trajectory {sc}, {y}'), cv)
                 for v, y in zip(w[1:], ('2025', '2030', '2035'))]))
    add(DL(1244, yrs1('2025'),
           CONV(S('256.32', 'kmwh', ('rec', 'W167_Q4', f'{q4}.expected_power_eur_per_MVA.2025'), 'expected 2025 cost'),
                'unit investment cost, probability-weighted (k EUR per MVA)'),
           CONV(S('253.88', 'kmwh', ('rec', 'W167_Q4', f'{q4}.expected_energy_eur_per_MWh_production_function.2025'),
                  'expected 2025 cost'), 'unit investment cost, probability-weighted (k EUR per MWh)')))
    # ---- 3.3 (merged 3.3 / 3.4) -------------------------------------------------------------------------------------
    add(DL(1255, n_dn, SINGLE('single'),
           *[SF(y, 'int', year_map(y), f'case files: Pmax of every WIND / PV generator equal in {YEAR_MAP[y][0]} and '
                                       f'{YEAR_MAP[y][1]} (four networks); network.py:215 reads <name>_<year>.json')
             for y in ('2025', '2028', '2030', '2031', '2034', '2035', '2037')]))
    add(DL(1262, SINGLE('single')))
    add(DL(1268, yrs1('2025'), yrs1('2030'), yrs1('2035')))
    for ln, node, tech, w in ((1270, 'case9', 'WIND', ('85', '130', '150')), (1271, 'case9', 'PV', ('47', '75', '95')),
                              (1273, 5, 'WIND', ('5', '40', '50', '50')), (1274, 5, 'PV', ('43', '60', '90')),
                              (1276, 7, 'WIND', ('7', '80', '90', '90')), (1277, 7, 'PV', ('20', '30', '50')),
                              (1279, 9, 'WIND', ('9', '20', '50', '50')), (1280, 9, 'PV', ('90', '90', '120'))):
        sp = []
        vals = list(w)
        if len(w) == 4:
            sp.append(SF(w[0], 'int', (lambda n: (lambda X: n if n in _dn_by_node(X) else None))(node),
                         'SRP1.json DistributionNetworks connection_node_id'))
            vals = list(w[1:])
        sp += [SF(v, 'int', res_sum(node, tech, y), f'sum of Pmax over the {tech} generators of '
                                                     f'{"case9" if node == "case9" else "the ADN at TN node " + str(node)} '
                                                     f'in {y} (case files; MW)') for v, y in zip(vals, ('2025', '2030', '2035'))]
        add(DL(ln, *sp))
    add(DL(1290, SF('32', 'int', n_flex_loads, 'case33_1/2/3_2025.json: loads per ADN, each with fl_reg = 1'),
           SF('1.25', 'raw', growth('Load'), 'case33_1/2/3_operational_data.xlsx sheet Growth Factors, Load (x 100)'),
           SF('3.00', 'raw', growth('Flexible Load'), 'the same workbooks, Flexible Load (x 100)'),
           ctx='moved from the former "Active Distribution Networks" subsection into the merged 3.3 paragraph that opens '
               '"Each ADN contains 32 Flexible Loads"; the growth rates are the ADN rates (all three ADN workbooks); the TN '
               'workbook (case9_operational_data.xlsx) has 1.00 % / 2.50 %, and the merged section now covers the TN too: '
               'read in its paragraph the meaning is unchanged (ADN demand), but the sentence no longer sits under an '
               'ADN-only heading'))
    add(DL(1291, SF('2025', 'int', B1R._year_in(SR, '2025'), 'SRP1.json Years: the base year of the profiles')))
    # ---- 3.4 storage parameters -------------------------------------------------------------------------------------
    add(DL(1298, CODE('0.97', 'CODE_SES', r'self\.eff_ch = ([0-9.]+)', uses=[('CODE_MCH', r'eff_ch')]),
           CODE('0.96', 'CODE_SES', r'self\.eff_dch = ([0-9.]+)', uses=[('CODE_MCH', r'eff_dch')])))
    add(DL(1299, CODE('10', 'CODE_DEFS', r'^ENERGY_STORAGE_MIN_ENERGY_STORED = ([0-9.]+)', 100.0),
           CODE('90', 'CODE_DEFS', r'^ENERGY_STORAGE_MAX_ENERGY_STORED = ([0-9.]+)', 100.0),
           CODE('50', 'CODE_DEFS', r'^ENERGY_STORAGE_RELATIVE_INIT_SOC = ([0-9.]+)', 100.0),
           S('15', 'int', ('rec', 'ESS_PARAMS', 'ageing.calendar_life_years'), 'SRP1_ESS_Params.json calendar_life_years')))
    add(DL(1300, CODE('0.05', 'CODE_MCH', r'^ESS_DAY_BALANCE_SLACK_FRACTION = ([0-9.]+)',
                      uses=[('CODE_MCH', r'e_capacity \* ESS_DAY_BALANCE_SLACK_FRACTION \+ EQUALITY_TOLERANCE')]),
           CODE('10^{3}', 'CODE_DEFS', r'^PENALTY_SHARED_ESS_BALANCE = ([0-9.e+-]+)',
                uses=[('CODE_MCH', r'base \* PENALTY_SHARED_ESS_BALANCE \* \(model\.slack_shared_es_soc_final_up')],
                desc='closure slack price c^Cl (EUR/MWh: base * 1e3 * slack_pu)', conv='penalty price (EUR/MWh)')))
    add(DL(1301, SF('0.25', 'raw', lambda X: B1R._lattice(X, 'P'), 'A1 campaign LATTICE_P_STEP (MVA)'),
           SF('0.5', 'raw', lambda X: B1R._lattice(X, 'E'), 'A1 campaign LATTICE_E_STEP (MWh)'),
           S('2', 'raw', ('rec', 'ESS_PARAMS', 'min_energy_to_power_factor'), 'SRP1_ESS_Params.json'),
           S('4', 'raw', ('rec', 'ESS_PARAMS', 'max_energy_to_power_factor'), 'SRP1_ESS_Params.json')))
    add(DL(1309, S('10,000', 'int', ('rec', 'ESS_PARAMS', 'ageing.calibration.cycles_n'), 'baseline calibration'),
           S('80', 'pct0', ('rec', 'ESS_PARAMS', 'ageing.calibration.reference_dod_d'), 'baseline depth of discharge (x 100)'),
           S('80', 'pct0', ('rec', 'ESS_PARAMS', 'ageing.calibration.eol_retention_r'), 'baseline end retention (x 100)')))
    add(DL(1310, S('0.985', 'raw', ('rec', 'ESS_PARAMS', 'ageing.calendar_retention_per_year'), 'SRP1_ESS_Params.json'),
           DER('80', der_retention80, 'calendar retention after 15 years without cycling, "approximately 80 %"'),
           S('15', 'int', ('rec', 'ESS_PARAMS', 'ageing.calendar_life_years'), 'SRP1_ESS_Params.json'), dict(soh070)))
    add(DL(1334, DER('8,000', der_c4('k_ln'), 'C4: T8 k x (-ln 0.70) = N x d of the datasheet (8,000 full cycles)'),
           DER('10,000', der_c4('N'), 'C4 entered cycles_n (e_c4 cell, frozen ext spec v3)'),
           DER('80', der_c4('dod_pct'), 'C4 entered depth of discharge (x 100)'),
           DER('8,000', der_c4('Nd'), 'N x d of the entered C4 statement')))
    # ---- 3.5 planning and evaluation settings -----------------------------------------------------------------------
    add(DL(1344, SF('0.5', 'raw', _alpha, '3 x 3 campaign spec candidates[*].interface_deviation_premium.alpha')))
    add(DL(1349, SINGLE('single'), SF('Two', 'int', lambda X: len(X.t['benchmark']['arms']), 'T6: the two NRF arms')))
    add(DL(1350, SF('three', 'int', lambda X: _all_same([X.fzp(f'{AF}.passive.n_starts'), X.fzp(f'{AF}.price_taker.n_starts')]),
                    'T6 n_starts (both arms)')))
    # ---- 4 opening, 4.1 ---------------------------------------------------------------------------------------------
    add(DL(1361, SINGLE('single'), *m3,
           SF('15', 'int', lambda X: sum(X.j[SR]['Years'].values()), 'SRP1.json Years: sum of the block lengths'), dict(two_pct)))
    add(DL(1362, NOT('0')))
    add(DL(1372, NOT('0')))
    add(DL(1373, S('fourteen', 'int', ('rec', 'W173', 'item2.x0.b_box.n_feasible'), 'W173: admissible unit neighbours of x = 0')))
    add(DL(1374, P('0.25', 'S', UNIT, [7]), P('1', 'E', UNIT, [7]), P('7', 'node', UNIT, [7]), P('2025', 'year', UNIT),
           S('253.5', 'k', 'ageing.rows[C2_calfade].value'), CONV(S('318.0', 'k', _cl(UNIT, 'I_other')), 'investment cost I'),
           S('64.4', 'k', _cl(REF4, 'd_gross')), P('5', 'node', C5, [5]), P('0.25', 'S', C5, [5]), P('0.5', 'E', C5, [5]),
           S('49.7', 'k', _cl(C5, 'd_gross'))))
    add(DL(1380, SINGLE('single')))
    for ln, claim, nodes, year, w in ((1393, C5, [5], 2025, ('0.25', '0.5', '2', '49.7', '12.6')),
                                      (1394, REF4, [7], 2025, ('0.25', '1', '4', '64.4', '13.1')),
                                      (1395, 'B:n7_2h_e1', [7], 2025, ('0.5', '1', '2', '114.1', '12.6')),
                                      (1396, 'B:n7_4h_e2', [7], 2025, ('0.5', '2', '4', '135.5', '13.4')),
                                      (1397, C59, [5, 9], 2025, ('0.25', '0.5', '2', '102.1', '12.6')),
                                      (1398, C57, [5, 7], 2030, ('0.25', '0.5', '2', '93.9', '13.4'))):
        add(ladder6(ln, claim, nodes, year, w))
    add(DL(1405, SF('2025', 'int', ex_lad([C5, REF4, 'B:n7_2h_e1', 'B:n7_4h_e2', C59]),
                    'investment_year of every tab:res_value_ladder plan except "5 and 7 (2030)"')))
    add(DL(1406, NOT('0')))
    add(DL(1413, P('two', 'nnodes', C59)))
    add(DL(1414, P('1', 'E', 'B:n7_2h_e1', [7]), P('2', 'dur', 'B:n7_2h_e1', [7]), P('4', 'dur', REF4, [7]),
           P('two', 'nnodes', C57), P('2030', 'year', C57), CONV(S('84.0', 'k', _cl(C57, 'd_net')), 'net of salvage')))
    # ---- 4.2 --------------------------------------------------------------------------------------------------------
    add(DL(1423, SF('7', 'int', lambda X: W177A._single_node(X, 'conservative_A64'), 'T3 conservative_A64 fit: its node')))
    add(DL(1424, P('2', 'dur', 'B:n7_2h_e1', [7]), P('4', 'dur', REF4, [7]), P('4', 'dur', REF4, [7]),
           SF('4', 'int', be_quarter, 'T3 objective_convention "e* = b + c/4 - p_cost/4"'),
           SF('4', 'int', be_quarter, 'the same (p/4)'),
           CONV(S('256,320', 'e10', 'break_even.p_cost_eur_per_MVA'), 'unit investment cost (EUR/MVA)'),
           S('six', 'int', 'break_even.conservative_A64.certified_only.n'),
           CONV(S('183,080', 'e10', 'break_even.conservative_A64.certified_only.breakeven_marginal_4h_energy_cost'),
                'break-even unit energy cost (EUR/MWh; fit to gross values)')))
    add(DL(1425, S('ten', 'int', 'break_even.conservative_A64.banded_mid.n'),
           CONV(S('171,310', 'e10', 'break_even.conservative_A64.breakeven_range[0]'), 'break-even unit energy cost (gross fit)'),
           CONV(S('192,560', 'e10', 'break_even.conservative_A64.breakeven_range[1]'), 'break-even unit energy cost (gross fit)'),
           DER('24', B1.der_breakeven_pct(0), 'energy-cost reduction to reach the range, lower end'),
           DER('33', B1.der_breakeven_pct(1), 'energy-cost reduction to reach the range, upper end'),
           CONV(S('253,880', 'e10', 'break_even.e_cost_eur_per_MWh'), 'unit investment cost (EUR/MWh)')))
    add(DL(1431, P('7', 'node', REF4, [7]), P('2025', 'year', REF4), P('4', 'dur', REF4, [7]), P('2', 'dur', 'B:n7_2h_e1', [7]),
           DER('2', der_two_tau, 'the fitting intervals add 2 tau to the bar')))
    add(DL(1437, P('4', 'dur', REF4, [7]), P('7', 'node', REF4, [7]),
           CONV(S('256,320', 'e10', 'break_even.p_cost_eur_per_MVA'), 'unit investment cost (EUR/MVA)')))
    add(DL(1454, dict(m15), S('14.5', 'k', _cl('H:m1.5:value_minus_I', 'd_gross')),
           S('18.8', 'k', _cl('H:m1.5:value_minus_I', 'gross_threshold_or_bar'))))
    add(DL(1455, S('1.75', 'raw', M175, 'fig2 data'), S('2', 'raw', M2, 'fig2 data'),
           S('14.2', 'k', _t10('H:m1.75:value_minus_I', 'd_gross')), S('46.3', 'k', _cl('H:m2:value_minus_I', 'd_gross'))))
    add(DL(1462, dict(m15)))
    J5 = 'J:e5_value_minus_I'
    add(DL(1467, SF('five', 'int', m2_positive_count, 'm = 2 ladder (H:m2, I:m2:e2, J:e3, J:e4, J:e5; 4 h at node 7): positive '
                                                     'margins'),
           P('4', 'dur', J5, [7]), P('7', 'node', J5, [7]), P('2025', 'year', J5),
           SF('1', 'raw', B1._e_steps, 'the E step between consecutive cells of the m = 2 ladder'),
           SF('41.0', 'k', lambda X: B1._incr(X, 'min'), 'T1 I/J increments: smallest'),
           SF('51.8', 'k', lambda X: B1._incr(X, 'max'), 'T1 I/J increments: largest')))
    add(DL(1472, P('4', 'dur', J5, [7]), P('7', 'node', J5, [7]), P('2025', 'year', J5)))
    for ln, c1, c2, w in ((1484, 'H:m2:value_minus_I', None, ('1', '46.3', '12.3')),
                          (1485, 'I:m2:e2_value_minus_I', 'I:m2:second_MWh', ('2', '92.8', '12.8', '46.5', '12.8')),
                          (1486, 'J:e3_value_minus_I', 'J:e2_to_e3', ('3', '133.8', '27.7', '41.0', '27.7')),
                          (1487, 'J:e4_value_minus_I', 'J:e3_to_e4', ('4', '176.6', '15.9', '42.8', '27.7')),
                          (1488, J5, 'J:e4_to_e5', ('5', '228.4', '26.5', '51.8', '26.5'))):
        bar1 = _cl(c1, 'uncertified_form_beside_A64_ruling4.bar') if c1 == 'J:e4_value_minus_I' else \
            _cl(c1, 'gross_threshold_or_bar')
        sp = [P(w[0], 'E', c1, [7]), S(w[1], 'k', _cl(c1, 'd_gross'), signed=True), S(w[2], 'k', bar1)]
        if c2:
            sp += [S(w[3], 'k', _cl(c2, 'd_gross'), signed=True), S(w[4], 'k', _cl(c2, 'gross_threshold_or_bar'))]
        add(DL(ln, *sp))
    RP = 'ref:5ca4f86c'
    add(DL(1507, P('two', 'nnodes', RP), P('2030', 'year', RP), P('0.25', 'S', RP, [5]), P('0.5', 'E', RP, [5]),
           P('5', 'node', RP, [5]), P('1', 'S', RP, [7]), P('3.5', 'E', RP, [7]), P('7', 'node', RP, [7]),
           S('61', 'int', ('rec', 'W173', 'item2.F2.b_box.n_feasible'), 'W173 F2: admissible unit neighbours'),
           S('17', 'int', ('rec', 'W173', 'item2.F2.b_box.n_evaluated_in_search'), 'W173 F2: evaluated in the search'),
           S('13', 'int', ('rec', 'W173', 'item2.F2.c_L_claims_box_neighbours.n'), 'W173 F2: re-evaluated (L claims)'),
           S('13', 'int', ('rec', 'W173', 'item2.F2.c_L_claims_box_neighbours.n'), 'the same 13')))
    # ---- 4.3 --------------------------------------------------------------------------------------------------------
    add(DL(1515, P('0.25', 'S', UNIT, [7]), P('1', 'E', UNIT, [7]), P('7', 'node', UNIT, [7]), P('2025', 'year', UNIT)))
    add(DL(1517, S('4.1', 'k', 'ageing.rows[no_ageing].value_minus_I'), S('31.6', 'k', 'ageing.rows[C2].value_minus_I'),
           S('64.4', 'k', 'ageing.rows[C2_calfade].value_minus_I')))
    add(DL(1518, dict(soh070), SF('31.6', 'k', lambda X: W177A._min_max_aged(X, 'value_minus_I', 'max'), 'T8 aged arms: smallest '
                                                                                                         'loss'),
           SF('73.7', 'k', lambda X: W177A._min_max_aged(X, 'value_minus_I', 'min'), 'T8 aged arms: largest loss')))
    add(DL(1525, dict(soh070)))
    for ln, arm, claim, w in ((1546, 'C2_calfade', 'E:n7_4h_e1_C2_calfade:value_minus_I', ('64.4', '13.1', '2035', '0.793', '0.86')),
                              (1548, 'C2', 'E:n7_4h_e1_C2:value_minus_I', ('31.6', '12.6', None, '0.888', '1.19')),
                              (1550, 'C3_unit', 'E:C3_unit_value_minus_I', ('73.7', '12.8', '2035', '0.786', '0.76')),
                              (1552, 'C3_midblock', 'E:n7_4h_e1_C3_midblock:value_minus_I', ('65.9', '12.9', '2035', '0.836', '0.76')),
                              (1554, 'C4', 'E:n7_4h_e1_C4:value_minus_I', ('45.2', '12.6', None, '0.834', '1.13')),
                              (1556, 'no_ageing', 'E:n7_4h_e1_no_ageing:value_minus_I', ('4.1', '12.9', None, '1.000', '1.29'))):
        sp = [S(w[0], 'k', f'ageing.rows[{arm}].value_minus_I', signed=True), S(w[1], 'k', _cl(claim, 'gross_threshold_or_bar'))]
        if w[2]:
            sp.append(S(w[2], 'int', f'ageing.rows[{arm}].floor_year_070'))
        sp += [S(w[3], 'r3', f'ageing.rows[{arm}].AE'), S(w[4], 'efc', f'ageing.rows[{arm}].EFC')]
        add(DL(ln, *sp))
    add(DL(1549, S('50', 'pct0', lambda X: W177A._retention(X, 'C3_unit'), 'T8 k of C3_unit: R = exp(-N D / k), N D = 8,000')))
    add(DL(1551, S('50', 'pct0', lambda X: W177A._retention(X, 'C3_midblock'), 'T8 k of C3_midblock')))
    add(DL(1553, SF('8,000', 'int', lambda X: round(W177A._arm(X, 'C4')['k'] * -math.log(0.70)), 'T8 k of C4 x (-ln 0.70)')))
    add(DL(1558, dict(soh050)))
    add(DL(1559, S('59.5', 'k', _t10('E:soh050:value_minus_I', 'd_gross'), signed=True),
           S('12.6', 'k', _t10('E:soh050:value_minus_I', 'gross_threshold_or_bar')),
           S('0.778', 'r3', ('rec', 'W155_SOH050', 'reports.e_soh050.AE'), 'W155 e_soh050 AE (W178b part (c))'),
           S('1.08', 'efc', ('rec', 'W155_SOH050', 'reports.e_soh050.EFC'), 'W155 e_soh050 EFC/day')))
    add(DL(1565, CONV(S('318.0', 'k', 'ageing.rows[C2_calfade].I'), 'investment cost I')))
    add(DL(1566, dict(soh070), dict(soh050)))
    # ---- 4.4 --------------------------------------------------------------------------------------------------------
    add(DL(1582, SF('three', 'int', lambda X: len(X.t['benchmark']['arms']) + 1, 'T6: coordinated + the two NRF arms')))
    add(DL(1583, S('13.9', 'pct', f'{BM}.benefit_relative', 'price-taker: (Q_pt - Q_coord) / Q_coord (the T6 relative)'),
           S('36.9', 'pct', f'{BM}.w160_additions.passive_arm_vs_coordinated_derived.benefit_passive_relative',
             'passive: (Q_passive - Q_coord) / Q_coord')))
    add(DL(1584, meur))
    add(DL(1606, SF('three', 'int', lambda X: _all_same([X.fzp(f'{AF}.passive.n_starts'), X.fzp(f'{AF}.price_taker.n_starts')]),
                    'T6 n_starts (both arms)')))
    add(DL(1624, S('77.2', 'pct', ('rec', 'FIG5_DATA', 'computed[share_TSO].value'), 'W165 fig5 data (W130 totals)'),
           S('22.8', 'pct', ('rec', 'FIG5_DATA', 'computed[share_DSO].value'), 'W165 fig5 data (W130 totals)'),
           S('0.6', 'gwh', f'{AF}.price_taker.curtailment_phase_A_by_start.warm_from_certified.by_agent.DSO.mwh_day_weighted'),
           S('807.6', 'gwh', f'{AF}.price_taker.curtailment_phase_A_by_start.warm_from_certified.by_agent.TSO.mwh_day_weighted')))
    add(DL(1628, S('four', 'int', f'{BM}.reverse_flow_interface_hours'),
           S('one', 'int', f'{BM}.sweep.sweep_passive_cold.n_blocks'),
           SF('twelve', 'int', B1._twelve, 'SRP1.json: 3 years x 4 days = 12 TN blocks'),
           S('eight', 'int', f'{BM}.sweep.sweep_price_taker_cold.n_blocks')))
    # ---- 4.5 --------------------------------------------------------------------------------------------------------
    add(DL(1636, DER('nine', B1R.der_nine, 'joint scenarios of the 3 x 3 instance'),
           SF('three', 'int', lambda X: X.j[C3X]['NumMarketScenarios'], '3 x 3 NumMarketScenarios'),
           SF('three', 'int', lambda X: B1R._op_scen(X.j[C3X]), '3 x 3 num_operation_scenarios'),
           SF('five', 'int', lambda X: len(X.j[C3X]['Years']), '3 x 3 Years (count)'),
           SF('three', 'int', lambda X: _all_same(X.j[C3X]['Years'].values()), '3 x 3 block length'),
           CONV(S('236.7', 'k', 'three_by_three.rows[0].value'), 'gross (3 x 3, production exit)'),
           CONV(S('318.0', 'k', 'three_by_three.rows[1].value'), 'investment cost I'),
           CONV(S('81.3', 'k', 'three_by_three.rows[2].value'), 'gross (3 x 3, production exit)'),
           CONV(S('18.0', 'k', 'three_by_three.rows[3].value'), 'gross (3 x 3 resolution)')))
    add(DL(1640, CONV(S('6.2', 'k', 'three_by_three.rows[5].value'), 'gross (post-hoc descent of Q(0), 3 x 3)'),
           CONV(S('236.7', 'k', 'three_by_three.rows[0].value'), 'gross (3 x 3, production exit)'),
           CONV(DER('230.5', der_v_adjusted, 'V - D (3 x 3)'), 'gross (3 x 3, V - D)'),
           DER('7', B1.der_one_minus_r(7, 'upper'), 'value lower than single-scenario, lower end'),
           DER('9', B1.der_one_minus_r(8, 'lower'), 'value lower than single-scenario, upper end'), SINGLE('single')))
    add(DL(1645, P('7', 'node', UNIT, [7]), P('0.25', 'S', UNIT, [7]), P('1', 'E', UNIT, [7]), P('2025', 'year', UNIT)))
    add(DL(1652, CONV(S('236.7', 'k', 'three_by_three.rows[0].value'), 'gross (3 x 3, production exit)')))
    add(DL(1653, CONV(S('318.0', 'k', 'three_by_three.rows[1].value'), 'investment cost I')))
    add(DL(1654, CONV(S('81.3', 'k', 'three_by_three.rows[2].value', signed=True), 'gross (3 x 3, production exit)')))
    add(DL(1655, CONV(S('18.0', 'k', 'three_by_three.rows[3].value'), 'gross (3 x 3 resolution)')))
    add(DL(1656, SINGLE('Single'), S('253.5', 'k', 'three_by_three.rows[6].value', 'V_SRP1 settled (T7 2 % row)')))
    add(DL(1657, S('0.934', 'r3', 'three_by_three.rows[7].value')))
    add(DL(1658, S('0.909', 'r3', 'three_by_three.rows[8].value')))
    add(DL(1667, SINGLE('single')))
    add(DL(1668, CONV(S('6.2', 'k', 'three_by_three.rows[5].value'), 'gross (post-hoc descent of Q(0), 3 x 3)')))
    # ---- 4.6 --------------------------------------------------------------------------------------------------------
    add(DL(1677, CONV(DER('54.5', B1.der_net_range('min'), 'net cost over the investment years, lower end'), 'net of salvage'),
           CONV(DER('64.4', B1.der_net_range('max'), 'net cost over the investment years, upper end'), 'net of salvage'),
           SF('three', 'int', lambda X: 1 + len(X.t['year_ladder']['per_year']), 'T4: 2025 + the two deferral years'),
           P('2030', 'year', 'yl_y2030'), P('2035', 'year', 'yl_y2035'), P('7', 'node', G[0], [7]),
           P('0.5', 'S', G[0], [7]), P('1', 'E', G[0], [7]), P('0.5', 'S', G[2], [7]), P('2', 'E', G[2], [7]),
           SF('17', 'int', lambda X: 17 if appE_claims_rows_17_20(X) else None, 'tab:appE_claims rows 17-20 printed = the four '
                                                                               'G claims (labels and gross Delta)'),
           SF('20', 'int', lambda X: 20 if appE_claims_rows_17_20(X) else None, 'the same pointer, upper end')))
    add(DL(1682, P('7', 'node', 'yl_y2030', [7]), P('0.25', 'S', 'yl_y2030', [7]), P('1', 'E', 'yl_y2030', [7])))
    add(DL(1699, NOT('0')))
    add(DL(1701, P('2035', 'year', 'yl_y2035'), P('2030', 'year', 'yl_y2030'),
           SF('12.3', 'k', lambda X: _all_same([X.fzp(f'{YL}.gross_v6.threshold'), X.fzp(f'{YL}.net_v6.threshold')]),
              'T4 gross_v6.threshold = net_v6.threshold (the one threshold of the 2035 - 2030 pair)',
              conv='threshold of the gross and of the net comparison (equal)')))
    add(DL(1706, SF('four', 'int', lambda X: len(X.t['discount']), 'T7 rows (discount rates)'),
           S('40.8', 'k', 'discount[0.0].value_minus_I'), S('114.6', 'k', 'discount[0.08].value_minus_I'),
           S('8', 'pct0', 'discount[0.08].rate'), P('2025', 'year', UNIT), dict(two_pct), dict(two_pct)))
    add(DL(1711, P('7', 'node', UNIT, [7]), P('0.25', 'S', UNIT, [7]), P('1', 'E', UNIT, [7]), P('2025', 'year', UNIT)))
    for ln, rate, w in ((1718, '0.0', ('0', '277.1', '40.8', '15.9')), (1719, '0.02', ('2', '253.5', '64.4', '13.1')),
                        (1720, '0.05', ('5', '225.2', '92.7', '13.1')), (1721, '0.08', ('8', '203.4', '114.6', '13.1'))):
        add(DL(ln, S(w[0], 'pct0', f'discount[{rate}].rate'), S(w[1], 'k', f'discount[{rate}].V'),
               S(w[2], 'k', f'discount[{rate}].value_minus_I', signed=True),
               S(w[3], 'k', f'discount[{rate}].threshold_a61_conservative')))
    add(DL(1728, CONV(S('318.0', 'k', 'discount[0.02].I'), 'investment cost I'), P('2025', 'year', UNIT)))
    add(DL(1730, dict(two_pct), dict(two_pct)))
    # ---- 4.7 --------------------------------------------------------------------------------------------------------
    add(DL(1739, SINGLE('single'), S('32', 'int', WC + ('result.totals.n_certified',)),
           S('42', 'int', WC + ('result.totals.n_cells',)), S('ten', 'int', WC + ('result.totals.n_uncertified',)),
           CONV(W3('15', 'N4', 15, 'x0: 14,970.62 EUR after the residual-rule stop N = 132'), 'gross (movement of Q)'),
           CONV(W3('21', 'N5', 21, 'unit: 20,806.32 EUR after N = 112'), 'gross (movement of Q)'),
           DER('three', der_three_continued, 'evaluations continued past a certificate'),
           W3('0.9', 'N25', 0.9, 'largest movement past a settling-rule certificate: 0.882 tau (a bound)')))
    add(DL(1743, LANG('two', '"two numerical outcomes" (the two panels)'),
           SF('132', 'int', lambda X: X.fig4['a.N'], 'W165 fig4 data a.N'),
           SF('15.0', 'k', lambda X: X.fig4_checks['a.N_and_slack_equal_summary']['detail'][1], 'fig4 data: Q(N) - Q(k*)'),
           SF('181', 'int', lambda X: X.fig4['a.k_star'], 'fig4 data a.k_star'), dict(m15)))
    add(DL(1749, LANG('two', '"two single-scenario evaluations" (panels a and b-c)'), SINGLE('single')))
    add(DL(1751, S('1.5', 'raw', ('rec', 'FIG4_DATA', 'values[b.h_f9eae48f.flex_price_multiplier].value'), 'fig4 data')))
    add(DL(1754, DER('60', der_window60, 'panel (b) shading: the last cycles of the window'), dict(tau2)))
    add(DL(1760, SF('three', 'int', lambda X: len(X.w173['item1']), 'W173 item1: s47, s51, s53'),
           SF('41', 'int', lambda X: sum(X.w173['item1'][k]['n_new_evaluations_total'] for k in X.w173['item1']), 'W173 sum'),
           SINGLE('single'), S('42', 'int', WC + ('result.totals.n_cells',)),
           DER('61.4', der_wall_h, 'wall time of the 42-evaluation campaign (220,876 s)'),
           S('8,570', 'int', WC + ('result.totals.cycles_total',)), S('25.6', 'raw', WC + ('result.totals.wall_s_per_cycle.median',)),
           W3('single', 'N45', 1, 'single-threaded (thread caps)')))
    # ---- 4.8 --------------------------------------------------------------------------------------------------------
    add(DL(1773, W3('10^{-5}', 'N33', 1e-5, 'RES availability slack (EQUALITY_TOLERANCE)'),
           W3('62', 'N35', 62, '61.91 MWh-equivalent (W109)')))
    add(DL(1781, one_inst))
    add(DL(1785, LANG('one', '"one 108-bus test system"'), DER('108', B1.der_108, 'buses'), SINGLE('single'), *m3))
    # ---- conclusions ------------------------------------------------------------------------------------------------
    add(DL(1795, two_stage))
    add(DL(1799, DER('108', B1.der_108, 'buses'), meur, pct122, P('0.25', 'S', UNIT, [7]), P('1', 'E', UNIT, [7]),
           S('64.4', 'k', _cl(REF4, 'd_gross')), CONV(S('318.0', 'k', _cl(UNIT, 'I_other')), 'investment cost I')))
    add(DL(1803, P('4', 'dur', REF4, [7]), P('7', 'node', REF4, [7]),
           SF('1.75', 'raw', pos_m('1.75'), 'multiplier with a positive, determinate margin'),
           SF('2', 'raw', pos_m('2'), 'multiplier with a positive, determinate margin')))
    # ---- Appendix A -------------------------------------------------------------------------------------------------
    add(DL(1911, LANG('one', '"serves one purpose"'),
           BOUND('4', _w169_max, 1e-5, 'bound 4 x 10^-5 on min(P^Ch, P^Dch) / rating at the certified points'),
           BOUND('10^{-5}', _w169_max, 4.0, 'the exponent of the same bound')))
    add(DL(1925, U('1', 'notation', 'iteration index k+1')))
    add(DL(1926, U('2', 'notation', 'alpha^2 (exponent)'), U('1', 'notation', 'iteration index k+1')))
    add(DL(1927, U('2', 'notation', 'alpha^2 (exponent)')))
    add(DL(1934, U('1', 'notation', 'iteration index k+1')))
    add(DL(1935, U('1', 'notation', 'iteration index k+1'), U('1', 'notation', 'iteration index k+1')))
    add(DL(1946, U('1', 'notation', 'iteration index k+1'), U('1', 'notation', 'iteration index k+1'),
           U('1', 'notation', 'iteration index k+1')))
    add(DL(1953, U('1', 'notation', 'iteration index k+1')))
    add(DL(1954, U('1', 'notation', 'iteration index k+1'), U('1', 'notation', 'iteration index k+1')))
    add(DL(1993, W3('4', 'N18', 4.0, 'settling_criterion.TAU = DELTA_R * R_REF / 4.0')))
    add(DL(2008, W3('three', 'N62', 3, 'the three SRP1 reference continuations')))
    add(DL(2023, CODE('100', 'CODE_SC1', r'^EPS0 = TAU / ([0-9.]+)', desc='EPS0 = TAU / 100 (a step below it carries no sign)'),
           W3('10', 'N8', 10.0, 'SWING_FLOOR = TAU / 10'),
           CODE('three', 'CODE_SC1', r'^K_EXCL = (\d+)', desc='K_EXCL: the first cycles after k0 not counted'),
           W3('2', 'N15', 2.0, '2 P_max = 60 (split from the token "2P_")'),
           CODE('100', 'CODE_SC1', r'^EPS0 = TAU / ([0-9.]+)', desc='EPS0 (stationary sub-case of M1)'),
           W3('2', 'N17', 2.0, '|last step| x 2 P_max <= tau, L_MONO = 60 = 2 x 30 (split from "2P_")')))
    add(DL(2029, W3('two', 'N73', 2, 'two certified evaluations'), W3('3', 'N28', 3.0, 'DETERMINACY_BAR_FACTOR (split from "3b")'),
           W3('2', 'N29', 2.0, 'DETERMINACY_TAU_MULTIPLE')))
    add(DL(2033, W3('three', 'N26', 3, 'bar = 3.0 x max(gap, slack)')))
    add(DL(2113, CODE('10^{-5}', 'CODE_DEFS', r'^EQUALITY_TOLERANCE = ([0-9.e+-]+)',
                      uses=[('CODE_MCH', r'e_capacity \* ESS_DAY_BALANCE_SLACK_FRACTION \+ EQUALITY_TOLERANCE')],
                      desc='the closure bound allowance (p.u. at baseMVA 100)'),
           ctx='moved from section 3.5 (storage parameters) to Appendix A settings; the same quantity, meaning unchanged'))
    add(DL(2114, CODE('10^{-4}', 'CODE_DEFS', r'^SMALL_TOLERANCE = ([0-9.e+-]+)',
                      uses=[('CODE_DEFS', r'^ESS_COMPLEMENTARITY_TOLERANCE = SMALL_TOLERANCE'),
                            ('CODE_MCH', r'<= ESS_COMPLEMENTARITY_TOLERANCE')], desc='eps^C'),
           CODE('10^{-5}', 'CODE_DEFS', r'^EPS_ESSO_THROUGHPUT = ([0-9.e+-]+)',
                uses=[('CODE_SESD', r'slack_penalty \+= EPS_ESSO_THROUGHPUT \* throughput')], desc='eps^E'),
           CODE('10^{3}', 'CODE_DEFS', r'^PENALTY_ESSO_SLACK = ([0-9.e+-]+)',
                uses=[('CODE_SESD', r'slack_penalty \+= PENALTY_ESSO_SLACK \* \(model\.slack_es_pnet_up')], desc='c^sigma',
                conv='penalty weight (solver-side, not in Q)'),
           ctx='moved from section 3.5 (storage parameters) to Appendix A settings; meaning unchanged'))
    # ---- Appendix B / C RES tables (moved) ----------------------------------------------------------------------------
    add(DL(2149, SF('2025', 'int', B1R._year_in(SR, '2025'), 'SRP1.json Years: the base year of the profiles')))
    tn_ctx = ('moved from section 3.3 (TN table) to Appendix B; caption now "for the multi-scenario representative years", '
              'which the column years are (SRP1__s53_3x3.json Years); values unchanged in meaning')
    add(DL(2158, *[yrs3(y) for y in ('2025', '2028', '2031', '2034', '2037')], ctx=tn_ctx))
    for ln, gid, w in ((2160, 4, ('4', '4', '20.00', '20.00', '30.00', '30.00', '50.00')),
                       (2161, 5, ('5', '6', '40.00', '40.00', '60.00', '60.00', '60.00')),
                       (2162, 6, ('6', '8', '25.00', '25.00', '40.00', '40.00', '40.00')),
                       (2163, 7, ('7', '4', '12.00', '12.00', '20.00', '20.00', '20.00')),
                       (2164, 8, ('8', '6', '20.00', '20.00', '30.00', '30.00', '50.00')),
                       (2165, 9, ('9', '8', '15.00', '15.00', '25.00', '25.00', '25.00'))):
        add(gen_row(ln, 'case9', gid, w, ctx=tn_ctx))
    for net, node, cap_ln, yl, rows in (
            ('case33_1', 5, 2344, 2348, ((2350, 2, ('2', '6', '40.00', '40.00', '50.00', '50.00', '50.00')),
                                         (2351, 3, ('3', '23', '15.00', '15.00', '25.00', '25.00', '25.00')),
                                         (2352, 4, ('4', '25', '20.00', '20.00', '20.00', '20.00', '30.00')),
                                         (2353, 5, ('5', '18', '8.00', '8.00', '15.00', '15.00', '35.00')))),
            ('case33_2', 7, 2416, 2420, ((2422, 2, ('2', '6', '40.00', '40.00', '50.00', '50.00', '50.00')),
                                         (2423, 3, ('3', '23', '40.00', '40.00', '40.00', '40.00', '40.00')),
                                         (2424, 4, ('4', '25', '10.00', '10.00', '20.00', '20.00', '20.00')),
                                         (2425, 5, ('5', '18', '10.00', '10.00', '10.00', '10.00', '30.00')))),
            ('case33_3', 9, 2487, 2491, ((2493, 2, ('2', '6', '20.00', '20.00', '50.00', '50.00', '50.00')),
                                         (2494, 3, ('3', '25', '30.00', '30.00', '30.00', '30.00', '30.00')),
                                         (2495, 4, ('4', '23', '30.00', '30.00', '30.00', '30.00', '50.00')),
                                         (2496, 5, ('5', '18', '30.00', '30.00', '30.00', '30.00', '40.00'))))):
        ctx = (f'moved from section 3.4 (ADN at TN node {node}) to Appendix C ({net}); caption unchanged; the column years '
               f'are the multi-scenario representative years; meaning unchanged')
        cn = SF(str(node), 'int', (lambda nn, nm: (lambda X: nn if _dn_by_node(X).get(nn) == nm else None))(node, net),
                f'SRP1.json DistributionNetworks: {net} connection_node_id')
        add(DL(cap_ln, dict(cn), ctx=ctx))
        add(DL(yl, *[yrs3(y) for y in ('2025', '2028', '2031', '2034', '2037')], ctx=ctx))
        for ln, gid, w in rows:
            add(gen_row(ln, net, gid, w, ctx=ctx))
    add(DL(2482, SF('9', 'int', lambda X: 9 if _dn_by_node(X).get(9) == 'case33_3' else None,
                    'SRP1.json DistributionNetworks: case33_3 connection_node_id'),
           ctx='subsection title capitalisation changed ("node" -> "Node"); number unchanged'))
    # ---- Appendix E -------------------------------------------------------------------------------------------------
    add(DL(2580, U('1', 'notation', '(1 + r) discount base')))
    add(DL(2582, U('1', 'notation', '(1 - SoH^Min) normalisation (the salvage equation is rated against the code by A1)')))
    add(DL(2589, U('0', 'notation', 'max{0, ...}')))
    add(DL(2597, NOT('0'),
           ctx='app:salvage_credit: the salvage equations and the paragraph moved from the section 4 opening to Appendix '
               'E; equations rated by A1'))
    lat_ctx = ('moved from section 3.5 (storage parameters) to Appendix E.3 with fig:investment_lattice; meaning unchanged')
    add(DL(2797, LANG('single', '"a single node"'), SF('2025', 'int', lattice_fig_year,
                                                         'fig_lattice_case.py CS / CE / BUDGET = the W167 2025 expected unit '
                                                         'costs and the budget'),
           DER('twelve', der_lattice('n'), 'admissible single-node plans in 2025'),
           DER('0.25', der_lattice('S'), 'the smallest admissible plan, S'), DER('0.5', der_lattice('E'), 'the smallest, E'),
           DER('one', der_lattice('unit'), 'admissible plans one unit step from x = 0'),
           S('fourteen', 'int', ('rec', 'W173', 'item2.x0.b_box.n_feasible'), 'W173: admissible unit neighbours of x = 0'),
           n_dn, SF('2025', 'int', x0_year(2025), 'W173 item2.x0.rows: neighbour installation years {2025, 2030}'),
           SF('2030', 'int', x0_year(2030), 'the same'), ctx=lat_ctx))
    add(DL(2803, LANG('single', '"single-node" (compound adjective)'),
           SF('2025', 'int', lattice_fig_year, 'fig_lattice_case.py: 2025 expected unit costs'), ctx=lat_ctx))
    add(DL(2841, S('1', 'raw', ('rec', 'BENCH_SPEC_V5', 'tie_breaker.decision.passive_dso'),
                   'benchmark spec v5 tie_breaker.decision.passive_dso (EUR/MWh)', conv='auxiliary DN objective coefficient '
                                                                                      '(EUR/MWh), excluded from Q'),
           ctx='moved from section 3.6 ("a minimum-curtailment term of 1 EUR/MWh") to Appendix E.4; meaning unchanged'))
    add(DL(2845, SF('three', 'int', lambda X: _all_same([X.fzp(f'{AF}.passive.n_starts'), X.fzp(f'{AF}.price_taker.n_starts')]),
                    'T6 n_starts'),
           LANG('three', '"three phases", enumerated in the sentence that follows'),
           ctx='the E.4 paragraph re-flowed (one sentence split, the pointer to section 3.5 dropped); meaning unchanged'))
    add(DL(2847, DER('864', der_864, 'network-hours of the DN re-evaluation'), n_dn,
           SF('twelve', 'int', B1._twelve, 'SRP1.json: 3 years x 4 days'),
           S('24', 'int', ('rec', SR, 'NumInstants'), 'SRP1.json NumInstants'),
           ctx='the E.4 paragraph re-flowed; meaning unchanged'))
    add(DL(2915, SF('three', 'int', lambda X: len(X.w173['item1']), 'W173 item1: s47, s51, s53')))
    set_ctx = ('moved from section 3.6 (Planning and Certification Settings) to Appendix E.6; the sentence is unchanged; '
               'its pointer "Certification (Subsection~\\ref{subsubsec:certification}) uses delta_R" now points at a '
               'subsection that no longer states tau = delta_R V / 4 (moved to eq:acceptance_tolerance, Appendix A): the '
               'numbers keep their meaning; the pointer is A4\'s scope')
    add(DL(2919, W3('0.07', 'N19', 0.07, 'settling_criterion.DELTA_R'), CONV(W3('259,375.33', 'N20', 259375.33, 'settling_criterion.R_REF'),
                                                                'gross (reference storage value Q(0) - Q(unit))'),
           S('4,539.07', 'raw', '$.constants.TAU', 'frozen constants.TAU (EUR)', conv='tolerance on Q (gross, EUR)'),
           W3('30', 'N16', 30, 'p_max'), ctx=set_ctx))
    add(DL(2920, SF('100', 'int', W178B.cap_after_record, 'frozen re-settling specs: fixed caps = N_old + 100'),
           LANG('one', 'pronoun ("a fresh one")'),
           SF('109', 'int', W178B.fresh_cap('after_first_k0'), 'dynamic caps min(k0 + 109, ceiling)'),
           SF('300', 'int', W178B.fresh_cap('ceiling'), 'dynamic cap ceiling'),
           SF('4', 'int', lambda X: _all_same([W178B._rs(X, 'delta0', 's47'), W178B._rs(X, 'delta0', 's51')]),
              'search specs s47 / s51 delta_0'),
           SF('20', 'int', lambda X: _all_same([W178B._rs(X, 'budgets', 's47')[0], W178B._rs(X, 'budgets', 's51')[0]]),
              'max_new_evaluations of s47 / s51'),
           SF('two', 'int', lambda X: W178B.n_searches_budget(X, 20), 'searches with budget 20'),
           SF('60', 'int', lambda X: W178B._rs(X, 'budgets', 's53')[0], 'max_new_evaluations of s53'),
           SF('30', 'int', lambda X: _all_same(W178B._rs(X, 'completion_cap', k) for k in ('s47', 's51', 's53')),
              'completion_cap (all three)'),
           SF('18,449.66', 'raw', lambda X: _all_same(W178B._rs(X, 'sigma_Q', k) for k in ('s47', 's51', 's53')),
              'sigma_Q (all three)', conv='search resolution on F (gross, EUR)'), ctx=set_ctx))
    add(DL(2921, SF('ten', 'int', lambda X: _all_same(W178B._rs(X, 'production_exit', k)[0] for k in ('s47', 's51', 's53')),
                    'required_consecutive_cycles (all three)'),
           SF('500', 'int', lambda X: _all_same(W178B._rs(X, 'production_exit', k)[1] for k in ('s47', 's51', 's53')),
              'cap (all three)'),
           SF('two', 'int', lambda X: W178B._rs(X, 'barrier_stop', 's53_barrier_stop')[0], 'barrier stop per poll'),
           LANG('one', '"in one poll"'),
           SF('three', 'int', lambda X: W178B._rs(X, 'barrier_stop', 's53_barrier_stop')[1], 'barrier stop overall'),
           SF('60', 'int', lambda X: _all_same(W178B._rs(X, 'max_polls', k) for k in ('s47', 's51', 's53')), 'max_polls'),
           U('single', 'compound adjective', '"single-node ladders"'),
           SF('two', 'int', W178B.n_searches_m2_cache, 'searches whose cache holds committed certified m = 2 records'),
           SF('2030', 'int', W178B.year_multi_node_s53, 'installation year of every multi-node entry of the s53 cache'),
           ctx=set_ctx))
    add(DL(2925, S('14', 'int', ('rec', 'W173', 'item1.s47.n_new_evaluations_total')),
           S('20', 'int', ('rec', 'W173', 'item1.s51.n_new_evaluations_total')),
           S('seven', 'int', ('rec', 'W173', 'item1.s53.n_new_evaluations_total'))))
    add(DL(2929, SF('10^{-6}', 'sci', w180d_relax, 'W180d conversion.relaxation_mwh_per_slack (IPOPT bound_relax_factor 1e-8 '
                                                   'x baseMVA 100 x 1 h)'),
           SF('16', 'int', w180d_row16('no'), 'W180d row16: E.21 row 16 = pb_y2025_n5 (= W181a e21_map[16])'),
           SF('1.8', 'raw', w180d_row16('S_mant'), 'W180d row16.S_mwh mantissa (1.7637e-6 MWh)'),
           SF('10^{-6}', 'sci', w180d_row16('S_exp'), 'W180d row16.S_mwh order of magnitude'),
           SF('5', 'raw', w180d_row16('F_mant'), 'W180d row16.S_over_e_av mantissa (5.039e-6)'),
           SF('10^{-6}', 'sci', w180d_row16('F_exp'), 'W180d row16.S_over_e_av order of magnitude'),
           ctx='moved from section 2 (storage operation in the network models) to Appendix E.6; the sentence is unchanged '
               'apart from its neighbours; meaning unchanged'))
    stat_ctx = ('moved within Appendix E.6 (after the settings paragraphs); one clause replaced ("so the rule bounds the '
                'stopping error ... at about tau, not well inside it" -> "so the accepted window ranges can lie close to tau, '
                'without providing a guaranteed bound"); the numbers keep their meaning')
    seq = campaign42_seq
    add(DL(2934, S('42', 'int', WC + ('result.totals.n_cells',)), SINGLE('single'),
           *[SF(w, 'int', seq_tok(seq, i), 'W153 cells (42) mapped to E.21 rows by W181a e21_map, compressed to ranges')
             for i, w in enumerate(('3', '11', '15', '22', '23', '26', '31', '34', '39', '42', '59'))],
           S('32', 'int', WC + ('result.totals.n_certified',)),
           S('24', 'int', WC + ('result.totals.branch_counts_certified.oscillatory',)),
           S('8', 'int', WC + ('result.totals.branch_counts_certified.monotone',)),
           S('10', 'int', WC + ('result.totals.n_uncertified',)),
           S('5', 'int', WC + ('result.totals.uncertified_by_cause.gap clause.n',)),
           S('2', 'int', WC + ('result.totals.uncertified_by_cause.lapse reset.n',)),
           S('3', 'int', WC + ('result.totals.uncertified_by_cause.growth test.n',)),
           S('174', 'raw', WC + ('result.totals.k_star.median',)),
           S('65', 'raw', WC + ('result.totals.cycles_after_k0_k_star_minus_N.median',)),
           S('21', 'int', WC + ('result.totals.cycles_after_k0_k_star_minus_N.min',)),
           S('89', 'int', WC + ('result.totals.cycles_after_k0_k_star_minus_N.max',)),
           S('0.86', 'efc', WC + ('result.totals.range_over_tau_at_k_star.median',)),
           SF('four', 'int', lambda X: len(added_later_rows(X)), 'cells_appended_w160 T10 rows, certified'),
           *[SF(w, 'int', seq_tok(added_later_rows, i), 'E.21 rows of the T10 cells (W181a e21_map)')
             for i, w in enumerate(('32', '33', '60', '61'))],
           S('Ten', 'int', 'at_or_above_0_95_tau.n_counted'),
           W3('46', 'N23', 46, 'certificates the tables use (49 - 1 superseded - 2 twins)'),
           W3('5', 'N24', 5, '1 - 0.95 (the >= 0.95 tau flag)'),
           *[SF(w, 'int', seq_tok(counted_rows, i), 'E.21 rows of the counted >= 0.95 tau certificates (W181a e21_map)')
             for i, w in enumerate(('2', '3', '5', '10', '11', '12', '14', '17', '20', '23'))],
           SF('36', 'int', beside_row, 'the "beside" cell (monotone) -> E.21 row'),
           SF('0.94', 'efc', beside_range, 'its range_over_tau (0.939)'),
           DER('three', der_three_continued, 'evaluations continued past a certificate'),
           *[SF(w, 'int', continued_row(i), f'E.21 row of {N25_ORDER[i]} (W163 N25; W181a e21_map)')
             for i, w in enumerate(('12', '14', '3'))],
           SF('0.004', 'raw', continued_over_tau('b_2a0ba8b2', 'v4 k* 173'), 'W163 N25 over_tau |movement| / tau'),
           SF('0.74', 'raw', continued_over_tau('b_4649234b', 'v5 k* 148'), 'W163 N25 over_tau'),
           SF('0.78', 'raw', continued_over_tau('d_c52e1670', 'v6 k* 150'), 'W163 N25 over_tau (v6 k* 150 reading)'),
           SF('0.88', 'raw', continued_over_tau('d_c52e1670', 'last-pair reading k* 148'), 'W163 N25 over_tau (k* 148 reading)'),
           S('21', 'int', WC + ('result.totals.non_clean_cycles_after_N_total',)),
           SF('10', 'int', lambda X: len(X.w153c['cells_with_non_clean_after_N']), 'W153 cells with non-clean cycles after N'),
           S('119', 'int', WC + ('result.totals.acceptable_clean_block_events_after_N_total',)),
           W3('ten', 'N59', 10, 'CLEAN_FACTOR 10'), ctx=stat_ctx))
    return L


# ======================================================================================================================
#  4. tokens of the '+' lines, the declarations applied, the negative control
# ======================================================================================================================
def added_lines(diff_u0):
    out = set()
    for m in re.finditer(r'^@@ -\d+(?:,\d+)? \+(\d+)(?:,(\d+))? @@', diff_u0, re.M):
        a, n = int(m.group(1)), (int(m.group(2)) if m.group(2) is not None else 1)
        out.update(range(a, a + n))
    return out


def removed_lines(diff_u0):
    out = set()
    for m in re.finditer(r'^@@ -(\d+)(?:,(\d+))? \+\d+(?:,\d+)? @@', diff_u0, re.M):
        a, n = int(m.group(1)), (int(m.group(2)) if m.group(2) is not None else 1)
        out.update(range(a, a + n))
    return out


COEF_RX = re.compile(r'^(\d+)([A-Za-z_].*)$')


def split_coefficients(toks):
    """math-mode alnum tokens with a leading integer coefficient (3b, 2P_): the digits become a number token."""
    n = 0
    for t in toks:
        if t['type'] == 'alnum' and t.get('in_math') and not t['excluded']:
            m = COEF_RX.match(t['text'])
            if m:
                d = m.group(1)
                t.update(split_from=t['text'], text=d, written=d, type='num', value=float(d), end0=t['col0'] + len(d),
                         tokenizer='W182b coefficient split')
                n += 1
    return n


def apply_declarations(X, toks, decls, lines):
    stale, claimed_twice, decl_out = [], [], []
    by_line = collections.defaultdict(list)
    for t in toks:
        by_line[t['line']].append(t)
    seen = collections.Counter(d['line'] for d in decls)
    dup = [ln for ln, n in seen.items() if n > 1]
    for d in decls:
        text = lines[d['line'] - 1]
        cand = [t for t in by_line[d['line']] if t['status'] != 'excluded']
        if len(cand) != len(d['specs']):
            stale.append({'line': d['line'], 'why': f"{len(d['specs'])} specs for {len(cand)} tokens",
                          'tokens': [t['written'] for t in cand], 'specs': [s['w'] for s in d['specs']]})
            continue
        bad = [(t['written'], s['w']) for t, s in zip(cand, d['specs']) if t['written'].lower() != s['w'].lower()]
        if bad:
            stale.append({'line': d['line'], 'why': 'written differs from the declaration', 'pairs': bad})
            continue
        res = []
        for t, s in zip(cand, d['specs']):
            if t['status'] is not None:
                claimed_twice.append(t['key'])
                continue
            try:
                r = evaluate(X, s, t, text)
            except Exception as e:  # noqa: BLE001 -- a declaration that cannot be evaluated is a run failure, listed
                stale.append({'line': d['line'], 'why': f'evaluation error {type(e).__name__}: {e}', 'written': t['written']})
                continue
            t.update(r)
            t['spec'] = s
            res.append({'written': t['written'], 'col': t['col'], 'status': r['status']})
        decl_out.append({'line': d['line'], 'tokens': res, 'note': d.get('note'), 'context': CTX.get(d['line'])})
    return stale, claimed_twice, decl_out, dup


NUMWORDS = dict(W164.NUMWORDS)
NUMWORDS.update(W174.EXTRA_WORDS)
WORD_OF = {}
for _w, _v in sorted(NUMWORDS.items(), key=lambda kv: len(kv[0])):
    WORD_OF.setdefault(_v, _w)


def perturb(t, down=False):
    """an altered copy of a token: the last digit +1 (9 -> 8), or -1 when `down` (0 -> 1); an exponent shifted by one; a
    number word replaced by its neighbour."""
    u = dict(t)
    if t['type'] == 'word':
        v = W174.wv3(t['text'])
        nv = v - 1 if (down and v - 1 in WORD_OF) else v + 1
        if nv not in WORD_OF:
            nv = v - 1
        u['text'] = u['written'] = WORD_OF[nv]
        u['value'] = float(nv)
        return u
    if t['type'] == 'sci':
        m = re.search(r'(-?)(\d+)', t['text'].split('^', 1)[1])
        e = int(m.group(1) + m.group(2))
        ne = e - 1 if (down or e < 0) else e + 1
        u['text'] = u['written'] = f'10^{{{ne}}}'
        u['value'] = 10.0 ** ne
        return u
    w = t['written']
    m = re.search(r'(\d)(?!.*\d)', w)
    dgt = int(m.group(1))
    nd = (dgt - 1 if dgt > 0 else 1) if down else (dgt + 1 if dgt < 9 else 8)
    nw = w[:m.start()] + str(nd) + w[m.end():]
    u['written'] = nw
    u['text'] = nw
    u['value'] = float(nw.replace(',', '').replace('+', ''))
    if t.get('sign') == '-':
        u['value'] = -abs(u['value'])
    return u


def negative_control(X, toks, lines):
    out = []
    for t in toks:
        if t['status'] not in ('match', 'approximate', 'derived') or 'spec' not in t:
            continue
        s = t['spec']
        u = perturb(t, down=(s['kind'] == 'bound'))
        try:
            r = evaluate(X, s, u, lines[t['line'] - 1])
            st = r['status']
        except Exception as e:  # noqa: BLE001 -- an evaluator that crashes on an altered number refuses it
            st = f'error {type(e).__name__}'
        out.append({'line': t['line'], 'col': t['col'], 'written': t['written'], 'altered': u['written'],
                    'status_altered': st, 'refused': st not in ('match', 'approximate', 'derived')})
    return out


def convention(t):
    """the objective convention of a rated token (monetary tokens); 'n/a' for counts, years, ratios of plans, etc."""
    s = t.get('spec') or {}
    if s.get('conv'):
        return s['conv']
    if t['status'] not in ('match', 'approximate', 'derived', 'MISMATCH', 'no counterpart'):
        return None
    src = str(s.get('src')) + ' ' + str(t.get('counterpart'))
    kind = s.get('kind')
    unit = t.get('unit')
    monetary = kind in ('k', 'meur') or unit in ('kEUR', 'MEUR') or (kind == 'derived' and unit in ('kEUR', 'MEUR'))
    if kind in ('pct',) and BM in src:
        return 'gross (ratio of gross Q values, settlement excluded)'
    if kind == 'derived' and 'benchmark.benefit' in src:
        return 'gross (Q = gross_operational_cost, settlement excluded)'
    if not monetary:
        return 'n/a'
    if any(k in src for k in ('d_net', 'M_net', 'D_net')):
        return 'net of the terminal salvage credit (Q_net = Q - salvage)'
    if 'salvage' in src:
        return 'terminal salvage credit'
    if src.strip().endswith('.I') or 'I_other' in src or '].I ' in src:
        return 'investment cost I'
    return 'gross (Q = gross_operational_cost, settlement excluded)'


# ======================================================================================================================
#  5. statement checks (claims in the hunks that bear on rated numbers), E.21 identity
# ======================================================================================================================
def statement_checks(X):
    out = []
    six = (C5, REF4, 'B:n7_2h_e1', 'B:n7_4h_e2', C59, C57)
    ok = all(X.CL[c]['gross_verdict'] == 'determinate' and X.cells[X.CL[c]['other_cell']]['status'] == 'certified' and
             X.cells[X.CL[c]['ref_cell']]['status'] == 'certified' for c in six)
    out.append({'line': 1404, 'text': 'tab:res_value_ladder note: "All listed evaluations are certified, and all cost increases '
                'are determinate."', 'rating': 'MATCH' if ok else 'MISMATCH',
                'evidence': {c: [X.CL[c]['gross_verdict'], X.cells[X.CL[c]['ref_cell']]['status'],
                                 X.cells[X.CL[c]['other_cell']]['status']] for c in six}})
    m2 = ('H:m2:value_minus_I', 'I:m2:e2_value_minus_I', 'J:e3_value_minus_I', 'J:e4_value_minus_I', 'J:e5_value_minus_I',
          'I:m2:second_MWh', 'J:e2_to_e3', 'J:e3_to_e4', 'J:e4_to_e5')
    ev = {c: [X.CL[c]['gross_verdict'], (X.CL[c].get('uncertified_form_beside_A64_ruling4') or {}).get('verdict')] for c in m2}
    ok = all(v[0] == 'determinate' and v[1] in (None, 'determinate') for v in ev.values())
    out.append({'line': 1500, 'text': 'tab:res_energy_ladder_m2 note: "All listed margins and increments are determinate." '
                '(the flagged 4 MWh row and its increments also in the uncertified form)', 'rating': 'MATCH' if ok else 'MISMATCH',
                'evidence': ev})
    it = X.w173['item1']
    rows = []
    ok = True
    for k in ('s47', 's51'):
        for p in it[k]['polls']:
            evaluated_from_dirs = p['n_new_evaluations'] > 0 and p['n_direction_points_feasible'] > 0 and \
                p['decision'] != 'stopped_for_review_completion_cap'
            rows.append([k, p['poll_index'], p['n_new_evaluations'], p['n_direction_points_feasible'], p['decision']])
            ok = ok and not evaluated_from_dirs
    tot = sum(r[2] for r in rows if r[4] != 'stopped_for_review_completion_cap')
    ok = ok and tot == int(it['s47']['n_new_evaluations_total']) + int(it['s51']['n_new_evaluations_total'])
    out.append({'line': 2925, 'text': '"All new evaluations in the variant-A searches came from unit-neighbour completion."',
                'rating': 'MATCH' if ok else 'MISMATCH', 'evidence': {'polls': rows, 'evaluated_total': tot,
                                                                      'note': 's51 poll 4 listed one admissible direction '
                                                                              'point, not evaluated (poll refused at the '
                                                                              'completion cap)'}})
    c = X.w173['item2']['F2']['c_L_claims_box_neighbours']
    det_better = [x for x in X.w173['item2']['F2'].get('L_claims_detail', []) if isinstance(x, dict) and
                  x.get('d_gross', 1) < 0 and x.get('gross_verdict') == 'determinate']
    ok = int(c['n']) == 13 and int(c['n_positive_plan_better']) + int(c['n_negative_neighbour_better']) == 13 and \
        all(n['gross_verdict'] != 'determinate' for n in c['negative']) and not det_better
    out.append({'line': 1507, 'text': '"None of those 13 was determinately better than the returned plan"',
                'rating': 'MATCH' if ok else 'MISMATCH', 'evidence': {k: c[k] for k in c}})
    ag = {r['arm']: r['verdict'] for r in X.t['ageing']['rows']}
    ok = ag.pop('no_ageing') == 'within resolution' and all(v == 'determinate' for v in ag.values()) and \
        X.A64['E:soh050:value_minus_I']['gross_verdict'] == 'determinate'
    out.append({'line': 1565, 'text': 'tab:res_ageing note: "w: margin within numerical resolution; all other margins are '
                'determinate."', 'rating': 'MATCH' if ok else 'MISMATCH', 'evidence': {'T8': {r['arm']: r['verdict'] for r in
                                                                                             X.t['ageing']['rows']},
                                                                                      'A64 e_soh050': X.A64['E:soh050:value_minus_I']['gross_verdict']}})
    tg = tn_growth(X)
    out.append({'line': 1290, 'text': 'growth rates 1.25 % / 3.00 % (ADN workbooks) -- TN workbook for comparison',
                'rating': 'INFO', 'evidence': {'case9_operational_data.xlsx Growth Factors': tg}})
    return out


def e21_rows(lines):
    a = next((i for i, ln in enumerate(lines) if '\\label{tab:appE_evaluations}' in ln), None)
    if a is None:
        return None
    b = next(i for i in range(a, len(lines)) if lines[i].startswith('\\end{longtable}'))
    return [ln for ln in lines[a:b] if re.match(r'^\d+ & ', ln)]


# ======================================================================================================================
#  6. guards, post-run, summary
# ======================================================================================================================
def guards_state():
    guards = {nm: {'counts': dict(g.counts), 'verify_0_failures': g.verify(0)} for nm, g in GUARDS}
    blocked = pickle.load is _blocked_load and pickle.loads is _blocked_loads
    pk = {'counts': {'w182b': dict(PICKLE_COUNTS), 'w178b': dict(W178B.PICKLE_COUNTS), 'w177b1r': dict(B1R.PICKLE_COUNTS)},
          'pickle_load_and_loads_blocked': blocked}
    pk['ok'] = blocked and all(v == {'load': 0, 'loads': 0} for v in pk['counts'].values())
    return guards, pk, all(not v['verify_0_failures'] for v in guards.values()) and pk['ok']


def post_run(out_dir):
    rels = [os.path.join(out_dir, OUT_NAMES[k]) for k in ('man', 'json', 'md', 'typing_json', 'typing_log')]
    rels += [os.path.join(out_dir, f) for f in sorted(os.listdir(os.path.join(REPO, out_dir))) if LAUNCH_LOG_RX.fullmatch(f)]
    for rel in rels:
        if not os.path.exists(os.path.join(REPO, rel)):
            _log(f'[W182b post-run PRECONDITION FAILED] {rel} missing')
            sys.exit(1)
    man = {rel: _sha(rel) for rel in rels}
    with open(os.path.join(REPO, out_dir, OUT_NAMES['post']), 'x', encoding='utf-8') as h:
        h.write(GRIO.dumps(man, indent=1, sort_keys=True) + '\n')
    _log(f"[W182b post-run] wrote {os.path.join(out_dir, OUT_NAMES['post'])}: " + ', '.join(f'{k} {v[:8]}' for k, v in man.items()))
    sys.exit(0)


def other_copies():
    r = subprocess.run(['ps', '-Ao', 'pid=,args='], capture_output=True, text=True).stdout.splitlines()
    me = os.getpid()
    out = []
    for ln in r:
        parts = ln.strip().split(None, 1)
        if len(parts) < 2 or int(parts[0]) == me:
            continue
        argv = parts[1].split()
        if argv and 'python' in os.path.basename(argv[0]) and any(a.endswith(SCRIPT_REL) for a in argv[1:3]) \
                and '--post-run' not in argv:
            out.append(ln.strip())
    return out


def _cell(s, n=230):
    s = str(s).replace('|', '\\|').replace('\n', ' ')
    return s if len(s) <= n else s[:n - 3] + '...'


def md_summary(o):
    s = o['summary']
    L = ['# W182b -- A2 number check, Overleaf edd7a63 (49 hunks from 8f051a2)', '',
         f"Generated by `{o['script']['path']}` (sha256 {o['script']['sha256'][:8]}), {o['utc']}; repo HEAD "
         f"{o['git_head'][:8]}; clone {o['clone']['head'][:8]}. Zero solves; pickle 0.", '',
         f"Objective convention: {o['objective_convention']}", '',
         f"Tokens on the '+' lines: {s['n_tokens']} ({s['n_added_lines']} '+' lines). A2 ratings: {s['by_rating']}. "
         f"Checker statuses: {s['by_status']}.", '',
         f"Negative control: {s['negative_control']['n_refused']} of {s['negative_control']['n']} perturbed rated tokens "
         f"refused.", '', '## Non-MATCH (MISMATCH, UNSOURCED)', '']
    nm = [t for t in o['tokens'] if t['rating'] in ('MISMATCH', 'UNSOURCED')]
    if not nm:
        L.append('none')
    for t in nm:
        L.append(f"- l. {t['line']}:{t['col']} `{t['written']}` {t['rating']} -- {_cell(t.get('counterpart'))}; "
                 f"{_cell(t.get('note'))}")
    L += ['', '## DERIVED (formula in the harness)', '', '| line | written | value at written precision | arithmetic | convention |',
          '|---|---|---|---|---|']
    for t in o['tokens']:
        if t['rating'] == 'DERIVED':
            L.append(f"| {t['line']} | {t['written']} | {_cell(t.get('value_at_written_precision'), 30)} | "
                     f"{_cell(t.get('counterpart'), 300)} | {t.get('convention')} |")
    L += ['', '## MATCH at a coarser written precision (approximate)', '']
    ap = [t for t in o['tokens'] if t['status'] == 'approximate']
    L += [f"- l. {t['line']} `{t['written']}` -> {t.get('value_at_written_precision')} ({_cell(t.get('counterpart'), 160)})"
          for t in ap] or ['none']
    L += ['', '## Statement checks', '']
    L += [f"- l. {x['line']}: {x['text']} -- **{x['rating']}**" for x in o['statement_checks']]
    L += ['', '## Moved lines (context)', '']
    for ln, c in sorted(o['contexts'].items(), key=lambda kv: int(kv[0])):
        L.append(f'- l. {ln}: {c}')
    L += ['', '## Every rated token', '', '| line | col | written | rating | status | class | convention | counterpart |',
          '|---|---|---|---|---|---|---|---|']
    for t in o['tokens']:
        L.append(f"| {t['line']} | {t['col']} | {t['written']} | {t['rating']} | {t['status']}"
                 f"{' (' + str(t.get('category')) + ')' if t.get('category') else ''} | {t['line_class'][:24]} | "
                 f"{t.get('convention') or ''} | {_cell(t.get('counterpart') or t.get('note'), 180)} |")
    L += ['', '## Integrity checks', ''] + [f"- {k}: {v}" for k, v in o['checks'].items()] + ['', f"failed: {o['failed']}", '']
    return '\n'.join(L) + '\n'


# ======================================================================================================================
#  7. main
# ======================================================================================================================
def main():
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--overleaf-commit', required=True)
    ap.add_argument('--expect', action='append', default=[], help='file.tex=<sha256> for every *.tex of the clone')
    ap.add_argument('--post-run', action='store_true')
    ap.add_argument('--trial-out-dir', default=None, help='trial runs only: an output directory outside the results')
    ap.add_argument('--launch-log', default=OUT_NAMES['log'])
    args = ap.parse_args()
    out_dir = args.trial_out_dir or OUT_DIR
    tag = 'W182b'
    if not LAUNCH_LOG_RX.fullmatch(args.launch_log):
        _log(f'[{tag} PRECONDITION FAILED] --launch-log {args.launch_log!r}')
        sys.exit(1)
    if args.overleaf_commit != DECLARED_COMMIT:
        _log(f'[{tag} PRECONDITION FAILED] --overleaf-commit {args.overleaf_commit} != declared {DECLARED_COMMIT}')
        sys.exit(1)
    if args.post_run:
        post_run(out_dir)
    t0 = time.time()
    pre = []
    others = other_copies()
    if others:
        pre.append(f'another copy of {SCRIPT_REL} is running: {others}')
    od = out_dir if os.path.isabs(out_dir) else os.path.join(REPO, out_dir)
    refused_logs = []
    if not os.path.isdir(od):
        pre.append(f'{out_dir} missing (the launch command creates it with mkdir)')
    else:
        present = sorted(os.listdir(od))
        extra = [f for f in present if not LAUNCH_LOG_RX.fullmatch(f)]
        if extra:
            pre.append(f'{out_dir} is not new: it holds {extra} (write-once)')
        refused_logs = [f for f in present if LAUNCH_LOG_RX.fullmatch(f) and f != args.launch_log]
    # ---- clone ----------------------------------------------------------------------------------------------------------
    mabs = os.path.join(REPO, CLONE)

    def cg(*a):
        r = subprocess.run(['git', '-C', mabs] + list(a), capture_output=True, text=True)
        return r.stdout if r.returncode == 0 else None
    head = (cg('rev-parse', 'HEAD') or '').strip()
    porcelain = (cg('status', '--porcelain', '--', '*.tex', BIB) or '').strip()
    if head != CLONE_HEAD:
        pre.append(f'clone HEAD {head} != declared {CLONE_HEAD}')
    expect = dict(e.split('=', 1) for e in args.expect if '=' in e)
    tex_now = sorted(f for f in os.listdir(mabs) if f.endswith('.tex'))
    tex_sha = {f: _sha(os.path.join(CLONE, f)) for f in tex_now}
    if sorted(expect) != tex_now or sorted(TEX_SHA) != tex_now:
        pre.append(f'--expect files {sorted(expect)} / declared {sorted(TEX_SHA)} != the clone *.tex {tex_now}')
    for f in tex_now:
        if expect.get(f) != tex_sha[f] or TEX_SHA.get(f) != tex_sha[f]:
            pre.append(f'{f}: sha256 {tex_sha[f]} != --expect {expect.get(f)} / declared {TEX_SHA.get(f)}')
    blob_eq = {}
    for f in tex_now + [BIB]:
        hb, hh = (cg('hash-object', f) or '').strip(), (cg('rev-parse', f'HEAD:{f}') or '').strip()
        blob_eq[f] = bool(hb) and hb == hh
    if porcelain:
        pre.append(f'clone *.tex / {BIB} not clean: {porcelain!r}')
    base_ok = (cg('rev-parse', '8f051a2') or '').strip() == BASE_COMMIT
    if not base_ok:
        pre.append('8f051a2 does not resolve to the declared base commit')
    # ---- inputs ---------------------------------------------------------------------------------------------------------
    all_inputs = dict(W177A.INPUTS)
    all_inputs.update(B1.EXTRA_INPUTS)
    all_inputs.update(B1R.EXTRA_INPUTS)
    all_inputs.update(W178B.EXTRA_INPUTS)
    all_inputs.update(EXTRA_INPUTS)
    inputs = {}
    for key, rel in all_inputs.items():
        if not os.path.exists(os.path.join(REPO, rel)):
            pre.append(f'capture path missing: {key} {rel}')
            continue
        tracked_in_clone = rel.startswith(CLONE + os.sep)
        inputs[key] = {'path': rel, 'sha256': _sha(rel),
                       'committed_clean': (True if tracked_in_clone else W160.W157.L132._committed_clean(rel)),
                       'last_commit': None if tracked_in_clone else W160._last_commit(rel)}
        if tracked_in_clone:
            r = subprocess.run(['git', '-C', mabs, 'status', '--porcelain', '--', os.path.relpath(rel, CLONE)],
                               capture_output=True, text=True).stdout.strip()
            inputs[key]['committed_clean'] = r == ''
        if not inputs[key]['committed_clean']:
            pre.append(f'{rel} not committed clean')
    if 'FROZEN_JSON' in inputs and inputs['FROZEN_JSON']['sha256'] != FZ_SHA:
        pre.append(f'frozen JSON sha256 != {FZ_SHA}')
    script_clean = W160.W157.L132._committed_clean(SCRIPT_REL)
    if pre:
        _log(f'[{tag} PRECONDITION FAILED] ' + ' | '.join(pre))
        sys.exit(1)
    _log(f'[{tag}] preconditions hold: clone {head[:8]} ({len(tex_now)} *.tex as declared); {len(inputs)} inputs committed '
         f'clean; script committed clean {script_clean}; earlier launch logs kept {refused_logs}')

    # ---- context -------------------------------------------------------------------------------------------------------
    X = W177A.Ctx(inputs)
    for k in B1.EXTRA_JSON + B1R.EXTRA_JSON + W178B.EXTRA_JSON + EXTRA_JSON:
        X.j[k] = _json(inputs[k]['path'])
    for k in B1R.EXTRA_TEXT + W178B.EXTRA_TEXT + EXTRA_TEXT:
        X.texts[k] = _text(inputs[k]['path'])
    X.texts['CODE_DEFS'] = _text(inputs['CODE_DEFS']['path'])
    X.texts['CODE_SES'] = _text(inputs['CODE_SES']['path'])
    X.texts['CODE_SESD'] = _text(inputs['CODE_SESD']['path'])
    X.texts['BE_CONV'] = X.t['break_even']['objective_convention']
    X.def_src = {'BE_CONV': f'{FZ_REL} tables.break_even.objective_convention'}
    main_text = _text(os.path.join(CLONE, 'main.tex'))
    lines = main_text.split('\n')
    X.main_lines = lines
    base_lines = (cg('show', f'{BASE_COMMIT}:main.tex') or '').split('\n')
    w181_lines = (cg('show', f'{W181A_COMMIT}:main.tex') or '').split('\n')
    e21_now, e21_w181 = e21_rows(lines), e21_rows(w181_lines)
    X.e21_rows_unchanged = bool(e21_now) and e21_now == e21_w181
    X.cells = dict(X.t['cells'])
    X.cells.update(X.t['cells_appended_w160'])

    # ---- the diff --------------------------------------------------------------------------------------------------------
    d0 = cg('diff', '-U0', BASE_COMMIT, CLONE_HEAD, '--', 'main.tex') or ''
    dd = cg('diff', BASE_COMMIT, CLONE_HEAD, '--', 'main.tex') or ''
    hunks = {'default_context': len(re.findall(r'^@@ ', dd, re.M)), 'zero_context': len(re.findall(r'^@@ ', d0, re.M))}
    plus = added_lines(d0)
    minus = removed_lines(d0)
    plus_dd = set()
    for m in re.finditer(r'^@@ -\d+(?:,\d+)? \+(\d+)(?:,(\d+))? @@(.*?)(?=^@@ |\Z)', dd, re.M | re.S):
        ln = int(m.group(1))
        for row in m.group(3).split('\n')[1:]:
            if row.startswith('+'):
                plus_dd.add(ln)
                ln += 1
            elif row.startswith('-'):
                continue
            elif row.startswith(' ') or row == '':
                ln += 1
    norm = lambda s: ' '.join(s.split())  # noqa: E731
    removed_norm = collections.defaultdict(list)
    for ln in sorted(minus):
        if ln - 1 < len(base_lines) and norm(base_lines[ln - 1]):
            removed_norm[norm(base_lines[ln - 1])].append(ln)

    # ---- tokens ---------------------------------------------------------------------------------------------------------
    doc = W164.Doc('main.tex', main_text)
    toks = []
    for t in W174.tokenize_v3(doc):
        if t['line'] not in plus:
            continue
        t['file'] = 'main.tex'
        t['status'] = None
        toks.append(t)
    n_split = split_coefficients(toks)
    for t in toks:
        if t['excluded']:
            t.update(status='excluded', category=f"structure: {t['excluded']}")
        elif t['type'] in W177A.IDENT_TYPES:
            t.update(status='excluded', category=f"identifier ({t['type']})")
    base_doc = W164.Doc('main.tex', '\n'.join(base_lines))
    removed_written = collections.Counter(t['written'] for t in W174.tokenize_v3(base_doc) if t['line'] in minus
                                          and not t['excluded'] and t['type'] not in W177A.IDENT_TYPES)
    added_written = collections.Counter(t['written'] for t in toks if t['status'] != 'excluded')
    decls = declarations()
    stale, claimed_twice, decl_out, dup_decl = apply_declarations(X, toks, decls, lines)
    unassigned = [t for t in toks if t['status'] is None]
    for t in toks:
        moved = removed_norm.get(norm(lines[t['line'] - 1]))
        t['line_class'] = (f'moved (identical text at 8f051a2 l. {moved[0]})' if moved else 'new or changed line')
        t['written_also_on_removed_lines'] = t['written'] in removed_written
        t['rating'] = RATING.get(t['status'], 'UNRATED')
        t['convention'] = convention(t)
        t['context'] = CTX.get(t['line'])
    controls = negative_control(X, toks, lines)
    stm = statement_checks(X)
    by_status = dict(collections.Counter(t['status'] for t in toks))
    by_rating = dict(collections.Counter(t['rating'] for t in toks))
    not_rated = dict(collections.Counter(t.get('category') or 'excluded' for t in toks if t['rating'] == 'NOT RATED'))
    by_class = {c: dict(collections.Counter(t['rating'] for t in toks if t['line_class'].startswith(c)))
                for c in ('moved', 'new')}
    removed_only = sorted((w, n) for w, n in removed_written.items() if w not in added_written)
    monetary_rated = [t for t in toks if t['rating'] in ('MATCH', 'DERIVED', 'MISMATCH', 'UNSOURCED')
                      and (t.get('unit') in ('kEUR', 'MEUR', 'EUR') or (t.get('spec') or {}).get('kind') in ('k', 'meur'))]
    checks = {
        'clone_head_is_declared': head == CLONE_HEAD,
        'every_clone_tex_sha_equals_declared_and_expect': all(TEX_SHA.get(f) == tex_sha[f] == expect.get(f) for f in tex_now),
        'every_clone_tex_and_bib_equals_its_HEAD_blob': all(blob_eq.values()),
        'clone_tex_and_bib_clean': porcelain == '',
        'main_tex_line_count': len(lines) - (1 if main_text.endswith('\n') else 0) == MAIN_LINES_N,
        'hunk_counts_as_declared': hunks == EXPECTED_HUNKS,
        'plus_lines_equal_under_both_contexts': plus == plus_dd,
        'base_and_w181_heads_readable': len(base_lines) > 100 and len(w181_lines) > 100,
        'e21_rows_unchanged_since_w181a_head': X.e21_rows_unchanged,
        'frozen_json_unchanged': _sha(FZ_REL) == FZ_SHA,
        'every_declaration_found_and_evaluated': not stale,
        'no_line_declared_twice': not dup_decl,
        'no_token_claimed_twice': not claimed_twice,
        'every_token_assigned': not unassigned,
        'statuses_in_vocabulary': all(t['status'] in STATUSES for t in toks),
        'every_input_committed_clean': all(v['committed_clean'] for v in inputs.values()),
        'negative_control_every_alteration_refused': bool(controls) and all(c['refused'] for c in controls),
        'every_monetary_rated_token_has_a_convention': all(t.get('convention') not in (None, 'n/a') for t in monetary_rated),
        'every_moved_line_with_rated_tokens_has_context': all(
            t['context'] for t in toks if t['line_class'].startswith('moved') and t['rating'] in ('MATCH', 'DERIVED',
                                                                                                  'MISMATCH', 'UNSOURCED')),
    }
    failed = sorted(k for k, v in checks.items() if v is not True)
    guards, pk, guards_ok = guards_state()
    code = 0 if not failed else 3
    if not guards_ok:
        code = 1
    for t in toks:
        for k in ('col0', 'end0'):
            t.pop(k, None)
        s = t.pop('spec', None)
        if s is not None:
            t['declared_kind'] = s.get('kind')
            t['declared_desc'] = s.get('desc') or s.get('why')
    mism = [t for t in toks if t['rating'] in ('MISMATCH', 'UNSOURCED')]
    out = {'schema': 'p515_s53_w182b_number_check', 'version': 1,
           'stage': 'P5.15 Addendum 82 v2, W182 item A2 -- every number changed or moved in the 49 hunks 8f051a2 -> edd7a63',
           'utc': datetime.now(timezone.utc).isoformat(), 'git_head': _git('rev-parse', 'HEAD'),
           'script': {'path': SCRIPT_REL, 'sha256': _sha(SCRIPT_REL), 'committed_clean': script_clean,
                      'imports': {os.path.basename(W178B.__file__): _sha(os.path.basename(W178B.__file__))}},
           'command_line': sys.argv, 'this_launch_log': args.launch_log,
           'earlier_launch_logs_kept': {f: _sha(os.path.join(out_dir, f)) for f in refused_logs},
           'clone': {'dir': CLONE, 'head': head, 'declared_commit': DECLARED_COMMIT, 'base_commit': BASE_COMMIT,
                     'tex_sha256': tex_sha, 'bib_sha256': _sha(os.path.join(CLONE, BIB)), 'tex_and_bib_equal_head_blob': blob_eq,
                     'porcelain_tex_and_bib': porcelain, 'edited': False, 'main_tex_lines': MAIN_LINES_N},
           'hunks': hunks, 'n_plus_lines': len(plus), 'n_minus_lines': len(minus),
           'objective_convention': X.fz.get('objective_convention'), 'frozen_tables': {'path': FZ_REL, 'sha256': FZ_SHA},
           'inputs': inputs, 'declarations_version': 'W182b (new, for edd7a63)', 'coefficient_tokens_split': n_split,
           'summary': {'n_tokens': len(toks), 'n_added_lines': len(plus), 'by_status': by_status, 'by_rating': by_rating,
                       'not_rated_by_category': not_rated, 'by_line_class': by_class,
                       'n_non_match': len(mism), 'n_declarations': len(decls),
                       'negative_control': {'n': len(controls), 'n_refused': sum(c['refused'] for c in controls)},
                       'removed_only_written_values': removed_only},
           'rating_map': RATING,
           'conventions': {'column': '1-based', 'line': '1-based, main.tex at edd7a63',
                           'statuses': "the checker's: match / approximate / derived / MISMATCH / no counterpart / "
                                       'unchecked (category) / excluded; evaluated now, never carried'},
           'contexts': {str(k): v for k, v in sorted(CTX.items())},
           'statement_checks': stm, 'negative_control': controls,
           'declarations_applied': decl_out, 'stale_declarations': stale, 'lines_declared_twice': dup_decl,
           'claimed_twice': claimed_twice, 'unassigned': [{'line': t['line'], 'written': t['written']} for t in unassigned],
           'checks': checks, 'failed': failed, 'tokens': toks, 'guards': guards, 'pickle_guard': pk, 'guards_ok': guards_ok,
           'exit_code': code, 'wall_s': time.time() - t0}

    def _clean(o):
        if isinstance(o, dict):
            return {str(k): _clean(v) for k, v in o.items() if not callable(v)}
        if isinstance(o, (list, tuple)):
            return [_clean(v) for v in o]
        if isinstance(o, float) and o != o:
            return 'NaN'
        if isinstance(o, set):
            return sorted(o)
        return o
    out = _clean(out)
    written = {}

    def wr(name, data):
        rel = os.path.join(out_dir, OUT_NAMES[name])
        full = rel if os.path.isabs(rel) else os.path.join(REPO, rel)
        with open(full, 'xb') as h:
            h.write(data if isinstance(data, bytes) else data.encode('utf-8'))
        with open(full, 'rb') as h:
            written[rel] = _sha_bytes(h.read())
    wr('json', GRIO.dumps(out, indent=1, sort_keys=True) + '\n')
    wr('md', md_summary(out))
    man = dict(written)
    man[SCRIPT_REL] = _sha(SCRIPT_REL)
    for v in inputs.values():
        man[v['path']] = v['sha256']
    for f in tex_now + [BIB]:
        man[os.path.join(CLONE, f)] = _sha(os.path.join(CLONE, f))
    wr('man', GRIO.dumps(man, indent=1, sort_keys=True) + '\n')
    # ---- log ------------------------------------------------------------------------------------------------------------
    _log(f"[{tag}] hunks {hunks}; '+' lines {len(plus)}; tokens {len(toks)}; coefficient tokens split {n_split}")
    _log(f'[{tag}] A2 ratings {by_rating}; statuses {by_status}')
    _log(f'[{tag}] not rated by category {not_rated}')
    for t in toks:
        if t['rating'] in ('MISMATCH', 'UNSOURCED') or t['status'] == 'approximate':
            _log(f"[{tag}]   l. {t['line']}:{t['col']} {t['rating']} ({t['status']}): {t['written']!r} -> "
                 f"{t.get('value_at_written_precision')!r} [{str(t.get('counterpart'))[:220]}] {str(t.get('note'))[:160]}")
    for x in stm:
        _log(f"[{tag}] STATEMENT l. {x['line']}: {x['rating']} -- {x['text'][:120]}")
    _log(f"[{tag}] negative control: {sum(c['refused'] for c in controls)} of {len(controls)} alterations refused; not "
         f"refused: {[(c['line'], c['written'], c['altered']) for c in controls if not c['refused']]}")
    for k, v in checks.items():
        _log(f'[{tag}] check {k}: {v}')
    if stale or claimed_twice or unassigned or dup_decl:
        _log(f'[{tag}] stale {stale}; claimed twice {claimed_twice}; declared twice {dup_decl}; '
             f"unassigned {[(t['line'], t['written']) for t in unassigned]}")
    if failed:
        _log(f'[{tag}] FAILED: {failed}')
    _log(f"[{tag}] wrote {', '.join(f'{k} {v[:8]}' for k, v in written.items())}")
    _log(f"[{tag}] guards {[(k, v['counts']['permitted_solve'], v['counts']['blocked_solve'], v['verify_0_failures']) for k, v in guards.items()]}; "
         f"pickle ok {pk['ok']}; exit {code}; wall {time.time() - t0:.1f} s")
    for _n, g in reversed(GUARDS):
        g.uninstall()
    pickle.load, pickle.loads = _PICKLE_ORIG
    sys.exit(code)


if __name__ == '__main__':
    main()
