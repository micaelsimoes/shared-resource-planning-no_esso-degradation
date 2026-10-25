"""
P5.15 Addendum 82 v2, W182 order item A1 (Planner task W182a) -- the four new equations of main.tex at Overleaf
edd7a63 (`eq:acceptance_tolerance`, `eq:interface_deviation`, `eq:interface_deviation_nonnegative`, `eq:soh_initial`)
and the appendix `app:salvage_credit`, audited against the code in the W166 form (item, printed equation, code line,
MATCH / M / L / H). "The presence of an equation is not evidence that the implementation changed -- rate what the code
does." ZERO SOLVES, NO MODEL BUILD, NO PICKLE, NO PRODUCTION IMPORT: main.tex is read with `git -C <clone> show
edd7a63:main.tex`; production source is read as text at the code head and parsed with `ast`, never executed; the
committed evaluation records are read with `json`; the shared-ESS cost workbook is read with openpyxl (read-only).

Per item:
  printed  -- the equation body located by its \\label at edd7a63 (line numbers), compared with the declared text;
  code     -- every cited code line verified by EXACT TEXT at the declared file:line of the code head, with its
              enclosing def (ast);
  symbols  -- a symbol-by-symbol correspondence, computed (not asserted): the printed body and the code expression are
              both parsed into a canonical form (signed terms / factors) and compared;
  records  -- where a committed record can test the reading numerically (eq:soh_initial: the recorded SoH chain of every
              committed evaluation record of the four certification / 3 x 3 campaign roots; app:salvage_credit: the
              printed formula recomputed from each record's ageing cells and the case inputs, against the recorded
              terminal salvage value);
  scope    -- L findings are scoped text scans of main.tex (what was searched is recorded);
  history  -- `git log -S` at the code head on the implementing line (first and every changing commit), and identity of
              the implementing files between every campaign run head and the code head;
  rating   -- computed: MISMATCH if any correspondence / anchor / record check fails (a refusal: the M / H judgement is
              then the report's); otherwise L if a scoped wording finding is confirmed; otherwise MATCH.
  For eq:acceptance_tolerance also: the planning-search code that ran (its run heads), and whether tau exists there.

Negative controls (in memory, each must be refused): an altered code line number, a sign swap in the deviation row,
a negated nonnegativity domain, an altered initial SoH literal, an altered tau denominator and delta_R literal, a
printed deviation equation with its sign swapped, a printed salvage discount exponent y^End - y^Inv, and a code
salvage product missing its remaining-life factor.

Guards: SolveProfileGuard(permitted=()) installed before any project import and verified at exactly 0; pickle.load /
loads blocked and counted; no production module imported (asserted over sys.modules); JSON written only through
gate_result_io. Refuses to run unless the output directory exists and holds nothing but launch.log, and when another
copy of this script is running. Read-only on the clone (no fetch / pull / checkout; `git show` and `rev-parse` only).

Command (repo root, canonical interpreter, attached, both streams captured):
  mkdir data/SRP1/Results/P515S53/w182a_equations_code && \\
  /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s53_w182a_equations_code.py \\
      > data/SRP1/Results/P515S53/w182a_equations_code/launch.log 2>&1
Optional: --out-dir <dir> (scratch trials only).
"""
import pickle
import sys
import os

REPO = os.path.dirname(os.path.abspath(__file__))
if REPO not in sys.path:
    sys.path.insert(0, REPO)
os.chdir(REPO)

from p513_solve_profile_guard import SolveProfileGuard  # noqa: E402

GUARD = SolveProfileGuard(permitted=(), label='P5.15 W182a equations vs code (never solves)').install()
PICKLE_COUNTS = {'load': 0, 'loads': 0}


def _blocked_load(*_a, **_k):
    PICKLE_COUNTS['load'] += 1
    raise RuntimeError('W182a: pickle.load called -- no model loads are permitted')


def _blocked_loads(*_a, **_k):
    PICKLE_COUNTS['loads'] += 1
    raise RuntimeError('W182a: pickle.loads called -- no model loads are permitted')


pickle.load, pickle.loads = _blocked_load, _blocked_loads

import argparse  # noqa: E402
import ast  # noqa: E402
import copy  # noqa: E402
import glob  # noqa: E402
import hashlib  # noqa: E402
import io  # noqa: E402
import json  # noqa: E402
import math  # noqa: E402
import re  # noqa: E402
import subprocess  # noqa: E402
from datetime import datetime, timezone  # noqa: E402

import gate_result_io as GRIO  # noqa: E402  (stdlib only; the shared writer that refuses string flags)

SCRIPT = os.path.basename(os.path.abspath(__file__))
DEFAULT_OUT_DIR = os.path.join('data', 'SRP1', 'Results', 'P515S53', 'w182a_equations_code')

# ------------------------------------------------------------------------------------------------------- identity
CLONE = os.path.join('manuscript', '6a67305f25e8348fb71380c3')
DECLARED_HEAD = 'edd7a63a6c2ea06c1db17e3b6941bde8aa6d6838'
BASE_HEAD = '8f051a2a0cd8b37378d7407699e38c3bf20a71e7'       # the author's marked copy (Addendum 82 v2)
MAIN_SHA_PREFIX = 'b08a3365607e56f2'                          # the critical review's main.tex fingerprint
CODE_HEAD = 'a17f6ba7cf643be915493f1e3a3421d2f51a6f8b'      # repo HEAD = code head declared by the Planner
LAST_PRODUCTION_COMMIT = '575bfec43ff47b03ed9f6ef955025cd35cfb28a2'

FORBIDDEN_MODULE_PREFIXES = ('shared_resources_planning', 'network', 'shared_energy_storage', 'model_construction_helpers',
                             'admm_', 'helper_functions', 'definitions', 'settling_criterion', 'p515_s47', 'p515_s51',
                             'p515_s53_f2', 'p515_s44', 'p515_s53_w142', 'pyomo.environ')

# production / rule files read as text at CODE_HEAD
SRC_FILES = ('settling_criterion.py', 'settling_criterion_v2.py', 'settling_criterion_v4.py', 'settling_criterion_v5.py',
             'settling_criterion_v6.py', 'p515_s53_w142_resettle_v6_hooks.py', 'p515_s53_w101_srp1_continuation_campaign.py',
             'model_construction_helpers.py', 'network.py', 'shared_resources_planning.py', 'shared_energy_storage_data.py',
             'shared_energy_storage_parameters.py')
PRODUCTION_FILES = ('model_construction_helpers.py', 'network.py', 'network_data.py', 'shared_resources_planning.py',
                    'shared_energy_storage_data.py', 'shared_energy_storage_parameters.py', 'shared_energy_storage.py',
                    'admm_parameters.py', 'definitions.py', 'helper_functions.py')
RULE_FILES = ('settling_criterion.py', 'settling_criterion_v2.py', 'settling_criterion_v3.py', 'settling_criterion_v4.py',
              'settling_criterion_v5.py', 'settling_criterion_v6.py')
CASE_FILES = ('data/SRP1/SRP1.json', 'data/SRP1/SharedESS/SRP1_ESS.xlsx', 'data/SRP1/SharedESS/SRP1_ESS_Params.json')

# the planning searches (manuscript App. E.21 "The three planning searches"): spec / results and launcher
SEARCHES = {
    'baseline_variant_A_s47': {'results': 'data/SRP1/Results/P515S47/campaign_s47_phase_b/campaign_results.json',
                               'harness': 'p515_s47_phase_b_record.py', 'convention_file': 'p515_s47_phase_b_record.py',
                               'convention_line': "OBJECTIVE_CONVENTION = ('Q(x) = certified_cost = gross_operational_cost"},
    'm2_variant_A_s51': {'results': 'data/SRP1/Results/P515S51/campaign_s51_f2_phase_b/campaign_results.json',
                         'harness': 'p515_s51_f2_phase_b.py', 'convention_file': 'p515_s51_f2_phase_b.py',
                         'convention_line': "'Q(x) = certified_cost = gross_operational_cost (settlement-excluded) AT m = 2"},
    'm2_variant_B_s53': {'results': 'data/SRP1/Results/P515S53/campaign_s53_f2_certificate_r1/campaign_results.json',
                         'harness': 'p515_s53_f2_certificate.py', 'convention_file': 'p515_s51_f2_phase_b.py',
                         'convention_line': "'Q(x) = certified_cost = gross_operational_cost (settlement-excluded) AT m = 2"},
}
SEARCH_SRC_SCAN = ('p515_s47_phase_b_record.py', 'p515_s51_f2_phase_b.py', 'p515_s53_f2_certificate.py',
                   'p515_s44_campaign_harness.py')
TAU_TOKENS = re.compile(r'settling_criterion|\bTAU\b|\bDELTA_R\b|\bR_REF\b')

# evaluation-record roots (the certification campaigns behind the tables, W166, plus the 3 x 3 pair)
P53 = os.path.join('data', 'SRP1', 'Results', 'P515S53')
RECORD_ROOTS = (os.path.join(P53, 'w142_resettle_v6'), os.path.join(P53, 'w142_resettle_ext_v6'),
                os.path.join(P53, 'w155_a64_cells'), os.path.join(P53, 'w90_3x3'))
CERT_ROOTS = RECORD_ROOTS[:3]

# ------------------------------------------------------------------------------------------------------- printed
PRINTED = {
    'eq:acceptance_tolerance': r'\tau = \frac{\delta_R\mathcal{V}}{4}',
    'eq:interface_deviation': r'P^{\text{I}}_{i,s,t} - \bar{P}^{\text{I}}_{i,t} = d^{+}_{i,s,t} - d^{-}_{i,s,t}',
    'eq:interface_deviation_nonnegative': r'd^{\pm}_{i,s,t} \ge 0',
    'eq:soh_initial': r'SoH_{e,y^{\text{Inv}},(y^{\text{Inv}})^{-}} = 1',
    'eq:salvage_credit': (r'C^\text{Salvage}(\boldsymbol{x}) = \sum_{e \in E^S} \sum_{y^\text{Inv} \in Y} '
                          r'\frac{\bar{c}^{\text{Inv,E}}_{y^\text{Inv}}\, E^{\text{Inv}}_{e,y^\text{Inv}}}'
                          r'{(1+r)^{\,y^\text{End}-y_0}} \;\frac{T^\text{Rem}_{e,y^\text{Inv}}}{T^\text{Cal}_{e,y^\text{Inv}}} '
                          r'\;\frac{SoH_{e,y^\text{Inv},y^\text{End}} - SoH^{\text{Min}}_e}{1 - SoH^{\text{Min}}_e}'),
    'eq:salvage_remaining_life': (r'T^\text{Rem}_{e,y^\text{Inv}} = \max\left\{0,\; y^\text{Inv} + '
                                  r'T^\text{Cal}_{e,y^\text{Inv}} - y^\text{End}\right\}'),
}
# prose the ratings rest on (exact substrings at edd7a63)
PROSE = {
    'tau_units': 'The tolerance is expressed in the same monetary units as $Q(\\boldsymbol{x})$.',
    'tau_values': ('uses $\\delta_R = 0.07$ on the reference value $\\mathcal{V} = 259{,}375.33$~\\euro{}, the value of '
                   'the reference unit in force when the rule was adopted, hence $\\tau = 4{,}539.07$~\\euro{}, kept frozen '
                   'thereafter'),
    'tau_C3': 'cycles, all after $k_0$, is at most $\\tau$;',
    'tau_C4': '$|t_{\\text{sum}}| \\le \\tau/2$',
    'tau_floors': 'a step smaller than $\\tau/100$ carries no sign, a swing smaller than $\\tau/10$ is treated as noise',
    'tau_M3': '(M3) $|\\text{last step}| \\cdot 2P_{\\max} \\le \\tau$',
    'tau_det': 'A difference is classified as determinate when its magnitude is at least $\\max\\{3b,2\\tau\\}$.',
    'searches_earlier_code': ('The three planning searches used earlier versions of the implementation without the '
                              'tightened interior-point tolerances used for numerical acceptance.'),
    'search_sigma': 'and the measured resolution $\\sigma_Q = 18{,}449.66$~\\euro{} of the recourse.',
    'dev_tso_common': ("Within each representative year--day block, the TSO's interface schedule is common to all "
                       "scenarios, while DSO exchanges may vary. Each DSO commits to its probability-weighted expected "
                       "exchange. Omitting the representative-year and representative-day indices"),
    'dev_reactive': 'Reactive-power deviations incur an analogous charge with the same numerical coefficient',
    'dev_no_tso': 'No deviation premium is imposed on the TSO.',
    'dev_single': ('With a single scenario, the interface exchange equals its expected value and the deviation-charge term '
                   'is omitted.'),
    'dev_nomenclature': ('Voltage, active and reactive power at the interface of DN $i$'),
    'soh_boundary': ('Let $y^{-}$ denote the boundary immediately before the block represented by year $y$. The state at '
                     'this boundary is the terminal state of the preceding block, or the initial state at installation '
                     'for a new cohort:'),
    'soh_new_cohort': 'Each newly installed cohort starts with unit state of health:',
    'salv_cbar': ('is the expected undiscounted energy cost per MWh in the installation year, $y^\\text{End}$ the end of '
                  'the planning horizon'),
    'salv_no_converter': 'the converter is not credited.',
    'salv_reported': ('The credit is reported, not optimized: the recourse does not see it, no plan is selected for it, '
                      'and the net value is $Q(\\boldsymbol{0}) - Q(\\boldsymbol{x}) + C^\\text{Salvage}(\\boldsymbol{x})$.'),
    'salv_net_form': 'the net form subtracts the terminal salvage credit~\\eqref{eq:salvage_credit}',
    'salv_rate': 'Costs are discounted at an annual rate of 2\\%.',
    'salv_after': ('The credit is calculated after the operational evaluation and does not affect dispatch or investment '
                   'selection.'),
}
# the predecessor text at 8f051a2 (the items existed there in another form)
PREDECESSOR = {
    'eq:acceptance_tolerance': '$\\tau = \\delta_R \\mathcal{V}/4$, with $\\mathcal{V}$ a reference storage value',
    'eq:interface_deviation': ('P^{\\text{I}}_{i,s,t} - \\bar{P}^{\\text{I}}_{i,t} = d^{+}_{i,s,t} - d^{-}_{i,s,t}, '
                               '\\qquad d^{\\pm}_{i,s,t} \\ge 0'),
    'eq:interface_deviation_nonnegative': 'd^{\\pm}_{i,s,t} \\ge 0',
    'eq:soh_initial': 'SoH_{e,y^\\text{Inv},y^\\text{Inv}-1} = 1',
    'app:salvage_credit': '\\label{eq:salvage_credit}',
}

# ------------------------------------------------------------------------------------------------------- code anchors
# id: (file, line at CODE_HEAD, exact stripped text, enclosing def or None for module level)
CODE = {
    # ---- tau
    'SC1_DELTA_R': ('settling_criterion.py', 55, 'DELTA_R = 0.07', None),
    'SC1_R_REF': ('settling_criterion.py', 56, 'R_REF = 259375.33', None),
    'SC1_TAU': ('settling_criterion.py', 57, 'TAU = DELTA_R * R_REF / 4.0', None),
    'SC1_EPS0': ('settling_criterion.py', 58, 'EPS0 = TAU / 100.0', None),
    'SC2_TAU': ('settling_criterion_v2.py', 69, 'TAU = SC1.TAU', None),
    'SC2_GAP': ('settling_criterion_v2.py', 74, 'GAP_BOUND = TAU / 2.0', None),
    'SC2_C3': ('settling_criterion_v2.py', 172, "a_parts.update({'range': rng, 'range_le_tau': rng <= TAU, 'band': [mn, mx]})",
               'evaluate'),
    'SC2_M3': ('settling_criterion_v2.py', 182, "'last_step_times_L_le_tau': last * self.l_mono <= TAU, 'band': None}",
               'evaluate'),
    'SC2_M2': ('settling_criterion_v2.py', 189,
               "b_parts.update({'range': rng, 'range_le_tau': rng <= TAU, 'band': [mn, mx], 'max_abs_dq': max(steps),",
               'evaluate'),
    'SC2_C4': ('settling_criterion_v2.py', 277, 'gap_ok = t_sum is not None and abs(t_sum) <= GAP_BOUND', 'observe'),
    'SC4_TAU': ('settling_criterion_v4.py', 62, 'TAU = SC2.TAU', None),
    'SC5_TAU': ('settling_criterion_v5.py', 52, 'TAU = SC4.TAU', None),
    'SC6_TAU': ('settling_criterion_v6.py', 57, 'TAU = SC5.TAU', None),
    'SC6_GAP': ('settling_criterion_v6.py', 62, 'GAP_BOUND = SC5.GAP_BOUND', None),
    'SC6_FLOOR': ('settling_criterion_v6.py', 80, 'SWING_FLOOR = TAU / 10.0                 # "Swing noise floor tau/10 (~454 EUR), adopted"', None),
    'SC6_DET': ('settling_criterion_v6.py', 363, 'return max(DETERMINACY_BAR_FACTOR * max(bar_r, bar_o), DETERMINACY_TAU_MULTIPLE * TAU)',
                'determinacy_threshold'),
    'HOOKS_V6': ('p515_s53_w142_resettle_v6_hooks.py', 30, 'import settling_criterion_v6 as SC6', None),
    'V_SRP1': ('p515_s53_w101_srp1_continuation_campaign.py', 119,
               'V_SRP1 = 259375.32654094696    # recert campaign_results R.R_tail (R_ref = 259,375.33 rounded)', None),
    # ---- interface deviation
    'DEV_P_RULE_DEF': ('model_construction_helpers.py', 1861, 'def row18_interface_deviation_p_rule(m, s_m, s_o, p):',
                       'row18_interface_deviation_p_rule'),
    'DEV_P_LHS': ('model_construction_helpers.py', 1866, 'return (m.pg_adn[s_m, s_o, p] - m.expected_interface_pf_p[p]',
                  'row18_interface_deviation_p_rule'),
    'DEV_P_RHS': ('model_construction_helpers.py', 1867, '== m.row18_dev_p_up[s_m, s_o, p] - m.row18_dev_p_down[s_m, s_o, p])',
                  'row18_interface_deviation_p_rule'),
    'DEV_Q_LHS': ('model_construction_helpers.py', 1873, 'return (m.qg_adn[s_m, s_o, p] - m.expected_interface_pf_q[p]',
                  'row18_interface_deviation_q_rule'),
    'DEV_Q_RHS': ('model_construction_helpers.py', 1874, '== m.row18_dev_q_up[s_m, s_o, p] - m.row18_dev_q_down[s_m, s_o, p])',
                  'row18_interface_deviation_q_rule'),
    'DEV_P_ROWS': ('model_construction_helpers.py', 1971, 'model.row18_dev_p_def = pe.Constraint(model.scenarios_market, '
                   'model.scenarios_operation, model.periods, rule=row18_interface_deviation_p_rule)',
                   'add_scenario_commitment_terms'),
    'DEV_DSO_ONLY': ('model_construction_helpers.py', 1948, 'if (not network.is_transmission) and premium_alpha > 0.0:',
                     'add_scenario_commitment_terms'),
    'DEV_SINGLE': ('model_construction_helpers.py', 1917, 'if n_scenarios == 1:', 'add_scenario_commitment_terms'),
    'DEV_CHARGE': ('model_construction_helpers.py', 1979,
                   'charge += (probability * model.row18_alpha * model.row18_premium[p] * network.baseMVA',
                   'add_scenario_commitment_terms'),
    'DEV_CHARGE_TERMS': ('model_construction_helpers.py', 1980,
                         '* (model.row18_dev_p_up[s_m, s_o, p] + model.row18_dev_p_down[s_m, s_o, p]',
                         'add_scenario_commitment_terms'),
    'DEV_IN_OBJ': ('model_construction_helpers.py', 1836, 'obj += model.row18_deviation_charge', 'objective_function_rule'),
    'DSO_IFACE_P': ('model_construction_helpers.py', 1297, 'return m.pg[ref_gen_idx, s_m, s_o, p] - shared_ess_p',
                    'interface_pf_p_distribution_def'),
    'DSO_IFACE_ESS': ('model_construction_helpers.py', 1293, 'm.shared_es_pnet[e, s_m0, s_o0, p]', 'interface_pf_p_distribution_def'),
    'TSO_IFACE_LOAD': ('model_construction_helpers.py', 1247, 'pc_adn = m.pc[adn_load_idx, s_m, s_o, p]',
                       'interface_pf_p_transmission_def'),
    'TSO_IFACE_DELTA': ('model_construction_helpers.py', 1259, 'pc_adn += m.interface_delta_p[dn, s_m0, s_o0, p]',
                        'interface_pf_p_transmission_def'),
    'ESS_PNET_SIGN': ('model_construction_helpers.py', 999,
                      'return m.shared_es_pnet[e, s_m0, s_o0, p] == m.shared_es_pch[e, s_m0, s_o0, p] - m.shared_es_pdch[e, s_m0, s_o0, p]',
                      'sess_pnet_rule'),
    'EXPECTED_P_DEF': ('model_construction_helpers.py', 2420, 'm.pg_adn[s_m, s_o, p]', 'dn_interface_expected_pf_p_def'),
    'EXPECTED_P_RULE': ('model_construction_helpers.py', 2443,
                        'return m.expected_interface_pf_p[p] == dn_interface_expected_pf_p_def(m, p, network)',
                        'dn_interface_expected_pf_p_rule'),
    'PG_ADN_EXPR': ('network.py', 444, 'model.pg_adn = pe.Expression(model.scenarios_market, model.scenarios_operation, '
                    'model.periods, rule=partial(interface_pf_p_distribution_def, network=network))', '_build_model'),
    'TSO_ALPHA_0': ('shared_resources_planning.py', 4298, 'premium_alpha=0.0,', 'create_transmission_network_model'),
    # ---- nonnegativity
    'NN_P_UP': ('model_construction_helpers.py', 1966, 'model.row18_dev_p_up = pe.Var(model.scenarios_market, '
                'model.scenarios_operation, model.periods, domain=pe.NonNegativeReals, initialize=0.0)',
                'add_scenario_commitment_terms'),
    'NN_P_DOWN': ('model_construction_helpers.py', 1967, 'model.row18_dev_p_down = pe.Var(model.scenarios_market, '
                  'model.scenarios_operation, model.periods, domain=pe.NonNegativeReals, initialize=0.0)',
                  'add_scenario_commitment_terms'),
    'NN_Q_UP': ('model_construction_helpers.py', 1968, 'model.row18_dev_q_up = pe.Var(model.scenarios_market, '
                'model.scenarios_operation, model.periods, domain=pe.NonNegativeReals, initialize=0.0)',
                'add_scenario_commitment_terms'),
    'NN_Q_DOWN': ('model_construction_helpers.py', 1969, 'model.row18_dev_q_down = pe.Var(model.scenarios_market, '
                  'model.scenarios_operation, model.periods, domain=pe.NonNegativeReals, initialize=0.0)',
                  'add_scenario_commitment_terms'),
    # ---- SoH initial
    'SOH_LOOP': ('shared_energy_storage_data.py', 687, 'for y in range(y_inv, max_tcal_norm):', '_build_subproblem'),
    'SOH_INIT': ('shared_energy_storage_data.py', 714, 'prev_soh = 1.00', '_build_subproblem'),
    'SOH_IF': ('shared_energy_storage_data.py', 715, 'if y > y_inv:', '_build_subproblem'),
    'SOH_PREV': ('shared_energy_storage_data.py', 716, 'prev_soh = model.es_soh_per_unit_cumul[y_inv, y - 1]', '_build_subproblem'),
    'SOH_CHAIN': ('shared_energy_storage_data.py', 724,
                  '== prev_soh * pe.exp(-model.es_D_per_unit[y_inv, y]) * (phi_cal ** num_years))', '_build_subproblem'),
    'SOH_FIXED_1': ('shared_energy_storage_data.py', 550, 'model.es_soh_per_unit_cumul.fix(1.00)', '_build_subproblem'),
    'SOH_POINT_DEFAULT': ('shared_energy_storage_data.py', 17, "AVAILABLE_ENERGY_SOH_POINT_DEFAULT = 'end'", None),
    'SOH_MID_INIT': ('shared_energy_storage_data.py', 622, 'prev_soh = 1.00', '_build_subproblem'),
    # ---- salvage
    'SALV_DEF': ('shared_energy_storage_data.py', 866, 'def _build_terminal_salvage_value_expression(shared_ess_data, model, shared_ess_idx):',
                 '_build_terminal_salvage_value_expression'),
    'SALV_TERMINAL_IDX': ('shared_energy_storage_data.py', 871, 'terminal_year_idx = len(shared_ess_data.years) - 1',
                          '_build_terminal_salvage_value_expression'),
    'SALV_E_RATED': ('shared_energy_storage_data.py', 883, 'e_rated = model.es_e_rated_per_unit[y_inv, terminal_year_idx]',
                     '_build_terminal_salvage_value_expression'),
    'SALV_E_AVAIL': ('shared_energy_storage_data.py', 884, 'e_available = model.es_e_available_per_unit[y_inv, terminal_year_idx]',
                     '_build_terminal_salvage_value_expression'),
    'SALV_ABOVE_EOL': ('shared_energy_storage_data.py', 890,
                       'usable_energy_above_eol = (e_available - min_soh * e_rated) / (1.00 - min_soh)',
                       '_build_terminal_salvage_value_expression'),
    'SALV_SUM': ('shared_energy_storage_data.py', 899, 'salvage_value += (', '_build_terminal_salvage_value_expression'),
    'SALV_DISC_SUM': ('shared_energy_storage_data.py', 911,
                      'terminal_years = sum(float(num_years) for num_years in shared_ess_data.years.values())',
                      '_get_terminal_discount_factor'),
    'SALV_DISC': ('shared_energy_storage_data.py', 912,
                  'return 1.00 / ((1.00 + shared_ess_data.discount_factor) ** terminal_years)', '_get_terminal_discount_factor'),
    'SALV_AGE': ('shared_energy_storage_data.py', 919,
                 'age_at_terminal = sum(float(value) for value in represented_years[investment_year_idx:])',
                 '_get_remaining_calendar_life'),
    'SALV_REM': ('shared_energy_storage_data.py', 920,
                 'remaining_life = max(float(shared_energy_storage.t_cal) - age_at_terminal, 0.00)', '_get_remaining_calendar_life'),
    'SALV_REM_FRAC': ('shared_energy_storage_data.py', 921,
                      'remaining_life_fraction = min(remaining_life / float(shared_energy_storage.t_cal), 1.00)',
                      '_get_remaining_calendar_life'),
    'SALV_CBAR': ('shared_energy_storage_data.py', 928,
                  "expected_cost += probability * shared_ess_data.cost_investment['energy'][scenario][year]",
                  '_get_expected_energy_investment_cost'),
    'SALV_NOT_IN_OBJ': ('shared_energy_storage_data.py', 853, 'expr=model.feasibility_penalty,', '_build_subproblem'),
    'SALV_EXPR': ('shared_energy_storage_data.py', 849, 'model.salvage_value = pe.Expression(expr=salvage_value)', '_build_subproblem'),
    'SALV_GET': ('shared_energy_storage_data.py', 148, 'salvage_value += pe.value(models[node_id].salvage_value)', 'get_salvage_value'),
    'SALV_NET': ('shared_resources_planning.py', 1306, 'net_operational_recourse = gross_operational_cost - terminal_salvage_value',
                 '_get_operational_recourse_components'),
    'SALV_PARAM_RECOVERY': ('shared_energy_storage_parameters.py', 14, 'self.energy_recovery_fraction = 1.00', '__init__'),
    'SALV_PARAM_FLOOR': ('shared_energy_storage_parameters.py', 15, 'self.recycling_floor_fraction = 0.00', '__init__'),
    'SALV_CBAR_READ': ('shared_energy_storage_data.py', 1743,
                       "num_scenarios, shared_ess_data.prob_market_scenarios = _get_operational_scenarios_info_from_excel_file(filename, 'Scenarios')",
                       '_read_shared_energy_storage_data_from_file'),
}

# git log -S anchors per item: (file, exact string)
HISTORY = {
    'eq:acceptance_tolerance': [('settling_criterion.py', 'TAU = DELTA_R * R_REF / 4.0'),
                                ('settling_criterion_v2.py', 'GAP_BOUND = TAU / 2.0'),
                                ('settling_criterion_v6.py', 'SWING_FLOOR = TAU / 10.0')],
    'eq:interface_deviation': [('model_construction_helpers.py', 'm.pg_adn[s_m, s_o, p] - m.expected_interface_pf_p[p]'),
                               ('model_construction_helpers.py', 'return m.pg[ref_gen_idx, s_m, s_o, p] - shared_ess_p')],
    'eq:interface_deviation_nonnegative': [('model_construction_helpers.py', 'model.row18_dev_p_up = pe.Var(')],
    'eq:soh_initial': [('shared_energy_storage_data.py', 'prev_soh = 1.00')],
    'app:salvage_credit': [('shared_energy_storage_data.py', 'def _get_remaining_calendar_life'),
                           ('shared_energy_storage_data.py',
                            'usable_energy_above_eol = (e_available - min_soh * e_rated) / (1.00 - min_soh)'),
                           ('shared_energy_storage_data.py', 'model.salvage_value = pe.Expression(expr=salvage_value)'),
                           ('shared_resources_planning.py',
                            'net_operational_recourse = gross_operational_cost - terminal_salvage_value')],
}

INPUTS = {}
_TRACKED = None


# ------------------------------------------------------------------------------------------------------- helpers
def _sha_bytes(b):
    return hashlib.sha256(b).hexdigest()


def _sha(path):
    with open(path, 'rb') as f:
        return _sha_bytes(f.read())


def _git(args, cwd=REPO, binary=False, check=True):
    r = subprocess.run(['git'] + list(args), cwd=cwd, check=check, capture_output=True, text=not binary)
    return r if not check else r.stdout


def _tracked(path):
    global _TRACKED
    if _TRACKED is None:
        _TRACKED = set(_git(['ls-files', '-z']).split('\0'))
    rel = os.path.relpath(os.path.abspath(path), REPO)
    return 'committed' if rel in _TRACKED else 'UNCOMMITTED'


def _reg(path):
    rel = os.path.relpath(os.path.abspath(path), REPO)
    if rel not in INPUTS:
        INPUTS[rel] = {'sha256': _sha(path), 'status': _tracked(path)}
    return INPUTS[rel]


def _reg_blob(label, data, status):
    INPUTS[label] = {'sha256': _sha_bytes(data), 'status': status}


def _read_json(path):
    _reg(path)
    with open(path, encoding='utf-8') as f:
        return json.load(f)


def _ck(name, ok, detail=None):
    return {'check': name, 'ok': bool(ok), 'detail': detail}


def _norm(s):
    return re.sub(r'\s+', ' ', s).strip()


def _lines_with(lines, needle):
    return [i + 1 for i, ln in enumerate(lines) if needle in ln]


def _flat_hits(lines, needle):
    """Line numbers where `needle` occurs in the whitespace-normalised text (a needle may span lines)."""
    starts, buf, pos = [], [], 0
    for i, ln in enumerate(lines):
        starts.append(pos)
        buf.append(ln)
        pos += len(ln) + 1
    flat = '\n'.join(buf)
    norm_map, out_chars = [], []
    prev_space = False
    for idx, ch in enumerate(flat):
        if ch.isspace():
            if prev_space:
                continue
            out_chars.append(' ')
            norm_map.append(idx)
            prev_space = True
        else:
            out_chars.append(ch)
            norm_map.append(idx)
            prev_space = False
    text = ''.join(out_chars)
    hits = []
    n = _norm(needle)
    k = text.find(n)
    while k >= 0:
        raw = norm_map[k]
        line = max(i for i, s in enumerate(starts) if s <= raw) + 1
        hits.append(line)
        k = text.find(n, k + 1)
    return hits


# ------------------------------------------------------------------------------------------------------- identity
def identity():
    head = _git(['rev-parse', 'HEAD'], cwd=CLONE).strip()
    status = _git(['status', '--porcelain', '--', 'main.tex'], cwd=CLONE).strip()
    base_is_ancestor = _git(['merge-base', '--is-ancestor', BASE_HEAD, DECLARED_HEAD], cwd=CLONE, check=False).returncode == 0
    n_commits = int(_git(['rev-list', '--count', f'{BASE_HEAD}..{DECLARED_HEAD}'], cwd=CLONE).strip())
    main_b = _git(['show', f'{DECLARED_HEAD}:main.tex'], cwd=CLONE, binary=True)
    base_b = _git(['show', f'{BASE_HEAD}:main.tex'], cwd=CLONE, binary=True)
    _reg_blob(f'{CLONE}@{DECLARED_HEAD[:7]}:main.tex', main_b, 'overleaf clone git object (read only)')
    _reg_blob(f'{CLONE}@{BASE_HEAD[:7]}:main.tex', base_b, 'overleaf clone git object (read only)')
    with open(os.path.join(CLONE, 'main.tex'), 'rb') as fh:
        wt_b = fh.read()
    repo_head = _git(['rev-parse', 'HEAD']).strip()
    watched = list(PRODUCTION_FILES) + list(RULE_FILES) + list(SRC_FILES) + list(CASE_FILES)
    dirty = _git(['status', '--porcelain', '--untracked-files=no', '--'] + watched).strip()
    head_eq_code = _git(['diff', '--quiet', CODE_HEAD, repo_head, '--'] + watched, check=False).returncode == 0
    last_prod = _git(['log', '-1', '--format=%H', CODE_HEAD, '--'] + list(PRODUCTION_FILES) + list(RULE_FILES)).strip()
    checks = [
        _ck('clone HEAD = declared edd7a63', head == DECLARED_HEAD, head),
        _ck('8f051a2 is an ancestor of edd7a63', base_is_ancestor),
        _ck('main.tex unmodified in the clone working tree', status == '', status),
        _ck('main.tex working copy = edd7a63 blob', wt_b == main_b),
        _ck(f'main.tex sha256 starts {MAIN_SHA_PREFIX} (critical review fingerprint)', _sha_bytes(main_b).startswith(MAIN_SHA_PREFIX),
            _sha_bytes(main_b)),
        _ck('code head a17f6ba7 = repo HEAD, or HEAD carries no change to a production / rule / case file since it '
            '(parallel W182 Workers commit report files)', repo_head == CODE_HEAD or head_eq_code, repo_head),
        _ck('production, rule, source and case files clean in the repo working tree', dirty == '', dirty),
        _ck('last commit touching production / rule files at the code head = 575bfec4', last_prod == LAST_PRODUCTION_COMMIT,
            last_prod),
    ]
    return ({'clone': CLONE, 'clone_head': head, 'base': BASE_HEAD, 'n_commits_base_to_head': n_commits,
             'main_sha256': _sha_bytes(main_b), 'main_n_lines': main_b.decode('utf-8').count('\n'),
             'repo_head': repo_head, 'repo_dirty_tracked': dirty.splitlines(), 'last_production_commit': last_prod,
             'checks': checks, 'ok': all(c['ok'] for c in checks)},
            main_b.decode('utf-8').split('\n'), base_b.decode('utf-8').split('\n'))


def code_sources():
    srcs = {}
    for f in SRC_FILES:
        b = _git(['show', f'{CODE_HEAD}:{f}'], binary=True)
        _reg_blob(f'repo@{CODE_HEAD[:8]}:{f}', b, 'committed (git object)')
        with open(f, 'rb') as fh:
            wt = fh.read()
        if wt != b:
            raise SystemExit(f'refusing [code-head check]: working-tree {f} differs from {CODE_HEAD[:8]}')
        srcs[f] = b.decode('utf-8')
    return srcs


# ------------------------------------------------------------------------------------------------------- AST helpers
def _ancestors(tree, lineno):
    chain = []

    def visit(node):
        for child in ast.iter_child_nodes(node):
            if not isinstance(child, (ast.stmt, ast.ExceptHandler)):
                continue
            if child.lineno <= lineno <= getattr(child, 'end_lineno', child.lineno):
                chain.append(child)
                visit(child)
                return
    visit(tree)
    return chain


def anchor_checks(srcs, ids, code=CODE):
    trees = {}
    out, checks = {}, []
    for key in ids:
        fname, line, text, func = code[key]
        lines = srcs[fname].split('\n')
        got = lines[line - 1].strip() if 0 < line <= len(lines) else None
        if fname not in trees:
            try:
                trees[fname] = ast.parse(srcs[fname])
            except SyntaxError:
                trees[fname] = None
        funcs = [n.name for n in _ancestors(trees[fname], line) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))] \
            if trees[fname] is not None else []
        out[key] = {'file': fname, 'line': line, 'text_at_line': got, 'enclosing_def': funcs[-1] if funcs else None}
        checks.append(_ck(f'{key}: {fname}:{line} reads "{text[:70]}"', got == text, got))
        checks.append(_ck(f'{key}: enclosing def {func}', (funcs[-1] if funcs else None) == func, funcs))
    return checks, out


def _func_node(src, name):
    tree = ast.parse(src)
    return next((n for n in ast.walk(tree) if isinstance(n, ast.FunctionDef) and n.name == name), None)


def _signed_terms(node, sign=1):
    """Flatten a +/- expression into [(sign, name)] with name = the attribute / subscript base name."""
    if isinstance(node, ast.BinOp) and isinstance(node.op, (ast.Add, ast.Sub)):
        return _signed_terms(node.left, sign) + _signed_terms(node.right, sign if isinstance(node.op, ast.Add) else -sign)
    if isinstance(node, ast.Subscript):
        node = node.value
    if isinstance(node, ast.Attribute):
        return [(sign, node.attr)]
    if isinstance(node, ast.Name):
        return [(sign, node.id)]
    return [(sign, ast.dump(node))]


def _canonical(lhs, rhs, lead):
    """lhs - rhs = 0 as {symbol: coefficient}, scaled so that `lead` has coefficient +1."""
    coef = {}
    for s, n in lhs:
        coef[n] = coef.get(n, 0) + s
    for s, n in rhs:
        coef[n] = coef.get(n, 0) - s
    k = coef.get(lead)
    if not k:
        return None
    return {n: c * k for n, c in coef.items()}


def code_equation(src, func_name):
    fn = _func_node(src, func_name)
    ret = next((n for n in ast.walk(fn) if isinstance(n, ast.Return)), None) if fn else None
    cmp_ = ret.value if ret is not None else None
    if not (isinstance(cmp_, ast.Compare) and len(cmp_.ops) == 1 and isinstance(cmp_.ops[0], ast.Eq)):
        return None
    return _signed_terms(cmp_.left), _signed_terms(cmp_.comparators[0])


def printed_equation(body, token_map):
    """'A - B = C - D' (LaTeX) -> ([(+1,a),(-1,b)], [(+1,c),(-1,d)]) through token_map."""
    if body.count('=') != 1:
        return None

    def side(s):
        terms = []
        s = s.strip()
        sign = 1
        for part in re.split(r'\s+(?=[+-]\s)', s):
            part = part.strip()
            if part.startswith('- '):
                sign, part = -1, part[2:].strip()
            elif part.startswith('+ '):
                sign, part = 1, part[2:].strip()
            else:
                sign = 1
            terms.append((sign, token_map.get(part, 'UNMAPPED:' + part)))
        return terms
    lhs, rhs = body.split('=')
    return side(lhs), side(rhs)


# ------------------------------------------------------------------------------------------------------- printed
def locate_label(lines, label):
    lab = [i + 1 for i, ln in enumerate(lines) if f'\\label{{{label}}}' in ln]
    if len(lab) != 1:
        return {'label': label, 'label_lines': lab, 'ok': False}
    ln = lab[0]
    if label.startswith('app:'):
        return {'label': label, 'label_line': ln, 'ok': True}
    begs = [i for i in range(1, ln) if '\\begin{equation}' in lines[i - 1]]
    if not begs:
        return {'label': label, 'label_line': ln, 'ok': False}
    beg = begs[-1]
    end = min(i for i in range(ln, len(lines) + 1) if '\\end{equation}' in lines[i - 1])
    body_lines = [lines[i - 1] for i in range(beg + 1, ln)]
    body = _norm(' '.join(body_lines)).rstrip(',').strip()
    return {'label': label, 'label_line': ln, 'begin_line': beg, 'end_line': end, 'body': body, 'ok': True}


def printed_block(lines, printed=PRINTED):
    out = {}
    for label, expected in printed.items():
        loc = locate_label(lines, label)
        loc['expected'] = expected
        loc['matches_declared'] = loc.get('body') == _norm(expected)
        out[label] = loc
    loc = locate_label(lines, 'app:salvage_credit')
    out['app:salvage_credit'] = loc
    return out


def prose_block(lines):
    return {k: _flat_hits(lines, v) for k, v in PROSE.items()}


# ------------------------------------------------------------------------------------------------------- item A: tau
def item_tau(srcs, pr, prose, main_lines, search_info, heads_identity, code=CODE):
    ids = ['SC1_DELTA_R', 'SC1_R_REF', 'SC1_TAU', 'SC1_EPS0', 'SC2_TAU', 'SC2_GAP', 'SC2_C3', 'SC2_M3', 'SC2_M2', 'SC2_C4',
           'SC4_TAU', 'SC5_TAU', 'SC6_TAU', 'SC6_GAP', 'SC6_FLOOR', 'SC6_DET', 'HOOKS_V6', 'V_SRP1']
    checks, anchors = anchor_checks(srcs, ids, code)
    p = pr['eq:acceptance_tolerance']
    checks.append(_ck('printed body at its label = declared', p.get('matches_declared'), p.get('body')))
    m = re.fullmatch(r'\\tau = \\frac\{\\delta_R\s*\\mathcal\{V\}\}\{(\d+(?:\.\d+)?)\}', p.get('body') or '')
    printed_num = {'delta_R', 'V'} if m else None
    printed_den = float(m.group(1)) if m else None
    # code: TAU = DELTA_R * R_REF / 4.0 parsed from the anchored line
    tau_line = srcs['settling_criterion.py'].split('\n')[code['SC1_TAU'][1] - 1].strip()
    code_num, code_den = None, None
    try:
        node = ast.parse(tau_line).body[0].value
        if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Div) and isinstance(node.right, ast.Constant):
            code_den = float(node.right.value)
            num = node.left
            if isinstance(num, ast.BinOp) and isinstance(num.op, ast.Mult):
                code_num = {getattr(num.left, 'id', None), getattr(num.right, 'id', None)}
    except SyntaxError:
        pass
    sym = {'delta_R': 'DELTA_R', 'V': 'R_REF'}
    checks.append(_ck('symbols: numerator {delta_R, V} <-> {DELTA_R, R_REF}',
                      printed_num is not None and code_num == {sym[s] for s in printed_num}, [printed_num, code_num]))
    checks.append(_ck('symbols: denominator 4 <-> 4.0', printed_den is not None and printed_den == code_den,
                      [printed_den, code_den]))

    def lit(key):
        ln = srcs['settling_criterion.py'].split('\n')[code[key][1] - 1].strip()
        try:
            return ast.literal_eval(ast.parse(ln).body[0].value)
        except Exception:
            return None
    d_r, r_ref = lit('SC1_DELTA_R'), lit('SC1_R_REF')
    tau_code = d_r * r_ref / code_den if (d_r is not None and r_ref is not None and code_den) else None
    pv = prose['tau_values']
    checks.append(_ck('values: printed delta_R = 0.07, V = 259,375.33, tau = 4,539.07 at l. %s' % pv,
                      len(pv) == 1 and d_r == 0.07 and r_ref == 259375.33 and tau_code is not None
                      and round(tau_code, 2) == 4539.07, {'delta_R': d_r, 'R_REF': r_ref, 'tau_code': tau_code}))
    checks.append(_ck('V = R_REF is V_SRP1 = 259,375.32654... rounded to the cent (W101 constant)',
                      abs(259375.32654094696 - r_ref) < 0.005 if r_ref else False, r_ref))
    # where tau enters: each printed use has its code line (anchors above) and its prose
    for k in ('tau_units', 'tau_C3', 'tau_C4', 'tau_floors', 'tau_M3', 'tau_det'):
        checks.append(_ck(f'prose {k} present once', len(prose[k]) == 1, prose[k]))
    # tau is not in any production module (it enters only the stop decision, through the hooks)
    prod_hits = {}
    for f in PRODUCTION_FILES:
        b = _git(['show', f'{CODE_HEAD}:{f}'], binary=True)
        _reg_blob(f'repo@{CODE_HEAD[:8]}:{f}', b, 'committed (git object)')
        prod_hits[f] = len(TAU_TOKENS.findall(b.decode('utf-8')))
    checks.append(_ck('tau / settling_criterion absent from every production module (outside Q and the model)',
                      all(v == 0 for v in prod_hits.values()), prod_hits))
    # the planning searches: run heads predate the rule; no tau in their code; resolution sigma_Q
    checks.append(_ck('search run heads: settling_criterion*.py absent, no tau token in the search code, '
                      'ab0bcbc9 (tau introduced) not an ancestor, resolution rule = max(bar_x + bar_inc, sigma_q)',
                      search_info['all_ok'], {k: v['summary'] for k, v in search_info['runs'].items()}))
    checks.append(_ck('prose: searches used earlier code; search resolution sigma_Q',
                      len(prose['searches_earlier_code']) == 1 and len(prose['search_sigma']) == 1,
                      [prose['searches_earlier_code'], prose['search_sigma']]))
    checks.append(_ck('certification run heads: settling_criterion*.py identical to the code head',
                      heads_identity['cert_rule_files_identical_all'], heads_identity['cert_rule_files_differ']))
    l_findings = []
    return {'item': 'eq:acceptance_tolerance', 'printed': p, 'anchors': anchors, 'checks': checks,
            'l_findings': l_findings, 'production_tau_hits': prod_hits,
            'values': {'delta_R': d_r, 'R_REF': r_ref, 'tau_code': tau_code, 'eps0': tau_code / 100 if tau_code else None,
                       'gap_bound': tau_code / 2 if tau_code else None, 'swing_floor': tau_code / 10 if tau_code else None}}


def search_runs():
    runs, ok = {}, True
    tau_intro = 'ab0bcbc9d099cea26883b0817b76ea8d2f67de78'
    for name, s in SEARCHES.items():
        cr = _read_json(s['results'])
        head = cr.get('git_head_at_run')
        tree = _git(['ls-tree', '--name-only', head]).split('\n')
        rule_present = sorted(f for f in tree if f.startswith('settling_criterion'))
        tok = {}
        for f in SEARCH_SRC_SCAN:
            r = _git(['show', f'{head}:{f}'], binary=True, check=False)
            if r.returncode != 0:
                tok[f] = None
                continue
            _reg_blob(f'repo@{head[:8]}:{f}', r.stdout, 'committed (git object)')
            tok[f] = len(TAU_TOKENS.findall(r.stdout.decode('utf-8')))
        pb = _git(['show', f'{head}:p515_s47_phase_b_record.py'])
        res_rule = 'return max(bar_x + bar_inc, sigma_q)' in pb
        conv_src = _git(['show', f'{head}:{s["convention_file"]}'])
        conv_ok = s['convention_line'] in conv_src
        if s['convention_file'] != s['harness']:
            conv_ok = conv_ok and 'OBJECTIVE_CONVENTION = F51.OBJECTIVE_CONVENTION' in _git(['show', f'{head}:{s["harness"]}'])
        salvage_excluded = conv_ok and ('excluded from F' in conv_src or 'EXCLUDED from F' in conv_src)
        anc = _git(['merge-base', '--is-ancestor', tau_intro, head], check=False).returncode == 0
        date = _git(['log', '-1', '--format=%aI', head]).strip()
        harness_present = tok.get(s['harness']) is not None
        run_ok = (not rule_present) and harness_present and all(v in (0, None) for v in tok.values()) and res_rule \
            and not anc
        ok = ok and run_ok
        runs[name] = {'results': s['results'], 'git_head_at_run': head, 'date': date, 'harness': s['harness'],
                      'settling_criterion_files_at_head': rule_present, 'tau_token_counts_at_head': tok,
                      'resolution_rule_sigma_q_present': res_rule, 'objective_convention_file': s['convention_file'],
                      'salvage_excluded_from_F_text': salvage_excluded,
                      'tau_intro_ab0bcbc9_is_ancestor': anc,
                      'summary': {'head': head[:8], 'rule_files': len(rule_present), 'tau_tokens': sum(v or 0 for v in tok.values()),
                                  'ok': run_ok}}
    return {'runs': runs, 'all_ok': ok, 'tau_introduced_in': tau_intro}


# ------------------------------------------------------------------------------------------------------- item B/C: deviation
DEV_TOKENS = {r'P^{\text{I}}_{i,s,t}': 'flow', r'\bar{P}^{\text{I}}_{i,t}': 'expectation',
              r'd^{+}_{i,s,t}': 'dev_up', r'd^{-}_{i,s,t}': 'dev_down'}
DEV_CODE_MAP_P = {'pg_adn': 'flow', 'expected_interface_pf_p': 'expectation', 'row18_dev_p_up': 'dev_up',
                  'row18_dev_p_down': 'dev_down'}
DEV_CODE_MAP_Q = {'qg_adn': 'flow', 'expected_interface_pf_q': 'expectation', 'row18_dev_q_up': 'dev_up',
                  'row18_dev_q_down': 'dev_down'}
ORIENTATION_WORDS = ('positive', 'import', 'export', 'withdraw', 'injection', 'inject', 'direction', 'sign convention',
                     'from the TN', 'from TN', 'TN-to-DN', 'TN--DN direction', 'into the DN', 'flowing')
EXCLUSION_WORDS = ('exclud', 'net of', 'excluding', 'without the storage', 'storage excluded')


def _scan_window(main_lines, lo, hi, words):
    hits = []
    for i in range(lo, hi + 1):
        ln = main_lines[i - 1]
        for w in words:
            if w.lower() in ln.lower():
                hits.append({'line': i, 'word': w, 'text': ln.strip()[:160]})
    return hits


def item_deviation(srcs, pr, prose, main_lines, code=CODE):
    ids = ['DEV_P_RULE_DEF', 'DEV_P_LHS', 'DEV_P_RHS', 'DEV_Q_LHS', 'DEV_Q_RHS', 'DEV_P_ROWS', 'DEV_DSO_ONLY', 'DEV_SINGLE',
           'DEV_CHARGE', 'DEV_CHARGE_TERMS', 'DEV_IN_OBJ', 'DSO_IFACE_P', 'DSO_IFACE_ESS', 'TSO_IFACE_LOAD', 'TSO_IFACE_DELTA',
           'ESS_PNET_SIGN', 'EXPECTED_P_DEF', 'EXPECTED_P_RULE', 'PG_ADN_EXPR', 'TSO_ALPHA_0']
    checks, anchors = anchor_checks(srcs, ids, code)
    p = pr['eq:interface_deviation']
    checks.append(_ck('printed body at its label = declared', p.get('matches_declared'), p.get('body')))
    pe_ = printed_equation(p.get('body') or '', DEV_TOKENS)
    printed_canon = _canonical(*pe_, 'flow') if pe_ else None
    mch = srcs['model_construction_helpers.py']
    ce_p = code_equation(mch, 'row18_interface_deviation_p_rule')
    ce_q = code_equation(mch, 'row18_interface_deviation_q_rule')
    code_canon_p = _canonical([(s, DEV_CODE_MAP_P.get(n, n)) for s, n in ce_p[0]],
                              [(s, DEV_CODE_MAP_P.get(n, n)) for s, n in ce_p[1]], 'flow') if ce_p else None
    code_canon_q = _canonical([(s, DEV_CODE_MAP_Q.get(n, n)) for s, n in ce_q[0]],
                              [(s, DEV_CODE_MAP_Q.get(n, n)) for s, n in ce_q[1]], 'flow') if ce_q else None
    checks.append(_ck('symbols / sign: printed P - Pbar - d+ + d- = 0 equals the code P row (pg_adn - expected - up + down)',
                      printed_canon is not None and printed_canon == code_canon_p, [printed_canon, code_canon_p]))
    checks.append(_ck('the Q row has the same signed form (prose: "an analogous charge")',
                      code_canon_q is not None and code_canon_q == code_canon_p and len(prose['dev_reactive']) == 1,
                      code_canon_q))
    # indices: code (s_m, s_o, p) <-> printed (s, t); year/day omitted in print (prose)
    fn = _func_node(mch, 'row18_interface_deviation_p_rule')
    args = [a.arg for a in fn.args.args] if fn else None
    checks.append(_ck('indices: rule over (s_m, s_o, p) = (s, t) with s = (m, o); year/day omitted in the print (prose)',
                      args == ['m', 's_m', 's_o', 'p'] and len(prose['dev_tso_common']) == 1, args))
    # orientation of the two interface copies (both TN -> DN positive) and storage exclusion
    dso_ret = _func_node(mch, 'interface_pf_p_distribution_def')
    dso_terms = _signed_terms(next(n for n in ast.walk(dso_ret) if isinstance(n, ast.Return)).value) if dso_ret else None
    tso_fn = _func_node(mch, 'interface_pf_p_transmission_def')
    tso_seg = ast.get_source_segment(mch, tso_fn) if tso_fn else ''
    pnet = code_equation(mch, 'sess_pnet_rule')
    checks.append(_ck('DSO copy = reference-generator injection into the DN minus the co-located ESS net power '
                      '(pg - shared_ess_p): TN -> DN positive, ESS excluded',
                      dso_terms == [(1, 'pg'), (-1, 'shared_ess_p')], dso_terms))
    checks.append(_ck('ESS net power = pch - pdch (charging positive), the scenario-free copy (s_m0, s_o0)',
                      pnet == ([(1, 'shared_es_pnet')], [(1, 'shared_es_pch'), (-1, 'shared_es_pdch')]), pnet))
    checks.append(_ck('TSO copy = the ADN load at the TN node (pc) + the scenario-free interface delta: TN -> DN positive',
                      'pc_adn = m.pc[adn_load_idx, s_m, s_o, p]' in tso_seg
                      and 'pc_adn += m.interface_delta_p[dn, s_m0, s_o0, p]' in tso_seg, None))
    # Pbar is the block's own expectation of the same flow (eq:commitment_deviation)
    ex = _func_node(mch, 'dn_interface_expected_pf_p_def')
    ex_seg = ast.get_source_segment(mch, ex) if ex else ''
    checks.append(_ck('Pbar = sum_s prob_m prob_o pg_adn[s, t] (the expectation of the same flow)',
                      'network.prob_market_scenarios[s_m] *' in ex_seg and 'network.prob_operation_scenarios[s_o] *' in ex_seg
                      and 'm.pg_adn[s_m, s_o, p]' in ex_seg, None))
    # scope: DSO only, multi-scenario only, inside Q
    checks.append(_ck('prose: no premium on the TSO; omitted at one scenario', len(prose['dev_no_tso']) == 1
                      and len(prose['dev_single']) == 1, [prose['dev_no_tso'], prose['dev_single']]))
    # the charge is symmetric in d+ and d- (so the orientation cannot change a number)
    ch_seg = '\n'.join(mch.split('\n')[1978:1981])
    sym = ('model.row18_dev_p_up[s_m, s_o, p] + model.row18_dev_p_down[s_m, s_o, p]' in ch_seg
           and 'model.row18_dev_q_up[s_m, s_o, p] + model.row18_dev_q_down[s_m, s_o, p]' in ch_seg)
    checks.append(_ck('charge sums d+ + d- (symmetric): a reversed orientation of P^I swaps d+ and d- and changes no number',
                      sym, None))
    checks.append(_ck('the ESS term in P^I is the scenario-free copy, so it cancels in P - Pbar (sum_s omega_s = 1): '
                      'excluding or including it changes no deviation', "m.shared_es_pnet[e, s_m0, s_o0, p]" in mch, None))
    # L findings (scoped scans)
    nom = prose['dev_nomenclature']
    sub = [i + 1 for i, ln in enumerate(main_lines) if '\\label{subsubsec:commitment}' in ln]
    sub_end = (p.get('label_line') or 0) + 40
    win_lo = sub[0] - 2 if sub else 0
    orient_hits = _scan_window(main_lines, win_lo, sub_end, ORIENTATION_WORDS)
    orient_nom = _scan_window(main_lines, nom[0], nom[0], ORIENTATION_WORDS) if len(nom) == 1 else None
    pi_lines = [i + 1 for i, ln in enumerate(main_lines) if 'P^{\\text{I}}' in ln]
    orient_pi = [h for i in pi_lines for h in _scan_window(main_lines, i, i, ORIENTATION_WORDS)]
    excl_hits = _scan_window(main_lines, win_lo, sub_end, EXCLUSION_WORDS) + \
        [h for i in pi_lines for h in _scan_window(main_lines, i, i, EXCLUSION_WORDS)]
    # whole-document context (reported, not deciding): lines carrying an orientation word AND 'interface' or 'DN'
    context_whole = [h for h in _scan_window(main_lines, 1, len(main_lines), ORIENTATION_WORDS + ('reverse',))
                     if 'interface' in h['text'].lower() or ' dn' in h['text'].lower()]
    l_findings = [
        {'finding': 'orientation of P^I (TN -> DN positive in both code copies) is not stated in main.tex',
         'confirmed': len(nom) == 1 and not orient_hits and not orient_nom and not orient_pi,
         'scope': {'nomenclature_line': nom, 'commitment_subsection_lines': [win_lo, sub_end],
                   'every_line_with_P^{\\text{I}}': pi_lines, 'words': list(ORIENTATION_WORDS)},
         'hits': {'subsection': orient_hits, 'nomenclature': orient_nom, 'P_I_lines': orient_pi},
         'context_whole_document_not_deciding': context_whole},
        {'finding': ('P^I in the code is the DN exchange net of the co-located shared-ESS schedule (pg - pnet_ESS); '
                     'main.tex does not say whether P^I includes the shared ESS at the interface node'),
         'confirmed': not excl_hits,
         'scope': {'commitment_subsection_lines': [win_lo, sub_end], 'every_line_with_P^{\\text{I}}': pi_lines,
                   'words': list(EXCLUSION_WORDS)}, 'hits': excl_hits},
    ]
    notes = ['units: the code row is in per unit on network.baseMVA (the charge multiplies by baseMVA, l. 1979); the '
             'print carries no unit -- notation only',
             'row 18 is inactive (rows deactivated, d+/- fixed at 0) at the DSO initialisation solve and activated with '
             'the settlement weight before the ADMM cycles (shared_resources_planning._set_row18_inactive_for_initialisation '
             '/ _activate_row18_with_settlement); not stated in main.tex; no reported number is taken at initialisation']
    return {'item': 'eq:interface_deviation', 'printed': p, 'anchors': anchors, 'checks': checks,
            'canonical': {'printed': printed_canon, 'code_p': code_canon_p, 'code_q': code_canon_q},
            'l_findings': l_findings, 'notes': notes}


def item_nonneg(srcs, pr, prose, main_lines, code=CODE):
    ids = ['NN_P_UP', 'NN_P_DOWN', 'NN_Q_UP', 'NN_Q_DOWN']
    checks, anchors = anchor_checks(srcs, ids, code)
    p = pr['eq:interface_deviation_nonnegative']
    checks.append(_ck('printed body at its label = declared', p.get('matches_declared'), p.get('body')))
    ok_print = (p.get('body') or '') == r'd^{\pm}_{i,s,t} \ge 0'
    lines = srcs['model_construction_helpers.py'].split('\n')
    domains = {}
    for key in ids:
        ln = lines[code[key][1] - 1].strip()
        try:
            call = ast.parse(ln).body[0].value
            kw = {k.arg: k.value for k in call.keywords}
            dom = kw.get('domain')
            domains[key] = dom.attr if isinstance(dom, ast.Attribute) else None
            domains[key + '_bounds_kw'] = 'bounds' in kw
        except Exception:
            domains[key] = None
    checks.append(_ck('symbols: d^{+/-} >= 0 <-> domain=pe.NonNegativeReals on the four deviation Vars (P and Q), no '
                      'other bound', ok_print and all(domains.get(k) == 'NonNegativeReals' and not domains.get(k + '_bounds_kw')
                                                       for k in ids), domains))
    return {'item': 'eq:interface_deviation_nonnegative', 'printed': p, 'anchors': anchors, 'checks': checks,
            'l_findings': [], 'domains': domains,
            'notes': ['at the DSO initialisation solve the pair is fixed at 0 with its row deactivated (consistent with >= 0)']}


# ------------------------------------------------------------------------------------------------------- item D: SoH
def item_soh(srcs, pr, prose, records, code=CODE):
    ids = ['SOH_LOOP', 'SOH_INIT', 'SOH_IF', 'SOH_PREV', 'SOH_CHAIN', 'SOH_FIXED_1', 'SOH_POINT_DEFAULT', 'SOH_MID_INIT']
    checks, anchors = anchor_checks(srcs, ids, code)
    p = pr['eq:soh_initial']
    checks.append(_ck('printed body at its label = declared', p.get('matches_declared'), p.get('body')))
    src = srcs['shared_energy_storage_data.py']
    tree = ast.parse(src)
    chain = _ancestors(tree, code['SOH_INIT'][1])
    loop = next((n for n in reversed(chain) if isinstance(n, ast.For)), None)
    struct = {}
    if loop is not None:
        struct['loop_iter'] = ast.get_source_segment(src, loop.iter)
        body = loop.body
        idx_init = next((i for i, s in enumerate(body) if isinstance(s, ast.Assign) and s.lineno == code['SOH_INIT'][1]), None)
        init = body[idx_init] if idx_init is not None else None
        struct['init_value'] = init.value.value if init is not None and isinstance(init.value, ast.Constant) else None
        nxt = body[idx_init + 1] if idx_init is not None and idx_init + 1 < len(body) else None
        struct['next_is_if_y_gt_y_inv'] = isinstance(nxt, ast.If) and ast.get_source_segment(src, nxt.test) == 'y > y_inv'
        struct['if_body'] = ast.get_source_segment(src, nxt.body[0]) if struct['next_is_if_y_gt_y_inv'] else None
        uses = [n for n in ast.walk(loop) if isinstance(n, ast.Name) and n.id == 'prev_soh' and isinstance(n.ctx, ast.Load)]
        struct['prev_soh_loaded_at'] = [n.lineno for n in uses]
    checks.append(_ck('symbols: SoH at the boundary before the installation block = 1 <-> prev_soh = 1.00 for y = y_inv '
                      '(loop starts at y_inv; reassigned only if y > y_inv)',
                      struct.get('loop_iter') == 'range(y_inv, max_tcal_norm)' and struct.get('init_value') == 1.0
                      and struct.get('next_is_if_y_gt_y_inv')
                      and struct.get('if_body') == 'prev_soh = model.es_soh_per_unit_cumul[y_inv, y - 1]'
                      and code['SOH_CHAIN'][1] in struct.get('prev_soh_loaded_at', []), struct))
    checks.append(_ck('prose: boundary = preceding block terminal state, or initial state for a new cohort; new cohort '
                      'starts at unit SoH', len(prose['soh_boundary']) == 1 and len(prose['soh_new_cohort']) == 1,
                      [prose['soh_boundary'], prose['soh_new_cohort']]))
    # records: every committed evaluation record's SoH chain
    n_init, n_prop, bad = 0, 0, []
    chain_dev = 0.0
    for r in records:
        att = r['rec'].get('ageing_trajectory_terminal') or {}
        for node, nd in (att.get('nodes') or {}).items():
            cells = {(c['y_inv'], c['y']): c for c in nd.get('cells', [])}
            for (yi, y), c in cells.items():
                if y == yi:
                    n_init += 1
                    if c['soh_prev_end'] != 1.0:
                        bad.append({'record': r['path'], 'node': node, 'y_inv': yi, 'y': y, 'soh_prev_end': c['soh_prev_end']})
                else:
                    n_prop += 1
                    prev = cells.get((yi, y - 1))
                    if prev is None or c['soh_prev_end'] != prev['soh_end']:
                        bad.append({'record': r['path'], 'node': node, 'y_inv': yi, 'y': y, 'soh_prev_end': c['soh_prev_end'],
                                    'prev_soh_end': prev['soh_end'] if prev else None})
                pred = c['soh_prev_end'] * math.exp(-c['D']) * c['phi_cal_in_model'] ** c['n_years']
                chain_dev = max(chain_dev, abs(pred - c['soh_end']))
    checks.append(_ck(f'records: soh_prev_end == 1.0 exactly at every installation cell ({n_init}) and == the previous '
                      f'block end at every later cell ({n_prop})', n_init > 0 and not bad, bad[:10]))
    return {'item': 'eq:soh_initial', 'printed': p, 'anchors': anchors, 'checks': checks, 'l_findings': [],
            'structure': struct, 'records': {'n_installation_cells': n_init, 'n_propagation_cells': n_prop,
                                             'context_eq_soh_chain_max_abs_dev': chain_dev, 'violations': bad},
            'notes': ['context (eq:soh_chain, not an A1 item): recorded soh_end = soh_prev_end exp(-D) phi^n to '
                      f'{chain_dev:.3e} at every recorded cell',
                      "the 'mid' available-energy variant (cell e_c3_midblock) uses the same literal 1.00 (l. 622)",
                      'outside a cohort\'s calendar window and at zero capacity the SoH variable is fixed at 1.0 (l. 550)']}


# ------------------------------------------------------------------------------------------------------- item E: salvage
def case_inputs():
    import openpyxl  # noqa: E402  (third-party reader; read_only, data_only)
    xlsx = 'data/SRP1/SharedESS/SRP1_ESS.xlsx'
    _reg(xlsx)
    with open(xlsx, 'rb') as fh:
        wb = openpyxl.load_workbook(io.BytesIO(fh.read()), read_only=True, data_only=True)
    sc = next(wb['Scenarios'].iter_rows(values_only=True))
    n = int(sc[1])
    probs = [float(x) for x in sc[2:2 + n]]
    rows = list(wb['Investment Cost, Energy'].iter_rows(values_only=True))
    years = [int(y) for y in rows[0][1:] if y is not None]
    cost = {s: {years[j]: float(rows[1 + s][1 + j]) for j in range(len(years))} for s in range(n)}
    srp1 = _read_json('data/SRP1/SRP1.json')
    ess_params = _read_json('data/SRP1/SharedESS/SRP1_ESS_Params.json')
    return {'probs': probs, 'cost_energy': cost, 'srp1_years': srp1['Years'], 'srp1_discount': srp1['DiscountFactor'],
            'salvage_params': ess_params['salvage_value'], 'ess_params_sha256': INPUTS['data/SRP1/SharedESS/SRP1_ESS_Params.json']['sha256']}


def printed_salvage(rec, inputs, case, exponent='end_minus_y0'):
    """eq:salvage_credit / eq:salvage_remaining_life as printed, from the record's terminal ageing cells."""
    years = case['Years']
    ylist = sorted(int(y) for y in years)
    y0 = ylist[0]
    y_end = y0 + sum(int(v) for v in years.values())
    r = case['DiscountFactor']
    att = rec['ageing_trajectory_terminal']
    t_cal = rec['ess_ageing_baseline']['calendar_life_years']
    rb = ((rec.get('ess_ageing_readback_terminal') or {}).get('expected') or {}).get('floor_row_lower')
    soh_min = rb if rb is not None else rec['ess_ageing_baseline']['minimum_soh']
    terminal_block = len(ylist) - 1
    total, per_node = 0.0, {}
    for node, nd in att['nodes'].items():
        s = 0.0
        for c in nd.get('cells', []):
            if c['y'] != terminal_block:
                continue
            y_inv = int(c['investment_year'])
            cbar = sum(p * inputs['cost_energy'][k][y_inv] for k, p in enumerate(inputs['probs']))
            expo = (y_end - y0) if exponent == 'end_minus_y0' else (y_end - y_inv)
            t_rem = max(0.0, y_inv + t_cal - y_end)
            s += cbar * c['e_rated'] / (1 + r) ** expo * (t_rem / t_cal) * (c['soh_end'] - soh_min) / (1 - soh_min)
        per_node[node] = s
        total += s
    return total, per_node, {'y0': y0, 'y_end': y_end, 'r': r, 't_cal': t_cal, 'soh_min': soh_min}


def salvage_reproduction(records, inputs, exponent='end_minus_y0'):
    rows, worst = [], 0.0
    for r in records:
        rec = r['rec']
        if not rec.get('ageing_trajectory_terminal'):
            rows.append({'record': r['path'], 'skipped': 'no ageing_trajectory_terminal'})
            continue
        di = rec.get('derived_instance')
        if di:
            case = r['instance']
        else:
            case = {'Years': inputs['srp1_years'], 'DiscountFactor': inputs['srp1_discount']}
        recorded = (rec.get('recourse_components') or {}).get('terminal_salvage_value')
        if recorded is None:
            rows.append({'record': r['path'], 'skipped': 'no recorded terminal_salvage_value'})
            continue
        printed, per_node, par = printed_salvage(rec, inputs, case, exponent)
        diff = abs(printed - recorded)
        node_diff = max(abs(per_node[n] - (rec['ageing_trajectory_terminal']['nodes'][n].get('salvage_value') or 0.0))
                        for n in per_node) if per_node else 0.0
        worst = max(worst, diff, node_diff)
        rows.append({'record': r['path'], 'candidate': rec.get('candidate_canonical'), 'eval_key': rec.get('eval_key'),
                     'recorded_salvage': recorded, 'printed_formula': printed, 'abs_diff': diff,
                     'max_node_abs_diff': node_diff, 'params': par,
                     'ess_params_sha256_in_child_equals_file': rec.get('ess_params_sha256_in_child') == inputs['ess_params_sha256']})
    tested = [x for x in rows if 'skipped' not in x]
    ok = bool(tested) and all(x['abs_diff'] <= 1e-6 and x['max_node_abs_diff'] <= 1e-6
                              and x['ess_params_sha256_in_child_equals_file'] for x in tested)
    nonzero = [x for x in tested if abs(x['recorded_salvage']) > 1.0]
    return {'rows': rows, 'n_tested': len(tested), 'n_nonzero_salvage': len(nonzero), 'worst_abs_diff_eur': worst,
            'tolerance_eur': 1e-6, 'ok': ok and len(nonzero) > 0}


def _product_factors(node):
    if isinstance(node, ast.BinOp) and isinstance(node.op, ast.Mult):
        return _product_factors(node.left) + _product_factors(node.right)
    if isinstance(node, ast.Attribute):
        return [node.attr]
    if isinstance(node, ast.Name):
        return [node.id]
    return [ast.dump(node)]


def item_salvage(srcs, pr, prose, main_lines, records, inputs, search_info, code=CODE, exponent='end_minus_y0'):
    ids = ['SALV_DEF', 'SALV_TERMINAL_IDX', 'SALV_E_RATED', 'SALV_E_AVAIL', 'SALV_ABOVE_EOL', 'SALV_SUM', 'SALV_DISC_SUM',
           'SALV_DISC', 'SALV_AGE', 'SALV_REM', 'SALV_REM_FRAC', 'SALV_CBAR', 'SALV_NOT_IN_OBJ', 'SALV_EXPR', 'SALV_GET',
           'SALV_NET', 'SALV_PARAM_RECOVERY', 'SALV_PARAM_FLOOR', 'SALV_CBAR_READ']
    checks, anchors = anchor_checks(srcs, ids, code)
    for lab in ('eq:salvage_credit', 'eq:salvage_remaining_life'):
        checks.append(_ck(f'printed {lab} at its label = declared', pr[lab].get('matches_declared'), pr[lab].get('body')))
    app = pr['app:salvage_credit']
    checks.append(_ck('app:salvage_credit label present once; both equations inside it (after the label)',
                      app.get('ok') and pr['eq:salvage_credit'].get('label_line', 0) > app.get('label_line', 10 ** 9)
                      and pr['eq:salvage_remaining_life'].get('label_line', 0) > app.get('label_line', 10 ** 9),
                      [app.get('label_line'), pr['eq:salvage_credit'].get('label_line')]))
    src = srcs['shared_energy_storage_data.py']
    fn = _func_node(src, '_build_terminal_salvage_value_expression')
    aug = next((n for n in ast.walk(fn) if isinstance(n, ast.AugAssign) and getattr(n.target, 'id', None) == 'salvage_value'),
               None) if fn else None
    factors = sorted(_product_factors(aug.value)) if aug is not None else None
    expected = sorted(['terminal_discount', 'energy_recovery_fraction', 'expected_unit_cost', 'remaining_life_fraction',
                       'residual_energy'])
    checks.append(_ck('symbols: code product = terminal_discount x energy_recovery_fraction x expected_unit_cost x '
                      'remaining_life_fraction x residual_energy <-> 1/(1+r)^(yEnd-y0) x [1] x cbar x Trem/Tcal x '
                      'E (SoH - SoHmin)/(1 - SoHmin)', factors == expected, factors))
    sp = inputs['salvage_params']
    checks.append(_ck('configuration: energy_recovery_fraction = 1.0 and recycling_floor_fraction = 0.0 in '
                      'SRP1_ESS_Params.json (the two code factors absent from the print are identities)',
                      sp.get('enabled') is True and sp.get('energy_recovery_fraction') == 1.0
                      and sp.get('recycling_floor_fraction') == 0.0, sp))
    checks.append(_ck('discount: code exponent = sum of block widths = y^End - y_0 (contiguous blocks); r = DiscountFactor '
                      '0.02 = printed 2 %', inputs['srp1_discount'] == 0.02 and len(prose['salv_rate']) == 1,
                      inputs['srp1_discount']))
    rep = salvage_reproduction(records, inputs, exponent)
    checks.append(_ck(f'records: printed formula reproduces the recorded terminal salvage of every committed evaluation '
                      f'record ({rep["n_tested"]} records, {rep["n_nonzero_salvage"]} with a non-zero credit; worst '
                      f'|diff| {rep["worst_abs_diff_eur"]:.3e} EUR)', rep['ok'], rep['worst_abs_diff_eur']))
    obj_seg = '\n'.join(src.split('\n')[code['SALV_EXPR'][1] - 1:code['SALV_NOT_IN_OBJ'][1] + 1])
    checks.append(_ck('not optimized: the ESSO objective is the feasibility penalty only; the salvage is a separate '
                      'Expression (no network objective references it)',
                      'expr=model.feasibility_penalty,' in obj_seg and 'salvage' not in srcs['model_construction_helpers.py']
                      and 'salvage' not in srcs['network.py'], None))
    checks.append(_ck('not selected for: every planning search excludes salvage from F (OBJECTIVE_CONVENTION at its run '
                      'head)', all(v['salvage_excluded_from_F_text'] for v in search_info['runs'].values()),
                      {k: v['salvage_excluded_from_F_text'] for k, v in search_info['runs'].items()}))
    for k in ('salv_cbar', 'salv_no_converter', 'salv_reported', 'salv_net_form', 'salv_after'):
        checks.append(_ck(f'prose {k} present once', len(prose[k]) == 1, prose[k]))
    # L finding: SoH indexed by the date y^End while SoH is indexed by representative-year blocks elsewhere
    yend_lines = [i + 1 for i, ln in enumerate(main_lines) if 'y^{\\text{End}}' in ln or 'y^\\text{End}' in ln]
    soh_chain = [i + 1 for i, ln in enumerate(main_lines) if '\\label{eq:soh_chain}' in ln]
    has_2040 = [i + 1 for i, ln in enumerate(main_lines) if '2040' in ln]
    terminal_idx_line = src.split('\n')[code['SALV_TERMINAL_IDX'][1] - 1].strip()
    l_findings = [
        {'finding': ('SoH_{e,y^Inv,y^End} indexes SoH by the horizon-end date; elsewhere SoH is indexed by '
                     'representative-year blocks (eq:soh_chain, end-of-block state). The code uses the end-of-block SoH '
                     'of the final representative block (terminal_year_idx = len(years) - 1), which under eq:soh_chain is '
                     'the state at y^End; main.tex does not state the mapping or the numeric y^End'),
         'confirmed': bool(soh_chain) and terminal_idx_line == code['SALV_TERMINAL_IDX'][2] and not has_2040,
         'scope': {'lines_with_y^End': yend_lines, 'eq:soh_chain_label_line': soh_chain,
                   'lines_with_2040': has_2040}},
    ]
    notes = ['the remaining-life fraction carries min(., 1) (l. 921), inactive in every evaluated case (T^Rem <= T^Cal)',
             'the cohort calendar window uses round(t_cal / width) blocks (l. 680); exact for t_cal = 15 and widths 5 and 3',
             'salvage sensitivities are added to the post-convergence investment sensitivities '
             '(shared_resources_planning._get_operational_sensitivities), which feed only the retired Benders cut; '
             'not on the DFO search path']
    return {'item': 'app:salvage_credit', 'printed': {k: pr[k] for k in ('app:salvage_credit', 'eq:salvage_credit',
                                                                           'eq:salvage_remaining_life')},
            'anchors': anchors, 'checks': checks, 'code_factors': factors, 'reproduction': rep,
            'l_findings': l_findings, 'notes': notes}


# ------------------------------------------------------------------------------------------------------- records / heads
def load_records():
    tracked = set(_git(['ls-files', '-z', '--'] + list(RECORD_ROOTS)).split('\0'))
    out = []
    for root in RECORD_ROOTS:
        for path in sorted(glob.glob(os.path.join(root, '**', 'evaluation_record.json'), recursive=True)):
            if path not in tracked:
                continue
            rec = _read_json(path)
            inst = None
            di = rec.get('derived_instance')
            if di:
                ij = _read_json(di['case_path'])
                if INPUTS[di['case_path']]['sha256'] != di['case_sha256']:
                    raise SystemExit(f'refusing [instance hash check]: {di["case_path"]}')
                inst = {'Years': ij['Years'], 'DiscountFactor': ij['DiscountFactor']}
            out.append({'path': path, 'rec': rec, 'instance': inst})
    return out


def run_heads_identity():
    tracked = set(_git(['ls-files', '-z', '--'] + list(RECORD_ROOTS)).split('\0'))
    heads = {}
    for root in RECORD_ROOTS:
        for path in sorted(glob.glob(os.path.join(root, '**', 'campaign_results.json'), recursive=True)):
            if path not in tracked:
                continue
            cr = _read_json(path)
            h = cr.get('git_head_at_run')
            if h:
                heads.setdefault(h, []).append(path)
    per, prod_differ, rule_differ = {}, {}, {}
    for h, paths in heads.items():
        cert = any(p.startswith(CERT_ROOTS) for p in paths)
        dp = [f for f in PRODUCTION_FILES + CASE_FILES
              if _git(['diff', '--quiet', h, CODE_HEAD, '--', f], check=False).returncode != 0]
        dr = [f for f in RULE_FILES if _git(['diff', '--quiet', h, CODE_HEAD, '--', f], check=False).returncode != 0] \
            if cert else None
        per[h] = {'campaign_results': paths, 'certification_root': cert, 'production_and_case_files_differ': dp,
                  'rule_files_differ': dr}
        if dp:
            prod_differ[h] = dp
        if cert and dr:
            rule_differ[h] = dr
    return {'n_heads': len(heads), 'per_head': per, 'production_identical_all': not prod_differ,
            'production_differ': prod_differ, 'cert_rule_files_identical_all': not rule_differ,
            'cert_rule_files_differ': rule_differ}


def history():
    out = {}
    for item, specs in HISTORY.items():
        rows = []
        for f, s in specs:
            log = _git(['log', '--format=%H|%aI|%s', '-S', s, CODE_HEAD, '--', f]).strip().split('\n')
            commits = [dict(zip(('commit', 'date', 'subject'), ln.split('|', 2))) for ln in log if ln]
            for c in commits:
                c['subject'] = c['subject'][:140]
            rows.append({'file': f, 'string': s, 'commits_newest_first': commits,
                         'introduced_in': commits[-1] if commits else None, 'n_commits': len(commits)})
        out[item] = rows
    return out


def predecessor(base_lines):
    return {k: _flat_hits(base_lines, v) for k, v in PREDECESSOR.items()}


# ------------------------------------------------------------------------------------------------------- rating
def rate(item):
    if not all(c['ok'] for c in item['checks']):
        return 'MISMATCH'
    if any(f['confirmed'] for f in item['l_findings']):
        return 'L'
    return 'MATCH'


def evaluate_all(srcs, main_lines, records, inputs, search_info, heads, code=CODE, printed=PRINTED, exponent='end_minus_y0'):
    pr = printed_block(main_lines, printed)
    prose = prose_block(main_lines)
    items = {
        'eq:acceptance_tolerance': item_tau(srcs, pr, prose, main_lines, search_info, heads, code),
        'eq:interface_deviation': item_deviation(srcs, pr, prose, main_lines, code),
        'eq:interface_deviation_nonnegative': item_nonneg(srcs, pr, prose, main_lines, code),
        'eq:soh_initial': item_soh(srcs, pr, prose, records, code),
        'app:salvage_credit': item_salvage(srcs, pr, prose, main_lines, records, inputs, search_info, code, exponent),
    }
    for v in items.values():
        v['rating'] = rate(v)
    return items, pr, prose


# ------------------------------------------------------------------------------------------------------- controls
def negative_controls(srcs, main_lines, records, inputs, search_info, heads):
    out = []

    def rec(name, items, key):
        r = items[key]['rating']
        out.append({'control': name, 'item': key, 'rating': r, 'refused': r == 'MISMATCH',
                    'failed_checks': [c['check'] for c in items[key]['checks'] if not c['ok']][:6]})

    def mut_src(fname, line, old, new):
        s = copy.deepcopy(srcs)
        lines = s[fname].split('\n')
        assert old in lines[line - 1], (fname, line, old)
        lines[line - 1] = lines[line - 1].replace(old, new)
        s[fname] = '\n'.join(lines)
        return s

    # 1 altered code line number
    code1 = dict(CODE)
    f, ln, t, fn = CODE['DEV_P_LHS']
    code1['DEV_P_LHS'] = (f, ln + 1, t, fn)
    rec('NC1 altered code line number (DEV_P_LHS 1866 -> 1867)',
        evaluate_all(srcs, main_lines, records, inputs, search_info, heads, code=code1)[0], 'eq:interface_deviation')
    # 2 sign swap in the deviation row (line text unchanged elsewhere; the anchor text check and the canonical form fail)
    s2 = mut_src('model_construction_helpers.py', 1866, 'm.pg_adn[s_m, s_o, p] - m.expected_interface_pf_p[p]',
                 'm.expected_interface_pf_p[p] - m.pg_adn[s_m, s_o, p]')
    code2 = dict(CODE)
    code2['DEV_P_LHS'] = ('model_construction_helpers.py', 1866,
                          'return (m.expected_interface_pf_p[p] - m.pg_adn[s_m, s_o, p]', 'row18_interface_deviation_p_rule')
    it2 = evaluate_all(s2, main_lines, records, inputs, search_info, heads, code=code2)[0]
    rec('NC2 sign swap in the code deviation row (anchor text re-declared to the swapped line)', it2,
        'eq:interface_deviation')
    out[-1]['canonical_check_failed'] = not next(c for c in it2['eq:interface_deviation']['checks']
                                                 if c['check'].startswith('symbols / sign'))['ok']
    out[-1]['refused'] = out[-1]['refused'] and out[-1]['canonical_check_failed']
    # 3 nonnegativity domain negated
    s3 = mut_src('model_construction_helpers.py', 1967, 'domain=pe.NonNegativeReals', 'domain=pe.Reals')
    code3 = dict(CODE)
    code3['NN_P_DOWN'] = (CODE['NN_P_DOWN'][0], 1967, CODE['NN_P_DOWN'][2].replace('pe.NonNegativeReals', 'pe.Reals'),
                          CODE['NN_P_DOWN'][3])
    it3 = evaluate_all(s3, main_lines, records, inputs, search_info, heads, code=code3)[0]
    rec('NC3 deviation Var domain pe.Reals (anchor re-declared)', it3, 'eq:interface_deviation_nonnegative')
    # 4 initial SoH literal altered
    s4 = mut_src('shared_energy_storage_data.py', 714, 'prev_soh = 1.00', 'prev_soh = 0.95')
    code4 = dict(CODE)
    code4['SOH_INIT'] = (CODE['SOH_INIT'][0], 714, 'prev_soh = 0.95', CODE['SOH_INIT'][3])
    rec('NC4 initial SoH literal 0.95 (anchor re-declared)',
        evaluate_all(s4, main_lines, records, inputs, search_info, heads, code=code4)[0], 'eq:soh_initial')
    # 5 tau denominator altered
    s5 = mut_src('settling_criterion.py', 57, 'TAU = DELTA_R * R_REF / 4.0', 'TAU = DELTA_R * R_REF / 2.0')
    code5 = dict(CODE)
    code5['SC1_TAU'] = ('settling_criterion.py', 57, 'TAU = DELTA_R * R_REF / 2.0', None)
    rec('NC5 tau denominator 2.0 (anchor re-declared)',
        evaluate_all(s5, main_lines, records, inputs, search_info, heads, code=code5)[0], 'eq:acceptance_tolerance')
    # 6 delta_R literal altered
    s6 = mut_src('settling_criterion.py', 55, 'DELTA_R = 0.07', 'DELTA_R = 0.08')
    code6 = dict(CODE)
    code6['SC1_DELTA_R'] = ('settling_criterion.py', 55, 'DELTA_R = 0.08', None)
    rec('NC6 delta_R literal 0.08 (anchor re-declared)',
        evaluate_all(s6, main_lines, records, inputs, search_info, heads, code=code6)[0], 'eq:acceptance_tolerance')
    # 7 printed deviation sign swapped (the declared text follows the swap: only the symbol comparison can refuse)
    ml7 = list(main_lines)
    lab = next(i for i, ln in enumerate(ml7) if '\\label{eq:interface_deviation}' in ln)
    k = max(i for i in range(lab) if 'd^{+}_{i,s,t}' in ml7[i])
    j = max(i for i in range(lab) if 'd^{-}_{i,s,t}' in ml7[i])
    ml7[k], ml7[j] = ml7[k].replace('d^{+}', 'd^{-}'), ml7[j].replace('d^{-}', 'd^{+}')
    pr7 = dict(PRINTED)
    pr7['eq:interface_deviation'] = r'P^{\text{I}}_{i,s,t} - \bar{P}^{\text{I}}_{i,t} = d^{-}_{i,s,t} - d^{+}_{i,s,t}'
    rec('NC7 printed deviation with d+ and d- swapped (declared text re-declared)',
        evaluate_all(srcs, ml7, records, inputs, search_info, heads, printed=pr7)[0], 'eq:interface_deviation')
    # 8 printed salvage discount exponent y^End - y^Inv (numerical reproduction must refuse)
    rec('NC8 salvage discounted over y^End - y^Inv instead of y^End - y_0',
        evaluate_all(srcs, main_lines, records, inputs, search_info, heads, exponent='end_minus_yinv')[0],
        'app:salvage_credit')
    # 9 code salvage product without its remaining-life factor
    s9 = mut_src('shared_energy_storage_data.py', 903, '* remaining_life_fraction', '* 1.00')
    rec('NC9 code salvage product without remaining_life_fraction',
        evaluate_all(s9, main_lines, records, inputs, search_info, heads)[0], 'app:salvage_credit')
    return out


# ------------------------------------------------------------------------------------------------------- main
def _other_copy_running():
    ps = subprocess.run(['ps', '-A', '-o', 'pid=,command='], capture_output=True, text=True, check=True).stdout
    me = os.getpid()
    found = []
    for ln in ps.splitlines():
        tok = ln.split()
        if len(tok) < 3 or int(tok[0]) == me:
            continue
        if os.path.basename(tok[1]).startswith('python') and any(os.path.basename(t) == SCRIPT for t in tok[2:4]):
            found.append(ln.strip())
    return found


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out-dir', default=DEFAULT_OUT_DIR)
    args = ap.parse_args()
    out_dir = args.out_dir
    started = datetime.now(timezone.utc).isoformat()
    entries = sorted(os.listdir(out_dir)) if os.path.isdir(out_dir) else None
    if entries is None or any(e != 'launch.log' for e in entries):
        raise SystemExit(f'refusing [fresh-dir check]: {out_dir} must exist and hold nothing but launch.log (found {entries})')
    others = _other_copy_running()
    if others:
        raise SystemExit(f'refusing [single-copy check]: another copy is running: {others}')
    required = [os.path.join(CLONE, 'main.tex')] + list(SRC_FILES) + list(PRODUCTION_FILES) + list(CASE_FILES) + \
        [s['results'] for s in SEARCHES.values()] + list(RECORD_ROOTS)
    missing = [p for p in required if not os.path.exists(p)]
    if missing:
        raise SystemExit(f'refusing [capture-path check]: missing sources {missing}')

    ident, main_lines, base_lines = identity()
    if not ident['ok']:
        raise SystemExit(f'refusing [identity check]: {[c for c in ident["checks"] if not c["ok"]]}')
    srcs = code_sources()
    records = load_records()
    inputs = case_inputs()
    search_info = search_runs()
    heads = run_heads_identity()
    items, pr, prose = evaluate_all(srcs, main_lines, records, inputs, search_info, heads)
    hist = history()
    pred = predecessor(base_lines)
    controls = negative_controls(srcs, main_lines, records, inputs, search_info, heads)
    controls_ok = all(c['refused'] for c in controls)

    ratings = {k: v['rating'] for k, v in items.items()}
    n_m = sum(1 for r in ratings.values() if r == 'MISMATCH')
    prediction = {'expected': 'at most one M row (the interface-deviation sign convention)', 'ratings': ratings,
                  'n_refused_for_M_or_H_judgement': n_m,
                  'held_at_most_one_M': n_m <= 1,
                  'the_named_M_found': ratings['eq:interface_deviation'] == 'MISMATCH'}

    bad_mods = sorted(m for m in sys.modules if m.startswith(FORBIDDEN_MODULE_PREFIXES))
    gfail = GUARD.verify(0)
    if bad_mods or gfail or PICKLE_COUNTS['load'] or PICKLE_COUNTS['loads'] or not controls_ok:
        raise SystemExit(f'harness fault: production modules {bad_mods}; guard {gfail}; pickle {PICKLE_COUNTS}; '
                         f'controls {[c for c in controls if not c["refused"]]}')
    result = {
        'schema': 'w182a_equations_code/1', 'stage': 'P5.15 W182a (Addendum 82 v2, W182 item A1)',
        'started_utc': started, 'interpreter': sys.executable, 'identity': ident,
        'objective_convention': ('Q = gross_operational_cost (settlement excluded); net = Q - terminal salvage credit; '
                                 'no Q computed here; salvage values compared in EUR'),
        'rating_rule': ('computed: MISMATCH if any anchor / correspondence / record check fails (refused; the M or H '
                        'judgement is the report\'s); else L if a scoped wording finding is confirmed; else MATCH'),
        'items': items, 'printed_locations': pr, 'prose_lines': prose, 'search_runs': search_info,
        'run_heads_identity': heads, 'history_git_log_S': hist, 'predecessor_at_8f051a2': pred,
        'case_inputs': {k: v for k, v in inputs.items() if k != 'cost_energy'},
        'n_evaluation_records': len(records),
        'negative_controls': controls, 'negative_controls_all_refused': controls_ok, 'prediction_A1': prediction,
        'guards': {'solve_profile_guard': {'label': GUARD.label, 'counts': GUARD.counts, 'verify_0_failures': gfail},
                   'pickle_counts': PICKLE_COUNTS, 'production_modules_imported': bad_mods},
        'finished_utc': datetime.now(timezone.utc).isoformat(),
    }
    out_json = os.path.join(out_dir, 'w182a_equations_code.json')
    with open(out_json, 'w', encoding='utf-8') as f:
        f.write(GRIO.dumps(result, indent=1, sort_keys=False) + '\n')
    manifest = {'outputs': {out_json: _sha(out_json)}, 'script': {SCRIPT: _sha(SCRIPT)},
                'inputs': {p: INPUTS[p] for p in sorted(INPUTS)}, 'n_inputs': len(INPUTS),
                'n_uncommitted_inputs': sum(1 for v in INPUTS.values() if v['status'] == 'UNCOMMITTED')}
    with open(os.path.join(out_dir, 'manifest_sha256.json'), 'w', encoding='utf-8') as f:
        f.write(GRIO.dumps(manifest, indent=1, sort_keys=True) + '\n')

    # ------------------------------------------------------------------------------------------------ console
    print(f'clone HEAD {ident["clone_head"]}; {ident["n_commits_base_to_head"]} commits after 8f051a2; main.tex sha256 '
          f'{ident["main_sha256"]} ({ident["main_n_lines"]} lines)')
    print(f'repo HEAD {ident["repo_head"]}; last production/rule commit {ident["last_production_commit"]}; '
          f'dirty tracked {ident["repo_dirty_tracked"]}')
    for k, v in items.items():
        print(f'{k}: {v["rating"]}')
        for c in v['checks']:
            print(f'    [{"ok" if c["ok"] else "FAIL"}] {c["check"]}' + ('' if c['ok'] else f' -- {c["detail"]}'))
        for lf in v['l_findings']:
            print(f'    L-finding confirmed={lf["confirmed"]}: {lf["finding"]}')
        for n in v.get('notes', []):
            print(f'    note: {n}')
    for k, v in pr.items():
        print(f'  printed {k}: label l. {v.get("label_line")} begin {v.get("begin_line")} end {v.get("end_line")}')
    rep = items['app:salvage_credit']['reproduction']
    for r in rep['rows']:
        if 'skipped' in r:
            print(f'  salvage {r["record"]}: skipped ({r["skipped"]})')
        elif abs(r['recorded_salvage']) > 1.0:
            print(f'  salvage {r["record"].split("/")[-3]}: recorded {r["recorded_salvage"]:.6f} printed '
                  f'{r["printed_formula"]:.6f} |diff| {r["abs_diff"]:.2e}')
    print(f'  salvage reproduction: {rep["n_tested"]} records, {rep["n_nonzero_salvage"]} non-zero, worst '
          f'{rep["worst_abs_diff_eur"]:.3e} EUR')
    for name, r in search_info['runs'].items():
        print(f'  search {name}: head {r["git_head_at_run"][:8]} {r["date"]} rule files {r["settling_criterion_files_at_head"]} '
              f'tau tokens {r["tau_token_counts_at_head"]} sigma_q rule {r["resolution_rule_sigma_q_present"]} '
              f'tau-intro ancestor {r["tau_intro_ab0bcbc9_is_ancestor"]}')
    print(f'  run heads: {heads["n_heads"]}; production+case identical to code head at all: '
          f'{heads["production_identical_all"]} {heads["production_differ"]}; rule files identical at all '
          f'certification heads: {heads["cert_rule_files_identical_all"]}')
    for item, rows in hist.items():
        for r in rows:
            intro = r['introduced_in']
            print(f'  history {item} [{r["file"]}] "{r["string"][:50]}": introduced {intro["commit"][:8]} {intro["date"][:10]} '
                  f'({r["n_commits"]} changing commits: {[c["commit"][:8] for c in r["commits_newest_first"]]})')
    print(f'  predecessor at 8f051a2 (line hits): {pred}')
    for c in controls:
        print(f'  control {c["control"]}: rating {c["rating"]} refused {c["refused"]}')
    print(f'prediction A1: {prediction}')
    print(f'inputs {len(INPUTS)} ({manifest["n_uncommitted_inputs"]} uncommitted)')
    print(f'guards: {result["guards"]}')
    print(f'output {out_json} sha256 {_sha(out_json)}')


if __name__ == '__main__':
    main()
