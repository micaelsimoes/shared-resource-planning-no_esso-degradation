"""P5.15 Addendum 44 ruling 5, task W74 item 2 -- FEASIBILITY CHECK for the fixed-configuration polish re-measurement.
READ-ONLY, ZERO SOLVES. Nothing is built, solved or re-run; no committed artifact is modified.

Question: can the six certified alpha-row cells (campaign s53_alpha_row_v25, spec 70965374) be re-polished "on the existing
certified cells, NO NEW ADMM RUNS"? That needs every TSO/DSO block at its certified terminal state. This script records:

  A  per cell, what is persisted for which agent (TSO / DSO / ESSO): presence, size, sha256, pair-manifest coverage and
     git tracking of `certified_models.pkl`, `esso_models_s39_D.pkl` (unpickled: node ids and component types -- the ESSO
     models; unpickling is not a solve), `results/FrozenSMOPF/*.pkl`, the production terminal workbook (sheet names and
     row counts), `hull_bound_detail.json` (the hull, per descriptor), `post_certification.json` (per-block f_i before /
     after, solve profile), and the spec's `post_certification` request;
  B  per cell and per block (20 TSO + 60 DSO), the IPOPT main log of the cell's run directory
     (data/SRP1/Results/P56A/evals/p515s44_s53_alpha_row_v25_<key16>_run/logs, git-ignored): number of solves, the
     classification of each solve by its starting multiplier norm (||z_L||_inf = 1.0 exactly = IPOPT's default, i.e. NO
     multiplier warm start: the initialisation solve and the hull-polish solve; every ADMM-cycle solve is multiplier-warm),
     the expectation n_solves = cycles_run + 2 (init + cycles + polish), and for the TERMINAL ADMM solve (second-to-last)
     and the POLISH solve (last) the objective scaling factor, final barrier mu, complementarity pair count, iterations,
     exit and final unscaled objective; the recovery-log solve counts; whether any committed manifest covers each log;
  C  the solve-count arithmetic of a six-cell polish from the committed solve profiles;
  D  which solve W73's TSO barrier-gap figures were read from: W73 took the LAST solve of each appended TSO log
     (p515_s53_alpha_row_reclass.parse_tso_log, `parts[-2]`) and ASSUMED it was the terminal ADMM cycle's; section B
     classifies it as the hull-polish solve. D re-evaluates W73's own formula on both solves, TSO only, with W73's own
     per-block weights (alpha_row_reclass.json, verified against the W73 manifest) cross-checked against the committed
     `admm_block_weight` of multiscenario_terminal.json; the polish-solve sum must reproduce W73's committed figure
     exactly (positive control). DSO gap estimates are NOT computed here (ruling 5's re-measurement task).

Formulas (preserved here, not only in prose):
  n_pairs = var_lb_only + 2 var_lb_ub + var_ub_only + ineq_lb_only + 2 ineq_lb_ub + ineq_ub_only   (problem-size lines of
            the solve; identical to W73 FORMULAS['tso_barrier_gap_estimate'])
  solve classes: #1 = initialisation, #2 .. #n-1 = ADMM cycles 1 .. cycles_run, #n = hull-polish primary attempt
  gap_w(solve) = sum over TSO blocks of weight_b x n_pairs x mu_last / obj_scale   (W73 FORMULAS['tso_barrier_gap_
            estimate'], unchanged; an order-of-magnitude barrier-offset estimate in base-objective EUR, weighted)
  polish solves per cell = n_blocks + retries_beyond_one_per_block (committed solve_profile); per block <= 3 attempts
            (primary + tier-1 + tier-2, network._run_smopf).

COMMANDS (repo root, canonical interpreter, attached, both streams, noclobber):
  set -o noclobber
  /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s53_w74_polish_feasibility.py --run \
      > data/SRP1/Results/P515S53/polish_feasibility_w74_launch.log 2>&1
  /Users/micaelsimoes/miniconda3/envs/opf_env_py311/bin/python -u p515_s53_w74_polish_feasibility.py --manifest
"""

import glob
import hashlib
import json
import os
import pickle
import re
import subprocess
import sys
import time
from datetime import datetime, timezone

REPO = os.path.dirname(os.path.abspath(__file__))
if REPO not in sys.path:
    sys.path.insert(0, REPO)

from p513_solve_profile_guard import SolveProfileGuard  # noqa: E402

GUARD = SolveProfileGuard(permitted=(), label='P5.15 W74 polish feasibility (read-only)').install()

ALPHA_ROOT = os.path.join('data', 'SRP1', 'Results', 'P515S53', 'alpha_row', 'campaign_s53_alpha_row_v25')
SPEC_REL = os.path.join(ALPHA_ROOT, 'campaign_spec_s53_alpha_row_v25_70965374.json')
CELLS = {'x0_a0p00': ('b123cd978794d690_x0_a0p00', 2), 'x0_a0p10': ('7516903c91153a29_x0_a0p10', 3),
         'x0_a0p25': ('62b46280a65f7744_x0_a0p25', 3), 'x0_a0p50': ('7d53b6f21b686a44_x0_a0p50', 1),
         'x0_a1p00': ('1bb2d63a07273887_x0_a1p00', 2), 'n7_4h_e1_a0p50': ('711fce9aa74d6878_n7_4h_e1_a0p50', 1)}
RUN_DIR = os.path.join('data', 'SRP1', 'Results', 'P56A', 'evals', 'p515s44_s53_alpha_row_v25_{key16}_run', 'logs')
NETWORK_OF = {'TSO': 'case9', 5: 'case33_1', 7: 'case33_2', 9: 'case33_3'}
YEARS = (2025, 2028, 2031, 2034, 2037)
DAYS = ('Spring', 'Summer', 'Autumn', 'Winter')
W73_JSON = os.path.join('data', 'SRP1', 'Results', 'P515S53', 'alpha_row', 'reclass_w73', 'alpha_row_reclass.json')
W73_MANIFEST = os.path.join('data', 'SRP1', 'Results', 'P515S53', 'alpha_row', 'reclass_w73',
                            'reclass_w73_manifest_sha256.json')

OUT = os.path.join('data', 'SRP1', 'Results', 'P515S53', 'polish_feasibility_w74')
OUT_JSON = os.path.join(OUT, 'polish_feasibility_w74.json')
OUT_MANIFEST = os.path.join(OUT, 'polish_feasibility_w74_manifest_sha256.json')
LAUNCH_LOG = os.path.join('data', 'SRP1', 'Results', 'P515S53', 'polish_feasibility_w74_launch.log')


def _abs(rel):
    return os.path.join(REPO, rel)


def _sha(rel):
    h = hashlib.sha256()
    with open(_abs(rel), 'rb') as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()


def _load(rel):
    with open(_abs(rel)) as handle:
        return json.load(handle)


def _tracked(rel):
    return subprocess.run(['git', 'ls-files', '--error-unmatch', rel], cwd=REPO, capture_output=True).returncode == 0


# ----------------------------------------------------------------------------------------------------------------------
def _file_entry(rel, manifest):
    present = os.path.isfile(_abs(rel))
    return {'path': rel, 'present': present, 'size_bytes': os.path.getsize(_abs(rel)) if present else None,
            'sha256': _sha(rel) if present else None,
            'in_pair_manifest': (rel in manifest) if present else False,
            'manifest_hash_matches': (manifest.get(rel) == _sha(rel)) if present and rel in manifest else None,
            'git_tracked': _tracked(rel) if present else False}


def inventory(label, eval_dir, pair, spec_entry):
    manifest = _load(os.path.join(ALPHA_ROOT, f'pair_{pair}_manifest_sha256.json'))
    ed = os.path.join(ALPHA_ROOT, 'evals', eval_dir)
    out = {'post_certification_request_in_spec': spec_entry.get('post_certification')}
    out['certified_models_pkl'] = _file_entry(os.path.join(ed, 'certified_models.pkl'), manifest)
    pkls_anywhere = sorted(os.path.relpath(p, REPO) for p in glob.glob(os.path.join(_abs(ed), '**', '*.pkl'),
                                                                       recursive=True))
    out['all_pkl_files_in_eval_dir'] = pkls_anywhere
    esso = _file_entry(os.path.join(ed, 'esso_models_s39_D.pkl'), manifest)
    if esso['present']:
        with open(_abs(esso['path']), 'rb') as handle:
            obj = pickle.load(handle)
        esso['content'] = {'type': type(obj).__name__,
                           'keys': sorted(str(k) for k in obj) if isinstance(obj, dict) else None,
                           'value_types': sorted({f'{type(v).__module__}.{type(v).__name__}' for v in obj.values()})
                           if isinstance(obj, dict) else None}
        del obj
    out['esso_models_pkl'] = esso
    out['frozen_smopf_pkls'] = [_file_entry(p, manifest) for p in pkls_anywhere if '/FrozenSMOPF/' in p]
    for fe in out['frozen_smopf_pkls']:
        m = re.search(r'(TSO|DSO)_(?:node(\d+)_)?(case\w+?)_(\d{4})_(\w+?)_cycle(\d+)\.pkl$', fe['path'])
        fe['parsed_from_name'] = ({'agent': m.group(1), 'node': m.group(2), 'network': m.group(3),
                                   'year': int(m.group(4)), 'day': m.group(5), 'cycle': int(m.group(6))} if m else None)
    wb = _file_entry(os.path.join(ed, 'results', 'SRP1_distributed_terminal.xlsx'), manifest)
    if wb['present']:
        import openpyxl
        book = openpyxl.load_workbook(_abs(wb['path']), read_only=True)
        wb['sheets'] = {n: {'max_row': book[n].max_row, 'max_column': book[n].max_column} for n in book.sheetnames}
        book.close()
    out['terminal_workbook'] = wb
    hb = _file_entry(os.path.join(ed, 'hull_bound_detail.json'), manifest)
    if hb['present']:
        rows = _load(hb['path'])
        hb['n_descriptors'] = len(rows)
        hb['descriptor_keys'] = sorted(rows[0]) if rows else None
        hb['by_channel'] = {c: sum(1 for r in rows if r['channel'] == c) for c in sorted({r['channel'] for r in rows})}
        del rows
    out['hull_bound_detail'] = hb
    pce = _file_entry(os.path.join(ed, 'post_certification.json'), manifest)
    pc = _load(pce['path'])
    hp = pc['hull_polish_full']
    pce['persisted_models'] = pc.get('persisted_models')
    pce['n_blocks'] = hp['n_blocks']
    pce['n_hull_descriptors'] = hp['n_hull_descriptors']
    pce['per_block_has_f_before_and_after'] = all(
        isinstance(b.get('weighted_base_objective_before'), float) and isinstance(b.get('weighted_base_objective_after'),
                                                                                 float) for b in hp['per_block'])
    pce['n_blocks_by_agent'] = {a: sum(1 for b in hp['per_block'] if b['agent'] == a)
                                for a in sorted({b['agent'] for b in hp['per_block']})}
    pce['solve_profile'] = hp['solve_profile']
    pce['gate_pass_raw'] = repr(hp['gate']['pass'])
    out['post_certification'] = pce
    for name in ('multiscenario_terminal.json', 'response_terminal.json', 'component_levels_terminal.json'):
        fe = _file_entry(os.path.join(ed, name), manifest)
        fe['top_level_keys'] = sorted(_load(fe['path'])) if fe['present'] else None
        out[name] = fe
    out['certified_tso_dso_models_persisted'] = out['certified_models_pkl']['present']
    return out


# ----------------------------------------------------------------------------------------------------------------------
def _solves(path):
    with open(_abs(path), errors='replace') as handle:
        txt = handle.read()
    segs = [s for s in re.split(r'(?=This is Ipopt version)', txt) if s.startswith('This is Ipopt')]
    out = []
    for s in segs:
        def f(p, cast=float, last=False):
            m = re.findall(p, s)
            return None if not m else cast(m[-1] if last else m[0])
        n = {k: f(rf'{v}:\s+(\d+)', int) for k, v in (
            ('var_lb_only', r'variables with only lower bounds'), ('var_lb_ub', r'variables with lower and upper bounds'),
            ('var_ub_only', r'variables with only upper bounds'),
            ('ineq_lb_only', r'inequality constraints with only lower bounds'),
            ('ineq_lb_ub', r'inequality constraints with lower and upper bounds'),
            ('ineq_ub_only', r'inequality constraints with only upper bounds'))}
        pairs = (None if any(v is None for v in n.values()) else
                 n['var_lb_only'] + 2 * n['var_lb_ub'] + n['var_ub_only'] + n['ineq_lb_only'] + 2 * n['ineq_lb_ub']
                 + n['ineq_ub_only'])
        out.append({'obj_scale': f(r'objective scaling factor = (\S+)'),
                    'z_L_inf_start': f(r'\|\|curr_z_L\|\|_inf = (\S+)'),
                    'mu_last': f(r'Current barrier parameter mu = (\S+)', float, True),
                    'iterations': f(r'Number of Iterations\.+: (\d+)', int, True),
                    'objective_unscaled_final': f(r'Objective\.+:\s+\S+\s+(\S+)', float, True),
                    'exit': f(r'EXIT: (.*)', str, True), 'n_var': f(r'Total number of variables\.+:\s+(\d+)', int),
                    'n_pairs': pairs, **n})
    return out


def log_survey(label, eval_dir, cycles_run, w73_manifest):
    key16 = eval_dir.split('_')[0]
    d = RUN_DIR.format(key16=key16)
    blocks = {}
    for agent in ('TSO', 5, 7, 9):
        net = NETWORK_OF[agent]
        for y in YEARS:
            for day in DAYS:
                rel = os.path.join(d, f'optim_log_{net}_{y}_{day}.log')
                block = f"{'TSO' if agent == 'TSO' else f'DSO{agent}'}|{y}|{day}"
                if not os.path.isfile(_abs(rel)):
                    blocks[block] = {'log': rel, 'present': False}
                    continue
                s = _solves(rel)
                n = len(s)
                recov = {suf: (len(_solves(os.path.join(d, f'optim_log_{net}_{y}_{day}_{suf}.log')))
                               if os.path.isfile(_abs(os.path.join(d, f'optim_log_{net}_{y}_{day}_{suf}.log'))) else 0)
                         for suf in ('recovery', 'recovery_tier2')}
                cold = [i + 1 for i, x in enumerate(s) if x['z_L_inf_start'] == 1.0]
                blocks[block] = {
                    'log': rel, 'present': True, 'sha256_at_read': _sha(rel),
                    'in_w73_manifest': rel in w73_manifest,
                    'w73_manifest_hash_matches': (w73_manifest.get(rel) == _sha(rel)) if rel in w73_manifest else None,
                    'n_solves': n, 'n_solves_minus_cycles_run': n - cycles_run,
                    'cold_multiplier_solves': cold,
                    'classification_holds': n == cycles_run + 2 and cold == [1, n],
                    'terminal_admm_solve': s[-2] if n >= 2 else None, 'polish_solve': s[-1] if n else None,
                    'recovery_log_solves': recov}
    return blocks


def _summ(vals):
    v = [x for x in vals if x is not None]
    return {'n': len(v), 'min': min(v) if v else None, 'max': max(v) if v else None,
            'distinct': len(set(v))}


def summarise_logs(blocks):
    out = {}
    for agent in ('TSO', 'DSO'):
        bl = [b for k, b in blocks.items() if k.startswith(agent) and b.get('present')]
        term = [b['terminal_admm_solve'] for b in bl]
        pol = [b['polish_solve'] for b in bl]
        out[agent] = {
            'n_logs': len(bl), 'classification_holds_all': all(b['classification_holds'] for b in bl),
            'n_in_any_committed_manifest': sum(1 for b in bl if b['in_w73_manifest']),
            'terminal_admm': {k: _summ([t[k] for t in term]) for k in ('obj_scale', 'mu_last', 'n_pairs')},
            'terminal_admm_exits': sorted({t['exit'] for t in term}),
            'polish': {k: _summ([p[k] for p in pol]) for k in ('obj_scale', 'mu_last', 'n_pairs')},
            'polish_primary_exits': sorted({p['exit'] for p in pol}),
            'obj_scale_terminal_equals_polish_blocks': sum(1 for t, p in zip(term, pol)
                                                           if t['obj_scale'] == p['obj_scale']),
            'mu_last_terminal_equals_polish_blocks': sum(1 for t, p in zip(term, pol) if t['mu_last'] == p['mu_last']),
            'n_pairs_terminal_equals_polish_blocks': sum(1 for t, p in zip(term, pol) if t['n_pairs'] == p['n_pairs']),
            'recovery_log_solves_total': sum(sum(b['recovery_log_solves'].values()) for b in bl)}
    return out


# ----------------------------------------------------------------------------------------------------------------------
def run():
    started = time.time()
    if os.path.exists(_abs(OUT_JSON)):
        raise SystemExit(f'REFUSED: output exists (write-once): {OUT_JSON}')
    spec = _load(SPEC_REL)
    by_label = {c['label']: c for c in spec['candidates']}
    w73 = _load(W73_MANIFEST)
    cells, inputs = {}, {SPEC_REL: _sha(SPEC_REL), W73_MANIFEST: _sha(W73_MANIFEST)}
    for label, (eval_dir, pair) in CELLS.items():
        rec_rel = os.path.join(ALPHA_ROOT, 'evals', eval_dir, 'evaluation_record.json')
        rec = _load(rec_rel)
        cycles = rec.get('cycles_run')
        inv = inventory(label, eval_dir, pair, by_label[label])
        blocks = log_survey(label, eval_dir, cycles, w73)
        cells[label] = {'eval_dir': eval_dir, 'eval_key': by_label[label]['eval_key'],
                        'candidate_key': by_label[label]['key'], 'status': rec.get('status'), 'cycles_run': cycles,
                        'inventory': inv, 'log_summary': summarise_logs(blocks), 'log_blocks': blocks}
        print(f"[W74-F] {label}: certified_models_pkl={inv['certified_models_pkl']['present']} "
              f"esso_pkl={inv['esso_models_pkl']['present']} frozen_pkls={len(inv['frozen_smopf_pkls'])} "
              f"workbook={inv['terminal_workbook']['present']} cycles={cycles} "
              f"TSO_cls={cells[label]['log_summary']['TSO']['classification_holds_all']} "
              f"DSO_cls={cells[label]['log_summary']['DSO']['classification_holds_all']}", flush=True)
    prof = {lab: c['inventory']['post_certification']['solve_profile'] for lab, c in cells.items()}
    n_blocks = {lab: c['inventory']['post_certification']['n_blocks'] for lab, c in cells.items()}
    arithmetic = {
        'blocks_per_cell': n_blocks, 'blocks_total': sum(n_blocks.values()),
        'primary_attempts_min': sum(n_blocks.values()),
        'attempts_max_3_per_block': 3 * sum(n_blocks.values()),
        'alpha_row_observed_retries': {lab: p['retries_beyond_one_per_block'] for lab, p in prof.items()},
        'alpha_row_observed_solves_total': sum(p['observed']['permitted_solve'] for p in prof.values()),
        'definition': 'per cell: n_blocks primary attempts + retries (<= 2 per block: tier-1 cold, tier-2 cold+adaptive)'}
    w73_json = _load(W73_JSON)
    if w73.get(W73_JSON) != _sha(W73_JSON):
        raise SystemExit(f'{W73_JSON} does not verify against the W73 manifest')
    inputs[W73_JSON] = _sha(W73_JSON)
    section_d = {}
    for label, c in cells.items():
        w73_blocks = w73_json['tso_unit_cell_investigation']['per_cell'][label]['blocks']
        ms_rel = os.path.join(ALPHA_ROOT, 'evals', c['eval_dir'], 'multiscenario_terminal.json')
        ms_blocks = _load(ms_rel)['blocks']
        inputs[ms_rel] = _sha(ms_rel)
        rows, sums = [], {'terminal_admm': 0.0, 'polish': 0.0}
        for block, b in c['log_blocks'].items():
            if not block.startswith('TSO'):
                continue
            wb = w73_blocks[block]
            w = wb['weight']
            row = {'block': block, 'weight_w73': w, 'weight_multiscenario_terminal': ms_blocks[block]['admm_block_weight'],
                   'w73_n_pairs': wb['n_pairs'], 'w73_mu_last': wb['mu_last'], 'w73_obj_scale': wb['obj_scale'],
                   'w73_iterations': wb['iterations']}
            for which, key in (('terminal_admm', 'terminal_admm_solve'), ('polish', 'polish_solve')):
                sv = b[key]
                g = w * sv['n_pairs'] * sv['mu_last'] / sv['obj_scale']
                row[which] = {'n_pairs': sv['n_pairs'], 'mu_last': sv['mu_last'], 'obj_scale': sv['obj_scale'],
                              'iterations': sv['iterations'], 'gap_w_eur': g}
                sums[which] += g
            row['w73_values_equal_polish_solve'] = (wb['n_pairs'] == row['polish']['n_pairs']
                                                    and wb['mu_last'] == row['polish']['mu_last']
                                                    and wb['obj_scale'] == row['polish']['obj_scale']
                                                    and wb['iterations'] == row['polish']['iterations'])
            rows.append(row)
        committed = w73_json['tso_unit_cell_investigation']['per_cell'][label]['TSO_barrier_gap_est_w_eur']
        section_d[label] = {
            'w73_committed_TSO_gap_w_eur': committed, 'recomputed_on_polish_solve': sums['polish'],
            'recomputed_on_terminal_admm_solve': sums['terminal_admm'],
            'positive_control_polish_reproduces_w73': abs(sums['polish'] - committed) <= 1e-9 * max(1.0, abs(committed)),
            'w73_values_equal_polish_solve_all_blocks': all(r['w73_values_equal_polish_solve'] for r in rows),
            'weights_w73_equal_multiscenario_terminal': all(r['weight_w73'] == r['weight_multiscenario_terminal']
                                                            for r in rows),
            'blocks': rows}
    d_pair = {k: section_d[k]['recomputed_on_terminal_admm_solve'] for k in ('x0_a0p50', 'n7_4h_e1_a0p50')}
    section_d['_x0_a0p50_minus_unit'] = {
        'w73_committed': (section_d['x0_a0p50']['w73_committed_TSO_gap_w_eur']
                          - section_d['n7_4h_e1_a0p50']['w73_committed_TSO_gap_w_eur']),
        'terminal_admm_solve': d_pair['x0_a0p50'] - d_pair['n7_4h_e1_a0p50']}
    payload = {'task': 'P5.15 Addendum 44 ruling 5 / W74 item 2: polish re-measurement feasibility (read-only)',
               'utc': datetime.now(timezone.utc).isoformat(),
               'git_head': subprocess.run(['git', 'rev-parse', 'HEAD'], cwd=REPO, capture_output=True,
                                          text=True).stdout.strip(),
               'script_sha256': _sha(os.path.basename(__file__)), 'spec': SPEC_REL, 'cells': cells,
               'solve_arithmetic': arithmetic, 'w73_solve_attribution': section_d, 'inputs_sha256': inputs}
    guard = GUARD.verify(0)
    payload['guard'] = {'counts': dict(GUARD.counts), 'verify_0_failures': guard}
    payload['wall_s'] = time.time() - started
    os.makedirs(_abs(OUT), exist_ok=True)
    with open(_abs(OUT_JSON), 'x') as handle:
        json.dump(payload, handle, indent=1)
    print(f"[W74-F] arithmetic {json.dumps(arithmetic)}", flush=True)
    for label, v in section_d.items():
        if label.startswith('_'):
            print(f'[W74-F] D {label}: {v}', flush=True)
        else:
            print(f"[W74-F] D {label}: W73 {v['w73_committed_TSO_gap_w_eur']:.2f} polish {v['recomputed_on_polish_solve']:.2f}"
                  f" terminal-ADMM {v['recomputed_on_terminal_admm_solve']:.2f} control "
                  f"{v['positive_control_polish_reproduces_w73']} w73==polish {v['w73_values_equal_polish_solve_all_blocks']}"
                  f" weights {v['weights_w73_equal_multiscenario_terminal']}", flush=True)
    print(f'[W74-F] wrote {OUT_JSON}; guard {dict(GUARD.counts)} verify(0) {guard}; wall {payload["wall_s"]:.1f}s',
          flush=True)
    GUARD.uninstall()
    sys.exit(0 if not guard else 1)


def manifest():
    if os.path.exists(_abs(OUT_MANIFEST)):
        raise SystemExit(f'REFUSED: {OUT_MANIFEST} exists (write-once)')
    out = _load(OUT_JSON)
    m = {OUT_JSON: _sha(OUT_JSON), LAUNCH_LOG: _sha(LAUNCH_LOG), os.path.basename(__file__): _sha(os.path.basename(__file__))}
    bad = [] if out['script_sha256'] == m[os.path.basename(__file__)] else [os.path.basename(__file__)]
    m.update(out['inputs_sha256'])
    for c in out['cells'].values():
        for b in c['log_blocks'].values():   # the IPOPT logs: NOT in any campaign manifest; hashed at read time here
            if b.get('present'):
                if _sha(b['log']) != b['sha256_at_read']:
                    bad.append(b['log'])
                m[b['log']] = b['sha256_at_read']
    if bad:
        raise SystemExit(f'REFUSED: changed since the run: {bad[:5]}')
    with open(_abs(OUT_MANIFEST), 'x') as handle:
        json.dump(m, handle, indent=1)
    failures = GUARD.verify(0)
    print(f'[W74-F] wrote {OUT_MANIFEST}: {len(m)} entries; guard verify(0) {failures}', flush=True)
    GUARD.uninstall()
    sys.exit(0 if not failures else 1)


if __name__ == '__main__':
    if sys.argv[1:] == ['--run']:
        run()
    elif sys.argv[1:] == ['--manifest']:
        manifest()
    else:
        raise SystemExit('usage: --run | --manifest')
